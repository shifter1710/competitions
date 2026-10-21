"""Small GTO report contract: the same filtered award rows drive every output."""
import re
from io import BytesIO
from urllib.parse import urlencode

from openpyxl import Workbook
from openpyxl.styles import Font

from src.gto import gto_stage_label
from src.gto import gto_status_label
from src.records import sanitize_spreadsheet_value
from src.storage.gto import GTO_AWARDS
from src.storage.helpers import GTO_STAGE_MAX
from src.storage.helpers import gto_year_max
from src.storage.helpers import GTO_YEAR_MIN

FILTER_KEYS = ('year', 'institute', 'group', 'stage', 'award', 'q')
FILTER_LABELS = {
    'year': 'Год записи ГТО',
    'institute': 'Институт',
    'group': 'Группа',
    'stage': 'Ступень',
    'award': 'Награда',
    'q': 'Студент или УИН',
}
EXPORT_COLUMNS = (
    'ID награждения',
    'ID студента',
    'Студент',
    'УИН',
    'Институт (текущий)',
    'Группа (текущая)',
    'Год записи ГТО',
    'Ступень',
    'Награда',
    'Дата награждения',
    'Сертификаты',
    'Защищённые ссылки',
)


def parse_gto_filters(args):
    filters = {key: '' for key in FILTER_KEYS}
    unknown = set(args) - set(FILTER_KEYS)
    if unknown:
        return filters, 'Неизвестный фильтр отчёта ГТО.'
    for key in FILTER_KEYS:
        values = args.getlist(key) if hasattr(args, 'getlist') else args.get(key, [])
        values = [values] if isinstance(values, str) else values
        if any(ord(char) < 32 or ord(char) == 127 for value in values for char in str(value)):
            return filters, f'Фильтр «{FILTER_LABELS[key]}» содержит недопустимые символы.'
        normalized = [str(value).strip() for value in values]
        if len(set(normalized)) > 1:
            return filters, f'Для фильтра «{FILTER_LABELS[key]}» укажите одно значение.'
        filters[key] = normalized[0] if normalized else ''
    for key, minimum, maximum in (('year', GTO_YEAR_MIN, gto_year_max()), ('stage', 1, GTO_STAGE_MAX)):
        value = filters[key]
        if value and (
            not re.fullmatch(r'[0-9]{1,4}' if key == 'year' else r'[0-9]{1,2}', value)
            or not minimum <= int(value) <= maximum
        ):
            return filters, f'Фильтр «{FILTER_LABELS[key]}»: укажите число от {minimum} до {maximum}.'
        if value:
            filters[key] = str(int(value))
    if filters['award'] and filters['award'] not in GTO_AWARDS:
        return filters, 'Указан недопустимый тип награды ГТО.'
    return filters, None


def gto_matches(row, filters):
    for key, column in (
        ('year', 'year'),
        ('institute', 'institute'),
        ('group', 'group_name'),
        ('stage', 'stage'),
        ('award', 'status'),
    ):
        if filters[key] and str(row.get(column) or '') != filters[key]:
            return False
    query = filters['q'].casefold()
    return not query or any(query in str(row.get(key) or '').casefold() for key in ('full_name', 'gto_uin'))


def build_gto_report(records, args):
    filters, error = parse_gto_filters(args)
    rows = [] if error else [dict(row) for row in records if gto_matches(row, filters)]
    query = urlencode([(key, filters[key]) for key in FILTER_KEYS if filters[key]])
    report_url = '/reports/gto' + ('?' + query if query else '')
    for row in rows:
        row['stage_label'] = gto_stage_label(row['stage'])
        row['status_label'] = gto_status_label(row['status'])
        row['date_label'] = '.'.join(reversed(row['awarded_on'].split('-'))) if row['awarded_on'] else ''
        row['student_url'] = f'/admin/people/{row["student_id"]}?' + urlencode({'tab': 'gto', 'return': report_url})
        row['certificates'] = [
            {**certificate, 'url': f'/gto/{row["student_id"]}/awards/{row["id"]}/certificates/{certificate["id"]}'}
            for certificate in row['certificates']
        ]
    metrics = {'students': len({row['student_id'] for row in rows}), 'awards': len(rows)}
    metrics.update({status: sum(row['status'] == status for row in rows) for status in GTO_AWARDS})
    options = {}
    for key, column in (('year', 'year'), ('institute', 'institute'), ('group', 'group_name')):
        values = {str(row.get(column) or '') for row in records} - {''}
        if filters[key]:
            values.add(filters[key])  # An unknown text criterion must not disappear/broaden.
        numeric = key == 'year' and not error and all(value.isascii() and value.isdigit() for value in values)
        options[key] = sorted(
            values, key=(lambda value: int(value)) if numeric else str.casefold, reverse=key == 'year'
        )
    return dict(
        filters=filters,
        error=error,
        rows=rows,
        metrics=metrics,
        options=options,
        report_url=report_url,
        export_url='/reports/gto/export.xlsx' + ('?' + query if query else ''),
        applied=[(FILTER_LABELS[key], filters[key]) for key in FILTER_KEYS if filters[key]],
    )


def build_gto_xlsx(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'ГТО'
    sheet.append(['Отчёт ГТО'])
    sheet.append(['Институт и группа — текущие сведения карточки студента.'])
    for label, value in report['applied']:
        safe_value = sanitize_spreadsheet_value(value)
        sheet.append([label, safe_value])
        if sheet.cell(sheet.max_row, 2).value != safe_value:
            raise ValueError('Условие нельзя без потерь записать в ячейку XLSX.')
    metrics = report['metrics']
    sheet.append(
        [
            'Уникальных студентов',
            metrics['students'],
            'Награждений',
            metrics['awards'],
            'Бронза',
            metrics['bronze'],
            'Серебро',
            metrics['silver'],
            'Золото',
            metrics['gold'],
        ]
    )
    sheet.append([])
    sheet.append(EXPORT_COLUMNS)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
    for row in report['rows']:
        values = [
            row['id'],
            row['student_id'],
            row['full_name'],
            '',
            row['institute'] or '',
            row['group_name'] or '',
            row['year'],
            row['stage_label'],
            row['status_label'],
            row['awarded_on'] or '',
            '\n'.join(c['filename'] for c in row['certificates']),
            '\n'.join(c['url'] for c in row['certificates']),
        ]
        safe_values = [sanitize_spreadsheet_value(value) for value in values]
        sheet.append(safe_values)
        if any(cell.value != value for cell, value in zip(sheet[sheet.max_row], safe_values) if isinstance(value, str)):
            raise ValueError('Текст слишком длинный для ячейки XLSX. Данные не изменены.')
        # Set the actual UIN value as a literal string, without altering its text.
        # In particular '=...' is never a formula and leading zeroes remain exact.
        cell = sheet.cell(sheet.max_row, 4)
        cell.value = row['gto_uin'] or ''
        if cell.value != (row['gto_uin'] or ''):
            raise ValueError('УИН слишком длинный для ячейки XLSX. Значение не изменено.')
        cell.data_type = 's'
        cell.number_format = '@'
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.auto_filter.ref = f'A{header_row}:L{sheet.max_row}'
    widths = {
        'A': 18,
        'B': 14,
        'C': 34,
        'D': 24,
        'E': 26,
        'F': 22,
        'G': 18,
        'H': 28,
        'I': 16,
        'J': 22,
        'K': 34,
        'L': 48,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
