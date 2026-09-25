import asyncio
import base64
import binascii
import hashlib
import hmac
import json
import logging
import re
import secrets
import shutil
import time
import uuid
import zipfile
from datetime import datetime
from datetime import timedelta
from datetime import timezone
from io import BytesIO
from pathlib import Path
from typing import Iterable
from typing import Sequence
from urllib.parse import urlencode

import pandas as pd
from jinja2 import Environment
from jinja2 import PackageLoader
from jinja2 import select_autoescape
from openpyxl.utils import get_column_letter
from pandas import isna
from sanic import html
from sanic import redirect
from sanic import Request
from sanic import Sanic
from sanic import text
from sanic.response import HTTPResponse
from sanic.response import json as json_response
from sanic.response import raw
from sanic_ext import render

from src.auth import hash_password
from src.auth import verify_password
from src.backup import run_backup
from src.models.competition import Competition
from src.models.custom_field import CustomField
from src.settings import settings
from src.storage.sqlite import SQLiteAdapter

jinja_env = Environment(
    loader=PackageLoader('src'),
    autoescape=select_autoescape(),
    enable_async=True,
)

# Синхронный клон окружения для error-страниц: хелперы unauthorized/forbidden
# синхронны (их зовут require_* из синхронных проверок ролей), а render()
# основной среды возвращает корутину (enable_async).
jinja_env_sync = Environment(
    loader=PackageLoader('src'),
    autoescape=select_autoescape(),
)

app = Sanic('SIBADI_competitions')

app.config.REQUEST_MAX_SIZE = 10 * 1024 * 1024
app.config.REQUEST_TIMEOUT = 60
app.config.RESPONSE_TIMEOUT = 60
# Ровно один доверенный proxy-хоп (nginx): приложение наружу напрямую не
# публикуется — deploy/nginx и docker-compose. nginx выставляет
# X-Forwarded-For через $proxy_add_x_forwarded_for, поэтому реальный IP
# клиента — всегда ПОСЛЕДНЯЯ запись заголовка, и клиент её подделать не
# может (свои поддельные значения nginx дописывает перед ней).
app.config.PROXIES_COUNT = 1

app.static(
    uri='/static',
    file_or_directory='src/static',
    name='static',
    directory_view=False,
)

app.ctx.storage = None

AUTH_ALLOWED_PATHS = {'/healthcheck', '/login'}
ADMIN_ROLE = 'admin'
EDITOR_ROLE = 'editor'
VIEWER_ROLE = 'viewer'
ATHLETE_ROLE = 'athlete'
MODERATOR_ROLES = {ADMIN_ROLE, EDITOR_ROLE}
WRITE_ROLES = {ADMIN_ROLE, EDITOR_ROLE, ATHLETE_ROLE}
KNOWN_ROLES = {ADMIN_ROLE, EDITOR_ROLE, VIEWER_ROLE, ATHLETE_ROLE}

# Очистка базы — явное админское действие в два шага (объём + фраза).
# См. docs/data-model-decisions.md «Очистка и выгрузка базы».
# Порядок пунктов — по возрастанию ущерба: вложения → записи → всё.
WIPE_SCOPES = ('attachments', 'records', 'all')
WIPE_MODES = ('scope', 'date')
WIPE_CONFIRM_PHRASE = 'УДАЛИТЬ'

INSECURE_SECRET_VALUES = {'', 'change-me', 'replace-with-random-string'}
# Слабые пароли посева учёток (M3): дефолты settings.py и заглушки .env.example.
WEAK_BOOTSTRAP_PASSWORDS = INSECURE_SECRET_VALUES | {'change-me-editor', 'change-me-viewer'}
# Имена переменных окружения для сообщений об ошибке старта (без значений!).
BOOTSTRAP_PASSWORD_ENV_VARS = {
    ADMIN_ROLE: 'AUTH_ADMIN_PASSWORD',
    EDITOR_ROLE: 'AUTH_EDITOR_PASSWORD',
    VIEWER_ROLE: 'AUTH_VIEWER_PASSWORD',
}
BOOTSTRAP_USERNAME_ENV_VARS = {
    ADMIN_ROLE: 'AUTH_ADMIN_USERNAME',
    EDITOR_ROLE: 'AUTH_EDITOR_USERNAME',
    VIEWER_ROLE: 'AUTH_VIEWER_USERNAME',
}
LOGIN_MAX_ATTEMPTS = 5
LOGIN_WINDOW_SECONDS = 180
LOGIN_LOCKOUT_SECONDS = 180
LOGIN_REJECT_STATUS = 429

login_failures: dict[str, list] = {}

logger = logging.getLogger(__name__)

AUDIT_PAGE_SIZE = 50
RECONCILE_PAGE_SIZE = 50
PRE_WIPE_ARCHIVES_SHOWN = 20

BASE_FIELD_SPECS: Sequence[dict[str, str]] = (
    {'key': 'student_name', 'label': 'ФИО'},
    {'key': 'student_sex', 'label': 'Пол'},
    {'key': 'institute', 'label': 'Институт'},
    {'key': 'group', 'label': 'Группа'},
    {'key': 'sport', 'label': 'Вид спорта'},
    {'key': 'date', 'label': 'Дата'},
    {'key': 'level', 'label': 'Уровень соревнований'},
    {'key': 'name', 'label': 'Название соревнований'},
    {'key': 'position', 'label': 'Место'},
    {'key': 'course', 'label': 'Курс'},
)

REQUIRED_IMPORT_COLUMNS: Sequence[str] = tuple(field['label'] for field in BASE_FIELD_SPECS)
INDEX_EXPORT_COLUMNS: Sequence[str] = tuple(field['label'] for field in BASE_FIELD_SPECS)

# Лёгкий реестр полей (решение 2026-09-13, docs/data-model-decisions.md
# «Реестр полей: лёгкая версия сейчас, полная запланирована»): у каждого
# базового поля настраиваются ТИП (text/number; link базовым недоступен)
# и ОБЯЗАТЕЛЬНОСТЬ. Дефолты воспроизводят сегодняшнее поведение валидации
# (числовые — Место и Курс; обязательные — ФИО, Дата, Курс). Синхронно
# с BASE_FIELD_SETTING_DEFAULTS в src/storage/sqlite.py.
BASE_FIELD_SETTING_DEFAULTS: dict[str, tuple[str, bool]] = {
    'student_name': ('text', True),
    'student_sex': ('text', False),
    'institute': ('text', False),
    'group': ('text', False),
    'sport': ('text', False),
    'date': ('text', True),
    'level': ('text', False),
    'name': ('text', False),
    'position': ('number', False),
    'course': ('number', True),
}
BASE_FIELD_LABELS: dict[str, str] = {field['key']: field['label'] for field in BASE_FIELD_SPECS}
# ФИО и Дата — ключ дедупликации и срезы отчётов: обязательность не
# отключается (в UI задизейблена с пояснением), тип фиксированный.
ALWAYS_REQUIRED_BASE_FIELDS: frozenset[str] = frozenset({'student_name', 'date'})
BASE_FIELD_VALUE_TYPES: Sequence[str] = ('text', 'number')

# Расширение отчётов (замечание №19, docs/data-model-decisions.md «Расширение
# отчётов»): срез × метрики × фильтры. Группировка всегда по данным записи.
REPORT_SLICES: Sequence[dict[str, str]] = (
    {'key': 'student', 'label': 'Студент'},
    {'key': 'group', 'label': 'Группа'},
    {'key': 'institute', 'label': 'Институт'},
    {'key': 'course', 'label': 'Курс'},
    {'key': 'sport', 'label': 'Вид спорта'},
    {'key': 'level', 'label': 'Уровень'},
    {'key': 'year', 'label': 'Год'},
)
REPORT_SLICE_KEYS = {item['key'] for item in REPORT_SLICES}
DEFAULT_REPORT_SLICE = 'student'
REPORT_METRIC_COLUMNS: Sequence[str] = ('Участий', 'Побед', 'Призовых')
# Колонки среза (без метрик) в HTML-таблице и выгрузке: для «студента» —
# как раньше (ФИО/пол/институт/группа/курс), для остальных — название среза.
REPORT_SLICE_COLUMNS: dict[str, tuple[str, ...]] = {
    'student': ('ФИО', 'Пол', 'Институт', 'Группа', 'Курс'),
    'group': ('Группа',),
    'institute': ('Институт',),
    'course': ('Курс',),
    'sport': ('Вид спорта',),
    'level': ('Уровень соревнований',),
    'year': ('Год',),
}
# Колонки с числовым значением среза (сортировка в HTML-таблице).
REPORT_SLICE_NUMERIC_KEYS = frozenset({'course', 'year'})
# Старые ссылки/закладки с выбором колонок продолжают работать.
LEGACY_REPORT_COLUMN_ALIASES = {'Количество участий': 'Участий'}
FIELD_TYPE_OPTIONS: Sequence[str] = ('text', 'number', 'date', 'url')

# №24 (docs/feedback-live.md): колонки, к которым можно привязать link-поле —
# базовые текстовые колонки таблицы. Сортировка/инлайн в целевой колонке
# остаются по её тексту, ссылка только оборачивает текст в <a>.
LINK_TARGETABLE_BASE_KEYS: Sequence[str] = (
    'student_name',
    'institute',
    'group',
    'sport',
    'level',
    'name',
)
USER_ROLES: Sequence[str] = (ADMIN_ROLE, EDITOR_ROLE, VIEWER_ROLE, ATHLETE_ROLE)
MIN_PASSWORD_LENGTH = 6
# Карточки студентов (Student Identity v1, Phase 1): допустимые значения
# поля «Пол» — только М/Ж или пусто (не указан).
STUDENT_SEX_OPTIONS: frozenset[str] = frozenset({'', 'М', 'Ж'})
# Импорт студентов из Excel (Student Identity v1, Phase 2.5): колонки файла
# (обязательна только «ФИО», остальные опциональны), лимит строк на файл и
# TTL staging-сессии предпросмотра (см. student_import_sessions).
STUDENT_IMPORT_COLUMNS: Sequence[str] = ('ФИО', 'Пол', 'Институт', 'Группа', 'Курс')
STUDENT_IMPORT_MAX_ROWS = 5000
STUDENT_IMPORT_SESSION_TTL_SECONDS = 2 * 60 * 60
# Резолвер ФИО (/api/athletes/search): лимит вариантов на запрос ПО KIND —
# до 8 карточек Students и до 8 легаси-подсказок. Карточки идут раньше, но
# легаси-квота не вытесняется (см. search_athletes).
ATHLETE_SEARCH_LIMIT = 8
ATTACHMENT_MAX_SIZE = 5 * 1024 * 1024
ATTACHMENT_EXTENSIONS = {
    'pdf': 'application/pdf',
    'png': 'image/png',
    'jpg': 'image/jpeg',
    'jpeg': 'image/jpeg',
}
PROFILE_FIELDS = ('student_name', 'student_sex', 'institute', 'group', 'course')
ATTACHMENT_SIGNATURES = {
    'application/pdf': b'%PDF-',
    'image/png': b'\x89PNG\r\n\x1a\n',
    'image/jpeg': b'\xff\xd8\xff',
}


DEFAULT_LEVELS = ('внутривузовские', 'межвузовские')

# Композиция главной по прототипу 02: серверные фильтры и пагинация реестра.
# per_page — из фиксированного набора; текущее поведение продукта показывало
# весь список сразу, поэтому дефолт — максимум набора.
INDEX_PER_PAGE_OPTIONS: Sequence[int] = (10, 25, 50, 100)
INDEX_DEFAULT_PER_PAGE = 100
INDEX_STATUS_FILTERS: Sequence[tuple[str, str]] = (
    ('approved', 'подтверждена'),
    ('pending', 'на проверке'),
    ('rejected', 'отклонена'),
)
INDEX_STATUS_KEYS = {key for key, _ in INDEX_STATUS_FILTERS}

# Управляемые справочники значений: записи остаются свободным текстом,
# справочник — только подсказки (docs/data-model-decisions.md,
# «Справочники значений»). Уровни живут в своей таблице на той же странице.
# Группы — иерархическая категория: существуют только с parent_id на
# институт, поэтому в одиночное создание не входят.
CATALOG_CATEGORIES: Sequence[str] = ('sport', 'institute')
GROUP_CATEGORY = 'group'
CATALOG_MANAGE_CATEGORIES: Sequence[str] = (*CATALOG_CATEGORIES, GROUP_CATEGORY)

# Переименование значений справочника (решение 2026-09-13): уровень живёт в
# levels, остальные — в catalog_values; label — для страницы-подтверждения.
CATALOG_RENAME_LABELS: dict[str, str] = {
    'level': 'Уровень',
    'sport': 'Вид спорта',
    'institute': 'Институт',
    GROUP_CATEGORY: 'Группа',
}


def seed_levels(storage: SQLiteAdapter):
    if not storage.get_level_names(include_inactive=True):
        for name in DEFAULT_LEVELS:
            storage.create_level(name)


def env_bootstrap_accounts() -> list[tuple[str, str, str]]:
    """Аккаунты, сеемые из AUTH_*-переменных: (роль, логин, пароль).

    admin и editor всегда настроены дефолтами; viewer с пустым логином
    или паролем не создаётся вовсе.
    """
    accounts = [
        (ADMIN_ROLE, settings.auth_admin_username, settings.auth_admin_password),
        (EDITOR_ROLE, settings.auth_editor_username, settings.auth_editor_password),
        (VIEWER_ROLE, settings.auth_viewer_username, settings.auth_viewer_password),
    ]
    return [(role, username, password) for role, username, password in accounts if username and password]


def seed_users(storage: SQLiteAdapter):
    for role, username, password in env_bootstrap_accounts():
        if storage.get_user(username) is None:
            storage.create_user(username, hash_password(password), role)


def find_weak_bootstrap_accounts(storage: SQLiteAdapter) -> list[str]:
    """Роли, чья учётка БУДЕТ посеяна из окружения со слабым паролем (M3).

    Посев не перезаписывает существующие учётки, поэтому слабый пароль в
    .env опасен только пока учётки с таким логином ещё нет.
    """
    return [
        role
        for role, username, password in env_bootstrap_accounts()
        if password in WEAK_BOOTSTRAP_PASSWORDS and storage.get_user(username) is None
    ]


@app.before_server_start
async def init_storage(app: Sanic, _):
    if not Sanic.test_mode and settings.auth_secret_key in INSECURE_SECRET_VALUES:
        raise RuntimeError('AUTH_SECRET_KEY is not configured: set it to a random value in the environment or .env')
    if app.ctx.storage is None:
        app.ctx.storage = SQLiteAdapter(settings.database_path)
    if not Sanic.test_mode:
        # Отказ старта вместо посева учёток со слабым паролем (M3) или с
        # недопустимым логином (L8). В сообщении — только имена переменных
        # окружения, значения (пароли) не раскрываются и не логируются.
        weak_roles = find_weak_bootstrap_accounts(app.ctx.storage)
        if weak_roles:
            weak_vars = ', '.join(BOOTSTRAP_PASSWORD_ENV_VARS[role] for role in weak_roles)
            raise RuntimeError(f'{weak_vars} is not configured: set a strong unique password in .env')
        for role, username, _password in env_bootstrap_accounts():
            if app.ctx.storage.get_user(username) is None and username_error(username):
                raise RuntimeError(
                    f'{BOOTSTRAP_USERNAME_ENV_VARS[role]} contains characters '
                    'that are not allowed in logins (: or control characters)'
                )
    seed_levels(app.ctx.storage)
    seed_users(app.ctx.storage)


def get_storage(app: Sanic) -> SQLiteAdapter:
    storage = getattr(app.ctx, 'storage', None)
    if storage is None:
        raise RuntimeError('Storage adapter is not initialized')
    return storage


def get_param(args: dict, key: str) -> str | None:
    raw_value = args.get(key, [])
    if raw_value:
        return raw_value[0]
    return None


def get_form_value(request: Request, key: str) -> str:
    value = request.form.get(key, '')
    if isinstance(value, list):
        return value[0]
    return value


def get_auth_user(request: Request) -> dict | None:
    return getattr(request.ctx, 'auth_user', None)


def get_current_user_id(request: Request) -> int | None:
    user = get_auth_user(request)
    if not user:
        return None
    try:
        record = get_storage(request.app).get_user(user['username'])
    except RuntimeError:
        return None
    return record['id'] if record else None


def log_audit_event(
    request: Request,
    action: str,
    details: dict | None = None,
    *,
    user_id: int | None = None,
    username: str | None = None,
) -> None:
    """Записать событие в журнал безопасности.

    Журнал — вспомогательный инструмент: сбой записи не должен ломать
    основной сценарий, поэтому любые ошибки глотаем с логом.
    """
    try:
        user = get_auth_user(request)
        if username is None:
            username = (user or {}).get('username') or ''
        if user_id is None and user:
            record = get_storage(request.app).get_user(user['username'])
            user_id = record['id'] if record else None
        get_storage(request.app).add_audit_event(
            user_id=user_id,
            username=username or '',
            action=action,
            details=json.dumps(details or {}, ensure_ascii=False),
        )
    except Exception:
        logger.exception('Failed to write audit event %s', action)


def create_auth_cookie_value(
    username: str,
    role: str,
    issued_at: int | None = None,
    pwd_ver: int = 0,
) -> str:
    if issued_at is None:
        issued_at = int(time.time())
    payload = f'{username}:{role}:{issued_at}:{pwd_ver}'
    signature = hmac.new(
        settings.auth_secret_key.encode(),
        payload.encode(),
        hashlib.sha256,
    ).hexdigest()
    token = f'{payload}:{signature}'
    return base64.urlsafe_b64encode(token.encode()).decode()


def decode_auth_token(token: str) -> dict | None:
    try:
        decoded = base64.urlsafe_b64decode(token.encode()).decode()
        username, role, issued_at, pwd_ver, signature = decoded.split(':', 4)
        issued_at = int(issued_at)
        pwd_ver = int(pwd_ver)
    except (ValueError, UnicodeDecodeError, binascii.Error):
        return None

    if role not in KNOWN_ROLES:
        return None
    if int(time.time()) - issued_at > settings.auth_session_ttl_seconds:
        return None

    expected_signature = hmac.new(
        settings.auth_secret_key.encode(),
        f'{username}:{role}:{issued_at}:{pwd_ver}'.encode(),
        hashlib.sha256,
    ).hexdigest()
    if not secrets.compare_digest(signature, expected_signature):
        return None

    return {'username': username, 'role': role, 'pwd_ver': pwd_ver}


def parse_auth_cookie(request: Request) -> dict | None:
    token = request.cookies.get(settings.auth_cookie_name)
    if not token:
        return None

    payload = decode_auth_token(token)
    if payload is None:
        return None

    try:
        storage = get_storage(request.app)
    except RuntimeError:
        storage = None
    if storage is not None:
        user = storage.get_user(payload['username'])
        if user is None or not user.get('active', 1):
            return None
        if int(user.get('pwd_ver', 0)) != payload['pwd_ver']:
            return None
        # id нужен middleware присутствия (№17): last_seen без лишнего запроса.
        return {'username': payload['username'], 'role': payload['role'], 'id': user['id']}

    return {'username': payload['username'], 'role': payload['role']}


def authenticate_user(request: Request, username: str, password: str) -> dict | None:
    user = get_storage(request.app).get_user(username)
    if user is None or not user['active']:
        return None
    if not verify_password(password, user['password_hash']):
        return None
    return {
        'id': user['id'],
        'username': user['username'],
        'role': user['role'],
        'pwd_ver': int(user.get('pwd_ver', 0)),
    }


def register_login_failure(ip: str):
    now = time.monotonic()
    attempts = [stamp for stamp in login_failures.get(ip, []) if now - stamp < LOGIN_WINDOW_SECONDS]
    attempts.append(now)
    login_failures[ip] = attempts
    return len(attempts)


def login_is_locked(ip: str) -> bool:
    now = time.monotonic()
    attempts = [stamp for stamp in login_failures.get(ip, []) if now - stamp < LOGIN_WINDOW_SECONDS]
    login_failures[ip] = attempts
    return len(attempts) >= LOGIN_MAX_ATTEMPTS


def clear_login_failures(ip: str):
    login_failures.pop(ip, None)


def username_error(username: str) -> str | None:
    """Недопустимые символы логина (L8): None — логин корректен.

    Двоеточие ломает разбор подписанной cookie «логин:роль:время:версия»,
    управляющие символы (включая DEL) в логине не нужны вовсе.
    """
    if ':' in username or any(ord(char) < 32 or ord(char) == 127 for char in username):
        return 'Логин не может содержать двоеточие и управляющие символы'
    return None


def request_is_secure(request: Request) -> bool:
    forwarded_proto = request.headers.get('x-forwarded-proto')
    if forwarded_proto:
        return forwarded_proto == 'https'
    return request.scheme == 'https'


def set_auth_cookie(request: Request, response, username: str, role: str, pwd_ver: int = 0):
    response.add_cookie(
        settings.auth_cookie_name,
        create_auth_cookie_value(username, role, pwd_ver=pwd_ver),
        httponly=True,
        samesite='Lax',
        secure=request_is_secure(request),
        path='/',
        max_age=settings.auth_session_ttl_seconds,
    )


def clear_auth_cookie(response):
    response.delete_cookie(settings.auth_cookie_name, path='/')


def create_csrf_token(request: Request) -> str:
    auth_cookie = request.cookies.get(settings.auth_cookie_name, '')
    return hmac.new(
        settings.auth_secret_key.encode(),
        b'csrf:' + auth_cookie.encode(),
        hashlib.sha256,
    ).hexdigest()


def request_csrf_is_valid(request: Request) -> bool:
    form_token = get_form_value(request, 'csrf_token')
    if not form_token:
        return False
    return secrets.compare_digest(form_token, create_csrf_token(request))


jinja_env.globals['csrf_token'] = create_csrf_token
jinja_env_sync.globals['csrf_token'] = create_csrf_token


def user_is_admin(request: Request) -> bool:
    user = get_auth_user(request)
    return bool(user and user['role'] == ADMIN_ROLE)


def user_can_write(request: Request) -> bool:
    user = get_auth_user(request)
    return bool(user and user['role'] in WRITE_ROLES)


def user_is_moderator(request: Request) -> bool:
    user = get_auth_user(request)
    return bool(user and user['role'] in MODERATOR_ROLES)


def user_is_athlete(request: Request) -> bool:
    user = get_auth_user(request)
    return bool(user and user['role'] == ATHLETE_ROLE)


def student_hashes_for_request(request: Request) -> list[str]:
    if not user_is_athlete(request):
        return []
    user_id = get_current_user_id(request)
    if user_id is None:
        return []
    storage = get_storage(request.app)
    names = storage.get_name_aliases(user_id)
    profile_name = storage.get_profile(user_id).get('student_name', '').strip()
    if profile_name and profile_name not in names:
        names = names + [profile_name]
    return [hashlib.sha256(name.strip().encode()).hexdigest() for name in names if name.strip()]


def user_owns_record(request: Request, review: dict) -> bool:
    if not user_is_athlete(request):
        return True
    if review.get('owner_id') == get_current_user_id(request):
        return True
    return review.get('student_id') in student_hashes_for_request(request)


# UI-ответы 401/403 (IA редизайна 2026-09): браузерной навигации (Accept:
# text/html) вместо голого text() отдаём редирект на вход / страницу
# «Доступ запрещён»; программным клиентам (fetch, API-тесты) поведение
# прежнее — строго text() с теми же статусами. CSRF-ответы middleware
# этими хелперами не затрагиваются.


def request_accepts_html(request: Request) -> bool:
    return 'text/html' in (request.headers.get('accept') or '')


def unauthorized(request: Request):
    if request_accepts_html(request):
        return redirect('/login')
    return text(body='Unauthorized', status=401)


def forbidden(request: Request):
    if request_accepts_html(request):
        return html(
            jinja_env_sync.get_template('error403.html').render(request=request),
            status=403,
        )
    return text(body='Forbidden', status=403)


def require_admin(request: Request):
    user = get_auth_user(request)
    if not user:
        return unauthorized(request)
    if user['role'] != ADMIN_ROLE:
        return forbidden(request)
    return None


def require_moderator(request: Request):
    user = get_auth_user(request)
    if not user:
        return unauthorized(request)
    if user['role'] not in MODERATOR_ROLES:
        return forbidden(request)
    return None


def require_writer(request: Request):
    user = get_auth_user(request)
    if not user:
        return unauthorized(request)
    if user['role'] not in WRITE_ROLES:
        return forbidden(request)
    return None


def build_redirect_with_message(*, message: str | None = None, error: str | None = None, url: str = '/'):
    params = {}
    if message:
        params['admin_message'] = message
    if error:
        params['admin_error'] = error
    query = urlencode(params)
    return redirect(f'{url}?{query}' if query else url)


def get_flash_args(request: Request) -> dict[str, str | None]:
    args = dict(request.args)
    return {
        'admin_message': get_param(args, 'admin_message'),
        'admin_error': get_param(args, 'admin_error'),
    }


def validate_import_columns(df: pd.DataFrame):
    missing_columns = [column for column in REQUIRED_IMPORT_COLUMNS if column not in df.columns]
    if missing_columns:
        raise ValueError(f'Missing required columns: {", ".join(missing_columns)}')


def clean_str(value) -> str:
    """Значение ячейки импорта/формы как строка; NaN/None → пустая строка.

    Замечание с живого импорта (docs/feedback-live.md «nan»): pandas
    превращает пустые ячейки Excel в NaN, и str(NaN) давал текст «nan»
    в Поле/Институте/Группе. Любая пустая ячейка теперь — пустое значение.
    """
    if value is None or isna(value):
        return ''
    return str(value).strip()


def normalize_position(value) -> int:
    if isna(value) or value == '':
        return 0
    return int(value)


def normalize_course(value) -> int:
    if isna(value):
        raise ValueError('Курс обязателен')
    return int(value)


def get_base_field_settings(storage: SQLiteAdapter | None) -> dict[str, dict]:
    """Эффективные настройки базовых полей: дефолты + строки field_settings.

    Настройки опциональны построчно: поле без строки в таблице действует
    по дефолту (п.4 постановки). ФИО и Дата всегда обязательны — их
    required принудительно True (ключ дедупликации и срезы отчётов).
    """
    merged = {
        key: {'value_type': value_type, 'required': required}
        for key, (value_type, required) in BASE_FIELD_SETTING_DEFAULTS.items()
    }
    if storage is not None:
        for key, setting in storage.get_field_settings().items():
            if key in merged and setting['value_type'] in BASE_FIELD_VALUE_TYPES:
                merged[key] = {'value_type': setting['value_type'], 'required': bool(setting['required'])}
    for key in ALWAYS_REQUIRED_BASE_FIELDS:
        # ФИО и Дата: обязательность не отключается, тип фиксированный —
        # даже если в таблице оказалась иная настройка.
        default_type = BASE_FIELD_SETTING_DEFAULTS[key][0]
        merged[key] = {'value_type': default_type, 'required': True}
    return merged


def get_allowed_link_targets(
    custom_fields: Sequence[CustomField],
    *,
    exclude_key: str | None = None,
) -> dict[str, str]:
    """№24: допустимые цели привязки link-поля — ключ → подпись колонки.

    Базовые текстовые колонки таблицы + активные кастомные текстовые поля
    (кроме самого link-поля).
    """
    targets = {key: BASE_FIELD_LABELS[key] for key in LINK_TARGETABLE_BASE_KEYS}
    for field in custom_fields:
        if field.field_type == 'text' and field.active and field.key != exclude_key:
            targets[field.key] = field.label
    return targets


def build_link_bindings(custom_fields: Sequence[CustomField]) -> dict[str, str]:
    """№24: карта «целевая колонка → ключ link-поля» для рендера таблицы.

    Целевая колонка показывает значение своей ячейки как гиперссылку на
    значение link-поля ЭТОЙ записи (если ссылка заполнена и это http/https).
    """
    bindings: dict[str, str] = {}
    for field in custom_fields:
        if field.field_type == 'url' and field.link_target and field.active:
            bindings[field.link_target] = field.key
    return bindings


def parse_link_target_form(
    request: Request,
    *,
    field_type: str,
    custom_fields: Sequence[CustomField],
    exclude_key: str | None = None,
) -> str | None:
    """Прочитать link_target из формы полей: NULL/пусто = не привязана.

    Значение принимается только у url-полей и только из допустимого списка
    целей — подделка разметки не уводит рендер в неизвестную колонку.
    """
    if field_type != 'url':
        return None
    raw_target = get_form_value(request, 'link_target').strip()
    if not raw_target:
        return None
    if raw_target not in get_allowed_link_targets(custom_fields, exclude_key=exclude_key):
        return None
    return raw_target


def parse_base_number_field(raw, *, label: str, setting: dict, empty_error: str, empty_value: int = 0):
    """number-поле базового реестра: пустое при required — отказ, без
    required — empty_value (как Место раньше давал 0)."""
    value = clean_str(raw)
    if not value:
        if setting['required']:
            raise ValueError(empty_error)
        return empty_value
    try:
        # Excel-ячейки приходят числами (в т.ч. float) — как раньше,
        # int() принимает и их; строка чистится.
        return int(value) if isinstance(raw, str) else int(raw)
    except (TypeError, ValueError):
        raise ValueError(f'Поле "{label}" должно быть числом') from None


def parse_base_text_field(raw, *, label: str, setting: dict):
    """text-поле базового реестра: строка без требования числа
    (кейс «Курс: Выпускник 2025/26», замечания №2/№7)."""
    value = clean_str(raw)
    if not value and setting['required']:
        raise ValueError(f'Поле "{label}" обязательно')
    return value


# Гибридный ввод дат-диапазонов (решение 2026-09-13, docs/data-model-decisions.md
# «Даты-диапазоны: визуально одно, под капотом два»). Одно поле ввода/колонка
# «Дата» разбирается на пару (date, date_to); формат решает парсер:
DATE_RANGE_SINGLE = re.compile(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})$')
# «25-27.06.2026» — границы в одном месяце
DATE_RANGE_ONE_MONTH = re.compile(r'^(\d{1,2})-(\d{1,2})\.(\d{1,2})\.(\d{4})$')
# «30.01-01.02.2026» — разные месяцы одного года
DATE_RANGE_ONE_YEAR = re.compile(r'^(\d{1,2})\.(\d{1,2})-(\d{1,2})\.(\d{1,2})\.(\d{4})$')
# «28.12.2025-02.01.2026» — полный формат (любые даты, включая разные годы)
DATE_RANGE_FULL = re.compile(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})-(\d{1,2})\.(\d{1,2})\.(\d{4})$')


def _range_date(day: str, month: str, year: str, label: str) -> datetime:
    try:
        return datetime(int(year), int(month), int(day))
    except ValueError:
        raise ValueError(f'Некорректная дата ({label})') from None


def _match_date_range(raw: str) -> tuple[datetime, datetime | None] | None:
    """Распознать формат диапазона (или однодневной даты). None — не узнано."""
    match = DATE_RANGE_FULL.match(raw)
    if match:
        return (
            _range_date(match.group(1), match.group(2), match.group(3), 'дата начала'),
            _range_date(match.group(4), match.group(5), match.group(6), 'дата окончания'),
        )
    match = DATE_RANGE_ONE_YEAR.match(raw)
    if match:
        year = match.group(5)
        return (
            _range_date(match.group(1), match.group(2), year, 'дата начала'),
            _range_date(match.group(3), match.group(4), year, 'дата окончания'),
        )
    match = DATE_RANGE_ONE_MONTH.match(raw)
    if match:
        month, year = match.group(3), match.group(4)
        return (
            _range_date(match.group(1), month, year, 'дата начала'),
            _range_date(match.group(2), month, year, 'дата окончания'),
        )
    match = DATE_RANGE_SINGLE.match(raw)
    if match:
        return _range_date(*match.groups(), 'дата'), None
    return None


def parse_date_value(value, manual_input: bool = False) -> tuple[datetime, datetime | None]:
    """Разобрать поле «Дата» в пару (date, date_to).

    Поддерживаются однодневное «25.06.2026», компактные диапазоны
    «25-27.06.2026» и «30.01-01.02.2026» и полный «25.06.2026-27.06.2026»
    — один парсер для ручного ввода и колонки «Дата» импорта. Excel-ячейка
    с типом дата приходит объектом datetime (однодневная, импорт). date_to
    раньше date — отказ (маршрут отдаёт 400).
    """
    if isinstance(value, datetime):
        return value, None

    raw = str(value or '').strip()
    if not raw:
        raise ValueError('Дата обязательна')

    parsed = _match_date_range(raw)
    if parsed is None:
        if manual_input:
            # Ручной ввод жёстче: только форматы дд.мм.гггг выше.
            raise ValueError(f'Некорректная дата: {raw}')
        # Легаси-прочтение импорта: текстовая ISO-ячейка Excel.
        try:
            return datetime.fromisoformat(raw), None
        except ValueError:
            raise ValueError(f'Некорректная дата: {raw}') from None

    date, date_to = parsed
    if date_to is not None and date_to < date:
        raise ValueError('Дата окончания не может быть раньше даты начала')
    return date, date_to


def format_date_range(date_from: datetime, date_to: datetime | None) -> str:
    """Компактное отображение диапазона (решение 2026-09-13): совпадающие
    месяц/год — «25-27.06.2026», совпадающий год — «30.01-01.02.2026»,
    разные годы — полный «28.12.2025-02.01.2026»; однодневное — «25.06.2026».
    Используется в таблице, экспорте (колонка «Дата» остаётся реимпортируемой)
    и предзаполнении инлайн-правки."""
    if not date_to or date_to <= date_from:
        return date_from.strftime('%d.%m.%Y')
    if (date_from.year, date_from.month) == (date_to.year, date_to.month):
        return f'{date_from.day:02d}-{date_to.day:02d}.{date_from.month:02d}.{date_from.year}'
    if date_from.year == date_to.year:
        return f'{date_from.day:02d}.{date_from.month:02d}-' f'{date_to.day:02d}.{date_to.month:02d}.{date_from.year}'
    return f'{date_from.strftime("%d.%m.%Y")}-{date_to.strftime("%d.%m.%Y")}'


# Компактное отображение дат-диапазонов в шаблонах (таблица реестра).
jinja_env.globals['format_date_range'] = format_date_range


def normalize_custom_field_key(label: str) -> str:
    normalized = re.sub(r'\W+', '_', label.strip().lower())
    normalized = normalized.strip('_')
    if not normalized:
        normalized = 'field'
    return normalized


def make_unique_custom_field_key(label: str, existing_fields: Sequence[CustomField]) -> str:
    base_key = normalize_custom_field_key(label)
    existing_keys = {field.key for field in existing_fields}
    candidate = base_key
    index = 2
    while candidate in existing_keys:
        candidate = f'{base_key}_{index}'
        index += 1
    return candidate


def checkbox_to_bool(value: str) -> bool:
    return value.lower() in {'1', 'true', 'on', 'yes'}


def parse_checkbox(request: Request, key: str) -> bool:
    return checkbox_to_bool(get_form_value(request, key))


def is_http_url(value: str) -> bool:
    """Ссылка строго со схемой http/https (без прочих схем вроде javascript:)."""
    return bool(re.fullmatch(r'https?://\S+', value))


def parse_custom_field_value(raw_value, field: CustomField) -> str:
    if isna(raw_value):
        raw_value = ''

    value = str(raw_value).strip()
    if not value:
        if field.required:
            raise ValueError(f'Поле "{field.label}" обязательно')
        return ''

    if field.field_type == 'number':
        return str(int(float(value)))
    if field.field_type == 'date':
        return datetime.strptime(value, settings.date_format).strftime(settings.date_format)
    if field.field_type == 'url':
        if not is_http_url(value):
            raise ValueError(f'Поле "{field.label}" должно быть ссылкой (http:// или https://)')
        return value
    return value


def extract_custom_field_values(record: dict, custom_fields: Sequence[CustomField]) -> dict[str, str]:
    extra_data = {}
    for field in custom_fields:
        raw_value = record.get(field.label, '')
        value = parse_custom_field_value(raw_value, field)
        extra_data[field.key] = value
    return extra_data


def extract_custom_field_values_from_request(
    request: Request,
    custom_fields: Sequence[CustomField],
) -> dict[str, str]:
    extra_data = {}
    for field in custom_fields:
        extra_data[field.key] = parse_custom_field_value(
            get_form_value(request, f'custom__{field.key}'),
            field,
        )
    return extra_data


def build_competition(
    record: dict,
    *,
    custom_fields: Sequence[CustomField],
    manual_input: bool = False,
    storage: SQLiteAdapter | None = None,
) -> Competition:
    field_settings = get_base_field_settings(storage)
    student_name = clean_str(record['ФИО'])
    if not student_name:
        raise ValueError('ФИО обязательно')

    date, date_to = parse_date_value(record['Дата'], manual_input)

    # required=1 у текстового базового поля — пустое значение = отказ
    # (как ФИО/дата сейчас). Числовые Место/Курс — в parse_base_number_field.
    level = parse_base_text_field(
        record['Уровень соревнований'],
        label=BASE_FIELD_LABELS['level'],
        setting=field_settings['level'],
    ).lower()

    competition = Competition(
        student_id=hashlib.sha256(student_name.encode()).hexdigest(),
        student_name=student_name,
        student_sex=parse_base_text_field(
            record['Пол'], label=BASE_FIELD_LABELS['student_sex'], setting=field_settings['student_sex']
        ),
        institute=parse_base_text_field(
            record['Институт'], label=BASE_FIELD_LABELS['institute'], setting=field_settings['institute']
        ),
        group=parse_base_text_field(
            record['Группа'], label=BASE_FIELD_LABELS['group'], setting=field_settings['group']
        ),
        date=date,
        date_to=date_to,
        sport=parse_base_text_field(
            record['Вид спорта'], label=BASE_FIELD_LABELS['sport'], setting=field_settings['sport']
        ),
        level=level,
        name=parse_base_text_field(
            record['Название соревнований'], label=BASE_FIELD_LABELS['name'], setting=field_settings['name']
        ),
        position=parse_base_number_field(
            record['Место'],
            label=BASE_FIELD_LABELS['position'],
            setting=field_settings['position'],
            empty_error='Поле "Место" обязательно',
        ),
        course=(
            parse_base_text_field(record['Курс'], label=BASE_FIELD_LABELS['course'], setting=field_settings['course'])
            if field_settings['course']['value_type'] != 'number'
            else parse_base_number_field(
                record['Курс'],
                label=BASE_FIELD_LABELS['course'],
                setting=field_settings['course'],
                empty_error='Курс обязателен',
            )
        ),
        extra_data=extract_custom_field_values(record, custom_fields),
    )
    if storage is not None:
        apply_catalog_canonical_values(storage, competition)
    return competition


def apply_catalog_canonical_values(storage: SQLiteAdapter, competition: Competition) -> None:
    """№1.1 (docs/feedback-live.md): значения справочников — в каноническом регистре.

    Если значение записи отличается от значения справочника ТОЛЬКО регистром,
    в запись подставляется каноническое написание из справочника. Исторические
    записи не трогаются — замена происходит только на пути создания записи
    (ручной ввод, импорт). Уровень по-прежнему нормализуется в lower
    (build_competition), а затем каноническое написание из levels имеет
    приоритет — так записи согласованы со справочником.

    Группа ищется внутри института: уникальность группы — по паре
    (институт, группа), поэтому каноническое подставляется только когда
    канонический институт найден в иерархии.
    """
    if competition.level:
        canonical = storage.find_level_canonical(competition.level)
        if canonical:
            competition.level = canonical
    for category in ('sport', 'institute'):
        value = getattr(competition, category)
        if value:
            canonical = storage.find_catalog_canonical(category, value)
            if canonical:
                setattr(competition, category, canonical)
    if competition.institute and competition.group:
        institute = storage.find_catalog_row('institute', competition.institute)
        if institute is not None:
            canonical = storage.find_catalog_canonical('group', competition.group, parent_id=institute['id'])
            if canonical:
                competition.group = canonical
    elif competition.group and not competition.institute:
        autofill_institute_from_group(storage, competition)


def autofill_institute_from_group(storage: SQLiteAdapter, competition: Competition) -> None:
    """№19а (docs/feedback-live.md): автозаполнение института по группе.

    Если имя группы принадлежит ровно одному институту иерархии справочников,
    подставляется канонический институт (и каноническое написание группы).
    Неизвестная группа или одно имя в разных институтах — институт остаётся
    пустым: свободный ввод не ломаем.
    """
    institute_value = storage.find_unique_group_institute(competition.group)
    if not institute_value:
        return
    competition.institute = institute_value
    institute = storage.find_catalog_row('institute', institute_value)
    if institute is not None:
        canonical = storage.find_catalog_canonical('group', competition.group, parent_id=institute['id'])
        if canonical:
            competition.group = canonical


def competition_to_export_row(
    competition: Competition,
    export_custom_fields: Sequence[CustomField],
) -> dict[str, str | int]:
    row = competition.model_dump(by_alias=True)
    row.pop('_id', None)
    row.pop('Время создания записи (UTC)', None)
    row.pop('extra_data', None)
    row.pop('Статус проверки', None)
    row.pop('Комментарий проверки', None)
    # Служебная стабильная связь с карточкой студента — внутренняя колонка,
    # в Excel-выгрузки (реестр/отчёты/обслуживание) не отдаётся: формат
    # выгрузок — часть контракта импорта/экспорта и не меняется.
    row.pop('student_ref_id', None)
    # Event Model, Wave 1 P1: служебные поля participation-identity
    # (дисциплина/результат/ссылка на событие) тоже внутренние — контракт
    # Excel-выгрузок не меняется, ключей в строке быть не должно.
    row.pop('discipline', None)
    row.pop('result', None)
    row.pop('calendar_event_id', None)
    # Даты-диапазоны: экспорт компактной строкой в существующей колонке
    # «Дата» — шаблон импорта не ломается, выгрузку можно импортировать
    # обратно (парсер диапазона работает на той же колонке).
    row['Дата'] = format_date_range(competition.date, competition.date_to)
    row.pop('Дата по', None)
    for field in export_custom_fields:
        row[field.label] = competition.extra_data.get(field.key, '')
    return row


def report_export_columns(slice_key: str = DEFAULT_REPORT_SLICE) -> list[str]:
    """Колонки выгрузки для среза: название среза + Участий/Побед/Призовых."""
    return [*REPORT_SLICE_COLUMNS[slice_key], *REPORT_METRIC_COLUMNS]


def parse_report_slice(args: dict) -> str:
    """Срез отчёта из GET-параметра group_by; отсутствие — «студент»."""
    value = get_param(args, 'group_by')
    return value if value else DEFAULT_REPORT_SLICE


def validate_report_filters(args: dict) -> str | None:
    for key in ('date_from', 'date_to'):
        value = get_param(args, key)
        if value:
            try:
                datetime.strptime(value, settings.date_format)
            except ValueError:
                return f'Некорректная дата в фильтре ({key}): {value}'

    position = get_param(args, 'position')
    if position and not re.fullmatch(r'[<>]\d+', position):
        return f'Некорректное значение места: {position}'

    group_by = get_param(args, 'group_by')
    if group_by and group_by not in REPORT_SLICE_KEYS:
        return f'Неизвестный срез отчёта: {group_by}'

    return None


def collect_custom_filters(request: Request) -> tuple[list[tuple[str, str, str]], str | None]:
    fields = {field.key: field for field in get_storage(request.app).get_custom_fields()}
    filters: list[tuple[str, str, str]] = []
    for key, raw_values in request.args.items():
        if not key.startswith('custom__'):
            continue
        field = fields.get(key.removeprefix('custom__'))
        value = raw_values[0].strip() if raw_values else ''
        if field is None or not value:
            continue
        if field.field_type == 'number':
            try:
                value = str(int(value))
            except ValueError:
                return [], f'Фильтр «{field.label}» должен быть числом'
        filters.append((field.key, field.field_type, value))
    return filters, None


def parse_index_filters(args: dict) -> tuple[dict[str, str], str | None]:
    """Серверные фильтры главной (прототип 02) из GET-параметров.

    ФИО — подстрока; институт/вид спорта/уровень — точное совпадение;
    даты — дд.мм.гггг (тот же формат, что в фильтрах отчёта). Возвращает
    сырые значения (для переотображения в форме) и текст ошибки при
    некорректной дате. Статус обрабатывает маршрут — он ролевой.
    """
    filters: dict[str, str] = {}
    for key in ('name', 'institute', 'sport', 'level'):
        filters[key] = (get_param(args, key) or '').strip()
    for key in ('date_from', 'date_to'):
        raw_value = (get_param(args, key) or '').strip()
        if raw_value:
            try:
                datetime.strptime(raw_value, settings.date_format)
            except ValueError:
                return {}, f'Некорректная дата в фильтре ({key}): {raw_value}'
        filters[key] = raw_value
    filters['status'] = ''
    return filters, None


def parse_index_page(args: dict) -> int:
    """Номер страницы главной; пусто/нечисло — первая, меньше 1 — первая."""
    try:
        return max(1, int(get_param(args, 'page') or 1))
    except ValueError:
        return 1


def parse_index_per_page(args: dict) -> int:
    """Размер страницы главной; вне допустимого набора — дефолт."""
    raw_value = get_param(args, 'per_page') or ''
    try:
        value = int(raw_value)
    except ValueError:
        return INDEX_DEFAULT_PER_PAGE
    return value if value in INDEX_PER_PAGE_OPTIONS else INDEX_DEFAULT_PER_PAGE


def pager_items(current: int, total: int) -> list[dict]:
    """Номера страниц пагинации с разрывами: края (первая/последняя) и
    окно ±1 вокруг текущей, между непоследовательными номерами — «…»."""
    if total <= 0:
        return []
    wanted = {1, total}
    for offset in (-1, 0, 1):
        candidate = current + offset
        if 1 <= candidate <= total:
            wanted.add(candidate)
    items: list[dict] = []
    previous: int | None = None
    for number in sorted(wanted):
        if previous is not None and number - previous > 1:
            items.append({'gap': True})
        items.append({'page': number})
        previous = number
    return items


def build_index_query(filters: dict[str, str], per_page: int) -> str:
    """GET-параметры главной без page: для ссылок пагинации (page добавит
    шаблон/маршрут) и селекта «Показывать по» (он ставит per_page сам).
    Пустые значения и дефолтный per_page в URL не попадают — ссылки чище."""
    params = {key: value for key, value in filters.items() if value}
    if per_page != INDEX_DEFAULT_PER_PAGE:
        params['per_page'] = str(per_page)
    return urlencode(params)


def get_report_rows(request: Request, slice_key: str):
    """Строки отчёта для среза с фильтрами из запроса.

    Для «студента» — StudentInfo (профиль записи + метрики), для остальных
    срезов — ReportSliceRow (значение поля записи + метрики).
    """
    args = dict(request.args)
    storage = get_storage(request.app)

    filter_kwargs = dict(
        date_from=get_param(args, 'date_from'),
        date_to=get_param(args, 'date_to'),
        position=get_param(args, 'position'),
        level=get_param(args, 'level'),
        name=get_param(args, 'name'),
        institute=get_param(args, 'institute') or '',
        group=get_param(args, 'group') or '',
        sport=get_param(args, 'sport') or '',
        custom_filters=collect_custom_filters(request)[0],
    )

    if slice_key == DEFAULT_REPORT_SLICE:
        return storage.get_filtered(**filter_kwargs)
    return storage.get_grouped_report(slice_key, **filter_kwargs)


# Присутствие пользователей (№17): last_seen обновляется не чаще раза в
# PRESENCE_THROTTLE_SECONDS на пользователя. Кэш в памяти экономит сам вызов
# storage; окончательный троттлинг — в UPDATE (storage.touch_user_seen),
# поэтому несколько воркеров не amplify-ят запись.
PRESENCE_THROTTLE_SECONDS = 60
PRESENCE_ONLINE_WINDOW = timedelta(minutes=5)

_last_seen_touch: dict[int, float] = {}


def reset_presence_tracking() -> None:
    """Сбросить кэш троттлинга (для тестов)."""
    _last_seen_touch.clear()


def touch_user_seen(request: Request) -> None:
    """Обновить last_seen аутентифицированного пользователя (№17).

    Активность — сам факт запроса с валидной сессией; сбой обновления не
    должен ломать запрос, поэтому ошибки глотаем с логом.
    """
    try:
        user = get_auth_user(request) or {}
        user_id = user.get('id')
        if not user_id:
            return
        now = time.monotonic()
        if now - _last_seen_touch.get(user_id, 0.0) < PRESENCE_THROTTLE_SECONDS:
            return
        get_storage(request.app).touch_user_seen(user_id)
        _last_seen_touch[user_id] = now
    except Exception:
        logger.exception('Failed to update last_seen_at')


@app.on_response
async def static_no_cache(request: Request, response: HTTPResponse):
    """Статика без версионирования в URL: заставляем браузер
    перепроверять файл по Last-Modified, иначе после деплоя дни висит
    старый CSS/JS (симптом: «поломанная тёмная тема»)."""
    if request.path.startswith('/static'):
        response.headers['Cache-Control'] = 'no-cache'


@app.on_request
async def authorize_request(request: Request):
    request.ctx.auth_user = parse_auth_cookie(request)

    if request.path.startswith('/static'):
        return None

    if request.path in AUTH_ALLOWED_PATHS:
        return None

    if get_auth_user(request):
        touch_user_seen(request)
        if request.method == 'POST' and request.path != '/login':
            if not request_csrf_is_valid(request):
                return text(body='CSRF token missing or invalid', status=403)
        return None

    if request.method == 'GET':
        return redirect('/login')

    return unauthorized(request)


@app.get('/login')
async def login_page(request: Request):
    if get_auth_user(request):
        return redirect('/')

    return await render(
        template_name=jinja_env.get_template('login.html'),
        context={
            'request': request,
            'error_message': request.args.get('error'),
        },
    )


@app.post('/login')
async def login(request: Request):
    # client_ip учитывает PROXIES_COUNT=1: за nginx это реальный IP клиента
    # (последняя запись X-Forwarded-For), а не адрес прокси.
    if login_is_locked(request.client_ip):
        return text(body='Слишком много неудачных попыток входа. Повторите позже', status=LOGIN_REJECT_STATUS)

    username = str(get_form_value(request, 'username')).strip()
    password = str(get_form_value(request, 'password'))

    user = authenticate_user(request, username, password)
    if user is None:
        register_login_failure(request.client_ip)
        log_audit_event(request, 'login_failed', username=username)
        response = await render(
            template_name=jinja_env.get_template('login.html'),
            context={
                'request': request,
                'error_message': 'Неверный логин или пароль',
            },
        )
        response.status = 401
        return response

    clear_login_failures(request.client_ip)
    log_audit_event(
        request,
        'login_success',
        user_id=user.get('id'),
        username=user['username'],
    )
    try:
        # Присутствие (№17): успешный вход фиксирует last_login_at
        # (и last_seen — вход тоже активность).
        get_storage(request.app).set_user_last_login(user['id'])
    except Exception:
        logger.exception('Failed to update last_login_at')
    response = redirect('/')
    set_auth_cookie(request, response, user['username'], user['role'], pwd_ver=user.get('pwd_ver', 0))
    return response


@app.post('/logout')
async def logout(request: Request):
    response = redirect('/login')
    clear_auth_cookie(response)
    return response


@app.get('/')
async def index(request: Request):
    storage = get_storage(request.app)
    custom_fields = storage.get_custom_fields()
    owner_filter = get_current_user_id(request) if user_is_athlete(request) else None
    profile_hashes = student_hashes_for_request(request)
    args = dict(request.args)

    # Серверная фильтрация реестра (прототип 02): GET-параметры рендерит,
    # фильтры живут в URL — ссылки шарятся. Некорректная дата — явная ошибка.
    index_filters, filter_error = parse_index_filters(args)
    if filter_error is not None:
        return text(body=filter_error, status=400)
    if user_is_moderator(request):
        raw_status = (get_param(args, 'status') or '').strip()
        if raw_status in INDEX_STATUS_KEYS:
            index_filters['status'] = raw_status

    per_page = parse_index_per_page(args)
    filter_kwargs = dict(
        owner_id=owner_filter,
        student_id_hashes=profile_hashes,
        name=index_filters['name'],
        institute=index_filters['institute'],
        sport=index_filters['sport'],
        level=index_filters['level'],
        review_status=index_filters['status'],
        date_from=index_filters['date_from'],
        date_to=index_filters['date_to'],
    )
    total_count = storage.count_competitions_filtered(**filter_kwargs)
    page_count = max(1, -(-total_count // per_page))
    # Страница вне диапазона зажимается на последнюю (как в журнале аудита).
    page = min(parse_index_page(args), page_count)
    competitions = storage.get_competitions_page(
        limit=per_page,
        offset=(page - 1) * per_page,
        **filter_kwargs,
    )

    def page_url(page_number: int) -> str:
        query = build_index_query(index_filters, per_page)
        separator = '&' if query else ''
        return f'/?{query}{separator}page={page_number}'

    pager = []
    for item in pager_items(page, page_count):
        if item.get('gap'):
            pager.append({'gap': True})
        else:
            number = item['page']
            pager.append({'label': number, 'url': page_url(number), 'active': number == page})

    return await render(
        template_name=jinja_env.get_template('index.html'),
        context={
            'request': request,
            'competitions': competitions,
            'custom_fields': custom_fields,
            # №24: карта «целевая колонка → ключ link-поля» для рендера.
            'link_bindings': build_link_bindings(custom_fields),
            'admin_custom_fields': storage.get_custom_fields(include_inactive=True),
            'field_type_options': FIELD_TYPE_OPTIONS,
            'can_write': user_can_write(request),
            'can_import': user_is_moderator(request),
            'is_moderator': user_is_moderator(request),
            'is_athlete': user_is_athlete(request),
            'is_admin': user_is_admin(request),
            'users': storage.list_users() if user_is_admin(request) else [],
            'user_roles': USER_ROLES,
            'levels': storage.get_level_names(),
            'sport_options': storage.list_catalog('sport'),
            'institute_options': storage.list_catalog('institute'),
            'groups_by_institute': storage.get_group_options_by_institute(),
            # «Статус»-колонка видна, если у пользователя есть хоть одна
            # неподтверждённая запись — независимо от страницы и фильтра.
            'has_unapproved': storage.count_competitions_filtered(
                owner_id=owner_filter,
                student_id_hashes=profile_hashes,
                unapproved_only=True,
            )
            > 0,
            'attachments_by_record': build_attachments_by_record(storage.get_attachments()),
            'admin_levels': storage.list_levels() if user_is_admin(request) else [],
            'current_username': (get_auth_user(request) or {}).get('username'),
            'admin_message': get_param(args, 'admin_message'),
            'admin_error': get_param(args, 'admin_error'),
            # Композиция главной (прототип 02): фильтры, счётчик, пагинация.
            'index_filters': index_filters,
            'status_filter_options': INDEX_STATUS_FILTERS,
            'index_filter_query': urlencode({key: value for key, value in index_filters.items() if value}),
            'index_is_filtered': any(index_filters.values()),
            'total_count': total_count,
            'shown_count': len(competitions),
            'page': page,
            'page_count': page_count,
            'per_page': per_page,
            'per_page_options': INDEX_PER_PAGE_OPTIONS,
            'pager_items': pager,
            'prev_url': page_url(page - 1) if page > 1 else None,
            'next_url': page_url(page + 1) if page < page_count else None,
        },
    )


@app.get('/reports')
async def reports_page(request: Request):
    if user_is_athlete(request):
        return forbidden(request)
    storage = get_storage(request.app)
    groups_by_institute = storage.get_group_options_by_institute()
    return await render(
        template_name=jinja_env.get_template('report.html'),
        context={
            'request': request,
            'custom_fields': storage.get_custom_fields(),
            'levels': storage.get_level_names(),
            'report_export_columns': report_export_columns(),
            'report_slice_options': REPORT_SLICES,
            'report_export_columns_by_slice': {
                slice_key: report_export_columns(slice_key) for slice_key in REPORT_SLICE_KEYS
            },
            'institute_options': storage.list_catalog('institute'),
            'sport_options': storage.list_catalog('sport'),
            'groups_by_institute': groups_by_institute,
            'group_options': sorted({group for groups in groups_by_institute.values() for group in groups}),
        },
    )


# Календарь соревнований (волна A, docs/feedback-live.md №23): план на год.
# Статус «запланировано/прошло» — АВТОМАТИЧЕСКИ по дате окончания (выводимой):
# date_to (или date) >= сегодня → «запланировано».
CALENDAR_MONTH_NAMES: Sequence[str] = (
    'Январь',
    'Февраль',
    'Март',
    'Апрель',
    'Май',
    'Июнь',
    'Июль',
    'Август',
    'Сентябрь',
    'Октябрь',
    'Ноябрь',
    'Декабрь',
)
CALENDAR_STATUS_PLANNED = 'planned'
CALENDAR_STATUS_PAST = 'past'
CALENDAR_STATUS_FILTERS: Sequence[tuple[str, str]] = (
    ('', 'Все'),
    (CALENDAR_STATUS_PLANNED, 'запланировано'),
    (CALENDAR_STATUS_PAST, 'прошло'),
)


def calendar_event_status(event: dict, today=None) -> str:
    end_raw = event.get('date_to') or event['date']
    end = datetime.fromisoformat(end_raw).date() if isinstance(end_raw, str) else end_raw
    current = today if today is not None else datetime.now().date()
    return CALENDAR_STATUS_PLANNED if end >= current else CALENDAR_STATUS_PAST


def decorate_calendar_event(event: dict, today=None) -> dict:
    """Готовые к шаблону поля: даты-объекты, компактный период, статус."""
    date_from = datetime.fromisoformat(event['date'])
    date_to = datetime.fromisoformat(event['date_to']) if event.get('date_to') else None
    decorated = dict(event)
    decorated['date_from'] = date_from
    decorated['date_to_obj'] = date_to
    decorated['date_label'] = format_date_range(date_from, date_to)
    decorated['status'] = calendar_event_status(event, today)
    return decorated


def group_calendar_events_by_month(events: Sequence[dict]) -> list[dict]:
    """Группировка карточек по месяцам начала (прототип 15): заголовок
    «Сентябрь 2026 · N», внутри — карточка соревнования."""
    groups: list[dict] = []
    index: dict[tuple[int, int], dict] = {}
    for event in events:
        date_from = event['date_from']
        key = (date_from.year, date_from.month)
        if key not in index:
            group = {
                'key': key,
                'label': f'{CALENDAR_MONTH_NAMES[date_from.month - 1]} {date_from.year}',
                'events': [],
            }
            index[key] = group
            groups.append(group)
        index[key]['events'].append(event)
    for group in groups:
        group['count'] = len(group['events'])
    return groups


def parse_calendar_event_form(request: Request) -> tuple[dict, str | None]:
    """Разобрать форму соревнования календаря. Возвращает (значения, ошибка).

    Период — одно гибридное поле, тот же парсер, что у записей реестра
    (решение «Даты-диапазоны: визуально одно, под капотом два»).
    """
    name = clean_str(get_form_value(request, 'name'))
    if not name:
        return {}, 'Название обязательно'
    try:
        date_from, date_to = parse_date_value(get_form_value(request, 'date'), manual_input=True)
    except ValueError as exc:
        return {}, str(exc)
    url = clean_str(get_form_value(request, 'url'))
    # Ссылка календаря рендерится как href: принимаем только http/https,
    # прочие схемы (javascript:, data:, vbscript:) — отказ.
    if url and not is_http_url(url):
        return {}, 'Ссылка должна начинаться с http:// или https://'
    return {
        'name': name,
        'date': date_from,
        'date_to': date_to,
        'level': clean_str(get_form_value(request, 'level')),
        'sport': clean_str(get_form_value(request, 'sport')),
        'url': url,
    }, None


@app.get('/calendar')
async def calendar_page(request: Request):
    # Решение по ролям: admin/editor — полный доступ, viewer — просмотр без
    # кнопок, athlete — 403 (план — внутренняя кухня).
    if user_is_athlete(request):
        return forbidden(request)
    storage = get_storage(request.app)
    args = dict(request.args)

    raw_status = (get_param(args, 'status') or '').strip()
    if raw_status not in {key for key, _ in CALENDAR_STATUS_FILTERS}:
        raw_status = ''
    sport_filter = (get_param(args, 'sport') or '').strip()

    today = datetime.now().date()
    events = [decorate_calendar_event(event, today) for event in storage.list_calendar_events(sport=sport_filter)]
    if raw_status:
        events = [event for event in events if event['status'] == raw_status]

    return await render(
        template_name=jinja_env.get_template('calendar.html'),
        context={
            'request': request,
            'month_groups': group_calendar_events_by_month(events),
            'events_count': len(events),
            'planned_count': sum(1 for event in events if event['status'] == CALENDAR_STATUS_PLANNED),
            'past_count': sum(1 for event in events if event['status'] == CALENDAR_STATUS_PAST),
            'can_manage': user_is_moderator(request),
            'status_filter': raw_status,
            'status_filter_options': CALENDAR_STATUS_FILTERS,
            'sport_filter': sport_filter,
            'sport_options': storage.list_catalog('sport'),
            'level_options': storage.get_level_names(),
            **get_flash_args(request),
        },
    )


@app.post('/calendar/new')
async def create_calendar_event(request: Request):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    values, error = parse_calendar_event_form(request)
    if error is not None:
        return text(body=error, status=400)

    get_storage(request.app).create_calendar_event(
        name=values['name'],
        date=values['date'].isoformat(),
        date_to=values['date_to'].isoformat() if values['date_to'] else None,
        level=values['level'],
        sport=values['sport'],
        url=values['url'],
    )
    return build_redirect_with_message(
        message=f'Соревнование «{values["name"]}» запланировано',
        url='/calendar',
    )


@app.post('/calendar/<event_id>/edit')
async def edit_calendar_event(request: Request, event_id: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_id = int(event_id)
    except ValueError:
        return text(body='Invalid event id', status=400)

    storage = get_storage(request.app)
    event = storage.get_calendar_event(numeric_id)
    if event is None:
        return text(body='Event not found', status=404)

    values, error = parse_calendar_event_form(request)
    if error is not None:
        return text(body=error, status=400)

    new_date = values['date'].isoformat()
    new_date_to = values['date_to'].isoformat() if values['date_to'] else None
    # P2: правка события — владелец полей участия; связанные записи (по
    # calendar_event_id) синхронизируются той же транзакцией (см.
    # SQLiteAdapter.update_calendar_event), synced — их число для аудита.
    synced_participations = storage.update_calendar_event(
        event_id=numeric_id,
        name=values['name'],
        date=new_date,
        date_to=new_date_to,
        level=values['level'],
        sport=values['sport'],
        url=values['url'],
    )
    log_audit_event(
        request,
        'calendar_event_edited',
        {
            'event_id': numeric_id,
            'name': {'old': event['name'], 'new': values['name']},
            'date': {'old': event['date'], 'new': new_date},
            'date_to': {'old': event.get('date_to'), 'new': new_date_to},
            'sport': {'old': event['sport'], 'new': values['sport']},
            'level': {'old': event['level'], 'new': values['level']},
            'synced_participations': synced_participations,
        },
    )
    # Правка со страницы соревнования (волна B, прототип 16) возвращает
    # внутрь события; из календаря — в календарь. next принимаем только
    # как путь внутрь /calendar/ — открытый редирект исключён.
    next_url = get_form_value(request, 'next')
    if next_url.startswith('/calendar/'):
        return build_redirect_with_message(message='Соревнование обновлено', url=next_url)
    return build_redirect_with_message(message='Соревнование обновлено', url='/calendar')


# Страница соревнования (волна B, прототип 16): участники = записи реестра
# по пресету события (name + date + date_to). Фильтр результата — GET,
# ссылки шарятся, как на главной.
CALENDAR_RESULT_FILTERS: Sequence[tuple[str, str]] = (
    ('', 'Все'),
    ('with', 'С результатом'),
    ('without', 'Без результата'),
)


def get_calendar_event_or_error(request: Request, event_id: str):
    """Общий разбор id события из URL: (event dict | None, response | None)."""
    try:
        numeric_id = int(event_id)
    except ValueError:
        return None, text(body='Invalid event id', status=400)
    event = get_storage(request.app).get_calendar_event(numeric_id)
    if event is None:
        return None, text(body='Event not found', status=404)
    return event, None


@app.get('/calendar/<event_id>')
async def calendar_event_page(request: Request, event_id: str):
    # Решение по ролям — как у календаря: admin/editor — полный доступ,
    # viewer — просмотр без кнопок, athlete — 403.
    if user_is_athlete(request):
        return forbidden(request)

    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    args = dict(request.args)
    result_filter = (get_param(args, 'result') or '').strip()
    if result_filter not in {key for key, _ in CALENDAR_RESULT_FILTERS}:
        result_filter = ''

    storage = get_storage(request.app)
    participants = storage.list_calendar_event_participants(event['id'])
    total_count = len(participants)
    no_result_count = sum(1 for participant in participants if participant['position'] == 0)
    if result_filter == 'with':
        participants = [participant for participant in participants if participant['position'] != 0]
    elif result_filter == 'without':
        participants = [participant for participant in participants if participant['position'] == 0]

    return await render(
        template_name=jinja_env.get_template('calendar_event.html'),
        context={
            'request': request,
            'event': decorate_calendar_event(event),
            'participants': participants,
            'total_count': total_count,
            'no_result_count': no_result_count,
            'shown_count': len(participants),
            'can_write': user_can_write(request),
            # Кнопка удаления участника — только admin (роут /competition/<id>/delete
            # админский): без флага кнопка в шаблоне не рисовалась вовсе.
            'is_admin': user_is_admin(request),
            # Управление файлом положения — модераторы (admin/editor);
            # can_write сюда не годится: он включает атлета.
            'can_manage': user_is_moderator(request),
            'result_filter': result_filter,
            'result_filter_options': CALENDAR_RESULT_FILTERS,
            'sport_options': storage.list_catalog('sport'),
            'level_options': storage.get_level_names(),
            # Настройки полей для строки ручного добавления (required/тип
            # Курса и Места — как в реестре; ФИО всегда обязательное).
            'base_field_settings': get_base_field_settings(storage),
            **get_flash_args(request),
        },
    )


@app.post('/calendar/<event_id>/participants')
async def add_calendar_event_participant(request: Request, event_id: str):
    # Участник добавляется как ОБЫЧНАЯ запись реестра (историчность —
    # docs/data-model-decisions.md): пресет события (name/date/date_to)
    # копируется в запись, никаких FK. Роли: admin/editor — полный доступ
    # (viewer пишет? нет), поэтому require_moderator; запись сразу approved.
    # При явном выборе карточки студента в резолвере ФИО форма присылает
    # student_ref_id: связь стабильная, но снимок полей остаётся снимком
    # формы (Excel/ручной ввод всегда приоритетнее карточки).
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    back_url = f'/calendar/{event["id"]}'
    student_name = clean_str(get_form_value(request, 'student_name'))
    if not student_name:
        return build_redirect_with_message(error='ФИО обязательно', url=back_url)

    event_date = datetime.fromisoformat(event['date'])
    event_date_to = datetime.fromisoformat(event['date_to']) if event.get('date_to') else None
    storage = get_storage(request.app)

    # Явный выбор карточки студента: непустое значение валидируется (существует
    # и активна); невалидное — отказ БЕЗ создания записи, молчаливо терять
    # выбранную админом связь нельзя.
    student_ref_id: int | None = None
    student_ref_raw = get_form_value(request, 'student_ref_id').strip()
    if student_ref_raw:
        try:
            student_ref_id = int(student_ref_raw)
        except ValueError:
            student_ref_id = None
        student = storage.get_student_by_id(student_ref_id) if student_ref_id is not None else None
        if student is None or not student['active']:
            return build_redirect_with_message(
                error='Выбранная карточка студента не найдена или неактивна. '
                'Сбросьте связь и выберите карточку заново.',
                url=back_url,
            )

    custom_fields = storage.get_custom_fields()
    record = {
        'ФИО': student_name,
        'Пол': get_form_value(request, 'student_sex'),
        'Институт': get_form_value(request, 'institute'),
        'Группа': get_form_value(request, 'group'),
        'Вид спорта': event['sport'],
        # Период — компактной строкой: тот же парсер, что у ручного ввода
        # (format_date_range реимпортируем обратно).
        'Дата': format_date_range(event_date, event_date_to),
        'Уровень соревнований': event['level'],
        'Название соревнований': event['name'],
        # Место опционально («можно дописать позже»): пусто → position = 0.
        'Место': get_form_value(request, 'position'),
        'Курс': get_form_value(request, 'course'),
    }

    try:
        competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
    except (TypeError, ValueError) as exc:
        return build_redirect_with_message(error=str(exc), url=back_url)

    competition.student_ref_id = student_ref_id
    # P2: участник события получает явную ссылку на событие — состав
    # участников читается по calendar_event_id, а не по пресету.
    competition.calendar_event_id = event['id']
    ensure_catalog_values(storage, [competition])
    storage.save_competitions(
        [competition],
        review_status='approved',
        owner_id=get_current_user_id(request),
    )
    return build_redirect_with_message(
        message=f'Участник «{student_name}» добавлен',
        url=back_url,
    )


@app.post('/calendar/<event_id>/delete')
async def delete_calendar_event(request: Request, event_id: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_id = int(event_id)
    except ValueError:
        return text(body='Invalid event id', status=400)

    storage = get_storage(request.app)
    event = storage.get_calendar_event(numeric_id)
    if event is None:
        return text(body='Event not found', status=404)

    # Удаление с участниками — отказ: сначала удалить участников (волна B
    # даст инструмент), счётчик подсказывает объём.
    participants = storage.count_calendar_event_participants(numeric_id)
    if participants > 0:
        return build_redirect_with_message(
            error=f'У соревнования «{event["name"]}» есть участники: {participants}. '
            'Сначала удалите участников, затем соревнование.',
            url='/calendar',
        )

    storage.delete_calendar_event(numeric_id)
    # Файл положения живёт вне БД: каталог события удаляется вместе с ним
    # (best-effort — как каталоги вложений записей).
    shutil.rmtree(calendar_regulation_dir(numeric_id), ignore_errors=True)
    log_audit_event(
        request,
        'calendar_event_deleted',
        {'event_id': numeric_id, 'name': event['name'], 'date': event['date']},
    )
    return build_redirect_with_message(message='Соревнование удалено', url='/calendar')


# Файл положения события календаря (решение 2026-09-22): один файл на событие
# (PDF/JPEG/PNG до 5 МБ, та же валидация, что у вложений), колонки
# calendar_events.regulation_* + файл в data/files/calendar/<event_id>/.
# Скачивание — все не-атлеты (как страница события), замена/удаление —
# модераторы.


@app.get('/calendar/<event_id>/regulation')
async def download_calendar_regulation(request: Request, event_id: str):
    # Права — как у страницы события: аноним перенаправляется на вход
    # middleware, атлету страница события недоступна — файл тоже.
    if user_is_athlete(request):
        return forbidden(request)
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    filename = event.get('regulation_filename') or ''
    stored_name = event.get('regulation_stored_name') or ''
    if not filename or not stored_name:
        return text(body='Regulation not found', status=404)
    file_path = regulation_source_path(event['id'], stored_name)
    if not file_path.is_file():
        return text(body='Not Found', status=404)

    extension = Path(filename).suffix.lower().lstrip('.')
    return raw(
        await asyncio.to_thread(file_path.read_bytes),
        headers={
            'content-type': ATTACHMENT_EXTENSIONS.get(extension, 'application/octet-stream'),
            'content-disposition': f'attachment; filename="{filename}"',
        },
    )


@app.post('/calendar/<event_id>/regulation')
async def upload_calendar_regulation(request: Request, event_id: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    back_url = f'/calendar/{event["id"]}'
    upload_file = request.files.get('regulation')
    if upload_file is None or not upload_file.body:
        return build_redirect_with_message(error='Выберите файл.', url=back_url)
    if len(upload_file.body) > ATTACHMENT_MAX_SIZE:
        return build_redirect_with_message(error='Файл больше 5 МБ', url=back_url)
    filename = (upload_file.name or 'regulation').rsplit('/', 1)[-1]
    content_type = detect_attachment_type(upload_file.body, filename)
    if content_type is None:
        return build_redirect_with_message(error='Допустимы только PDF, JPEG и PNG', url=back_url)

    # Различаем первое прикрепление и замену до загрузки (flash-текст).
    had_regulation = bool(event.get('regulation_filename'))

    # Порядок как у вложений: сначала новый файл на диск, затем колонки БД,
    # затем best-effort удаление прежнего файла (БД уже согласована).
    extension = Path(filename).suffix.lower().lstrip('.') or 'bin'
    stored_name = f'{uuid.uuid4().hex}.{extension}'
    event_dir = calendar_regulation_dir(event['id'])
    event_dir.mkdir(parents=True, exist_ok=True)
    (event_dir / stored_name).write_bytes(upload_file.body)

    old_stored_name = event.get('regulation_stored_name') or ''
    get_storage(request.app).set_calendar_regulation(event['id'], filename, stored_name)

    if old_stored_name:
        try:
            regulation_source_path(event['id'], old_stored_name).unlink(missing_ok=True)
        except OSError:
            logger.warning('Failed to remove old regulation file of event %s', event['id'], exc_info=True)

    log_audit_event(
        request,
        'calendar_regulation_uploaded',
        {
            'event_id': event['id'],
            'event_name': event['name'],
            'filename': filename,
            'replaced': bool(old_stored_name),
        },
    )
    return build_redirect_with_message(
        message='Положение заменено' if had_regulation else 'Положение прикреплено',
        url=back_url,
    )


@app.post('/calendar/<event_id>/regulation/delete')
async def delete_calendar_regulation(request: Request, event_id: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    back_url = f'/calendar/{event["id"]}'
    filename = event.get('regulation_filename') or ''
    stored_name = event.get('regulation_stored_name') or ''
    if not stored_name:
        return build_redirect_with_message(error='Положение не прикреплено.', url=back_url)

    get_storage(request.app).clear_calendar_regulation(event['id'])
    shutil.rmtree(calendar_regulation_dir(event['id']), ignore_errors=True)
    log_audit_event(
        request,
        'calendar_regulation_deleted',
        {'event_id': event['id'], 'event_name': event['name'], 'filename': filename},
    )
    return build_redirect_with_message(message='Положение удалено', url=back_url)


# --- Импорт участников события календаря из Excel. ---
#
# Массовое добавление участников В КОНТЕКСТЕ события: пресет соревнования
# (название/даты/уровень/вид спорта) берётся из события и в Excel НЕ
# повторяется — обязательное значение только ФИО. Паттерн импорта студентов
# (Phase 2.5): разбор → staging-сессия в памяти → предпросмотр с
# кандидатами-карточками → явные решения построчно или батчем; в БД не
# пишется НИЧЕГО до подтверждения. Приоритет снимка участия: ЯВНОЕ значение
# Excel > карточка Student (явно выбранная или единственный точный кандидат
# для автозаполнения пол/институт/группа/курс) > пусто; карточка Student при
# этом НЕ меняется. Подтверждённые участия получают student_ref_id — это
# только запись стабильной связи, runtime identity (Phase 3) не включается.
# Права: страницы/действия импорта — модераторы (admin/editor); «создать
# карточку студента» — только admin (конвенция Phase 1/2.5: Students —
# admin-only), editor видит выбор существующей карточки и пропуск.

EVENT_PARTICIPANT_IMPORT_BASE_COLUMNS: Sequence[str] = ('ФИО', 'Пол', 'Институт', 'Группа', 'Курс', 'Место')
EVENT_PARTICIPANT_IMPORT_MAX_ROWS = 5000
EVENT_IMPORT_SESSION_TTL_SECONDS = 2 * 60 * 60
# Поиск карточки для строки предпросмотра («Найти студента»): лимит
# результатов и максимальная длина запроса.
EVENT_IMPORT_STUDENT_SEARCH_LIMIT = 8
EVENT_IMPORT_STUDENT_SEARCH_QUERY_MAX_LENGTH = 100

# Staging живёт в памяти (прецедент student_import_sessions): однопроцессное
# приложение, рестарт просто теряет сессии — безопасно, файл загрузят заново.
# Один файл на пользователя; сессия привязана к владельцу и событию; ФИО из
# файла не логируются.
event_import_sessions: dict[str, dict] = {}


def event_import_custom_fields(custom_fields: Sequence[CustomField]) -> list[CustomField]:
    """Custom-поля импорта участников события: активные с show_in_template,
    КРОМЕ url-полей (ссылки уровня записи/события к снимку участия не
    относятся). Признак — field_type, не label.

    Сами url-поля в системе остаются (реестр, обычный импорт/экспорт,
    link_target) — исключение касается только этого workflow. Фильтр
    применяется на входе каждого пути (шаблон, разбор, предпросмотр, правки,
    валидация, добавление) и обязан доходить до валидации
    (build_event_participant_competition): иначе гипотетическое
    ОБЯЗАТЕЛЬНОЕ url-поле заблокировало бы все строки импорта участников.
    """
    return [field for field in custom_fields if field.active and field.show_in_template and field.field_type != 'url']


def event_participant_import_columns(custom_fields: Sequence[CustomField]) -> list[str]:
    """Колонки файла/шаблона импорта участников: базовые + активные
    custom-поля с show_in_template (генерически, по существующим флагам)."""
    return [
        *EVENT_PARTICIPANT_IMPORT_BASE_COLUMNS,
        *[field.label for field in custom_fields if field.show_in_template],
    ]


def normalize_import_position(value) -> str:
    """Место из ячейки Excel как строка — та же нормализация, что у курса
    (числовые ячейки pandas отдаёт float'ами: 1 → «1»). Пусто → «»."""
    return normalize_import_course(value)


def parse_event_participant_rows(
    df: pd.DataFrame,
    custom_fields: Sequence[CustomField],
) -> tuple[list[dict], list[str], list[str]]:
    """Разобранные строки файла импорта участников + неизвестные колонки +
    колонки custom-полей, присутствующие в файле.

    Строка staging: row_number = index + 2 (заголовок — строка 1), статус
    pending/error. Пол — normalize_student_sex («м» → «М», недопустимое —
    ошибка строки, а не молчаливая правка); курс и место — числовая
    нормализация Excel. Дубли строк в файле здесь НЕ помечаются: группы
    пересчитываются при каждом рендере (правки строк меняют состав).
    """
    import_columns = set(event_participant_import_columns(custom_fields))
    columns = [str(column) for column in df.columns]
    unknown_columns = [column for column in columns if column not in import_columns]
    custom_by_label = {field.label: field for field in custom_fields if field.show_in_template}
    custom_columns = [column for column in columns if column in custom_by_label]
    rows: list[dict] = []
    for index, record in df.iterrows():
        full_name = clean_str(record.get('ФИО'))
        sex = normalize_student_sex(clean_str(record.get('Пол')))
        error = None
        if not full_name:
            error = 'Пустое ФИО.'
        elif sex is None:
            error = 'Пол должен быть «М» или «Ж».'
        rows.append(
            {
                'row_number': int(index) + 2,
                'full_name': full_name,
                'sex': sex or '',
                'institute': clean_str(record.get('Институт')),
                'group': clean_str(record.get('Группа')),
                'course': normalize_import_course(record.get('Курс')),
                'position': normalize_import_position(record.get('Место')),
                'custom': {custom_by_label[column].key: clean_str(record.get(column)) for column in custom_columns},
                'status': 'error' if error else 'pending',
                'error': error,
                # Явный выбор карточки админом (select-student/create-student).
                'student_id': None,
                # Явное решение «без карточки» (clear-student): подавляет
                # автозаполнение единственным точным кандидатом — участие
                # добавляется без связи, сопоставить можно позже.
                'student_forced_none': False,
                # Карточка, реально использованная добавленной записью.
                'added_student_id': None,
            }
        )
    return rows, unknown_columns, custom_columns


EVENT_PARTICIPANT_SNAPSHOT_KEYS: Sequence[str] = ('full_name', 'sex', 'institute', 'group', 'course', 'position')


def event_import_duplicate_rows(rows: Sequence[dict]) -> dict[int, list[int]]:
    """Полностью одинаковые строки файла: номер строки → номера всех строк группы.

    «Одинаковость» — полное равенство снимка импортируемых колонок
    (strip + casefold, включая custom-значения): «Иванов|100 м|1» и
    «Иванов|200 м|3» — ДВА корректных участия, а точная копия строки —
    повод спросить админа. Молчаливого skip нет: обе строки группы
    попадают в «Требуют решения», пока ОБЕ не разобраны (см.
    event_import_pending_duplicates — семантика «оставить одну / создать
    обе / исправить / пропустить» из §17 достигается действиями).
    """
    groups: dict[tuple, list[int]] = {}
    for row in rows:
        key = (
            tuple((row[column] or '').strip().casefold() for column in EVENT_PARTICIPANT_SNAPSHOT_KEYS),
            tuple(sorted((key.strip(), (value or '').strip().casefold()) for key, value in row['custom'].items())),
        )
        groups.setdefault(key, []).append(row['row_number'])
    return {row_number: numbers for numbers in groups.values() if len(numbers) > 1 for row_number in numbers}


def event_import_pending_duplicates(rows: Sequence[dict]) -> dict[int, list[int]]:
    """Дубли среди НЕРАЗОБРАННЫХ (pending) строк сессии.

    Это БЛОКИРУЮЩАЯ часть event_import_duplicate_rows: повтор-строка требует
    решения, только пока её не разобранная копия тоже висит в предпросмотре.
    После явного решения над копией (пропуск/добавление) или её ошибки
    строка перестаёт блокироваться и становится добавляемой — «создать обе»
    выполняется добавлением первой (вторая после этого разблокируется).
    """
    return event_import_duplicate_rows([row for row in rows if row['status'] == 'pending'])


def sweep_event_import_sessions() -> None:
    """Удалить просроченные сессии (TTL); вызывается при каждом обращении."""
    now = time.time()
    expired = [
        token
        for token, session in event_import_sessions.items()
        if now - session['created_at'] > EVENT_IMPORT_SESSION_TTL_SECONDS
    ]
    for token in expired:
        event_import_sessions.pop(token, None)


def get_event_import_session(request: Request, token: str, event_id: int):
    """Сессия предпросмотра для запроса: (session, None) или (None, redirect).

    Неизвестный/просроченный токен, чужая сессия и сессия другого события —
    одно и то же поведение: возврат на страницу загрузки с flash.
    """
    sweep_event_import_sessions()
    session = event_import_sessions.get(token)
    if session is None or session['user_id'] != get_current_user_id(request) or session['event_id'] != event_id:
        return None, build_redirect_with_message(
            error='Время сессии предпросмотра истекло. Загрузите файл заново.',
            url=f'/calendar/{event_id}/participants/import',
        )
    return session, None


def replace_event_import_session(user_id: int, token: str, session: dict) -> None:
    """Сохранить новую сессию предпросмотра; прежняя сессия того же
    пользователя заменяется (один файл на пользователя)."""
    sweep_event_import_sessions()
    for existing_token, existing in list(event_import_sessions.items()):
        if existing['user_id'] == user_id:
            event_import_sessions.pop(existing_token, None)
    event_import_sessions[token] = session


def event_import_preview_url(event_id: int, token: str) -> str:
    return f'/calendar/{event_id}/participants/import/preview/{token}'


def event_import_row(session: dict, raw_row_number: str):
    """Строка сессии по номеру из URL: (row, None) или (None, ответ 400)."""
    row_number, error = parse_reconcile_int(raw_row_number, 'row number')
    if error is not None:
        return None, error
    for row in session['rows']:
        if row['row_number'] == row_number:
            return row, None
    return None, text(body='Invalid row number', status=400)


def event_participant_preset(event: dict) -> dict:
    """Пресет события в колонки записи: sport/date/date_to/level/name.

    Тот же пресет, что у ручного добавления участника (POST
    /calendar/<id>/participants): период — компактной строкой, парсер
    ручного ввода понимает её обратно.
    """
    event_date = datetime.fromisoformat(event['date'])
    event_date_to = datetime.fromisoformat(event['date_to']) if event.get('date_to') else None
    return {
        'Вид спорта': event['sport'],
        'Дата': format_date_range(event_date, event_date_to),
        'Уровень соревнований': event['level'],
        'Название соревнований': event['name'],
    }


def event_import_effective_student(storage: SQLiteAdapter, row: dict, candidates: Sequence[dict]) -> dict | None:
    """Карточка для автозаполнения/связи: явный выбор админа, иначе при
    отсутствии решения «без карточки» — единственный точный кандидат.

    Полные тёзки (≥2) и 0 кандидатов НЕ выбираются автоматически — участие
    просто добавляется без связи (сопоставить можно позже). Явное решение
    «без карточки» (student_forced_none, clear-student) подавляет
    автозаполнение единственным кандидатом. Выбранная карточка проверяется
    на активность. Возвращаемая карточка нормализована: id — в 'id' (у
    кандидатов find_student_cards ключ 'student_id', у get_student_by_id —
    'id').
    """
    if row.get('student_id') is not None:
        student = storage.get_student_by_id(row['student_id'])
        if student is not None and student['active']:
            return student
        return None
    if row.get('student_forced_none'):
        return None
    if len(candidates) == 1:
        candidate = dict(candidates[0])
        candidate['id'] = candidate['student_id']
        return candidate
    return None


def event_import_row_record(
    storage: SQLiteAdapter,
    row: dict,
    preset: dict,
    student: dict | None,
    custom_fields: Sequence[CustomField],
    file_groups: dict[str, list[dict]] | None = None,
) -> tuple[dict, list[str], dict | None]:
    """Снимок строки для build_competition: Excel > Student > пусто.

    Возвращает (record, автозаполненные-из-карточки-колонки, hints).
    Карточка Student только читается (master не мутируется). Если институт
    пуст и группа заполнена (и карточка института не дала) — подсказки
    student_catalog_hints: сначала карта группа→институты ТЕКУЩЕГО файла,
    затем справочник (справочники при этом только читаются).
    """
    record = {
        'ФИО': row['full_name'],
        'Пол': row['sex'],
        'Институт': row['institute'],
        'Группа': row['group'],
        'Курс': row['course'],
        'Место': row['position'],
        **preset,
    }
    record.update(
        {field.label: row['custom'].get(field.key, '') for field in custom_fields if field.key in row['custom']}
    )
    autofilled: list[str] = []
    if student is not None:
        for column, student_key in (
            ('Пол', 'sex'),
            ('Институт', 'institute'),
            ('Группа', 'group_name'),
            ('Курс', 'course'),
        ):
            value = str(student.get(student_key) or '').strip()
            if value and not str(record.get(column) or '').strip():
                record[column] = value
                autofilled.append(column)
    hints = None
    if not str(record.get('Институт') or '').strip() and str(record.get('Группа') or '').strip():
        hints = student_catalog_hints(storage, '', record['Группа'], file_groups)
        if hints.get('institute'):
            record['Институт'] = hints['institute']
            record['Группа'] = hints.get('group') or record['Группа']
    return record, autofilled, hints


def build_event_participant_competition(
    record: dict,
    custom_fields: Sequence[CustomField],
    storage: SQLiteAdapter,
) -> tuple[Competition | None, str | None]:
    """Финальная валидация строки тем же механизмом, что ручной ввод
    (build_competition: field_settings, required custom-поля, число места).
    Возвращает (competition, None) или (None, текст ошибки)."""
    try:
        return build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage), None
    except (TypeError, ValueError) as exc:
        return None, str(exc)


def event_import_existing_participations(
    existing_participants: Sequence[dict],
    row: dict,
    student: dict | None,
) -> list[dict]:
    """Уже существующие участия того же человека в этом событии.

    Совпадение — ФИО (casefold) ИЛИ выбранная карточка (student_ref_id).
    Только информация для предпросмотра («уже участвует»), НЕ блокировка:
    несколько участий одного Student в событии — нормальный случай
    (разные дисциплины), решение всегда за админом.
    """
    target = (row['full_name'] or '').strip().casefold()
    student_ref = student['id'] if student else None
    return [
        participant
        for participant in existing_participants
        if (participant['student_name'] or '').strip().casefold() == target
        or (student_ref is not None and participant.get('student_ref_id') == student_ref)
    ]


def event_import_row_display(record: dict, competition: Competition | None) -> dict:
    """Значения строки для таблицы предпросмотра: из ПОСТРОЕННОЙ записи
    (ровно то, что будет вставлено — с канонизацией справочников), при
    ошибке валидации — из record после автозаполнения."""
    if competition is not None:
        return {
            'sex': competition.student_sex,
            'institute': competition.institute,
            'group': competition.group,
            'course': competition.course,
            'position': competition.position,
        }
    return {
        'sex': record.get('Пол', ''),
        'institute': record.get('Институт', ''),
        'group': record.get('Группа', ''),
        'course': record.get('Курс', ''),
        'position': record.get('Место', ''),
    }


def event_import_preview_context(
    storage: SQLiteAdapter,
    event: dict,
    session: dict,
    custom_fields: Sequence[CustomField],
) -> dict:
    """Контекст предпросмотра импорта участников.

    Всё производное (кандидаты, автозаполнение, подсказки, дубли файла,
    существующие участия, валидация, категории, счётчики) пересчитывается
    при каждом рендере; в сессии — строки файла, явные решения и
    фиксируемые рендером признаки had_student_id/had_validation_error
    (см. event_import_bulk_rows).
    """
    duplicates = event_import_duplicate_rows(session['rows'])
    # Блокирующие дубли — только среди не разобранных строк (§17): после
    # явного решения над копией (пропуск/добавление) строка разблокируется.
    pending_duplicates = event_import_pending_duplicates(session['rows'])
    file_groups = student_file_group_institutes(session['rows'])
    preset = event_participant_preset(event)
    existing_participants = storage.list_calendar_event_participants(event['id'])
    groups: dict[str, list[dict]] = {'ready': [], 'review': [], 'error': [], 'resolved': []}
    resolved_counts = {'added': 0, 'skipped': 0}
    pending = 0
    linked_ready = 0
    unlinked_ready = 0
    for row in session['rows']:
        status = row['status']
        view = {
            **row,
            'candidates': [],
            'effective_student': None,
            'auto_preselected': False,
            'autofilled': [],
            'hints': None,
            'validation_error': None,
            'display': None,
            'custom_display': {},
            'duplicate_rows': duplicates.get(row['row_number']),
            'existing': [],
        }
        if status == 'pending':
            pending += 1
            view['candidates'] = storage.find_student_candidates(row['full_name'])
            student = event_import_effective_student(storage, row, view['candidates'])
            view['effective_student'] = student
            view['auto_preselected'] = student is not None and row.get('student_id') is None
            record, autofilled, hints = event_import_row_record(
                storage, row, preset, student, custom_fields, file_groups
            )
            view['autofilled'] = autofilled
            view['hints'] = hints
            competition, validation_error = build_event_participant_competition(record, custom_fields, storage)
            view['validation_error'] = validation_error
            view['display'] = event_import_row_display(record, competition)
            view['custom_display'] = {
                field.label: (
                    competition.extra_data.get(field.key, '')
                    if competition is not None
                    else record.get(field.label, '')
                )
                for field in custom_fields
                if field.key in row['custom']
            }
            view['existing'] = event_import_existing_participations(existing_participants, row, student)
            # Классификация рендера фиксируется в строке сессии: bulk-commit
            # берёт только «готовые» строки и отменяется целиком, если их
            # состав изменился ПОСЛЕ рендера. После правки строки следующий
            # рендер пересчитает признаки естественно.
            row['had_student_id'] = student['id'] if student else None
            row['had_validation_error'] = bool(validation_error)
            # Сопоставление с карточкой НЕ блокирует импорт: «Готова» — любая
            # валидная строка без не разобранной строки-дубля (карточка при
            # этом может подобраться автоматически — единственный точный
            # кандидат). «Требуют решения» — только неразобранные дубли,
            # ошибочные — только ошибки валидации.
            blocked_by_duplicate = len(pending_duplicates.get(row['row_number'], [])) > 1
            if validation_error:
                category = 'error'
            elif not blocked_by_duplicate:
                category = 'ready'
                if student is None:
                    unlinked_ready += 1
                else:
                    linked_ready += 1
            else:
                category = 'review'
        elif status == 'error':
            category = 'error'
        else:
            category = 'resolved'
            if status in resolved_counts:
                resolved_counts[status] += 1
        groups[category].append(view)
    return {
        'token': session['token'],
        'event': decorate_calendar_event(event),
        'filename': session['filename'],
        'unknown_columns': session['unknown_columns'],
        'custom_columns': session['custom_columns'],
        'custom_keys': {field.label: field.key for field in custom_fields},
        'groups': groups,
        'total': len(session['rows']),
        'pending': pending,
        'ready_count': len(groups['ready']),
        'review_count': len(groups['review']),
        'error_count': len(groups['error']),
        'resolved_count': len(groups['resolved']),
        'added_count': resolved_counts['added'],
        'skipped_count': resolved_counts['skipped'],
        # Для confirm-текста bulk: сколько готовых строк уйдут со связью и
        # сколько — без карточки.
        'linked_ready': linked_ready,
        'unlinked_ready': unlinked_ready,
    }


def event_import_bulk_rows(session: dict) -> list[dict]:
    """Строки массового добавления: ровно группа «Готовы к добавлению» на
    момент последнего рендера предпросмотра — счётчик кнопки совпадает с
    числом вставляемых строк.

    Признаки had_validation_error фиксирует контекст предпросмотра; строки
    «Требуют решения» (не разобранные дубли файла) в батч НЕ входят и отмену
    не вызывают — повторная проверка на момент клика отменяет всё, только
    если состав «готовых» строк изменился. Сопоставление с карточкой
    обязательным НЕ является: строка без карточки добавляется с
    student_ref_id IS NULL."""
    pending_duplicates = event_import_pending_duplicates(session['rows'])
    return [
        row
        for row in session['rows']
        if row['status'] == 'pending'
        and not row.get('had_validation_error')
        and len(pending_duplicates.get(row['row_number'], [])) <= 1
    ]


def event_import_counters(session: dict) -> dict[str, int]:
    """Итоги сессии для аудита/сообщений (без содержимого файла)."""
    counters = {'added': 0, 'skipped': 0, 'errors': 0}
    for row in session['rows']:
        if row['status'] == 'error':
            counters['errors'] += 1
        elif row['status'] in counters:
            counters[row['status']] += 1
    counters['total'] = len(session['rows'])
    return counters


def event_import_upload_error(df: pd.DataFrame) -> str | None:
    """Ошибки файла на уровне колонок/объёма; None — файл пригоден."""
    if 'ФИО' not in {str(column) for column in df.columns}:
        return 'В файле нет обязательной колонки «ФИО». Скачайте шаблон и заполните его.'
    if len(df) > EVENT_PARTICIPANT_IMPORT_MAX_ROWS:
        return f'В файле больше {EVENT_PARTICIPANT_IMPORT_MAX_ROWS} строк. Разбейте список на части.'
    if df.empty:
        return 'В файле нет строк с данными.'
    return None


@app.get('/calendar/<event_id>/participants/import')
async def event_participants_import_page(request: Request, event_id: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error
    return await render(
        template_name=jinja_env.get_template('calendar_event_import.html'),
        context={
            'request': request,
            'event': decorate_calendar_event(event),
            **get_flash_args(request),
        },
    )


@app.get('/calendar/<event_id>/participants/import/template')
async def export_event_participants_template(request: Request, event_id: str):
    """Пустой шаблон: ФИО|Пол|Институт|Группа|Курс|Место + активные
    custom-поля с show_in_template (кроме url-полей — см.
    event_import_custom_fields). Без колонок события (пресет берётся из
    события) и без каких-либо данных."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    custom_fields = event_import_custom_fields(get_storage(request.app).get_custom_fields())
    df = pd.DataFrame(columns=event_participant_import_columns(custom_fields))
    buffer = BytesIO()
    await asyncio.to_thread(df.to_excel, buffer, index=False)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="Шаблон_участники_{now_str}.xlsx"',
        },
    )


@app.post('/calendar/<event_id>/participants/import')
async def event_participants_import_upload(request: Request, event_id: str):
    """Загрузка файла: разбор → staging в памяти → предпросмотр.

    Ошибки файла — возврат на страницу загрузки с flash; в БД не пишется
    ничего.
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error

    back_url = f'/calendar/{event["id"]}/participants/import'
    upload_file = request.files.get('file')
    if upload_file is None or not upload_file.body:
        return build_redirect_with_message(error='Выберите файл.', url=back_url)
    if not (upload_file.name or '').lower().endswith('.xlsx'):
        return build_redirect_with_message(error='Файл должен быть в формате .xlsx.', url=back_url)

    try:
        df = await asyncio.to_thread(pd.read_excel, io=upload_file.body)
    except Exception:
        return build_redirect_with_message(
            error='Не удалось прочитать файл. Проверьте, что это не повреждённый .xlsx.',
            url=back_url,
        )

    upload_error = event_import_upload_error(df)
    if upload_error is not None:
        return build_redirect_with_message(error=upload_error, url=back_url)

    custom_fields = event_import_custom_fields(get_storage(request.app).get_custom_fields())
    rows, unknown_columns, custom_columns = await asyncio.to_thread(parse_event_participant_rows, df, custom_fields)

    token = secrets.token_urlsafe(16)
    replace_event_import_session(
        get_current_user_id(request),
        token,
        {
            'token': token,
            'user_id': get_current_user_id(request),
            'created_at': time.time(),
            'event_id': event['id'],
            'filename': upload_file.name or '',
            'unknown_columns': unknown_columns,
            'custom_columns': custom_columns,
            'rows': rows,
        },
    )
    return redirect(event_import_preview_url(event['id'], token))


def event_import_row_search(
    request: Request,
    storage: SQLiteAdapter,
    session: dict,
) -> tuple[int | None, str, list[dict], bool]:
    """Поиск карточки для строки предпросмотра (GET-параметры find/q).

    Возвращает (номер строки, запрос, результаты, усечён ли список). Только
    чтение: в сессию ничего не пишет и не влияет на had_student_id/bulk.
    Некорректные find/q (нет параметра, не число, строка не существует/не
    pending, пустой или слишком длинный запрос) молча игнорируются — обычный
    предпросмотр без блока поиска.
    """
    raw_find = request.args.get('find', '')
    if not raw_find:
        return None, '', [], False
    row_number, error = parse_reconcile_int(raw_find, 'row number')
    if error is not None:
        return None, '', [], False
    row = next((row for row in session['rows'] if row['row_number'] == row_number), None)
    if row is None or row['status'] != 'pending':
        return None, '', [], False

    query = request.args.get('q', '').strip()
    if not query or len(query) > EVENT_IMPORT_STUDENT_SEARCH_QUERY_MAX_LENGTH:
        return row_number, '', [], False
    results = storage.search_student_candidates(query, limit=EVENT_IMPORT_STUDENT_SEARCH_LIMIT)
    truncated = len(results) >= EVENT_IMPORT_STUDENT_SEARCH_LIMIT
    return row_number, query, results, truncated


@app.get('/calendar/<event_id>/participants/import/preview/<token>')
async def event_participants_import_preview(request: Request, event_id: str, token: str):
    """Предпросмотр: НОЛЬ записей в БД — только чтение кандидатов/справочников.

    GET-параметры find/q — поиск карточки для конкретной строки («Найти
    студента»): только чтение, в сессию ничего не пишет и не влияет на
    bulk-логику; некорректные find/q молча игнорируются."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error
    session, session_error = get_event_import_session(request, token, event['id'])
    if session_error is not None:
        return session_error

    storage = get_storage(request.app)
    search_row, search_query, search_results, search_truncated = event_import_row_search(request, storage, session)
    context = event_import_preview_context(
        storage, event, session, event_import_custom_fields(storage.get_custom_fields())
    )
    return await render(
        template_name=jinja_env.get_template('calendar_event_import_preview.html'),
        context={
            'request': request,
            'is_admin': user_is_admin(request),
            **context,
            'search_row': search_row,
            'search_query': search_query,
            'search_results': search_results,
            'search_truncated': search_truncated,
            **get_flash_args(request),
        },
    )


def resolve_event_import_row(request: Request, event_id: str, token: str, row_number: str):
    """Общие проверки действия над строкой: событие, сессия, строка."""
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return None, None, None, error
    session, session_error = get_event_import_session(request, token, event['id'])
    if session_error is not None:
        return None, None, None, session_error
    row, row_error = event_import_row(session, row_number)
    if row_error is not None:
        return None, None, None, row_error
    return event, session, row, None


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/select-student')
async def event_import_row_select_student(request: Request, event_id: str, token: str, row_number: str):
    """Явный выбор карточки для строки: пометка ТОЛЬКО в сессии предпросмотра.

    Автозаполнение и категория пересчитаются при следующем рендере (строка
    переходит в «Готовы к добавлению»); ничего не создаётся.
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    student_id, id_error = parse_reconcile_int(get_form_value(request, 'student_id'), 'student id')
    if id_error is not None:
        return id_error
    storage = get_storage(request.app)
    student = storage.get_student_by_id(student_id)
    if student is None:
        return build_redirect_with_message(
            error='Студент не найден.',
            url=event_import_preview_url(event['id'], token),
        )
    if not student['active']:
        return build_redirect_with_message(
            error=f'Карточка #{student_id} неактивна — выберите активную.',
            url=event_import_preview_url(event['id'], token),
        )

    row['student_id'] = student_id
    row['student_forced_none'] = False
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: выбрана карточка #{student_id} — {student["full_name"]}.',
        url=event_import_preview_url(event['id'], token),
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/clear-student')
async def event_import_row_clear_student(request: Request, event_id: str, token: str, row_number: str):
    """Сбросить карточку строки: участие будет добавлено без связи.

    Явное решение «без карточки» (подавляет автозаполнение единственным
    точным кандидатом): пометка ТОЛЬКО в сессии предпросмотра, категория
    пересчитается при следующем рендере; запись появится в сопоставлении
    (/admin/people/reconcile) и её можно связать с карточкой позже.
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    row['student_id'] = None
    row['student_forced_none'] = True
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: карточка сброшена — участие будет добавлено без связи.',
        url=event_import_preview_url(event['id'], token),
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/skip')
async def event_import_row_skip(request: Request, event_id: str, token: str, row_number: str):
    """Пропустить строку (в т.ч. ошибочную — как быстрый способ убрать её
    из виду)."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] not in ('pending', 'error'):
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    row['status'] = 'skipped'
    row['added_student_id'] = None
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]} пропущена.',
        url=event_import_preview_url(event['id'], token),
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/edit')
async def event_import_row_edit(request: Request, event_id: str, token: str, row_number: str):
    """Правка строки в сессии: ФИО/пол/институт/группа/курс/место/custom.

    Некорректные данные — строка становится ошибочной; после правки ФИО
    кандидаты и автозаполнение пересчитаются при следующем рендере (строка
    мигрирует между группами естественно). Изменённое ФИО сбрасывает ОБА
    решения по карточке — явный выбор и «без карточки» (оба относились к
    прежнему ФИО и могли устареть): для нового ФИО кандидаты подбираются
    заново, решение принимают отдельным действием.
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] not in ('pending', 'error'):
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    storage = get_storage(request.app)
    custom_fields = event_import_custom_fields(storage.get_custom_fields())
    # Правка ФИО сбрасывает явный выбор карточки и решение «без карточки»:
    # оба решения относятся к прежнему ФИО и могли стать ошибочными
    # (устаревшее молчаливо переносить нельзя). Кандидаты и автозаполнение
    # пересчитаются при следующем рендере.
    new_name = get_form_value(request, 'full_name').strip()
    if new_name.casefold() != (row['full_name'] or '').casefold():
        row['student_id'] = None
        row['student_forced_none'] = False
    row['full_name'] = new_name
    row['sex'] = get_form_value(request, 'sex').strip()
    row['institute'] = get_form_value(request, 'institute').strip()
    row['group'] = get_form_value(request, 'group').strip()
    row['course'] = get_form_value(request, 'course').strip()
    row['position'] = get_form_value(request, 'position').strip()
    # Все импортируемые custom-поля (в т.ч. отсутствовавшие в файле —
    # обязательное поле можно заполнить при правке); url-поля сюда не входят.
    for field in custom_fields:
        row['custom'][field.key] = get_form_value(request, f'custom__{field.key}').strip()
    error = None
    if not row['full_name']:
        error = 'Пустое ФИО.'
    elif row['sex'] not in STUDENT_SEX_OPTIONS:
        error = 'Пол должен быть «М» или «Ж».'
    else:
        # Быстрая проверка правки тем же путём, что предпросмотр: с карточкой
        # строки (явный выбор/единственный кандидат) — иначе пустой курс,
        # который подставился бы из карточки, ложно проваливал бы правку.
        candidates = storage.find_student_candidates(row['full_name'])
        student = event_import_effective_student(storage, row, candidates)
        record, _, _ = event_import_row_record(
            storage,
            row,
            event_participant_preset(event),
            student,
            custom_fields,
            student_file_group_institutes(session['rows']),
        )
        _, error = build_event_participant_competition(record, custom_fields, storage)
    if error is not None:
        row['status'] = 'error'
        row['error'] = error
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} обновлена, но данные некорректны: {error}',
            url=event_import_preview_url(event['id'], token),
        )
    row['status'] = 'pending'
    row['error'] = None
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]} обновлена.',
        url=event_import_preview_url(event['id'], token),
    )


def event_import_add_row_participation(
    request: Request, storage: SQLiteAdapter, event: dict, session: dict, row: dict
) -> str | None:
    """Добавить участие по строке с её текущей карточкой (явный выбор или
    единственный кандидат) — или без карточки, если её нет. Возвращает текст
    ошибки или None.

    Кандидаты и валидация пересчитываются на момент клика — предпросмотр
    мог устареть. Вставка — обычным путём записи реестра (пресет события,
    approved, student_ref_id подобранной карточки или NULL — сопоставить
    можно позже)."""
    candidates = storage.find_student_candidates(row['full_name'])
    student = event_import_effective_student(storage, row, candidates)
    custom_fields = event_import_custom_fields(storage.get_custom_fields())
    record, _, _ = event_import_row_record(
        storage,
        row,
        event_participant_preset(event),
        student,
        custom_fields,
        student_file_group_institutes(session['rows']),
    )
    competition, validation_error = build_event_participant_competition(record, custom_fields, storage)
    if validation_error is not None:
        return validation_error
    competition.student_ref_id = student['id'] if student else None
    # P2: участие явно связывается с событием (участники читаются по ссылке,
    # не по пресету).
    competition.calendar_event_id = event['id']
    ensure_catalog_values(storage, [competition])
    storage.save_competitions(
        [competition],
        review_status='approved',
        owner_id=get_current_user_id(request),
    )
    row['status'] = 'added'
    row['student_id'] = student['id'] if student else None
    row['added_student_id'] = student['id'] if student else None
    log_audit_event(
        request,
        'event_participant_imported',
        {'event_id': event['id'], 'student_ref_id': row['added_student_id']},
    )
    return None


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/add')
async def event_import_row_add(request: Request, event_id: str, token: str, row_number: str):
    """Добавить одну строку («готовую»): участие создаётся сразу, решение
    фиксируется в сессии («Решено: добавлен · карточка #N»)."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] == 'error':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} содержит ошибку — исправьте её.',
            url=event_import_preview_url(event['id'], token),
        )
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    storage = get_storage(request.app)
    add_error = await asyncio.to_thread(event_import_add_row_participation, request, storage, event, session, row)
    if add_error is not None:
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]}: {add_error}',
            url=event_import_preview_url(event['id'], token),
        )
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: участник добавлен.',
        url=event_import_preview_url(event['id'], token),
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/row/<row_number>/create-student')
async def event_import_row_create_student(request: Request, event_id: str, token: str, row_number: str):
    """Создать карточку по строке и СРАЗУ добавить связанного участника.

    Одно явное решение админа: форма предзаполнена Excel-значениями и
    проверена перед отправкой. Сначала валидируется БУДУЩЕЕ участие (с
    данными будущей карточки), затем создаётся карточка (audit
    student_created, source='event-participant-import') и участие с
    student_ref_id новой карточки. Доступно только admin — editor создаёт
    карточки через раздел Students (admin-only, конвенция Phase 1/2.5).
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    event, session, row, error = resolve_event_import_row(request, event_id, token, row_number)
    if error is not None:
        return error
    if row['status'] == 'error':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} содержит ошибку — исправьте её.',
            url=event_import_preview_url(event['id'], token),
        )
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=event_import_preview_url(event['id'], token),
        )

    storage = get_storage(request.app)
    full_name = get_form_value(request, 'full_name').strip()
    sex = get_form_value(request, 'sex').strip()
    institute = get_form_value(request, 'institute').strip()
    group_name = get_form_value(request, 'group').strip()
    course = get_form_value(request, 'course').strip()
    form_error = None
    if not full_name:
        form_error = 'Пустое ФИО.'
    elif sex not in STUDENT_SEX_OPTIONS:
        form_error = 'Пол должен быть «М» или «Ж».'
    if form_error is not None:
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]}: {form_error}',
            url=event_import_preview_url(event['id'], token),
        )

    file_groups = student_file_group_institutes(session['rows'])
    hints = student_catalog_hints(storage, institute, group_name, file_groups)

    # Валидация будущего участия ДО создания карточки: провал не должен
    # оставлять карточку без участия.
    future_student = {
        'id': 0,
        'full_name': full_name,
        'sex': sex,
        'institute': hints['institute'],
        'group_name': hints['group'],
        'course': course,
    }
    custom_fields = event_import_custom_fields(storage.get_custom_fields())
    record, _, _ = event_import_row_record(
        storage,
        row,
        event_participant_preset(event),
        future_student,
        custom_fields,
        file_groups,
    )
    competition, validation_error = build_event_participant_competition(record, custom_fields, storage)
    if validation_error is not None:
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]}: {validation_error}',
            url=event_import_preview_url(event['id'], token),
        )

    student_id = storage.create_student(full_name, sex, hints['institute'], hints['group'], course)
    log_audit_event(
        request,
        'student_created',
        {'student_id': student_id, 'full_name': full_name, 'source': 'event-participant-import'},
    )
    competition.student_ref_id = student_id
    # P2: участие явно связывается с событием (участники читаются по ссылке,
    # не по пресету).
    competition.calendar_event_id = event['id']
    ensure_catalog_values(storage, [competition])
    storage.save_competitions(
        [competition],
        review_status='approved',
        owner_id=get_current_user_id(request),
    )
    row['status'] = 'added'
    row['student_id'] = student_id
    row['student_forced_none'] = False
    row['added_student_id'] = student_id
    log_audit_event(
        request,
        'event_participant_imported',
        {'event_id': event['id'], 'student_ref_id': student_id},
    )
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: студент «{full_name}» создан, участник добавлен.',
        url=event_import_preview_url(event['id'], token),
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/bulk-commit')
async def event_import_bulk_commit(request: Request, event_id: str, token: str):
    """Добавить «готовые» строки одной транзакцией.

    В батч входят ТОЛЬКО строки «Готовы к добавлению» на момент последнего
    рендера (фиксированный признак had_validation_error); неразобранные
    строки в батч не входят. Предпроверка: эффективная карточка каждой
    строки батча пересчитывается на момент клика и сравнивается с
    зафиксированной рендером (had_student_id — id или None) — если состав
    «готовых» изменился, не вставляется ничего (строки перейдут в правильные
    группы при следующем рендере). Строка без карточки добавляется с
    student_ref_id IS NULL (сопоставить можно позже). Сбой вставки — откат
    всего батча вместе со справочниками (import_competitions).
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error
    session, session_error = get_event_import_session(request, token, event['id'])
    if session_error is not None:
        return session_error

    storage = get_storage(request.app)
    custom_fields = event_import_custom_fields(storage.get_custom_fields())
    bulk_rows = event_import_bulk_rows(session)
    file_groups = student_file_group_institutes(session['rows'])
    preset = event_participant_preset(event)
    prepared: list[tuple[dict, Competition, dict | None]] = []
    for row in bulk_rows:
        candidates = storage.find_student_candidates(row['full_name'])
        student = event_import_effective_student(storage, row, candidates)
        actual_ref = student['id'] if student else None
        if actual_ref != row.get('had_student_id'):
            return build_redirect_with_message(
                error='Список карточек изменился — обновите страницу и проверьте строки.',
                url=event_import_preview_url(event['id'], token),
            )
        record, _, _ = event_import_row_record(storage, row, preset, student, custom_fields, file_groups)
        competition, validation_error = build_event_participant_competition(record, custom_fields, storage)
        if validation_error is not None:
            return build_redirect_with_message(
                error=f'Строка {row["row_number"]}: {validation_error}',
                url=event_import_preview_url(event['id'], token),
            )
        competition.student_ref_id = actual_ref
        # P2: участие явно связывается с событием (участники читаются по
        # ссылке, не по пресету).
        competition.calendar_event_id = event['id']
        prepared.append((row, competition, student))

    if prepared:
        await asyncio.to_thread(
            storage.import_competitions,
            [competition for _, competition, _ in prepared],
            owner_id=get_current_user_id(request),
        )
    for row, _, student in prepared:
        row['status'] = 'added'
        row['student_id'] = student['id'] if student else None
        row['added_student_id'] = student['id'] if student else None
        log_audit_event(
            request,
            'event_participant_imported',
            {'event_id': event['id'], 'student_ref_id': row['added_student_id']},
        )
    return build_redirect_with_message(
        message=f'Добавлено участников из Excel: {len(prepared)}.',
        url=f'/calendar/{event["id"]}',
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/finish')
async def event_import_finish(request: Request, event_id: str, token: str):
    """Завершить импорт: итоговое audit-событие с счётчиками (без содержимого
    файла), сессия удаляется.

    Неразобранные pending-строки блокируют завершение; ошибочные — нет
    (учтены как errors, их можно пропустить по одной).
    """
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error
    session, session_error = get_event_import_session(request, token, event['id'])
    if session_error is not None:
        return session_error

    pending = sum(1 for row in session['rows'] if row['status'] == 'pending')
    if pending:
        return build_redirect_with_message(
            error=f'Завершение недоступно: не разобрано строк: {pending}.',
            url=event_import_preview_url(event['id'], token),
        )

    counters = event_import_counters(session)
    log_audit_event(
        request,
        'event_participants_import_completed',
        {'event_id': event['id'], 'event_name': event['name'], 'counters': counters},
    )
    event_import_sessions.pop(token, None)
    return build_redirect_with_message(
        message=(f'Импорт завершён: добавлено {counters["added"]}, ' f'пропущено {counters["skipped"]}.'),
        url=f'/calendar/{event["id"]}',
    )


@app.post('/calendar/<event_id>/participants/import/preview/<token>/discard')
async def event_import_discard(request: Request, event_id: str, token: str):
    """Отменить импорт: сессия удаляется без аудита. Уже добавленные
    подтверждением участники остаются (добавление было явным)."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    event, error = get_calendar_event_or_error(request, event_id)
    if error is not None:
        return error
    session, session_error = get_event_import_session(request, token, event['id'])
    if session_error is not None:
        return session_error

    added = sum(1 for row in session['rows'] if row['status'] == 'added')
    event_import_sessions.pop(token, None)
    if added:
        return build_redirect_with_message(
            message=f'Импорт отменён. Добавленные участники ({added}) сохранены.',
            url=f'/calendar/{event["id"]}',
        )
    return build_redirect_with_message(message='Импорт отменён.', url=f'/calendar/{event["id"]}')


@app.get('/admin')
async def admin_page(request: Request):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    return await render(
        template_name=jinja_env.get_template('admin.html'),
        context={
            'request': request,
            'can_import': user_is_moderator(request),
            'is_admin': user_is_admin(request),
            **get_flash_args(request),
        },
    )


@app.get('/admin/data-map')
async def admin_data_map_page(request: Request):
    """Карта данных: таблицы базы, связи и потоки (постер для владельца)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    return await render(
        template_name=jinja_env.get_template('admin_data_map.html'),
        context={'request': request},
    )


@app.get('/admin/data-guide')
async def admin_data_guide_page(request: Request):
    """Инструкция простыми словами: что делает каждый раздел,
    откуда данные и куда попадают (для администратора-тренера)."""
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    return await render(
        template_name=jinja_env.get_template('admin_data_guide.html'),
        context={'request': request},
    )


@app.get('/admin/import')
async def admin_import_page(request: Request):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    return await render(
        template_name=jinja_env.get_template('admin_import.html'),
        context={
            'request': request,
            'pending_queue_count': storage.count_import_queue('pending'),
            'is_admin': user_is_admin(request),
            **get_flash_args(request),
        },
    )


# Конфликт-режим импорта (№6/№8, docs/data-model-decisions.md «Конфликт-режим
# импорта: очередь на подтверждение»): разбора строка-кандидат показывается
# рядом с похожей существующей записью; решение — принять/пропустить.
QUEUE_FIELD_ORDER: Sequence[str] = (
    'student_name',
    'student_sex',
    'institute',
    'group',
    'course',
    'sport',
    'date',
    'date_to',
    'level',
    'name',
    'position',
)


def queue_candidate_view(payload: dict) -> list[tuple[str, str]]:
    """Поля кандидата из очереди как пары (название, значение) для страницы."""
    view = []
    for key in QUEUE_FIELD_ORDER:
        value = payload.get(key)
        if key in ('date', 'date_to') and value:
            try:
                value = format_date_range(
                    datetime.fromisoformat(payload['date']),
                    datetime.fromisoformat(payload['date_to']) if payload.get('date_to') else None,
                )
            except (TypeError, ValueError):
                pass
        view.append((BASE_FIELD_LABELS.get(key, key), '' if value is None else str(value)))
    return view


def queue_edit_values(payload: dict) -> dict:
    """Предзаполнение инлайн-формы ручного решения: значения кандидата,
    даты — в формате дд.мм.гггг (парсер ручного ввода понимает и диапазон
    «25-27.06.2026» в одном поле)."""
    date_display = ''
    if payload.get('date'):
        try:
            date_display = format_date_range(
                datetime.fromisoformat(payload['date']),
                datetime.fromisoformat(payload['date_to']) if payload.get('date_to') else None,
            )
        except (TypeError, ValueError):
            date_display = ''
    return {
        'student_name': '' if payload.get('student_name') is None else str(payload.get('student_name')),
        'student_sex': '' if payload.get('student_sex') is None else str(payload.get('student_sex')),
        'institute': '' if payload.get('institute') is None else str(payload.get('institute')),
        'group': '' if payload.get('group') is None else str(payload.get('group')),
        'course': '' if payload.get('course') is None else str(payload.get('course')),
        'sport': '' if payload.get('sport') is None else str(payload.get('sport')),
        'date': date_display,
        'level': '' if payload.get('level') is None else str(payload.get('level')),
        'name': '' if payload.get('name') is None else str(payload.get('name')),
        'position': '' if payload.get('position') is None else str(payload.get('position')),
    }


def queue_field_changes(old: Competition, new: Competition) -> dict:
    """Изменения полей old→new для аудита замены существующей записи."""
    changes = {}
    for key in QUEUE_FIELD_ORDER:
        old_value = getattr(old, key, None)
        new_value = getattr(new, key, None)
        if key in ('date', 'date_to'):
            old_value = old_value.isoformat() if old_value else None
            new_value = new_value.isoformat() if new_value else None
        if old_value != new_value:
            changes[key] = {
                'old': '' if old_value is None else str(old_value),
                'new': '' if new_value is None else str(new_value),
            }
    return changes


def resolve_queue_entry(request: Request, entry_id: str):
    """Общая проверка действия над записью очереди: id и статус pending."""
    try:
        numeric_id = int(entry_id)
    except ValueError:
        return None, text(body='Invalid queue entry id', status=400)
    entry = get_storage(request.app).get_import_queue_entry(numeric_id)
    if entry is None:
        return None, text(body='Queue entry not found', status=404)
    if entry['status'] != 'pending':
        return None, text(body='Queue entry already resolved', status=409)
    return entry, None


@app.get('/admin/import-queue')
async def admin_import_queue_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    entries = []
    for entry in storage.list_import_queue('pending'):
        matched = None
        if entry['matched_record_id']:
            matched = storage.get_competition_by_id(int(entry['matched_record_id']))
        entries.append(
            {
                **entry,
                'candidate': queue_candidate_view(entry['payload']),
                'edit_values': queue_edit_values(entry['payload']),
                'matched': matched,
            }
        )
    return await render(
        template_name=jinja_env.get_template('admin_import_queue.html'),
        context={
            'request': request,
            'queue_entries': entries,
            **get_flash_args(request),
        },
    )


@app.post('/admin/import-queue/<entry_id>/accept')
async def accept_import_queue_entry(request: Request, entry_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    entry, error = resolve_queue_entry(request, entry_id)
    if error is not None:
        return error
    storage = get_storage(request.app)
    competition = Competition.model_validate(entry['payload'])

    # Строка вставляется только по явному решению админа; если пока она
    # лежала в очереди, такая запись появилась другим путём — не дублируем.
    duplicate_now = any(
        competition_duplicate_key(comp) == competition_duplicate_key(competition) for comp in storage.get_competitions()
    )
    if duplicate_now:
        storage.set_import_queue_status(entry['id'], 'skipped')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {'decision': 'skipped', 'queue_id': entry['id'], 'reason': 'duplicate_already_exists'},
        )
        return redirect(to='/admin/import-queue?admin_message=Запись+уже+существует,+кандидат+пропущен')

    ensure_catalog_values(storage, [competition])
    storage.save_competitions([competition], review_status='approved', owner_id=entry['created_by'])
    storage.set_import_queue_status(entry['id'], 'accepted')
    log_audit_event(
        request,
        'import_conflict_resolved',
        {
            'decision': 'accepted',
            'queue_id': entry['id'],
            'matched_record_id': entry['matched_record_id'],
        },
    )
    return redirect(to='/admin/import-queue?admin_message=Кандидат+принят+и+добавлен+в+реестр')


@app.post('/admin/import-queue/<entry_id>/skip')
async def skip_import_queue_entry(request: Request, entry_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    entry, error = resolve_queue_entry(request, entry_id)
    if error is not None:
        return error
    storage = get_storage(request.app)
    storage.set_import_queue_status(entry['id'], 'skipped')
    log_audit_event(
        request,
        'import_conflict_resolved',
        {'decision': 'skipped', 'queue_id': entry['id'], 'matched_record_id': entry['matched_record_id']},
    )
    return redirect(to='/admin/import-queue?admin_message=Кандидат+пропущен')


# Замена существующей записи данными кандидата (№25): историчность обеспечивает
# аудит old→new по полям; record_id существующей сохраняется, её владелец и
# review-статус не меняются (update_competition не трогает эти колонки).
@app.post('/admin/import-queue/<entry_id>/replace')
async def replace_import_queue_entry(request: Request, entry_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    entry, error = resolve_queue_entry(request, entry_id)
    if error is not None:
        return error
    if not entry['matched_record_id']:
        return text(body='Queue entry has no matched record', status=400)
    storage = get_storage(request.app)
    existing = storage.get_competition_by_id(int(entry['matched_record_id']))
    if existing is None:
        return text(body='Matched record not found', status=404)

    competition = Competition.model_validate(entry['payload'])
    ensure_catalog_values(storage, [competition])
    changes = queue_field_changes(existing, competition)
    storage.update_competition(int(entry['matched_record_id']), competition)
    storage.set_import_queue_status(entry['id'], 'replaced')
    log_audit_event(
        request,
        'import_conflict_resolved',
        {
            'decision': 'replaced',
            'queue_id': entry['id'],
            'matched_record_id': entry['matched_record_id'],
            'changes': changes,
        },
    )
    return redirect(to='/admin/import-queue?admin_message=Существующая+запись+заменена+данными+кандидата')


# Ручное решение (№25, №27в): админ правит поля кандидата инлайн, кроме ФИО
# и даты — они задают конфликт и всегда берутся из payload очереди; после
# сохранения кандидат вставляется как обычный accepted.
@app.post('/admin/import-queue/<entry_id>/edit')
async def edit_import_queue_entry(request: Request, entry_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    entry, error = resolve_queue_entry(request, entry_id)
    if error is not None:
        return error
    storage = get_storage(request.app)
    custom_fields = storage.get_custom_fields()
    # №27в: ФИО и дата задают конфликт с существующей записью — из формы не
    # берутся (подмена игнорируется), только из payload очереди. Остальные
    # поля приходят из формы. Дату форматируем в дд.мм.гггг (с диапазоном),
    # так её понимает ручной парсер build_competition.
    payload = entry['payload']
    record = {
        'ФИО': '' if payload.get('student_name') is None else str(payload['student_name']),
        'Пол': get_form_value(request, 'student_sex'),
        'Институт': get_form_value(request, 'institute'),
        'Группа': get_form_value(request, 'group'),
        'Вид спорта': get_form_value(request, 'sport'),
        'Дата': queue_edit_values(payload)['date'],
        'Уровень соревнований': get_form_value(request, 'level'),
        'Название соревнований': get_form_value(request, 'name'),
        'Место': get_form_value(request, 'position'),
        'Курс': get_form_value(request, 'course'),
    }
    try:
        competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
    except (TypeError, ValueError) as exc:
        return text(body=f'Invalid row data: {exc}', status=400)

    # Тот же анти-дубликат, что и у «Принять»: строка вставляется только
    # по явному решению админа.
    duplicate_now = any(
        competition_duplicate_key(comp) == competition_duplicate_key(competition) for comp in storage.get_competitions()
    )
    if duplicate_now:
        storage.set_import_queue_status(entry['id'], 'skipped')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {'decision': 'skipped', 'queue_id': entry['id'], 'reason': 'duplicate_already_exists', 'edited': True},
        )
        return redirect(to='/admin/import-queue?admin_message=Запись+уже+существует,+кандидат+пропущен')

    ensure_catalog_values(storage, [competition])
    storage.save_competitions([competition], review_status='approved', owner_id=entry['created_by'])
    storage.set_import_queue_status(entry['id'], 'accepted')
    log_audit_event(
        request,
        'import_conflict_resolved',
        {
            'decision': 'accepted',
            'edited': True,
            'queue_id': entry['id'],
            'matched_record_id': entry['matched_record_id'],
        },
    )
    return redirect(to='/admin/import-queue?admin_message=Кандидат+принят+с+ручными+правками')


@app.get('/admin/fields')
async def admin_fields_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    return await render(
        template_name=jinja_env.get_template('admin_fields.html'),
        context={
            'request': request,
            'admin_custom_fields': storage.get_custom_fields(include_inactive=True),
            'field_type_options': FIELD_TYPE_OPTIONS,
            # №24: варианты привязки link-полей к колонкам таблицы.
            'link_target_base_options': [(key, BASE_FIELD_LABELS[key]) for key in LINK_TARGETABLE_BASE_KEYS],
            'link_target_custom_options': [
                (field.key, field.label) for field in storage.get_custom_fields() if field.field_type == 'text'
            ],
            'base_field_settings': get_base_field_settings(storage),
            'base_field_labels': BASE_FIELD_LABELS,
            'always_required_base_fields': ALWAYS_REQUIRED_BASE_FIELDS,
            'base_field_value_types': BASE_FIELD_VALUE_TYPES,
            **get_flash_args(request),
        },
    )


def parse_base_field_settings_form(request: Request) -> dict[str, tuple[str, bool]]:
    """Собрать настройки базовых полей из формы /admin/fields/base.

    Заблокированные поля (ФИО, Дата) игнорируют форму: обязательность
    не отключается, тип фиксированный — защита от подделки разметки.
    """
    settings_form: dict[str, tuple[str, bool]] = {}
    for key in BASE_FIELD_SETTING_DEFAULTS:
        default_type, default_required = BASE_FIELD_SETTING_DEFAULTS[key]
        if key in ALWAYS_REQUIRED_BASE_FIELDS:
            settings_form[key] = (default_type, True)
            continue
        value_type = get_form_value(request, f'type_{key}')
        if value_type not in BASE_FIELD_VALUE_TYPES:
            value_type = default_type
        settings_form[key] = (value_type, parse_checkbox(request, f'required_{key}'))
    return settings_form


@app.post('/admin/fields/base')
async def update_base_field_settings(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    new_settings = parse_base_field_settings_form(request)
    old_settings = get_base_field_settings(storage)
    changed = {
        key: {
            'old': old_settings[key],
            'new': {'value_type': value_type, 'required': required},
        }
        for key, (value_type, required) in new_settings.items()
        if old_settings[key]['value_type'] != value_type or old_settings[key]['required'] != required
    }
    if changed:
        storage.update_field_settings(new_settings)
        log_audit_event(request, 'field_settings_changed', {'fields': changed})
    return redirect(to='/admin/fields?admin_message=Настройки+базовых+полей+сохранены')


@app.get('/admin/levels')
async def admin_levels_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    # Уровни переехали на страницу «Справочники»; старый URL остаётся
    # рабочим, flash-параметры сохраняются в редиректе.
    query = request.query_string
    return redirect(f'/admin/catalogs?{query}' if query else '/admin/catalogs')


def build_catalog_entries(storage: SQLiteAdapter, category: str) -> list[dict]:
    """Значения справочника для страницы админки: каждое со счётчиком записей."""
    return [
        {**row, 'records_count': storage.count_records_using(category, row['value'])}
        for row in storage.list_catalog_all(category)
    ]


@app.get('/admin/catalogs')
async def admin_catalogs_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    return await render(
        template_name=jinja_env.get_template('admin_catalogs.html'),
        context={
            'request': request,
            'sport_values': build_catalog_entries(storage, 'sport'),
            'institute_tree': storage.list_catalog_tree(),
            'institute_options': build_institute_options(storage),
            'admin_levels': [
                {**level, 'records_count': storage.count_records_using('level', level['name'])}
                for level in storage.list_levels()
            ],
            **get_flash_args(request),
        },
    )


def resolve_catalog_value(request: Request, category: str, value_id: str):
    """Общая проверка для действий над значением справочника: категория и id."""
    if category not in CATALOG_MANAGE_CATEGORIES:
        return None, text(body='Unknown catalog', status=404)
    try:
        numeric_value_id = int(value_id)
    except ValueError:
        return None, text(body='Invalid value id', status=400)
    row = get_storage(request.app).get_catalog_value(numeric_value_id)
    if row is None or row['category'] != category:
        return None, build_redirect_with_message(error='Значение не найдено', url='/admin/catalogs')
    return row, None


@app.post('/admin/catalogs/<category>')
async def create_catalog_value(request: Request, category: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    if category not in CATALOG_CATEGORIES:
        return text(body='Unknown catalog', status=404)

    value = get_form_value(request, 'value').strip()
    if not value:
        return build_redirect_with_message(error='Значение обязательно', url='/admin/catalogs')

    # №1.1: регистровый дубль — отказ с подсказкой канонического написания.
    existing = get_storage(request.app).find_catalog_canonical(category, value)
    if existing is not None and existing != value:
        return build_redirect_with_message(
            error=f'Значение «{value}» уже есть: «{existing}» (отличается только регистром)',
            url='/admin/catalogs',
        )

    get_storage(request.app).add_catalog_value(category, value)
    return build_redirect_with_message(message='Значение добавлено', url='/admin/catalogs')


@app.post('/admin/catalogs/<category>/<value_id>/hide')
async def hide_catalog_value(request: Request, category: str, value_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    row, error = resolve_catalog_value(request, category, value_id)
    if error is not None:
        return error

    get_storage(request.app).hide_catalog_value(row['id'])
    return build_redirect_with_message(message='Значение скрыто из подсказок', url='/admin/catalogs')


@app.post('/admin/catalogs/<category>/<value_id>/unhide')
async def unhide_catalog_value(request: Request, category: str, value_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    row, error = resolve_catalog_value(request, category, value_id)
    if error is not None:
        return error

    get_storage(request.app).unhide_catalog_value(row['id'])
    return build_redirect_with_message(message='Значение снова видно в подсказках', url='/admin/catalogs')


@app.post('/admin/catalogs/<category>/<value_id>/delete')
async def delete_catalog_value(request: Request, category: str, value_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    row, error = resolve_catalog_value(request, category, value_id)
    if error is not None:
        return error

    storage = get_storage(request.app)
    parent_value = None
    if category == GROUP_CATEGORY and row.get('parent_id'):
        parent = storage.get_catalog_value(row['parent_id'])
        parent_value = parent['value'] if parent else None
    records_count = storage.count_records_using(category, row['value'], parent_value=parent_value)
    if records_count > 0:
        # Значение висит на записях — историчность неприкосновенна, удалять нельзя.
        return text(
            body=f'Нельзя удалить значение «{row["value"]}»: на нём записей — {records_count}',
            status=400,
        )

    if category == 'institute':
        child_groups = storage.count_child_groups(row['id'])
        if child_groups:
            # Группы без родителя не бывают: не оставляем сирот, удаление
            # института — только после удаления его групп.
            return text(
                body=f'У института «{row["value"]}» есть группы ({child_groups}) — сначала удалите их',
                status=400,
            )

    storage.delete_catalog_value(row['id'])
    return build_redirect_with_message(message='Значение удалено', url='/admin/catalogs')


@app.post('/admin/catalogs/institute/<institute_id>/group')
async def create_catalog_group(request: Request, institute_id: str):
    """Добавить группу в институт (иерархия справочника, №15)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    try:
        numeric_institute_id = int(institute_id)
    except ValueError:
        return text(body='Invalid value id', status=400)

    storage = get_storage(request.app)
    institute = storage.get_catalog_value(numeric_institute_id)
    if institute is None or institute['category'] != 'institute':
        return build_redirect_with_message(error='Институт не найден', url='/admin/catalogs')

    value = get_form_value(request, 'value').strip()
    if not value:
        return build_redirect_with_message(error='Название группы обязательно', url='/admin/catalogs')

    # №1.1: регистровый дубль внутри этого института — отказ с подсказкой.
    existing = storage.find_catalog_canonical('group', value, parent_id=institute['id'])
    if existing is not None and existing != value:
        return build_redirect_with_message(
            error=f'Группа «{value}» уже есть: «{existing}» (отличается только регистром)',
            url='/admin/catalogs',
        )

    storage.add_catalog_value(GROUP_CATEGORY, value, parent_id=institute['id'])
    return build_redirect_with_message(
        message=f'Группа «{value}» добавлена в институт «{institute["value"]}»',
        url='/admin/catalogs',
    )


def resolve_rename_target(request: Request, category: str, value_id: str):
    """Строка справочника для переименования: уровень — по id в levels,
    остальные категории — по id в catalog_values. Для группы вторым
    элементом сразу её институт (переименование затрагивает пару)."""
    if category not in CATALOG_RENAME_LABELS:
        return None, None, text(body='Unknown catalog', status=404)
    try:
        numeric_value_id = int(value_id)
    except ValueError:
        return None, None, text(body='Invalid value id', status=400)

    storage = get_storage(request.app)
    if category == 'level':
        level = next((item for item in storage.list_levels() if item['id'] == numeric_value_id), None)
        if level is None:
            return None, None, build_redirect_with_message(error='Значение не найдено', url='/admin/catalogs')
        # уровни хранят имя в «name»; дальше работаем с унифицированным «value»
        return {'id': level['id'], 'value': level['name'], 'active': level.get('active')}, None, None

    row = storage.get_catalog_value(numeric_value_id)
    if row is None or row['category'] != category:
        return None, None, build_redirect_with_message(error='Значение не найдено', url='/admin/catalogs')
    parent_value = None
    if category == GROUP_CATEGORY:
        parent = storage.get_catalog_value(row['parent_id']) if row.get('parent_id') else None
        if parent is None or parent['category'] != 'institute':
            return None, None, build_redirect_with_message(error='Институт группы не найден', url='/admin/catalogs')
        parent_value = parent['value']
    return row, parent_value, None


def catalog_rename_conflict(storage: SQLiteAdapter, category: str, row: dict, new_name: str) -> bool:
    """Конфликт нового имени с существующим значением категории (для группы —
    в том же институте). Проверка до выполнения, чтобы вернуть ошибку рано;
    авторитетная — внутри storage.rename_catalog_value, в транзакции."""
    if category == 'level':
        return new_name in storage.get_level_names(include_inactive=True)
    if category == GROUP_CATEGORY:
        return any(
            item['value'] == new_name and item['parent_id'] == row.get('parent_id')
            for item in storage.list_catalog_all(GROUP_CATEGORY)
        )
    return any(item['value'] == new_name for item in storage.list_catalog_all(category))


@app.get('/admin/catalogs/<category>/<value_id>/rename')
async def catalog_rename_page(request: Request, category: str, value_id: str):
    """Шаг 2 переименования без JS: подтверждение с числом записей и галочкой
    «обновить записи» (шаг 1 — GET-форма со старым значением в списке)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    row, parent_value, error = resolve_rename_target(request, category, value_id)
    if error is not None:
        return error

    new_name = (get_param(dict(request.args), 'new_name') or '').strip()
    if not new_name:
        return build_redirect_with_message(error='Новое имя обязательно', url='/admin/catalogs')
    if new_name == row['value']:
        return build_redirect_with_message(error='Новое имя совпадает со старым', url='/admin/catalogs')
    storage = get_storage(request.app)
    if catalog_rename_conflict(storage, category, row, new_name):
        return build_redirect_with_message(error=f'«{new_name}» уже есть в справочнике', url='/admin/catalogs')

    return await render(
        template_name=jinja_env.get_template('admin_catalog_rename.html'),
        context={
            'request': request,
            'category': category,
            'category_label': CATALOG_RENAME_LABELS[category],
            'value_id': row['id'],
            'old_name': row['value'],
            'new_name': new_name,
            'parent_value': parent_value,
            'records_count': storage.count_records_using(category, row['value'], parent_value=parent_value),
        },
    )


@app.post('/admin/catalogs/<category>/<value_id>/rename')
async def rename_catalog_value(request: Request, category: str, value_id: str):
    """Переименование значения справочника (решение 2026-09-13, по образцу merge).

    Без confirm — JSON-превью, сколько записей затронет. С confirm —
    выполнение: справочник и (по галочке update_records) записи одной
    транзакцией в storage; каждое выполнение — audit catalog_value_renamed
    (кто, что, во что, сколько записей обновлено).
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    row, parent_value, error = resolve_rename_target(request, category, value_id)
    if error is not None:
        return error

    new_name = get_form_value(request, 'new_name').strip()
    if not new_name:
        return build_redirect_with_message(error='Новое имя обязательно', url='/admin/catalogs')
    if new_name == row['value']:
        return build_redirect_with_message(error='Новое имя совпадает со старым', url='/admin/catalogs')
    storage = get_storage(request.app)
    if catalog_rename_conflict(storage, category, row, new_name):
        return build_redirect_with_message(error=f'«{new_name}» уже есть в справочнике', url='/admin/catalogs')

    if not checkbox_to_bool(get_form_value(request, 'confirm')):
        return json_response(
            {
                'preview': True,
                'category': category,
                'old': row['value'],
                'new': new_name,
                'records': storage.count_records_using(category, row['value'], parent_value=parent_value),
            }
        )

    update_records = checkbox_to_bool(get_form_value(request, 'update_records'))
    try:
        records_updated = storage.rename_catalog_value(
            category,
            row['value'],
            new_name,
            parent_value=parent_value,
            update_records=update_records,
        )
    except ValueError as exc:
        return build_redirect_with_message(error=str(exc), url='/admin/catalogs')

    audit_details = {
        'category': category,
        'old': row['value'],
        'new': new_name,
        'update_records': update_records,
        'records_updated': records_updated,
    }
    if category == GROUP_CATEGORY:
        audit_details['parent'] = parent_value
    log_audit_event(request, 'catalog_value_renamed', audit_details)

    message = f'«{row["value"]}» → «{new_name}»'
    message += f': обновлено записей — {records_updated}' if update_records else ' (записи не тронуты)'
    return build_redirect_with_message(message=message, url='/admin/catalogs')


def build_institute_options(storage: SQLiteAdapter) -> list[dict]:
    """Активные институты плоским списком id+значение (форма переноса группы)."""
    return [{'id': row['id'], 'value': row['value']} for row in storage.list_catalog_all('institute') if row['active']]


def resolve_group_move_target(request: Request, value_id: str, target_institute_id: str):
    """Группа и целевой институт переноса: ((group, target), None) или
    (None, redirect с ошибкой на /admin/catalogs)."""
    try:
        numeric_group_id = int(value_id)
        numeric_target_id = int(target_institute_id)
    except (TypeError, ValueError):
        return None, text(body='Invalid value id', status=400)
    storage = get_storage(request.app)
    group = storage.get_catalog_value(numeric_group_id)
    if group is None or group['category'] != GROUP_CATEGORY or not group.get('parent_id'):
        return None, build_redirect_with_message(error='Группа или институт не найдены', url='/admin/catalogs')
    target = storage.get_catalog_value(numeric_target_id)
    if target is None or target['category'] != 'institute':
        return None, build_redirect_with_message(error='Группа или институт не найдены', url='/admin/catalogs')
    if target['id'] == group['parent_id']:
        return None, build_redirect_with_message(
            error='Группа уже относится к выбранному институту',
            url='/admin/catalogs',
        )
    conflict = storage.find_catalog_row(GROUP_CATEGORY, group['value'], parent_id=target['id'])
    if conflict is not None:
        return None, build_redirect_with_message(
            error=(
                f'В институте «{target["value"]}» уже есть группа «{group["value"]}» '
                '— переименуйте одну из групп перед переносом'
            ),
            url='/admin/catalogs',
        )
    return (group, target), None


@app.get('/admin/catalogs/group/<value_id>/move')
async def catalog_group_move_page(request: Request, value_id: str):
    """Перенос группы в другой институт (решение 2026-09-22), шаг 2 без JS:
    страница-подтверждение со счётчиками затрагиваемых данных и галочкой."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    target_institute_id = (get_param(dict(request.args), 'target_institute_id') or '').strip()
    if not target_institute_id:
        return build_redirect_with_message(error='Выберите институт для переноса', url='/admin/catalogs')
    resolved, error = resolve_group_move_target(request, value_id, target_institute_id)
    if error is not None:
        return error
    group, target = resolved

    storage = get_storage(request.app)
    old_institute = storage.get_catalog_value(group['parent_id'])
    if old_institute is None:
        return build_redirect_with_message(error='Группа или институт не найдены', url='/admin/catalogs')

    return await render(
        template_name=jinja_env.get_template('admin_catalog_group_move.html'),
        context={
            'request': request,
            'value_id': group['id'],
            'group_value': group['value'],
            'old_institute': old_institute['value'],
            'target_institute': target['value'],
            'target_institute_id': target['id'],
            'records_count': storage.count_records_using('group', group['value'], parent_value=old_institute['value']),
            'students_count': storage.count_students_using_group_pair(group['value'], old_institute['value']),
        },
    )


@app.post('/admin/catalogs/group/<value_id>/move')
async def catalog_group_move(request: Request, value_id: str):
    """Выполнение переноса группы: справочник + записи + карточки студентов
    одной транзакцией в storage; подтверждение — обязательной галочкой."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    target_institute_id = get_form_value(request, 'target_institute_id').strip()
    resolved, error = resolve_group_move_target(request, value_id, target_institute_id)
    if error is not None:
        return error
    group, target = resolved

    # Без галочки — назад на страницу подтверждения, ничего не меняя.
    back_url = f'/admin/catalogs/group/{value_id}/move?target_institute_id={target_institute_id}'
    if not checkbox_to_bool(get_form_value(request, 'confirm')):
        return redirect(to=back_url)

    storage = get_storage(request.app)
    old_institute = storage.get_catalog_value(group['parent_id'])
    if old_institute is None:
        return build_redirect_with_message(error='Группа или институт не найдены', url='/admin/catalogs')
    try:
        counts = storage.move_catalog_group(group['id'], target['id'])
    except ValueError as exc:
        return build_redirect_with_message(error=str(exc), url='/admin/catalogs')

    log_audit_event(
        request,
        'catalog_group_moved',
        {
            'group': group['value'],
            'old_institute': old_institute['value'],
            'new_institute': target['value'],
            'students_updated': counts['students_updated'],
            'competitions_updated': counts['competitions_updated'],
            'group_id': group['id'],
        },
    )
    return build_redirect_with_message(
        message=(
            f'Группа «{group["value"]}» перенесена: «{old_institute["value"]}» → «{target["value"]}»; '
            f'обновлено записей — {counts["competitions_updated"]}, '
            f'карточек студентов — {counts["students_updated"]}'
        ),
        url='/admin/catalogs',
    )


def ru_plural(number: int, one: str, few: str, many: str) -> str:
    """Русская плюрализация: 1 минуту / 2 минуты / 5 минут."""
    mod_100 = number % 100
    mod_10 = number % 10
    if mod_100 in range(11, 15):
        return many
    if mod_10 == 1:
        return one
    if mod_10 in range(2, 5):
        return few
    return many


def format_login_moment(raw_value) -> str:
    """Последний вход для подсказки: дд.мм.гггг чч:мм или «—»."""
    if not raw_value:
        return '—'
    try:
        moment = datetime.fromisoformat(str(raw_value)).replace(tzinfo=timezone.utc).astimezone()
    except ValueError:
        return '—'
    return moment.strftime('%d.%m.%Y %H:%M')


def describe_presence(last_seen_raw, *, now: datetime | None = None) -> dict:
    """Колонка «Активность» на /admin/users (№17).

    «Онлайн» = last_seen в пределах PRESENCE_ONLINE_WINDOW (сессии
    stateless-cookie, точного списка живых сессий нет — приближение).
    Остальное — человекочитаемое «был(а): …»; храним в UTC, показываем
    локальное время сервера.
    """
    if not last_seen_raw:
        return {'online': False, 'label': '—'}
    try:
        last_seen = datetime.fromisoformat(str(last_seen_raw)).replace(tzinfo=timezone.utc).astimezone()
    except ValueError:
        return {'online': False, 'label': '—'}

    local_now = now or datetime.now(timezone.utc).astimezone()
    delta = max(timedelta(0), local_now - last_seen)
    if delta <= PRESENCE_ONLINE_WINDOW:
        return {'online': True, 'label': 'онлайн'}

    minutes = int(delta.total_seconds() // 60)
    if minutes < 60:
        return {'online': False, 'label': f'{minutes} {ru_plural(minutes, "минуту", "минуты", "минут")} назад'}

    if last_seen.date() == local_now.date():
        return {'online': False, 'label': f'сегодня {last_seen.strftime("%H:%M")}'}
    if last_seen.date() == (local_now - timedelta(days=1)).date():
        return {'online': False, 'label': f'вчера {last_seen.strftime("%H:%M")}'}
    return {'online': False, 'label': last_seen.strftime('%d.%m.%Y %H:%M')}


@app.get('/admin/users')
async def admin_users_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    users = [
        {
            **user,
            'records_count': storage.count_records_by_owner(user['id']),
            'presence': describe_presence(user.get('last_seen_at')),
            'last_login_label': format_login_moment(user.get('last_login_at')),
        }
        for user in storage.list_users()
    ]
    return await render(
        template_name=jinja_env.get_template('admin_users.html'),
        context={
            'request': request,
            'users': users,
            'user_roles': USER_ROLES,
            'current_username': (get_auth_user(request) or {}).get('username'),
            **get_flash_args(request),
        },
    )


@app.get('/admin/students')
async def admin_students_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    return await render(
        template_name=jinja_env.get_template('admin_students.html'),
        context={
            'request': request,
            **get_flash_args(request),
        },
    )


@app.get('/profile')
async def profile_page(request: Request):
    return await render(
        template_name=jinja_env.get_template('profile.html'),
        context={
            'request': request,
            'is_athlete': user_is_athlete(request),
        },
    )


@app.get('/template/empty.xlsx')
async def export_empty_template(request: Request):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    custom_fields = [field for field in storage.get_custom_fields() if field.show_in_template]
    columns = list(REQUIRED_IMPORT_COLUMNS) + [field.label for field in custom_fields]

    df = pd.DataFrame(columns=columns)
    buffer = BytesIO()
    df.to_excel(buffer, index=False)
    apply_template_date_format(buffer, columns)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="Шаблон_соревнования_{now_str}.xlsx"',
        },
    )


def apply_template_date_format(buffer: BytesIO, columns: Sequence[str]) -> None:
    from openpyxl import load_workbook

    if 'Дата' not in columns:
        buffer.seek(0)
        return
    date_column_letter = get_column_letter(columns.index('Дата') + 1)
    buffer.seek(0)
    workbook = load_workbook(buffer)
    sheet = workbook.worksheets[0]
    for cell in sheet[date_column_letter]:
        cell.number_format = 'DD.MM.YYYY'
    sheet.column_dimensions[date_column_letter].width = 14
    buffer.seek(0)
    workbook.save(buffer)
    buffer.seek(0)


def build_index_dataframe(competitions, export_custom_fields) -> pd.DataFrame:
    df = pd.DataFrame.from_records([competition_to_export_row(comp, export_custom_fields) for comp in competitions])
    df = df.reindex(columns=list(INDEX_EXPORT_COLUMNS) + [field.label for field in export_custom_fields])
    for column in df.columns:
        if df[column].dtype == object:
            df[column] = df[column].map(sanitize_spreadsheet_value)
    return df


@app.get('/export/index')
async def export_index(request: Request):
    if user_is_athlete(request):
        return forbidden(request)
    storage = get_storage(request.app)
    competitions = storage.get_competitions()
    export_custom_fields = [field for field in storage.get_custom_fields() if field.show_in_export]
    df = await asyncio.to_thread(build_index_dataframe, competitions, export_custom_fields)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    filename = f'Отчет_{now_str}.xlsx'
    buffer = BytesIO()
    await asyncio.to_thread(df.to_excel, buffer, index=False)
    buffer.seek(0)

    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="{filename}"',
        },
    )


def import_record_value_is_empty(value) -> bool:
    """Пустая ячейка импорта: None/NaN или пробельная строка."""
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return str(value).strip() == ''


def autofill_record_from_known_athlete(
    record: dict,
    storage: SQLiteAdapter,
    cache: dict[str, dict],
) -> None:
    """№23доп (docs/feedback-live.md): автозаполнение пустых полей строки импорта.

    Для строки с пустыми Пол/Институт/Группа/Курс — данные известного атлета
    (профиль или последняя запись, точное ФИО). Заполненные поля строки НЕ
    перезаписываются: запись — исторический факт. Заполняется ДО
    build_competition и подсчёта дублей: дубль-ключ считается после
    автозаполнения, и пустой Курс проходит валидацию только после неё.
    """
    name = clean_str(record.get('ФИО', ''))
    if not name:
        return
    key = name.lower()
    if key not in cache:
        cache[key] = storage.find_athlete_fields(name)
    known = cache[key]
    if not known:
        return
    for column, result_key in (
        ('Пол', 'sex'),
        ('Институт', 'institute'),
        ('Группа', 'group'),
    ):
        if import_record_value_is_empty(record.get(column)) and known.get(result_key):
            record[column] = known[result_key]
    if import_record_value_is_empty(record.get('Курс')) and known.get('course'):
        record['Курс'] = known['course']


def build_import_competitions(df: pd.DataFrame, custom_fields, storage: SQLiteAdapter) -> list[Competition]:
    competitions = []
    known_athletes: dict[str, dict] = {}
    for index, row in df.iterrows():
        record = row.to_dict()
        try:
            autofill_record_from_known_athlete(record, storage, known_athletes)
            competitions.append(build_competition(record, custom_fields=custom_fields, storage=storage))
        except (TypeError, ValueError) as exc:
            raise ValueError(f'строка {index + 2}: {exc}') from exc
    return competitions


def competition_duplicate_key(competition: Competition) -> tuple:
    return (
        competition.student_name,
        competition.date.date().isoformat(),
        competition.sport,
        competition.name,
    )


def ensure_catalog_values(storage: SQLiteAdapter, competitions: Iterable[Competition]):
    """Новое значение вида спорта/института из записи или импорта попадает в справочник.

    Как с уровнями: справочник только собирает подсказки и ничего не
    перезаписывает в записях (docs/data-model-decisions.md). Дубликаты
    игнорируются на стороне storage. Пара институт→группа попадает в иерархию:
    группа кладётся под свой институт (уникальность (parent, value)).

    Для одиночных записей (ручной ввод). Массовый импорт пишет справочники
    батчем той же семантикой — SQLiteAdapter._ensure_import_catalogs.
    """
    for competition in competitions:
        for category, value in (('sport', competition.sport), ('institute', competition.institute)):
            if value:
                storage.add_catalog_value(category, value)
        if competition.institute and competition.group:
            storage.ensure_catalog_pair(competition.institute, competition.group)


def competition_partial_key(competition: Competition) -> tuple:
    """Похожесть (№6/№8): ФИО + дата_начала. Совпадение пары при несовпадении
    полного ключа дедупликации — кандидат в очередь подтверждения."""
    return (competition.student_name, competition.date.date().isoformat())


def split_import_competitions(
    competitions: Sequence[Competition],
    existing_competitions: Iterable[Competition],
) -> tuple[list[Competition], int, list[tuple[Competition, int | None]]]:
    """Разложить строки импорта на вставку / дубли / конфликты (№6/№8).

    Точный дубль (ФИО+дата+вид спорта+название) — пропуск как раньше.
    Похожая строка (совпадают ФИО+дата_начала, полный ключ другой) —
    НЕ вставляется и попадает в очередь: (конфликт, matched_record_id)
    похожей существующей записи. Конфликты между строками одного файла —
    тоже в очередь: первая строка вставляется, похожая на неё вторая —
    в очередь с matched_record_id=None (похожая запись появится после
    вставки первой).
    """
    existing_by_full_key = {competition_duplicate_key(comp): comp for comp in existing_competitions}
    existing_partial_keys = {competition_partial_key(comp) for comp in existing_competitions}
    new_competitions = []
    conflicts: list[tuple[Competition, int | None]] = []
    skipped_duplicates = 0
    seen_full_keys = set()
    inserted_partial_keys = set()
    for competition in competitions:
        full_key = competition_duplicate_key(competition)
        if full_key in existing_by_full_key or full_key in seen_full_keys:
            skipped_duplicates += 1
            continue
        partial_key = competition_partial_key(competition)
        if partial_key in existing_partial_keys or partial_key in inserted_partial_keys:
            matched = None
            if partial_key in existing_partial_keys:
                matched = next(
                    (
                        int(comp.record_id)
                        for comp in existing_competitions
                        if competition_partial_key(comp) == partial_key and comp.record_id
                    ),
                    None,
                )
            conflicts.append((competition, matched))
            continue
        seen_full_keys.add(full_key)
        inserted_partial_keys.add(partial_key)
        new_competitions.append(competition)
    return new_competitions, skipped_duplicates, conflicts


def build_import_summary(inserted: int, skipped_duplicates: int, conflicts: int, *, is_admin: bool) -> str:
    """Сводка импорта реестра (IA-редизайн 2026-09): admin видит адресата
    разбора, editor — что спорные строки переданы администратору (очередь
    ему недоступна)."""
    summary = f'Импортировано записей: {inserted}'
    if skipped_duplicates:
        summary += f'. Пропущено дублей: {skipped_duplicates}'
    if conflicts:
        if is_admin:
            summary += f'. На подтверждение: {conflicts}.'
        else:
            summary += f'. На подтверждение: {conflicts} — передано администратору.'
    return summary


@app.post('/')
async def upload(request: Request):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    upload_file = request.files.get('file')
    if not upload_file or not upload_file.body:
        return text(body='No file uploaded', status=400)

    storage = get_storage(request.app)
    custom_fields = storage.get_custom_fields()

    try:
        df = await asyncio.to_thread(pd.read_excel, io=upload_file.body)
        validate_import_columns(df)
    except ValueError as exc:
        return text(body=str(exc), status=400)

    try:
        competitions = await asyncio.to_thread(build_import_competitions, df, custom_fields, storage)
    except (TypeError, ValueError) as exc:
        return text(body=f'Invalid row data: {exc}', status=400)

    existing = await asyncio.to_thread(storage.get_competitions)
    new_competitions, skipped_duplicates, conflicts = split_import_competitions(competitions, existing)

    # Одна транзакция и один COMMIT на весь импорт, в отдельном потоке:
    # построчные коммиты справочников (WAL+fsync ~4 раза на строку) и запись
    # в event loop замораживали приложение на всё время импорта (QA №1).
    # Дубли отсеяны до транзакции; справочники storage пополняет значениями
    # из вставленных записей (решение 2026-09-13 по инциденту rename).
    await asyncio.to_thread(
        storage.import_competitions,
        new_competitions,
        owner_id=get_current_user_id(request),
    )

    # Конфликт-режим импорта (решение 2026-09-13, docs/data-model-decisions.md
    # «Конфликт-режим импорта: очередь на подтверждение»): похожие строки
    # не вставляются молча и не отклоняют импорт — ждут явного решения админа.
    created_by = get_current_user_id(request)
    for conflict, matched_record_id in conflicts:
        await asyncio.to_thread(
            storage.add_import_queue_entry,
            conflict.model_dump(mode='json', by_alias=False),
            matched_record_id=matched_record_id,
            created_by=created_by,
        )

    return text(
        body=build_import_summary(
            len(new_competitions),
            skipped_duplicates,
            len(conflicts),
            is_admin=user_is_admin(request),
        )
    )


def get_athlete_profile_defaults(request: Request, storage: SQLiteAdapter) -> dict:
    if not user_is_athlete(request):
        return {}
    user_id = get_current_user_id(request)
    if user_id is None:
        return {}
    return storage.get_profile(user_id)


def read_profile_payload(request: Request) -> tuple[dict, str | None]:
    profile = {}
    for field in PROFILE_FIELDS:
        value = get_form_value(request, field).strip()
        if value:
            profile[field] = value
    if profile.get('student_sex') not in (None, '', 'М', 'Ж'):
        return {}, 'Пол должен быть М или Ж'
    if 'course' in profile:
        try:
            profile['course'] = str(int(profile['course']))
        except ValueError:
            return {}, 'Курс должен быть числом'
    return profile, None


@app.get('/api/profile')
async def get_profile(request: Request):
    if get_auth_user(request) is None:
        return unauthorized(request)
    user_id = get_current_user_id(request)
    profile = get_storage(request.app).get_profile(user_id) if user_id else {}
    return json_response({'profile': profile})


@app.post('/api/profile')
async def save_profile(request: Request):
    if get_auth_user(request) is None:
        return unauthorized(request)
    user_id = get_current_user_id(request)
    if user_id is None:
        return text(body='Unknown user', status=400)
    profile, error = read_profile_payload(request)
    if error:
        return text(body=error, status=400)
    storage = get_storage(request.app)
    storage.set_profile(user_id, profile)
    if profile.get('student_name'):
        storage.add_name_alias(user_id, profile['student_name'])
    return json_response({'profile': profile})


@app.get('/api/students')
async def list_students(request: Request):
    # Полный список ФИО раскрывает персональные данные других студентов —
    # доступен только admin/editor (см. docs/data-model-decisions.md).
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    names = get_storage(request.app).get_student_names()
    return json_response(names)


@app.get('/api/athletes/search')
async def search_athletes(request: Request):
    # №23доп (docs/feedback-live.md): резолвер атлета. Подсказки по ФИО
    # раскрывают персональные данные других студентов — только admin/editor
    # (docs/data-model-decisions.md); атлету и наблюдателю — 403. У атлета
    # своя автоподстановка из собственного профиля (/api/profile).
    #
    # Источник — ОБЪЕДИНЕНИЕ: активные карточки студентов (таблица students,
    # находятся и без истории участий) + легаси-известные атлеты (профили и
    # последние записи). Каждый вариант помечен kind: 'student' | 'legacy';
    # при точном совпадении ФИО (strip + casefold) карточка Student приоритетна
    # — легаси-дубль не показывается, данные берутся из карточки. Разные
    # написания остаются обоими вариантами. Слияния/авто-связывания нет:
    # kind 'student' лишь позволяет вызывающей стороне записать student_ref_id
    # ПОСЛЕ явного выбора человеком.
    #
    # Лимит — ПО KIND (до ATHLETE_SEARCH_LIMIT карточек и до столько же
    # легаси-вариантов): карточки идут раньше, но НЕ вытесняют легаси-подсказки
    # при полном лимите совпадающих карточек — легаси-люди без карточек должны
    # оставаться досягаемыми и в резолвере главной (клиент фильтрует карточки).
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error
    query = str(request.args.get('q', '')).strip()
    if not query:
        return json_response([])

    storage = get_storage(request.app)

    def combined_search() -> list[dict]:
        students = storage.search_student_suggestions(query, limit=ATHLETE_SEARCH_LIMIT)
        legacy = storage.search_athletes(query, limit=ATHLETE_SEARCH_LIMIT)
        student_names = {(item['name'] or '').strip().casefold() for item in students}
        result = [
            {
                'kind': 'student',
                'student_id': item['student_id'],
                'name': item['name'],
                **{key: item[key] for key in ('sex', 'institute', 'group', 'course') if item.get(key)},
            }
            for item in students
        ]
        result.extend(
            {
                'kind': 'legacy',
                'student_id': None,
                **{key: value for key, value in item.items() if key != 'name' and value},
                'name': item['name'],
            }
            for item in legacy
            if (item['name'] or '').strip().casefold() not in student_names
        )
        return result

    athletes = await asyncio.to_thread(combined_search)
    return json_response(athletes)


@app.get('/api/students/lookup')
async def lookup_student(request: Request):
    # Атлету нельзя отдавать чужие ФИО — только факт точного совпадения.
    user = get_auth_user(request)
    if user is None:
        return unauthorized(request)
    if user['role'] not in MODERATOR_ROLES and user['role'] != ATHLETE_ROLE:
        return forbidden(request)

    name = str(request.args.get('name', '')).strip()
    if not name:
        return text(body='Параметр name обязателен', status=400)

    # Тот же хеш, что использует привязка записей к кабинету атлета.
    name_hash = hashlib.sha256(name.encode()).hexdigest()
    records_count = get_storage(request.app).count_records_by_student_hash(name_hash)
    return json_response({'found': records_count > 0})


@app.post('/competition')
async def add_competition(request: Request):
    auth_error = require_writer(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    custom_fields = storage.get_custom_fields()
    record = {
        'ФИО': get_form_value(request, 'student_name'),
        'Пол': get_form_value(request, 'student_sex'),
        'Институт': get_form_value(request, 'institute'),
        'Группа': get_form_value(request, 'group'),
        'Вид спорта': get_form_value(request, 'sport'),
        'Дата': get_form_value(request, 'date'),
        'Уровень соревнований': get_form_value(request, 'level'),
        'Название соревнований': get_form_value(request, 'name'),
        'Место': get_form_value(request, 'position'),
        'Курс': get_form_value(request, 'course'),
    }
    record.update({field.label: get_form_value(request, f'custom__{field.key}') for field in custom_fields})

    profile_defaults = get_athlete_profile_defaults(request, storage)
    for form_key, profile_key in (
        ('ФИО', 'student_name'),
        ('Пол', 'student_sex'),
        ('Институт', 'institute'),
        ('Группа', 'group'),
        ('Курс', 'course'),
    ):
        if not str(record.get(form_key, '')).strip() and profile_key in profile_defaults:
            record[form_key] = profile_defaults[profile_key]

    try:
        competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
    except (TypeError, ValueError) as exc:
        return text(body=f'Invalid row data: {exc}', status=400)

    ensure_catalog_values(storage, [competition])
    review_status = 'pending' if user_is_athlete(request) else 'approved'
    storage.save_competitions(
        [competition],
        review_status=review_status,
        owner_id=get_current_user_id(request),
    )
    return redirect(to='/')


@app.post('/competition/<record_id>')
async def update_competition(request: Request, record_id: str):
    auth_error = require_writer(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_id = int(record_id)
    except ValueError:
        return text(body='Invalid record id', status=400)

    storage = get_storage(request.app)
    existing_review = storage.get_competition_review(numeric_id)
    if existing_review is not None and not user_owns_record(request, existing_review):
        return forbidden(request)

    custom_fields = storage.get_custom_fields()
    # P2: снимок существующей записи — источник event-owned полей и
    # discipline/result (форма их не присылает). Записи со ссылкой на
    # событие календаря (calendar_event_id) этими полями не владеют:
    # название/вид спорта/уровень/дата берутся из записи-события (правка —
    # только со страницы события), подделка формы их изменить не может.
    # NULL-записи правятся прежним путём — все поля из формы.
    existing = storage.get_competition_by_id(numeric_id)
    record = {
        'ФИО': get_form_value(request, 'student_name'),
        'Пол': get_form_value(request, 'student_sex'),
        'Институт': get_form_value(request, 'institute'),
        'Группа': get_form_value(request, 'group'),
        'Вид спорта': get_form_value(request, 'sport'),
        'Дата': get_form_value(request, 'date'),
        'Уровень соревнований': get_form_value(request, 'level'),
        'Название соревнований': get_form_value(request, 'name'),
        'Место': get_form_value(request, 'position'),
        'Курс': get_form_value(request, 'course'),
    }
    record.update({field.label: get_form_value(request, f'custom__{field.key}') for field in custom_fields})
    if existing is not None and existing.calendar_event_id is not None:
        record['Название соревнований'] = existing.name
        record['Вид спорта'] = existing.sport
        record['Уровень соревнований'] = existing.level
        record['Дата'] = format_date_range(existing.date, existing.date_to)

    try:
        competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
    except (TypeError, ValueError) as exc:
        return text(body=f'Invalid row data: {exc}', status=400)

    # Серверная часть полей, которых нет в форме: discipline/result —
    # внутренние поля participation-identity (P1), их модельные None не
    # должны затирать сохранённые значения.
    if existing is not None:
        competition.discipline = existing.discipline
        competition.result = existing.result

    storage.update_competition(numeric_id, competition)
    # Роль решает статус после правки: модератор (admin/editor)
    # подтверждает, атлет отправляет на повторную проверку — независимо
    # от прежнего статуса записи.
    if user_is_moderator(request):
        storage.set_competition_review(numeric_id, 'approved')
        # Правка модератора = авто-подтверждение, фиксируем как record_approved.
        log_audit_event(request, 'record_approved', {'record_id': numeric_id, 'auto': True})
    else:
        storage.set_competition_review(numeric_id, 'pending')
    return redirect(to='/')


@app.post('/competition/<record_id>/review/<decision>')
async def review_competition(request: Request, record_id: str, decision: str):
    auth_error = require_moderator(request)
    if auth_error is not None:
        return auth_error

    if decision not in ('approve', 'reject'):
        return text(body='Unknown review decision', status=400)

    try:
        numeric_id = int(record_id)
    except ValueError:
        return text(body='Invalid record id', status=400)

    if get_storage(request.app).get_competition_review(numeric_id) is None:
        return text(body='Record not found', status=404)

    status = 'approved' if decision == 'approve' else 'rejected'
    comment = get_form_value(request, 'comment').strip() if decision == 'reject' else ''
    get_storage(request.app).set_competition_review(numeric_id, status, comment)
    log_audit_event(
        request,
        'record_approved' if decision == 'approve' else 'record_rejected',
        {'record_id': numeric_id, 'comment': comment},
    )
    return redirect(to='/')


@app.post('/competition/<record_id>/delete')
async def delete_competition(request: Request, record_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_id = int(record_id)
    except ValueError:
        return text(body='Invalid record id', status=400)

    storage = get_storage(request.app)
    # ФИО до удаления — в аудит удаления записи (после удаления взять неоткуда).
    student_name = storage.get_competition_student_name(numeric_id)
    # Удаление записи тянет за собой её вложения (M4): сначала транзакция БД
    # (строки вложений + запись), затем файлы — тот же порядок, что у
    # perform_wipe. Файлы удаляются best-effort: остаток каталога — только
    # предупреждение в лог, БД уже согласована.
    storage.delete_competition_with_attachments(numeric_id)
    remove_attachment_record_dirs([numeric_id])
    if (files_dir() / str(numeric_id)).exists():
        logger.warning('Attachment directory of record %s still exists after record deletion', numeric_id)
    log_audit_event(request, 'record_deleted', {'record_id': numeric_id, 'student_name': student_name})
    return redirect(to='/')


@app.post('/admin/fields')
async def create_custom_field(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    label = get_form_value(request, 'label').strip()
    field_type = get_form_value(request, 'field_type').strip() or 'text'
    if not label:
        return build_redirect_with_message(error='Название поля обязательно', url='/admin/fields')
    if field_type not in FIELD_TYPE_OPTIONS:
        return build_redirect_with_message(error='Недопустимый тип поля', url='/admin/fields')

    try:
        storage.create_custom_field(
            key=make_unique_custom_field_key(label, storage.get_custom_fields(include_inactive=True)),
            label=label,
            field_type=field_type,
            required=parse_checkbox(request, 'required'),
            show_in_table=parse_checkbox(request, 'show_in_table'),
            show_in_export=parse_checkbox(request, 'show_in_export'),
            show_in_template=parse_checkbox(request, 'show_in_template'),
            sort_order=int(get_form_value(request, 'sort_order') or 0),
            link_target=parse_link_target_form(
                request, field_type=field_type, custom_fields=storage.get_custom_fields()
            ),
        )
    except Exception as exc:
        return build_redirect_with_message(error=f'Не удалось создать поле: {exc}', url='/admin/fields')
    return build_redirect_with_message(message='Поле добавлено', url='/admin/fields')


@app.post('/admin/fields/<field_id>')
async def update_custom_field(request: Request, field_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_field_id = int(field_id)
    except ValueError:
        return text(body='Invalid field id', status=400)

    label = get_form_value(request, 'label').strip()
    field_type = get_form_value(request, 'field_type').strip() or 'text'
    if not label:
        return build_redirect_with_message(error='Название поля обязательно', url='/admin/fields')
    if field_type not in FIELD_TYPE_OPTIONS:
        return build_redirect_with_message(error='Недопустимый тип поля', url='/admin/fields')

    try:
        sort_order = int(get_form_value(request, 'sort_order') or 0)
    except ValueError:
        return build_redirect_with_message(error='Порядок должен быть числом', url='/admin/fields')

    storage = get_storage(request.app)
    # №24: старое значение привязки нужно для audit-события при изменении.
    old_field = next(
        (field for field in storage.get_custom_fields(include_inactive=True) if field.field_id == numeric_field_id),
        None,
    )
    new_link_target = parse_link_target_form(
        request,
        field_type=field_type,
        custom_fields=storage.get_custom_fields(),
        exclude_key=old_field.key if old_field else None,
    )
    storage.update_custom_field(
        field_id=numeric_field_id,
        label=label,
        field_type=field_type,
        required=parse_checkbox(request, 'required'),
        show_in_table=parse_checkbox(request, 'show_in_table'),
        show_in_export=parse_checkbox(request, 'show_in_export'),
        show_in_template=parse_checkbox(request, 'show_in_template'),
        sort_order=sort_order,
        active=parse_checkbox(request, 'active'),
        link_target=new_link_target,
    )
    if old_field is not None and old_field.link_target != new_link_target:
        log_audit_event(
            request,
            'field_settings_changed',
            {
                'fields': {
                    old_field.key: {
                        'link_target': {'old': old_field.link_target, 'new': new_link_target},
                    }
                }
            },
        )
    return build_redirect_with_message(message='Настройки поля сохранены', url='/admin/fields')


@app.post('/admin/fields/<field_id>/delete')
async def delete_custom_field(request: Request, field_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_field_id = int(field_id)
    except ValueError:
        return text(body='Invalid field id', status=400)

    storage = get_storage(request.app)
    storage.disable_custom_field(numeric_field_id)
    return build_redirect_with_message(message='Поле отключено', url='/admin/fields')


@app.post('/admin/users')
async def create_user(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    username = get_form_value(request, 'username').strip()
    password = get_form_value(request, 'password')
    role = get_form_value(request, 'role').strip()
    if not username:
        return build_redirect_with_message(error='Имя пользователя обязательно', url='/admin/users')
    invalid_username_message = username_error(username)
    if invalid_username_message:
        return build_redirect_with_message(error=invalid_username_message, url='/admin/users')
    if role not in USER_ROLES:
        return build_redirect_with_message(error='Недопустимая роль', url='/admin/users')
    if len(password) < MIN_PASSWORD_LENGTH:
        return build_redirect_with_message(
            error=f'Пароль должен быть не короче {MIN_PASSWORD_LENGTH} символов', url='/admin/users'
        )
    if storage.get_user(username) is not None:
        return build_redirect_with_message(error='Пользователь уже существует', url='/admin/users')

    storage.create_user(username, hash_password(password), role)
    return build_redirect_with_message(message='Пользователь добавлен', url='/admin/users')


@app.post('/admin/users/<user_id>/password')
async def reset_user_password(request: Request, user_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_user_id = int(user_id)
    except ValueError:
        return text(body='Invalid user id', status=400)

    password = get_form_value(request, 'password')
    if len(password) < MIN_PASSWORD_LENGTH:
        return build_redirect_with_message(
            error=f'Пароль должен быть не короче {MIN_PASSWORD_LENGTH} символов', url='/admin/users'
        )

    storage = get_storage(request.app)
    target_user = storage.get_user_by_id(numeric_user_id)
    if target_user is None:
        return build_redirect_with_message(error='Пользователь не найден', url='/admin/users')
    storage.set_user_password(numeric_user_id, hash_password(password))
    log_audit_event(
        request,
        'password_changed',
        {'target_user_id': numeric_user_id, 'target_username': target_user['username']},
    )
    return build_redirect_with_message(message='Пароль обновлён', url='/admin/users')


@app.post('/admin/users/<user_id>/active')
async def toggle_user_active(request: Request, user_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_user_id = int(user_id)
    except ValueError:
        return text(body='Invalid user id', status=400)

    storage = get_storage(request.app)
    user = storage.get_user_by_id(numeric_user_id)
    if user is None:
        return build_redirect_with_message(error='Пользователь не найден', url='/admin/users')
    if user['username'] == (get_auth_user(request) or {}).get('username'):
        return build_redirect_with_message(error='Нельзя отключить собственную учётную запись', url='/admin/users')

    storage.set_user_active(numeric_user_id, not user['active'])
    return build_redirect_with_message(message='Статус пользователя изменён', url='/admin/users')


def count_active_admins(storage: SQLiteAdapter) -> int:
    return sum(1 for user in storage.list_users() if user['role'] == ADMIN_ROLE and user['active'])


@app.post('/admin/users/<user_id>/delete')
async def delete_user(request: Request, user_id: str):
    # Удаление аккаунта: записи — исторические факты, остаются с owner_id NULL;
    # псевдонимы и привязка кабинета исчезают вместе с аккаунтом. Подтверждение —
    # ввод логина в поле (docs/data-model-decisions.md «Удаление пользователей»).
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_user_id = int(user_id)
    except ValueError:
        return text(body='Invalid user id', status=400)

    storage = get_storage(request.app)
    target_user = storage.get_user_by_id(numeric_user_id)
    if target_user is None:
        return build_redirect_with_message(error='Пользователь не найден', url='/admin/users')
    if target_user['username'] == (get_auth_user(request) or {}).get('username'):
        return build_redirect_with_message(error='Нельзя удалить собственную учётную запись', url='/admin/users')
    if target_user['role'] == ADMIN_ROLE and target_user['active'] and count_active_admins(storage) <= 1:
        return build_redirect_with_message(
            error='Нельзя удалить последнего активного администратора', url='/admin/users'
        )

    confirm_username = get_form_value(request, 'confirm_username').strip()
    if not confirm_username or confirm_username != target_user['username']:
        return build_redirect_with_message(
            error='Для подтверждения введите логин удаляемого пользователя',
            url='/admin/users',
        )

    detached = storage.delete_user(numeric_user_id)
    log_audit_event(
        request,
        'user_deleted',
        {
            'target_user_id': numeric_user_id,
            'target_username': target_user['username'],
            'records_detached': detached,
        },
    )
    return build_redirect_with_message(
        message=f'Пользователь «{target_user["username"]}» удалён. Отвязано записей: {detached}',
        url='/admin/users',
    )


def build_attachments_by_record(all_attachments: list[dict]) -> dict[int, list[dict]]:
    grouped: dict[int, list[dict]] = {}
    for attachment in all_attachments:
        grouped.setdefault(attachment['record_id'], []).append(attachment)
    return grouped


def attachments_dir() -> Path:
    base = Path(settings.data_folder) / 'files'
    base.mkdir(parents=True, exist_ok=True)
    return base


def calendar_regulation_dir(event_id: int) -> Path:
    """Каталог файла положения события: data/files/calendar/<event_id>
    (общий корень data/files — вложения и положения бэкапятся вместе)."""
    return Path(settings.data_folder) / 'files' / 'calendar' / str(event_id)


def regulation_source_path(event_id: int, stored_name: str) -> Path:
    """Путь к файлу положения события на диске по служебному имени."""
    return calendar_regulation_dir(event_id) / stored_name


def detect_attachment_type(body: bytes, filename: str) -> str | None:
    extension = Path(filename).suffix.lower().lstrip('.')
    expected_type = ATTACHMENT_EXTENSIONS.get(extension)
    if expected_type is None:
        return None
    signature = ATTACHMENT_SIGNATURES[expected_type]
    return expected_type if body.startswith(signature) else None


@app.post('/competition/<record_id>/attachments')
async def upload_attachment(request: Request, record_id: str):
    auth_error = require_writer(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_id = int(record_id)
    except ValueError:
        return text(body='Invalid record id', status=400)

    storage = get_storage(request.app)
    review = storage.get_competition_review(numeric_id)
    if review is None:
        return text(body='Record not found', status=404)
    if not user_owns_record(request, review):
        return forbidden(request)

    upload_file = request.files.get('file')
    if not upload_file or not upload_file.body:
        return text(body='No file uploaded', status=400)
    if len(upload_file.body) > ATTACHMENT_MAX_SIZE:
        return text(body='Файл больше 5 МБ', status=400)

    filename = (upload_file.name or 'attachment').rsplit('/', 1)[-1]
    content_type = detect_attachment_type(upload_file.body, filename)
    if content_type is None:
        return text(body='Допустимы только PDF, JPEG и PNG', status=400)

    extension = Path(filename).suffix.lower().lstrip('.') or 'bin'
    stored_name = f'{uuid.uuid4().hex}.{extension}'
    record_dir = attachments_dir() / str(numeric_id)
    record_dir.mkdir(parents=True, exist_ok=True)
    (record_dir / stored_name).write_bytes(upload_file.body)

    storage.create_attachment(
        record_id=numeric_id,
        filename=filename,
        stored_name=stored_name,
        content_type=content_type,
        size=len(upload_file.body),
        uploaded_by=get_current_user_id(request),
    )
    return text(body='Файл загружен')


@app.get('/attachment/<attachment_id>')
async def download_attachment(request: Request, attachment_id: str):
    if get_auth_user(request) is None:
        return unauthorized(request)

    try:
        numeric_attachment_id = int(attachment_id)
    except ValueError:
        return text(body='Invalid attachment id', status=400)

    storage = get_storage(request.app)
    attachment = storage.get_attachment(numeric_attachment_id)
    if attachment is None:
        return text(body='Not Found', status=404)
    review = storage.get_competition_review(attachment['record_id'])
    if review is not None and not user_owns_record(request, review):
        return forbidden(request)

    file_path = attachments_dir() / str(attachment['record_id']) / attachment['stored_name']
    if not file_path.is_file():
        return text(body='Not Found', status=404)

    return raw(
        await asyncio.to_thread(file_path.read_bytes),
        headers={
            'content-type': attachment['content_type'],
            'content-disposition': f'attachment; filename="{attachment["stored_name"]}"',
        },
    )


@app.post('/attachment/<attachment_id>/delete')
async def delete_attachment(request: Request, attachment_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_attachment_id = int(attachment_id)
    except ValueError:
        return text(body='Invalid attachment id', status=400)

    storage = get_storage(request.app)
    attachment = storage.get_attachment(numeric_attachment_id)
    if attachment is None:
        return text(body='Not Found', status=404)

    file_path = attachments_dir() / str(attachment['record_id']) / attachment['stored_name']
    file_path.unlink(missing_ok=True)
    storage.delete_attachment(numeric_attachment_id)
    return redirect(to='/')


@app.post('/admin/levels')
async def create_level(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    name = get_form_value(request, 'name').strip()
    if not name:
        return build_redirect_with_message(error='Название уровня обязательно', url='/admin/catalogs')
    # Уровень всегда хранится в нижнем регистре (так нормализуются импорт
    # и записи) — ручное добавление приведено к тому же виду. Именно ручное
    # добавление без нормализации было источником регистровых дублей (№1,
    # docs/feedback-live.md).
    name = name.lower()
    existing_levels = get_storage(request.app).get_level_names(include_inactive=True)
    # №1.1: дубль (в том числе регистровый — имя уже в lower) — отказ.
    if any(item.lower() == name for item in existing_levels):
        return build_redirect_with_message(error='Такой уровень уже существует', url='/admin/catalogs')

    get_storage(request.app).create_level(name)
    return build_redirect_with_message(message='Уровень добавлен', url='/admin/catalogs')


@app.post('/admin/levels/<level_id>/delete')
async def disable_level(request: Request, level_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_level_id = int(level_id)
    except ValueError:
        return text(body='Invalid level id', status=400)

    get_storage(request.app).disable_level(numeric_level_id)
    return build_redirect_with_message(message='Уровень скрыт из списков', url='/admin/catalogs')


@app.post('/admin/levels/<level_id>/hard-delete')
async def hard_delete_level(request: Request, level_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_level_id = int(level_id)
    except ValueError:
        return text(body='Invalid level id', status=400)

    storage = get_storage(request.app)
    level = next((item for item in storage.list_levels() if item['id'] == numeric_level_id), None)
    if level is None:
        return build_redirect_with_message(error='Уровень не найден', url='/admin/catalogs')

    records_count = storage.count_records_using('level', level['name'])
    if records_count > 0:
        return build_redirect_with_message(
            error='У уровня есть записи — скройте или переименуйте', url='/admin/catalogs'
        )

    storage.hard_delete_level(numeric_level_id)
    log_audit_event(
        request,
        'catalog_value_deleted',
        {'category': 'level', 'value': level['name'], 'level_id': numeric_level_id},
    )
    return build_redirect_with_message(message='Уровень удалён', url='/admin/catalogs')


@app.post('/admin/users/<user_id>/alias')
async def add_user_alias(request: Request, user_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    try:
        numeric_user_id = int(user_id)
    except ValueError:
        return text(body='Invalid user id', status=400)

    storage = get_storage(request.app)
    if storage.get_user_by_id(numeric_user_id) is None:
        return build_redirect_with_message(error='Пользователь не найден', url='/admin/users')

    name = get_form_value(request, 'name').strip()
    if not name:
        return build_redirect_with_message(error='ФИО обязательно', url='/admin/users')

    target_user = storage.get_user_by_id(numeric_user_id)
    storage.add_name_alias(numeric_user_id, name)
    log_audit_event(
        request,
        'alias_added',
        {
            'target_user_id': numeric_user_id,
            'target_username': target_user['username'] if target_user else '',
            'name': name,
        },
    )
    return build_redirect_with_message(message=f'ФИО «{name}» привязано к аккаунту', url='/admin/users')


# --- Карточки студентов (Student Identity v1, Phase 1 — фундамент). ---
#
# Отдельная таблица актуальных данных студента (ФИО/пол/институт/группа/курс)
# и его псевдонимов ФИО. В Phase 1 карточки живут ИЗОЛИРОВАННО: записи
# соревнований, аккаунты, импорт, отчёты и merge про них не знают
# (student_ref_id не пишется, авто-создания/авто-связей нет). Правка карточки
# меняет ТОЛЬКО актуальные «профильные» данные, исторические записи не
# перезаписываются.


def parse_student_form(request: Request) -> tuple[dict, str | None]:
    """Поля карточки студента из формы. None-ошибка — текст для редиректа.

    Поле формы группы называется «group» (HTML-конвенция) и маппится в
    group_name хранилища. Все значения — свободный текст со strip().
    """
    full_name = get_form_value(request, 'full_name').strip()
    sex = get_form_value(request, 'sex').strip()
    institute = get_form_value(request, 'institute').strip()
    group_name = get_form_value(request, 'group').strip()
    course = get_form_value(request, 'course').strip()
    if not full_name:
        return {}, 'Укажите ФИО студента.'
    if sex not in STUDENT_SEX_OPTIONS:
        return {}, 'Пол может быть «М», «Ж» или не указан.'
    return {
        'full_name': full_name,
        'sex': sex,
        'institute': institute,
        'group_name': group_name,
        'course': course,
    }, None


def student_field_changes(student: dict, form: dict) -> dict:
    """Diff изменённых полей карточки для аудита: поле → {old, new}."""
    return {
        field: {'old': student[field], 'new': form[field]}
        for field in ('full_name', 'sex', 'institute', 'group_name', 'course')
        if student[field] != form[field]
    }


def resolve_student_id(request: Request, student_id: str):
    """Числовой id карточки из URL: (id, None) или (None, ответ 400)."""
    try:
        return int(student_id), None
    except ValueError:
        return None, text(body='Invalid student id', status=400)


@app.get('/admin/people')
async def admin_people_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    search = (get_param(dict(request.args), 'q') or '').strip()
    students = storage.list_students(search=search)
    total = len(storage.list_students())
    return await render(
        template_name=jinja_env.get_template('admin_people.html'),
        context={
            'request': request,
            'students': students,
            'q': search,
            'total': total,
            'reconcile_unlinked': storage.count_student_reconciliation()['records_unlinked'],
            **get_flash_args(request),
        },
    )


@app.post('/admin/people')
async def admin_person_create(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    form, form_error = parse_student_form(request)
    if form_error is not None:
        return build_redirect_with_message(error=form_error, url='/admin/people')

    storage = get_storage(request.app)
    student_id = storage.create_student(
        form['full_name'],
        form['sex'],
        form['institute'],
        form['group_name'],
        form['course'],
    )
    log_audit_event(request, 'student_created', {'student_id': student_id, 'full_name': form['full_name']})
    return build_redirect_with_message(
        message=f'Студент «{form["full_name"]}» добавлен.',
        url=f'/admin/people/{student_id}',
    )


# --- Сопоставление данных (Student Identity v1, Phase 2). ---
#
# Ручное заполнение стабильных связей student_ref_id у СУЩЕСТВУЮЩИХ записей
# и аккаунтов атлетов. Кандидаты на странице — только ПРЕДЛОЖЕНИЯ, ничего
# не связывается автоматически. Привязка меняет ТОЛЬКО student_ref_id:
# снимки данных в записях, легаси-ключ sha256(ФИО), кабинет атлета,
# отчёты, импорт/экспорт и календарь работают как раньше (Phase 3 не
# начата). Регистрируется ДО /admin/people/<student_id>, чтобы статический
# путь /admin/people/reconcile не разбирался как id карточки.


def parse_reconcile_int(raw: str, label: str):
    """Целое из строки формы/URL: (значение, None) или (None, ответ 400).

    Пустая строка — тоже некорректный id (поле обязательное).
    """
    raw = (raw or '').strip()
    if not raw:
        return None, text(body=f'Invalid {label}', status=400)
    try:
        return int(raw), None
    except ValueError:
        return None, text(body=f'Invalid {label}', status=400)


def reconcile_back_url(raw: str) -> str:
    """Возврат внутрь раздела сопоставления: открытые редиректы исключены,
    посторонние пути заменяются вкладкой записей."""
    raw = (raw or '').strip()
    if raw.startswith('/admin/people/reconcile'):
        return raw
    return '/admin/people/reconcile'


def reconcile_student_error(storage: SQLiteAdapter, student_id: int, code: str) -> str:
    """Flash-текст ошибки привязки по коду хранилища (карточка/активность)."""
    student = storage.get_student_by_id(student_id)
    if student is None:
        return 'Студент не найден.'
    if code == 'student_inactive':
        return f'Студент «{student["full_name"]}» неактивен: привязка возможна только к активным студентам.'
    return 'Студент не найден.'


def reconcile_link_error(
    storage: SQLiteAdapter,
    record_ids: list[int],
    student_id: int,
    code: str,
) -> str:
    """Flash-текст ошибки привязки записей (одиночной и массовой)."""
    if code in ('student_not_found', 'student_inactive'):
        return reconcile_student_error(storage, student_id, code)
    if code == 'records_not_found':
        return 'Запись не найдена.'
    if len(record_ids) == 1:
        return f'Запись №{record_ids[0]} уже привязана к студенту.'
    return 'Ничего не привязано: часть выбранных записей уже привязана. Обновите список и повторите.'


def reconcile_create_url(*, record_id: int | None = None, user_id: int | None = None, back: str) -> str:
    """URL страницы создания студента из записи/профиля с возвратом."""
    params = {'back': back}
    if record_id is not None:
        params['record_id'] = record_id
    if user_id is not None:
        params['user_id'] = user_id
    query = urlencode(params)
    if record_id is not None:
        return f'/admin/people/reconcile/create?{query}'
    return f'/admin/people/reconcile/users/create?{query}'


@app.get('/admin/people/reconcile')
async def admin_reconcile_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    q = (get_param(dict(request.args), 'q') or '').strip()
    try:
        page = max(1, int(get_param(dict(request.args), 'page') or 1))
    except ValueError:
        page = 1

    counters = storage.count_student_reconciliation()
    total = storage.count_unlinked_competitions(search=q)
    pages = max(1, -(-total // RECONCILE_PAGE_SIZE))
    page = min(page, pages)
    records = storage.list_unlinked_competitions(
        search=q, limit=RECONCILE_PAGE_SIZE, offset=(page - 1) * RECONCILE_PAGE_SIZE
    )
    # Кандидаты — только для текущей страницы (точное совпадение ФИО).
    for record in records:
        record['candidates'] = storage.find_student_candidates(record['student_name'])
        # Даты — datetime для format_date_range в шаблоне.
        record['date'] = datetime.fromisoformat(record['date']) if record['date'] else None
        record['date_to'] = datetime.fromisoformat(record['date_to']) if record['date_to'] else None
    active_students = [student for student in storage.list_students() if student['active']]

    query = urlencode({'q': q}) if q else ''

    def page_url(page_number: int) -> str:
        return (
            f'/admin/people/reconcile?{query}&page={page_number}'
            if query
            else f'/admin/people/reconcile?page={page_number}'
        )

    return await render(
        template_name=jinja_env.get_template('admin_reconcile.html'),
        context={
            'request': request,
            'counters': counters,
            'q': q,
            'records': records,
            'active_students': active_students,
            'page': page,
            'pages': pages,
            'total': total,
            'page_size': RECONCILE_PAGE_SIZE,
            'back_url': page_url(page),
            'prev_url': page_url(page - 1) if page > 1 else None,
            'next_url': page_url(page + 1) if page < pages else None,
            **get_flash_args(request),
        },
    )


@app.get('/admin/people/reconcile/users')
async def admin_reconcile_users_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    counters = storage.count_student_reconciliation()
    users = storage.list_unlinked_athlete_users()
    # Кандидаты — по ФИО из профиля аккаунта; псевдонимы аккаунта НЕ участвуют.
    for user in users:
        user['candidates'] = storage.find_student_candidates(user['profile_data'].get('student_name') or '')

    return await render(
        template_name=jinja_env.get_template('admin_reconcile_users.html'),
        context={
            'request': request,
            'counters': counters,
            'users': users,
            'back_url': '/admin/people/reconcile/users',
            **get_flash_args(request),
        },
    )


@app.post('/admin/people/reconcile/link')
async def admin_reconcile_link(request: Request):
    """Привязка записей к карточке: одна точка для одиночной (кнопка
    кандидата) и массовой (галочки + список студентов) привязки."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    raw_ids = [value for value in request.form.getlist('record_ids[]') if str(value).strip()]
    numeric_ids: list[int] = []
    for raw_id in raw_ids:
        numeric_id, id_error = parse_reconcile_int(str(raw_id), 'record id')
        if id_error is not None:
            return id_error
        numeric_ids.append(numeric_id)

    back = reconcile_back_url(get_form_value(request, 'back'))
    if not numeric_ids:
        return build_redirect_with_message(error='Не выбрано ни одной записи.', url=back)
    student_id_raw = get_form_value(request, 'student_id').strip()
    if not student_id_raw:
        return build_redirect_with_message(error='Выберите студента.', url=back)
    student_id, id_error = parse_reconcile_int(student_id_raw, 'student id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    count, error = storage.link_competitions(numeric_ids, student_id)
    if error is not None:
        return build_redirect_with_message(
            error=reconcile_link_error(storage, numeric_ids, student_id, error), url=back
        )

    student = storage.get_student_by_id(student_id)
    full_name = student['full_name'] if student else ''
    if len(numeric_ids) == 1 and count == 1:
        log_audit_event(
            request,
            'competition_linked_to_student',
            {'record_id': numeric_ids[0], 'old_ref': None, 'new_ref': student_id, 'student_id': student_id},
        )
        message = f'Запись №{numeric_ids[0]} привязана к студенту «{full_name}».'
    else:
        # Одна запись аудита на всю массовую привязку.
        log_audit_event(
            request,
            'student_records_bulk_linked',
            {'student_id': student_id, 'record_ids': numeric_ids, 'count': count},
        )
        message = f'Привязано записей: {count} — студент «{full_name}».'
    return build_redirect_with_message(message=message, url=back)


@app.get('/admin/people/reconcile/create')
async def admin_reconcile_create_page(request: Request):
    """Создание карточки из записи: поля предзаполнены снимком записи."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    record_id, id_error = parse_reconcile_int(get_param(dict(request.args), 'record_id') or '', 'record id')
    if id_error is not None:
        return id_error
    back = reconcile_back_url(get_param(dict(request.args), 'back') or '')

    storage = get_storage(request.app)
    record = storage.get_competition_by_id(record_id)
    found, student_ref = storage.get_competition_student_ref(record_id)
    if record is None or not found:
        return build_redirect_with_message(error='Запись не найдена.', url=back)
    if student_ref is not None:
        return build_redirect_with_message(
            error=f'Запись №{record_id} уже привязана к студенту.',
            url=back,
        )

    return await render(
        template_name=jinja_env.get_template('admin_reconcile_create.html'),
        context={
            'request': request,
            'source': 'record',
            'record': record,
            'record_id': record_id,
            'form': {
                'full_name': record.student_name,
                'sex': record.student_sex or '',
                'institute': record.institute or '',
                'group_name': record.group or '',
                'course': '' if record.course is None else str(record.course),
            },
            'back_url': back,
            **get_flash_args(request),
        },
    )


@app.post('/admin/people/reconcile/create')
async def admin_reconcile_create_record(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    record_id, id_error = parse_reconcile_int(get_form_value(request, 'record_id'), 'record id')
    if id_error is not None:
        return id_error
    back = reconcile_back_url(get_form_value(request, 'back'))
    form, form_error = parse_student_form(request)
    if form_error is not None:
        return build_redirect_with_message(
            error=form_error,
            url=reconcile_create_url(record_id=record_id, back=back),
        )

    storage = get_storage(request.app)
    record = storage.get_competition_by_id(record_id)
    found, student_ref = storage.get_competition_student_ref(record_id)
    if record is None or not found:
        return build_redirect_with_message(error='Запись не найдена.', url=back)
    if student_ref is not None:
        return build_redirect_with_message(
            error=f'Запись №{record_id} уже привязана к студенту.',
            url=back,
        )

    student_id = storage.create_student(
        form['full_name'],
        form['sex'],
        form['institute'],
        form['group_name'],
        form['course'],
    )
    log_audit_event(
        request,
        'student_created',
        {'student_id': student_id, 'full_name': form['full_name'], 'source': 'reconciliation-record'},
    )
    _, link_error = storage.link_competitions([record_id], student_id)
    if link_error is not None:
        # Гонка: карточка уже создана (и остаётся), запись связал кто-то другой.
        if link_error == 'already_linked':
            message = (
                f'Студент «{form["full_name"]}» создан, но запись №{record_id} уже была привязана другим действием.'
            )
        else:
            message = f'Студент «{form["full_name"]}» создан, но запись №{record_id} не найдена.'
        return build_redirect_with_message(message=message, url=f'/admin/people/{student_id}')
    log_audit_event(
        request,
        'competition_linked_to_student',
        {'record_id': record_id, 'old_ref': None, 'new_ref': student_id, 'student_id': student_id},
    )
    return build_redirect_with_message(
        message=f'Студент «{form["full_name"]}» создан, запись №{record_id} привязана к нему.',
        url=f'/admin/people/{student_id}',
    )


@app.post('/admin/people/reconcile/unlink')
async def admin_reconcile_unlink(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    record_id, id_error = parse_reconcile_int(get_form_value(request, 'record_id'), 'record id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    found, student_ref = storage.get_competition_student_ref(record_id)
    if not found:
        return build_redirect_with_message(error='Запись не найдена.', url='/admin/people/reconcile')
    if student_ref is None:
        return build_redirect_with_message(
            error=f'Запись №{record_id} не привязана к студенту.',
            url='/admin/people/reconcile',
        )

    old_ref = storage.unlink_competition(record_id)
    if old_ref is None:
        return build_redirect_with_message(
            error=f'Запись №{record_id} не привязана к студенту.',
            url='/admin/people/reconcile',
        )
    student = storage.get_student_by_id(old_ref)
    full_name = student['full_name'] if student else ''
    log_audit_event(
        request,
        'competition_unlinked_from_student',
        {'record_id': record_id, 'old_ref': old_ref},
    )
    return build_redirect_with_message(
        message=f'Запись №{record_id} отвязана от студента «{full_name}».',
        url=f'/admin/people/{old_ref}',
    )


@app.post('/admin/people/reconcile/relink')
async def admin_reconcile_relink(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    record_id, id_error = parse_reconcile_int(get_form_value(request, 'record_id'), 'record id')
    if id_error is not None:
        return id_error
    student_id, id_error = parse_reconcile_int(get_form_value(request, 'student_id'), 'student id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    old_ref, error = storage.relink_competition(record_id, student_id)
    if error is not None:
        if error == 'records_not_found':
            return build_redirect_with_message(error='Запись не найдена.', url='/admin/people/reconcile')
        if error in ('student_not_found', 'student_inactive'):
            # Возврат к карточке, где кнопка: прежняя связь записи ещё цела.
            _, current_ref = storage.get_competition_student_ref(record_id)
            url = f'/admin/people/{current_ref}' if current_ref else '/admin/people/reconcile'
            return build_redirect_with_message(error=reconcile_student_error(storage, student_id, error), url=url)

    student = storage.get_student_by_id(student_id)
    full_name = student['full_name'] if student else ''
    log_audit_event(
        request,
        'competition_relinked',
        {'record_id': record_id, 'old_ref': old_ref, 'new_ref': student_id, 'student_id': student_id},
    )
    return build_redirect_with_message(
        message=f'Запись №{record_id} перепривязана на студента «{full_name}».',
        url=f'/admin/people/{student_id}',
    )


@app.get('/admin/people/reconcile/users/create')
async def admin_reconcile_user_create_page(request: Request):
    """Создание карточки из профиля аккаунта атлета: псевдонимы аккаунта
    НЕ переносятся — у карточки будут только данные профиля."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    user_id, id_error = parse_reconcile_int(get_param(dict(request.args), 'user_id') or '', 'user id')
    if id_error is not None:
        return id_error
    back = reconcile_back_url(get_param(dict(request.args), 'back') or '')

    storage = get_storage(request.app)
    user = storage.get_unlinked_athlete_user(user_id)
    if user is None:
        legacy_user = storage.get_user_by_id(user_id)
        if legacy_user is None:
            return build_redirect_with_message(error='Пользователь не найден.', url=back)
        if legacy_user['role'] != 'athlete':
            return build_redirect_with_message(error='Пользователь не является атлетом.', url=back)
        return build_redirect_with_message(
            error=f'Аккаунт {legacy_user["username"]} уже привязан к студенту.',
            url=back,
        )

    profile = user['profile_data']
    return await render(
        template_name=jinja_env.get_template('admin_reconcile_create.html'),
        context={
            'request': request,
            'source': 'user',
            'user': user,
            'user_id': user_id,
            'form': {
                'full_name': str(profile.get('student_name') or ''),
                'sex': str(profile.get('student_sex') or ''),
                'institute': str(profile.get('institute') or ''),
                'group_name': str(profile.get('group') or ''),
                'course': str(profile.get('course') or ''),
            },
            'back_url': back,
            **get_flash_args(request),
        },
    )


@app.post('/admin/people/reconcile/users/create')
async def admin_reconcile_create_user(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    user_id, id_error = parse_reconcile_int(get_form_value(request, 'user_id'), 'user id')
    if id_error is not None:
        return id_error
    back = reconcile_back_url(get_form_value(request, 'back'))
    form, form_error = parse_student_form(request)
    if form_error is not None:
        return build_redirect_with_message(
            error=form_error,
            url=reconcile_create_url(user_id=user_id, back=back),
        )

    storage = get_storage(request.app)
    user = storage.get_unlinked_athlete_user(user_id)
    if user is None:
        legacy_user = storage.get_user_by_id(user_id)
        if legacy_user is None:
            return build_redirect_with_message(error='Пользователь не найден.', url=back)
        if legacy_user['role'] != 'athlete':
            return build_redirect_with_message(error='Пользователь не является атлетом.', url=back)
        return build_redirect_with_message(
            error=f'Аккаунт {legacy_user["username"]} уже привязан к студенту.',
            url=back,
        )

    student_id = storage.create_student(
        form['full_name'],
        form['sex'],
        form['institute'],
        form['group_name'],
        form['course'],
    )
    log_audit_event(
        request,
        'student_created',
        {'student_id': student_id, 'full_name': form['full_name'], 'source': 'reconciliation-user'},
    )
    _, link_error = storage.link_user(user_id, student_id)
    if link_error is not None:
        if link_error == 'user_already_linked':
            message = (
                f'Студент «{form["full_name"]}» создан, но аккаунт {user["username"]} уже был '
                'привязан другим действием.'
            )
        else:
            message = f'Студент «{form["full_name"]}» создан, но аккаунт {user["username"]} недоступен для привязки.'
        return build_redirect_with_message(message=message, url=f'/admin/people/{student_id}')
    log_audit_event(
        request,
        'user_linked_to_student',
        {'user_id': user_id, 'username': user['username'], 'old_ref': None, 'new_ref': student_id},
    )
    return build_redirect_with_message(
        message=f'Студент «{form["full_name"]}» создан, аккаунт {user["username"]} привязан к нему.',
        url=f'/admin/people/{student_id}',
    )


@app.post('/admin/people/reconcile/users/link')
async def admin_reconcile_user_link(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    user_id, id_error = parse_reconcile_int(get_form_value(request, 'user_id'), 'user id')
    if id_error is not None:
        return id_error
    back = reconcile_back_url(get_form_value(request, 'back'))
    student_id_raw = get_form_value(request, 'student_id').strip()
    if not student_id_raw:
        return build_redirect_with_message(error='Выберите студента.', url=back)
    student_id, id_error = parse_reconcile_int(student_id_raw, 'student id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    _, error = storage.link_user(user_id, student_id)
    if error is not None:
        if error == 'user_not_found':
            return build_redirect_with_message(error='Пользователь не найден.', url=back)
        if error == 'user_not_athlete':
            return build_redirect_with_message(error='Пользователь не является атлетом.', url=back)
        if error == 'user_already_linked':
            user = storage.get_user_by_id(user_id)
            username = user['username'] if user else ''
            return build_redirect_with_message(
                error=f'Аккаунт {username} уже привязан к студенту.',
                url=back,
            )
        return build_redirect_with_message(error=reconcile_student_error(storage, student_id, error), url=back)

    user = storage.get_user_by_id(user_id)
    username = user['username'] if user else ''
    student = storage.get_student_by_id(student_id)
    full_name = student['full_name'] if student else ''
    log_audit_event(
        request,
        'user_linked_to_student',
        {'user_id': user_id, 'username': username, 'old_ref': None, 'new_ref': student_id},
    )
    return build_redirect_with_message(
        message=f'Аккаунт {username} привязан к студенту «{full_name}».',
        url=back,
    )


@app.post('/admin/people/reconcile/users/unlink')
async def admin_reconcile_user_unlink(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    user_id, id_error = parse_reconcile_int(get_form_value(request, 'user_id'), 'user id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    user = storage.get_user_by_id(user_id)
    if user is None:
        return build_redirect_with_message(error='Пользователь не найден.', url='/admin/people/reconcile/users')
    old_ref = storage.unlink_user(user_id)
    if old_ref is None:
        return build_redirect_with_message(
            error=f'Аккаунт {user["username"]} не привязан к студенту.',
            url='/admin/people/reconcile/users',
        )
    student = storage.get_student_by_id(old_ref)
    full_name = student['full_name'] if student else ''
    log_audit_event(
        request,
        'user_unlinked_from_student',
        {'user_id': user_id, 'username': user['username'], 'old_ref': old_ref, 'new_ref': None},
    )
    return build_redirect_with_message(
        message=f'Аккаунт {user["username"]} отвязан от студента «{full_name}».',
        url=f'/admin/people/{old_ref}',
    )


@app.post('/admin/people/reconcile/users/relink')
async def admin_reconcile_user_relink(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    user_id, id_error = parse_reconcile_int(get_form_value(request, 'user_id'), 'user id')
    if id_error is not None:
        return id_error
    student_id, id_error = parse_reconcile_int(get_form_value(request, 'student_id'), 'student id')
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    old_ref, error = storage.relink_user(user_id, student_id)
    if error is not None:
        if error == 'user_not_found':
            return build_redirect_with_message(error='Пользователь не найден.', url='/admin/people/reconcile/users')
        if error == 'user_not_athlete':
            return build_redirect_with_message(
                error='Пользователь не является атлетом.',
                url='/admin/people/reconcile/users',
            )
        return build_redirect_with_message(
            error=reconcile_student_error(storage, student_id, error), url=f'/admin/people/{student_id}'
        )

    user = storage.get_user_by_id(user_id)
    username = user['username'] if user else ''
    student = storage.get_student_by_id(student_id)
    full_name = student['full_name'] if student else ''
    log_audit_event(
        request,
        'user_relinked_to_student',
        {'user_id': user_id, 'username': username, 'old_ref': old_ref, 'new_ref': student_id},
    )
    return build_redirect_with_message(
        message=f'Аккаунт {username} перепривязан на студента «{full_name}».',
        url=f'/admin/people/{student_id}',
    )


# --- Импорт студентов из Excel (Student Identity v1, Phase 2.5). ---
#
# Загрузка списка студентов с предпросмотром: файл разбирается в память
# (staging-сессия), карточки создаются ТОЛЬКО явным подтверждением строк —
# предпросмотр не пишет в БД ничего. Кандидаты — только ПРЕДЛОЖЕНИЯ:
# точное совпадение с существующей карточкой не создаёт и не связывает
# ничего автоматически. Импорт трогает ТОЛЬКО students + audit_log:
# записи соревнований, аккаунты и student_ref_id не меняются, справочники
# только читаются (подсказки института/группы), ничего в них не сеется.
# Регистрируется ДО /admin/people/<student_id>, чтобы статические пути
# /admin/people/import и /admin/people/import/template не разбирались
# как id карточки.


def normalize_import_course(value) -> str:
    """Курс из ячейки Excel как строка: 2.0 → «2», NaN/None → «».

    Числовые ячейки pandas отдаёт float'ами («2» в Excel → 2.0); целые
    выводятся без дробной части, остальное (текст, нецелые) — как str().
    """
    if value is None or isna(value):
        return ''
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def normalize_student_sex(value: str) -> str | None:
    """Пол из ячейки Excel в каноническом виде: '' → '', м/м → «М», ж/ж → «Ж».

    upper() в Python работает с кириллицей («м» → «М»); латинские m/M не
    совпадают с кириллическими «М»/«Ж» и отвергаются естественно. Всё
    остальное — None (ошибка строки, а не молчаливая правка).
    """
    if value == '':
        return ''
    if value.upper() == 'М':
        return 'М'
    if value.upper() == 'Ж':
        return 'Ж'
    return None


def parse_student_import_rows(df: pd.DataFrame) -> tuple[list[dict], list[str]]:
    """Разобранные строки файла импорта студентов + неизвестные колонки.

    Возвращает строки staging: row_number = index + 2 (заголовок — строка 1)
    и статус pending/error. Значения — clean_str (NaN → пустая строка), курс —
    normalize_import_course, пол — normalize_student_sex (регистр «м»/«ж»
    приводится к «М»/«Ж»). Ошибки строки: пустое ФИО, недопустимый Пол
    (ошибка, а не молчаливая правка). Дубли ФИО в файле здесь НЕ помечаются
    — группы пересчитываются при каждом использовании (правки строк меняют
    их состав естественно).
    """
    columns = [str(column) for column in df.columns]
    unknown_columns = [column for column in columns if column not in STUDENT_IMPORT_COLUMNS]
    rows: list[dict] = []
    for index, record in df.iterrows():
        full_name = clean_str(record.get('ФИО'))
        sex = normalize_student_sex(clean_str(record.get('Пол')))
        institute = clean_str(record.get('Институт'))
        group = clean_str(record.get('Группа'))
        course = normalize_import_course(record.get('Курс'))
        error = None
        if not full_name:
            error = 'Пустое ФИО.'
        elif sex is None:
            error = 'Пол должен быть «М» или «Ж».'
        rows.append(
            {
                'row_number': int(index) + 2,
                'full_name': full_name,
                'sex': sex,
                'institute': institute,
                'group': group,
                'course': course,
                'status': 'error' if error else 'pending',
                'error': error,
                'student_id': None,
            }
        )
    return rows, unknown_columns


def student_import_duplicate_rows(rows: Sequence[dict]) -> dict[int, list[int]]:
    """Одинаковые ФИО в файле: номер строки → номера всех строк группы.

    Группировка по casefold(strip(ФИО)); возвращаются только строки из групп
    с БОЛЕЕ одной строкой. Каждому члену группы виден весь список строк.
    """
    by_name: dict[str, list[int]] = {}
    for row in rows:
        name = (row['full_name'] or '').strip()
        if name:
            by_name.setdefault(name.casefold(), []).append(row['row_number'])
    return {row_number: numbers for numbers in by_name.values() if len(numbers) > 1 for row_number in numbers}


def student_file_group_institutes(rows: Sequence[dict]) -> dict[str, list[dict]]:
    """Карта группа → институты из строк файла импорта (по самой сессии).

    Ключ — casefold(strip(группа)); значение — институты этой группы,
    уникальные по casefold, в порядке появления (сохраняется первое
    написание). Строки без группы или института не участвуют; статус строки
    не важен — карта строится из ВСЕХ строк сессии при каждом рендере,
    поэтому правка строки-«донора» меняет подсказки остальных естественно.
    """
    institutes: dict[str, list[dict]] = {}
    for row in rows:
        group = (row['group'] or '').strip()
        institute = (row['institute'] or '').strip()
        if not group or not institute:
            continue
        bucket = institutes.setdefault(group.casefold(), [])
        if not any(item['cf'] == institute.casefold() for item in bucket):
            bucket.append({'cf': institute.casefold(), 'raw': institute})
    return institutes


def canonical_group_in_institute(storage: SQLiteAdapter, institute: str, group: str) -> str:
    """Каноническое написание группы внутри института (уникальность группы —
    по паре институт+группа); без изменений, если пара справочнику неизвестна."""
    institute_row = storage.find_catalog_row('institute', institute)
    if institute_row is None:
        return group
    canonical = storage.find_catalog_canonical('group', group, parent_id=institute_row['id'])
    return canonical if canonical else group


def student_catalog_group_hints(
    storage: SQLiteAdapter,
    hints: dict,
    group: str,
    file_groups: dict[str, list[dict]] | None = None,
) -> dict:
    """Институт не указан, группа указана: институт — из ТЕКУЩЕГО файла
    (file_groups, если группа встречается в нём ровно с одним институтом),
    иначе из справочника (если группа принадлежит ровно одному институту);
    иначе предупреждение без изменения значений.

    Файл важнее справочника: импорт списков одной группы обычно идёт из
    файла факультета. Расхождение со справочником — предупреждение, значение
    остаётся файловым.
    """
    file_entry = (file_groups or {}).get(group.casefold())
    if file_entry is not None:
        if len(file_entry) > 1:
            hints['warnings'].append('Институт не определён: в файле группа относится к разным институтам')
            return hints
        raw_institute = file_entry[0]['raw']
        institute_value = storage.find_catalog_canonical('institute', raw_institute) or raw_institute
        hints['institute'] = institute_value
        hints['institute_autofilled'] = True
        hints['institute_source'] = 'file'
        hints['group'] = canonical_group_in_institute(storage, institute_value, group)
        catalog_owner = storage.find_unique_group_institute(group)
        if catalog_owner and catalog_owner.casefold() != institute_value.casefold():
            hints['warnings'].append('Требует внимания: справочник относит группу к другому институту')
        return hints
    institute_value = storage.find_unique_group_institute(group)
    if not institute_value:
        hints['warnings'].append(
            'Институт не определён: группа отсутствует в справочнике или относится к нескольким институтам'
        )
        return hints
    hints['institute'] = institute_value
    hints['institute_autofilled'] = True
    hints['institute_source'] = 'catalog'
    hints['group'] = canonical_group_in_institute(storage, institute_value, group)
    return hints


def student_catalog_pair_hints(storage: SQLiteAdapter, hints: dict, institute: str, group: str) -> dict:
    """Институт и группа указаны: канонизация обоих по справочнику; группа,
    однозначно принадлежащая ДРУГОМУ институту — предупреждение, значения
    не переписываются."""
    canonical_institute = storage.find_catalog_canonical('institute', institute)
    if canonical_institute:
        hints['institute'] = canonical_institute
        hints['group'] = canonical_group_in_institute(storage, canonical_institute, group)
    owner = storage.find_unique_group_institute(group)
    if owner and owner.lower() != hints['institute'].lower():
        hints['warnings'].append('Требует внимания: группа относится к другому институту')
    return hints


def student_catalog_hints(
    storage: SQLiteAdapter,
    institute: str,
    group: str,
    file_groups: dict[str, list[dict]] | None = None,
) -> dict:
    """Подсказки справочников для строки импорта студентов — ТОЛЬКО ЧТЕНИЕ.

    Возвращает значения (возможно канонизированные по регистру), признак
    автозаполнения института (institute_source: 'file' — из текущего файла,
    'catalog' — из справочника) и предупреждения. Импорт студентов НИЧЕГО
    не пишет в справочники (в отличие от импорта записей соревнований).
    """
    hints = {
        'institute': institute,
        'group': group,
        'institute_autofilled': False,
        'institute_source': None,
        'warnings': [],
    }

    if not institute and group:
        return student_catalog_group_hints(storage, hints, group, file_groups=file_groups)
    if institute and group:
        return student_catalog_pair_hints(storage, hints, institute, group)
    if institute:
        canonical_institute = storage.find_catalog_canonical('institute', institute)
        if canonical_institute:
            hints['institute'] = canonical_institute
    return hints


# Staging предпросмотра живёт в памяти (прецедент login_failures): в БД не
# пишется НИЧЕГО до подтверждения строк. Приложение однопроцессное, поэтому
# словарь процесса — корректное хранилище; рестарт просто теряет сессии
# (безопасно — админ загрузит файл заново). Один файл на админа: новая
# загрузка заменяет прежнюю сессию. ФИО из файла не логируются.
student_import_sessions: dict[str, dict] = {}


def sweep_student_import_sessions() -> None:
    """Удалить просроченные сессии (TTL); вызывается при каждом обращении."""
    now = time.time()
    expired = [
        token
        for token, session in student_import_sessions.items()
        if now - session['created_at'] > STUDENT_IMPORT_SESSION_TTL_SECONDS
    ]
    for token in expired:
        student_import_sessions.pop(token, None)


def get_student_import_session(request: Request, token: str):
    """Сессия предпросмотра для запроса: (session, None) или (None, redirect).

    Неизвестный/просроченный токен и чужая сессия — одно и то же поведение:
    возврат на страницу загрузки с flash о истечении сессии.
    """
    sweep_student_import_sessions()
    session = student_import_sessions.get(token)
    if session is None or session['user_id'] != get_current_user_id(request):
        return None, build_redirect_with_message(
            error='Время сессии предпросмотра истекло. Загрузите файл заново.',
            url='/admin/people/import',
        )
    return session, None


def student_import_preview_url(token: str) -> str:
    return f'/admin/people/import/preview/{token}'


def student_import_row(session: dict, raw_row_number: str):
    """Строка сессии по номеру из URL: (row, None) или (None, ответ 400)."""
    row_number, error = parse_reconcile_int(raw_row_number, 'row number')
    if error is not None:
        return None, error
    for row in session['rows']:
        if row['row_number'] == row_number:
            return row, None
    return None, text(body='Invalid row number', status=400)


def student_import_bulk_rows(session: dict) -> list[dict]:
    """Строки массового создания: pending без дублей ФИО в файле и БЕЗ
    кандидатов на момент последнего рендера предпросмотра (признак
    had_candidates фиксируется контекстом предпросмотра) — ровно группа
    «Новые». Неразобранные строки с кандидатами в батч НЕ входят и
    отмену не вызывают; повторная проверка в bulk-create отменяет всё,
    только если у «чистой» строки совпадения ПОЯВИЛИСЬ с момента рендера."""
    duplicates = student_import_duplicate_rows(session['rows'])
    return [
        row
        for row in session['rows']
        if row['status'] == 'pending' and row['row_number'] not in duplicates and not row.get('had_candidates')
    ]


def student_import_preview_context(storage: SQLiteAdapter, session: dict) -> dict:
    """Контекст предпросмотра. Всё производное (кандидаты, подсказки
    справочников, группы дублей, категории, счётчики) пересчитывается при
    каждом рендере; в сессии хранятся строки файла, решения и фиксируемый
    рендером признак had_candidates (см. student_import_bulk_rows)."""
    duplicates = student_import_duplicate_rows(session['rows'])
    # Карта группа → институты текущего файла строится один раз на рендер —
    # и предпросмотр, и создание используют один и тот же источник данных.
    file_groups = student_file_group_institutes(session['rows'])
    groups: dict[str, list[dict]] = {'new': [], 'review': [], 'error': [], 'resolved': []}
    resolved_counts = {'created': 0, 'reused_existing': 0, 'skipped': 0}
    pending = 0
    for row in session['rows']:
        status = row['status']
        if status == 'pending':
            pending += 1
        elif status in resolved_counts:
            resolved_counts[status] += 1
        view = {**row, 'candidates': [], 'hints': None, 'duplicate_rows': duplicates.get(row['row_number'])}
        if status == 'pending':
            view['candidates'] = storage.find_student_candidates(row['full_name'])
            view['hints'] = student_catalog_hints(storage, row['institute'], row['group'], file_groups)
            # Классификация рендера фиксируется в строке сессии: массовое
            # создание (bulk-create) берёт только «чистые» строки (без
            # кандидатов на момент предпросмотра) и отменяется целиком,
            # только если у такой строки совпадения появились ПОЗЖЕ. После
            # правки строки следующий рендер пересчитает признак естественно.
            row['had_candidates'] = bool(view['candidates'])
            category = 'review' if (view['candidates'] or view['duplicate_rows']) else 'new'
        elif status == 'error':
            category = 'error'
        else:
            category = 'resolved'
        groups[category].append(view)
    return {
        'token': session['token'],
        'filename': session['filename'],
        'unknown_columns': session['unknown_columns'],
        'groups': groups,
        'total': len(session['rows']),
        'pending': pending,
        'new_count': len(groups['new']),
        'review_count': len(groups['review']),
        'error_count': len(groups['error']),
        'resolved_count': len(groups['resolved']),
        'created_count': resolved_counts['created'],
        'reused_count': resolved_counts['reused_existing'],
        'skipped_count': resolved_counts['skipped'],
    }


def student_import_counters(session: dict) -> dict[str, int]:
    """Итоги сессии для аудита/сообщений: created/reused_existing/skipped/
    errors/total (по статусам строк)."""
    counters = {'created': 0, 'reused_existing': 0, 'skipped': 0, 'errors': 0}
    for row in session['rows']:
        if row['status'] == 'error':
            counters['errors'] += 1
        elif row['status'] in counters:
            counters[row['status']] += 1
    counters['total'] = len(session['rows'])
    return counters


def create_imported_student(
    request: Request,
    storage: SQLiteAdapter,
    row: dict,
    file_groups: dict[str, list[dict]] | None = None,
) -> int:
    """Создать карточку по строке импорта с подсказками справочников.

    Создаётся ровно то, что показано в предпросмотре: институт/группа —
    после канонизации и автозаполнения по группе (student_catalog_hints,
    с той же картой группа→институты текущего файла — file_groups).
    """
    hints = student_catalog_hints(storage, row['institute'], row['group'], file_groups)
    student_id = storage.create_student(
        row['full_name'],
        row['sex'],
        hints['institute'],
        hints['group'],
        row['course'],
    )
    row['status'] = 'created'
    row['student_id'] = student_id
    log_audit_event(
        request,
        'student_created',
        {'student_id': student_id, 'full_name': row['full_name'], 'source': 'excel-import'},
    )
    return student_id


@app.get('/admin/people/import')
async def admin_student_import_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    return await render(
        template_name=jinja_env.get_template('admin_people_import.html'),
        context={
            'request': request,
            **get_flash_args(request),
        },
    )


@app.get('/admin/people/import/template')
async def export_student_import_template(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    df = pd.DataFrame(columns=list(STUDENT_IMPORT_COLUMNS))
    buffer = BytesIO()
    await asyncio.to_thread(df.to_excel, buffer, index=False)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="Шаблон_студенты_{now_str}.xlsx"',
        },
    )


def student_import_upload_error(df: pd.DataFrame) -> str | None:
    """Ошибки файла на уровне колонок/объёма; None — файл пригоден."""
    if 'ФИО' not in {str(column) for column in df.columns}:
        return 'В файле нет обязательной колонки «ФИО». Скачайте шаблон и заполните его.'
    if len(df) > STUDENT_IMPORT_MAX_ROWS:
        return f'В файле больше {STUDENT_IMPORT_MAX_ROWS} строк. Разбейте список на части.'
    if df.empty:
        return 'В файле нет строк с данными.'
    return None


def replace_student_import_session(user_id: int, token: str, session: dict) -> None:
    """Сохранить новую сессию предпросмотра; прежняя сессия того же админа
    заменяется (один файл на админа)."""
    sweep_student_import_sessions()
    for existing_token, existing in list(student_import_sessions.items()):
        if existing['user_id'] == user_id:
            student_import_sessions.pop(existing_token, None)
    student_import_sessions[token] = session


@app.post('/admin/people/import')
async def admin_student_import_upload(request: Request):
    """Загрузка файла: разбор → staging в памяти → предпросмотр.

    Ошибки файла — возврат на страницу загрузки с flash; в БД не пишется
    ничего.
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    upload_file = request.files.get('file')
    if upload_file is None or not upload_file.body:
        return build_redirect_with_message(error='Выберите файл.', url='/admin/people/import')
    if not (upload_file.name or '').lower().endswith('.xlsx'):
        return build_redirect_with_message(
            error='Файл должен быть в формате .xlsx.',
            url='/admin/people/import',
        )

    try:
        df = await asyncio.to_thread(pd.read_excel, io=upload_file.body)
    except Exception:
        return build_redirect_with_message(
            error='Не удалось прочитать файл. Проверьте, что это не повреждённый .xlsx.',
            url='/admin/people/import',
        )

    upload_error = student_import_upload_error(df)
    if upload_error is not None:
        return build_redirect_with_message(error=upload_error, url='/admin/people/import')

    rows, unknown_columns = await asyncio.to_thread(parse_student_import_rows, df)

    token = secrets.token_urlsafe(16)
    user_id = get_current_user_id(request)
    replace_student_import_session(
        user_id,
        token,
        {
            'token': token,
            'user_id': user_id,
            'created_at': time.time(),
            'filename': upload_file.name or '',
            'unknown_columns': unknown_columns,
            'rows': rows,
        },
    )
    return redirect(student_import_preview_url(token))


@app.get('/admin/people/import/preview/<token>')
async def admin_student_import_preview(request: Request, token: str):
    """Предпросмотр: НОЛЬ записей в БД — только чтение кандидатов/справочников."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error

    context = student_import_preview_context(get_storage(request.app), session)
    return await render(
        template_name=jinja_env.get_template('admin_people_import_preview.html'),
        context={
            'request': request,
            **context,
            **get_flash_args(request),
        },
    )


@app.post('/admin/people/import/preview/<token>/bulk-create')
async def admin_student_import_bulk_create(request: Request, token: str):
    """Создать «новых» (pending без кандидатов и дублей на момент рендера)
    одной транзакцией.

    Неразобранные строки с кандидатами в батч не входят — их решает админ
    по отдельности. Предпроверка: кандидаты каждой строки БАТЧА
    пересчитываются на момент клика — если у любой «чистой» строки
    появились совпадения с момента рендера, не создаётся ничего (строки
    перейдут в «Требуют решения» при следующем рендере).
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error

    storage = get_storage(request.app)
    bulk_rows = student_import_bulk_rows(session)
    # Повторная проверка на момент клика: совпадения появились у «чистой»
    # строки (их не было при рендере) — отмена целиком.
    if any(storage.find_student_candidates(row['full_name']) for row in bulk_rows):
        return build_redirect_with_message(
            error='Ничего не создано: у части строк появились совпадения с существующими студентами. '
            'Проверьте раздел «Требуют решение».',
            url=student_import_preview_url(token),
        )

    # Значения — как в предпросмотре: с канонизацией и автозаполнением
    # института по группе (справочники при этом только читаются). Карта
    # группа→институты — та же функция от строк сессии, что и в рендере.
    file_groups = student_file_group_institutes(session['rows'])
    prepared = []
    for row in bulk_rows:
        hints = student_catalog_hints(storage, row['institute'], row['group'], file_groups)
        prepared.append((row['full_name'], row['sex'], hints['institute'], hints['group'], row['course']))
    student_ids = await asyncio.to_thread(storage.create_students, prepared)
    for row, student_id in zip(bulk_rows, student_ids):
        row['status'] = 'created'
        row['student_id'] = student_id
        log_audit_event(
            request,
            'student_created',
            {'student_id': student_id, 'full_name': row['full_name'], 'source': 'excel-import'},
        )
    return build_redirect_with_message(
        message=f'Создано студентов: {len(student_ids)}.',
        url=student_import_preview_url(token),
    )


@app.post('/admin/people/import/preview/<token>/row/<row_number>/create')
async def admin_student_import_row_create(request: Request, token: str, row_number: str):
    """Явное создание одной строки. Разрешено и при 100% совпадении —
    полные тёзки бывают; переспрашивать кандидатов не нужно (кнопку нажали
    осознанно)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error
    row, row_error = student_import_row(session, row_number)
    if row_error is not None:
        return row_error
    if row['status'] == 'error':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} содержит ошибку — исправьте её.',
            url=student_import_preview_url(token),
        )
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=student_import_preview_url(token),
        )

    storage = get_storage(request.app)
    create_imported_student(request, storage, row, student_file_group_institutes(session['rows']))
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: студент «{row["full_name"]}» создан.',
        url=student_import_preview_url(token),
    )


@app.post('/admin/people/import/preview/<token>/row/<row_number>/use-existing')
async def admin_student_import_row_use_existing(request: Request, token: str, row_number: str):
    """Использовать существующую карточку: пометка ТОЛЬКО в сессии предпросмотра.

    Ничего не создаётся и не связывается — записи и аккаунты не трогаются
    (связи — отдельный ручной workflow сопоставления, Phase 2).
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error
    row, row_error = student_import_row(session, row_number)
    if row_error is not None:
        return row_error
    if row['status'] == 'error':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} содержит ошибку — исправьте её.',
            url=student_import_preview_url(token),
        )
    if row['status'] != 'pending':
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=student_import_preview_url(token),
        )

    student_id, id_error = parse_reconcile_int(get_form_value(request, 'student_id'), 'student id')
    if id_error is not None:
        return id_error
    storage = get_storage(request.app)
    if storage.get_student_by_id(student_id) is None:
        return build_redirect_with_message(
            error='Студент не найден.',
            url=student_import_preview_url(token),
        )

    row['status'] = 'reused_existing'
    row['student_id'] = student_id
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]}: использован существующий студент #{student_id}.',
        url=student_import_preview_url(token),
    )


@app.post('/admin/people/import/preview/<token>/row/<row_number>/skip')
async def admin_student_import_row_skip(request: Request, token: str, row_number: str):
    """Пропустить строку (в т.ч. ошибочную — как быстрый способ убрать её
    из виду)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error
    row, row_error = student_import_row(session, row_number)
    if row_error is not None:
        return row_error
    if row['status'] not in ('pending', 'error'):
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=student_import_preview_url(token),
        )

    row['status'] = 'skipped'
    row['student_id'] = None
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]} пропущена.',
        url=student_import_preview_url(token),
    )


@app.post('/admin/people/import/preview/<token>/row/<row_number>/edit')
async def admin_student_import_row_edit(request: Request, token: str, row_number: str):
    """Правка строки в сессии. В отличие от очереди конфликтов импорта
    записей, ФИО здесь редактируемо — им можно привести строку к написанию
    существующего студента. Некорректные данные — строка становится
    ошибочной; кандидаты/подсказки пересчитаются при следующем рендере."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error
    row, row_error = student_import_row(session, row_number)
    if row_error is not None:
        return row_error
    if row['status'] not in ('pending', 'error'):
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} уже разобрана.',
            url=student_import_preview_url(token),
        )

    row['full_name'] = get_form_value(request, 'full_name').strip()
    row['sex'] = get_form_value(request, 'sex').strip()
    row['institute'] = get_form_value(request, 'institute').strip()
    row['group'] = get_form_value(request, 'group').strip()
    row['course'] = get_form_value(request, 'course').strip()
    error = None
    if not row['full_name']:
        error = 'Пустое ФИО.'
    elif row['sex'] not in STUDENT_SEX_OPTIONS:
        error = 'Пол должен быть «М» или «Ж».'
    if error is not None:
        row['status'] = 'error'
        row['error'] = error
        return build_redirect_with_message(
            error=f'Строка {row["row_number"]} обновлена, но данные некорректны: {error}',
            url=student_import_preview_url(token),
        )
    row['status'] = 'pending'
    row['error'] = None
    return build_redirect_with_message(
        message=f'Строка {row["row_number"]} обновлена.',
        url=student_import_preview_url(token),
    )


@app.post('/admin/people/import/preview/<token>/finish')
async def admin_student_import_finish(request: Request, token: str):
    """Завершить импорт: одно audit-событие с итогами, сессия удаляется.

    Ошибочные строки завершению НЕ мешают (они учтены как errors) — блокируют
    только неразобранные pending-строки.
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error

    pending = sum(1 for row in session['rows'] if row['status'] == 'pending')
    if pending:
        return build_redirect_with_message(
            error=f'Завершение недоступно: не разобрано строк: {pending}.',
            url=student_import_preview_url(token),
        )

    counters = student_import_counters(session)
    log_audit_event(request, 'student_import_completed', counters)
    student_import_sessions.pop(token, None)
    return build_redirect_with_message(
        message=(
            f'Импорт завершён: создано {counters["created"]}, '
            f'использовано существующих {counters["reused_existing"]}, '
            f'пропущено {counters["skipped"]}.'
        ),
        url='/admin/people',
    )


@app.post('/admin/people/import/preview/<token>/discard')
async def admin_student_import_discard(request: Request, token: str):
    """Отменить импорт: сессия удаляется без аудита. Уже созданные
    подтверждением карточки остаются (создание было явным)."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    session, session_error = get_student_import_session(request, token)
    if session_error is not None:
        return session_error

    created = sum(1 for row in session['rows'] if row['status'] == 'created')
    student_import_sessions.pop(token, None)
    if created:
        return build_redirect_with_message(
            message=f'Импорт отменён. Созданные карточки ({created}) сохранены.',
            url='/admin/people',
        )
    return build_redirect_with_message(message='Импорт отменён.', url='/admin/people')


@app.get('/admin/people/<student_id>')
async def admin_person_card(request: Request, student_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    student = storage.get_student_by_id(numeric_id)
    if student is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    linked_records = storage.list_linked_records(numeric_id, limit=50)
    for record in linked_records:
        # Даты — datetime для format_date_range в шаблоне.
        record['date'] = datetime.fromisoformat(record['date']) if record['date'] else None
        record['date_to'] = datetime.fromisoformat(record['date_to']) if record['date_to'] else None
    return await render(
        template_name=jinja_env.get_template('admin_person.html'),
        context={
            'request': request,
            'student': student,
            'aliases': storage.list_student_aliases(numeric_id),
            'linked_records_count': storage.linked_records_count(numeric_id),
            'linked_records': linked_records,
            'linked_users': storage.linked_athlete_users(numeric_id),
            'active_students': [item for item in storage.list_students() if item['active']],
            **get_flash_args(request),
        },
    )


@app.post('/admin/people/<student_id>/edit')
async def admin_person_edit(request: Request, student_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    student = storage.get_student_by_id(numeric_id)
    if student is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    form, form_error = parse_student_form(request)
    if form_error is not None:
        return build_redirect_with_message(error=form_error, url=f'/admin/people/{numeric_id}')

    changes = student_field_changes(student, form)
    if not storage.update_student(
        numeric_id,
        form['full_name'],
        form['sex'],
        form['institute'],
        form['group_name'],
        form['course'],
    ):
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    log_audit_event(
        request,
        'student_updated',
        {'student_id': numeric_id, 'full_name': form['full_name'], 'changed': changes},
    )
    return build_redirect_with_message(
        message=f'Данные студента «{form["full_name"]}» сохранены.',
        url=f'/admin/people/{numeric_id}',
    )


@app.post('/admin/people/<student_id>/active')
async def admin_person_toggle_active(request: Request, student_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    student = storage.get_student_by_id(numeric_id)
    if student is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')

    deactivate = bool(student['active'])
    storage.set_student_active(numeric_id, not deactivate)
    if deactivate:
        log_audit_event(request, 'student_deactivated', {'student_id': numeric_id, 'full_name': student['full_name']})
        message = f'Студент «{student["full_name"]}» помечен неактивным.'
    else:
        log_audit_event(request, 'student_activated', {'student_id': numeric_id, 'full_name': student['full_name']})
        message = f'Студент «{student["full_name"]}» снова активен.'
    return build_redirect_with_message(message=message, url=f'/admin/people/{numeric_id}')


@app.post('/admin/people/<student_id>/alias')
async def admin_person_add_alias(request: Request, student_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error

    name = get_form_value(request, 'name').strip()
    if not name:
        return build_redirect_with_message(error='Укажите псевдоним ФИО.', url=f'/admin/people/{numeric_id}')

    storage = get_storage(request.app)
    if storage.get_student_by_id(numeric_id) is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    if not storage.add_student_alias(numeric_id, name):
        return build_redirect_with_message(
            error=f'Псевдоним «{name}» уже есть у этого студента.',
            url=f'/admin/people/{numeric_id}',
        )
    log_audit_event(request, 'student_alias_added', {'student_id': numeric_id, 'name': name})
    return build_redirect_with_message(
        message=f'Псевдоним «{name}» добавлен.',
        url=f'/admin/people/{numeric_id}',
    )


@app.post('/admin/people/<student_id>/alias/<alias_id>/delete')
async def admin_person_remove_alias(request: Request, student_id: str, alias_id: str):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error
    try:
        numeric_alias_id = int(alias_id)
    except ValueError:
        return text(body='Invalid alias id', status=400)

    storage = get_storage(request.app)
    if storage.get_student_by_id(numeric_id) is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    # Имя псевдонима нужно для сообщения и аудита — берём до удаления.
    alias = next(
        (item for item in storage.list_student_aliases(numeric_id) if item['id'] == numeric_alias_id),
        None,
    )
    if alias is None:
        return build_redirect_with_message(error='Псевдоним не найден.', url=f'/admin/people/{numeric_id}')

    storage.remove_student_alias(numeric_id, numeric_alias_id)
    log_audit_event(
        request,
        'student_alias_removed',
        {'student_id': numeric_id, 'alias_id': numeric_alias_id, 'name': alias['name']},
    )
    return build_redirect_with_message(
        message=f'Псевдоним «{alias["name"]}» удалён.',
        url=f'/admin/people/{numeric_id}',
    )


@app.post('/admin/people/<student_id>/delete')
async def admin_person_delete(request: Request, student_id: str):
    """Полное удаление ОШИБОЧНО созданной карточки (hard delete, admin).

    Возможно только для карточки без связей: storage блокирует удаление
    при привязанных записях/аккаунтах и слитых карточках, не меняя ни
    одной строки. Записи соревнований вместе с карточкой не удаляются.
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    numeric_id, id_error = resolve_student_id(request, student_id)
    if id_error is not None:
        return id_error

    storage = get_storage(request.app)
    # Снимок до удаления: ФИО нужно для сообщения и аудита после того,
    # как строка исчезла (паттерн admin_person_remove_alias).
    student = storage.get_student_by_id(numeric_id)
    if student is None:
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')

    code, blockers = storage.delete_student(numeric_id)
    if code == 'not_found':
        return build_redirect_with_message(error='Студент не найден.', url='/admin/people')
    if code == 'blocked':
        # Отказ не пишется в аудит (конвенция остальных проверок ролей/связей).
        facts = []
        if blockers.get('records'):
            facts.append(f'записей — {blockers["records"]}')
        if blockers.get('athlete_users'):
            facts.append(f'аккаунтов атлета — {blockers["athlete_users"]}')
        if blockers.get('merged_children'):
            facts.append(f'объединённых карточек — {blockers["merged_children"]}')
        return build_redirect_with_message(
            error=(
                f'Удалить карточку нельзя: к ней привязано {", ".join(facts)}. '
                'Отвязайте их на этой странице или деактивируйте карточку.'
            ),
            url=f'/admin/people/{numeric_id}',
        )

    log_audit_event(
        request,
        'student_deleted',
        {'student_id': numeric_id, 'full_name': student['full_name']},
    )
    return build_redirect_with_message(
        message=f'Карточка «{student["full_name"]}» (#{numeric_id}) удалена.',
        url='/admin/people',
    )


@app.post('/admin/students/merge')
async def merge_students(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    from_name = get_form_value(request, 'from_name').strip()
    to_name = get_form_value(request, 'to_name').strip()
    confirm = checkbox_to_bool(get_form_value(request, 'confirm'))
    rewrite_names = checkbox_to_bool(get_form_value(request, 'rewrite_names'))
    if not from_name or not to_name:
        return text(body='Укажите оба ФИО', status=400)
    if from_name == to_name:
        return text(body='ФИО совпадают', status=400)

    from_hash = hashlib.sha256(from_name.encode()).hexdigest()
    to_hash = hashlib.sha256(to_name.encode()).hexdigest()
    storage = get_storage(request.app)
    affected = storage.count_records_by_student_hash(from_hash)
    if not confirm:
        return json_response(
            {
                'preview': True,
                'from_name': from_name,
                'to_name': to_name,
                'records_to_merge': affected,
            }
        )

    merged = storage.merge_students(from_hash, to_hash, new_name=to_name if rewrite_names else None)
    aliases_updated = storage.carry_name_aliases(from_name, to_name)
    log_audit_event(
        request,
        'students_merged',
        {
            'from_name': from_name,
            'to_name': to_name,
            'merged': merged,
            'rewrite_names': rewrite_names,
            'aliases_updated': aliases_updated,
        },
    )
    return json_response(
        {
            'merged': merged,
            'rewritten_names': rewrite_names,
            'aliases_updated': aliases_updated,
        }
    )


def backups_dir() -> Path:
    return Path(settings.data_folder) / 'backups'


def files_dir() -> Path:
    return Path(settings.data_folder) / 'files'


def format_size(num_bytes: int | float) -> str:
    """Человекочитаемый размер: 12,4 ГБ / 5 МБ / 512 Б (запятая как разделитель)."""
    size = float(num_bytes)
    for unit in ('Б', 'КБ', 'МБ', 'ГБ', 'ТБ'):
        if size < 1024 or unit == 'ТБ':
            if unit == 'Б':
                return f'{int(size)} Б'
            rounded = f'{size:.1f}'.replace('.', ',')
            if rounded.endswith(',0'):
                rounded = rounded[:-2]
            return f'{rounded} {unit}'
        size /= 1024
    return f'{int(size)} Б'


def file_size(path: Path) -> int:
    try:
        return path.stat().st_size
    except OSError:
        return 0


def dir_size(path: Path) -> int:
    if not path.is_dir():
        return 0
    return sum(item.stat().st_size for item in path.rglob('*') if item.is_file())


def format_timestamp(timestamp: float) -> str:
    return datetime.fromtimestamp(timestamp).strftime('%d.%m.%Y %H:%M')


def collect_maintenance_stats() -> dict:
    """Панель состояния обслуживания (замечания №10 и №16б из FEEDBACK.md).

    Только чтение: размеры БД (файл + WAL/SHM), вложений, бэкапов, место на
    диске тома с data/, последний бэкап и список архивов очисток.
    """
    data_root = Path(settings.data_folder)
    db_path = Path(settings.database_path)

    db_files = []
    for suffix, label in (('', 'файл БД'), ('-wal', 'WAL'), ('-shm', 'SHM')):
        candidate = db_path.with_name(db_path.name + suffix)
        size = file_size(candidate)
        if size > 0 or not suffix:
            db_files.append({'label': label, 'size': size, 'size_label': format_size(size)})
    db_total = sum(item['size'] for item in db_files)

    attachments_size = dir_size(data_root / 'files')
    backups_size = dir_size(data_root / 'backups')

    disk_root = data_root if data_root.is_dir() else db_path.parent if db_path.parent.is_dir() else Path('.')
    usage = shutil.disk_usage(disk_root)
    disk_percent = round(usage.used / usage.total * 100, 1) if usage.total else 0

    last_backup = None
    pre_wipe_archives = []
    root = data_root / 'backups'
    if root.is_dir():
        backup_files = [item for item in root.rglob('*') if item.is_file()]
        if backup_files:
            latest = max(backup_files, key=lambda item: item.stat().st_mtime)
            stat = latest.stat()
            last_backup = {
                'name': latest.name,
                'subdir': str(latest.parent.relative_to(root)) if latest.parent != root else '',
                'time': format_timestamp(stat.st_mtime),
                'size_label': format_size(stat.st_size),
            }
        pre_wipe_files = sorted(
            (item for item in root.iterdir() if item.is_file() and item.name.startswith('pre-wipe-')),
            key=lambda item: item.stat().st_mtime,
            reverse=True,
        )[:PRE_WIPE_ARCHIVES_SHOWN]
        pre_wipe_archives = [
            {
                'name': item.name,
                'time': format_timestamp(item.stat().st_mtime),
                'size_label': format_size(item.stat().st_size),
            }
            for item in pre_wipe_files
        ]

    return {
        'db_files': db_files,
        'db_total_label': format_size(db_total),
        'attachments_size_label': format_size(attachments_size),
        'backups_size_label': format_size(backups_size),
        'disk': {
            'percent': disk_percent,
            'used_label': format_size(usage.used),
            'total_label': format_size(usage.total),
            'free_label': format_size(usage.free),
            'caption': f'Диск: {format_size(usage.used)} из {format_size(usage.total)} занято',
        },
        'last_backup': last_backup,
        'pre_wipe_archives': pre_wipe_archives,
    }


@app.get('/admin/maintenance')
async def admin_maintenance_page(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error
    storage = get_storage(request.app)
    stats = await asyncio.to_thread(collect_maintenance_stats)
    return await render(
        template_name=jinja_env.get_template('admin_maintenance.html'),
        context={
            'request': request,
            'wipe_confirm_phrase': WIPE_CONFIRM_PHRASE,
            'records_count': storage.count_competitions(),
            'attachments_count': storage.count_attachments(),
            'stats': stats,
            **get_flash_args(request),
        },
    )


def sanitize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    for column in df.columns:
        if df[column].dtype == object:
            df[column] = df[column].map(sanitize_spreadsheet_value)
    return df


def build_users_dataframe(users: Sequence[dict]) -> pd.DataFrame:
    rows = [
        {
            'Логин': user['username'],
            'Роль': user['role'],
            'Активен': 'да' if user.get('active') else 'нет',
            'Псевдонимы ФИО': ', '.join(user.get('name_aliases') or []),
        }
        for user in users
    ]
    return sanitize_dataframe(pd.DataFrame.from_records(rows, columns=['Логин', 'Роль', 'Активен', 'Псевдонимы ФИО']))


def build_levels_dataframe(levels: Sequence[dict]) -> pd.DataFrame:
    rows = [{'Уровень': level['name'], 'Статус': 'активен' if level.get('active') else 'скрыт'} for level in levels]
    return sanitize_dataframe(pd.DataFrame.from_records(rows, columns=['Уровень', 'Статус']))


def build_sports_dataframe(sport_names: Sequence[str]) -> pd.DataFrame:
    return sanitize_dataframe(
        pd.DataFrame.from_records([{'Вид спорта': name} for name in sport_names], columns=['Вид спорта'])
    )


def build_database_export_frames(storage: SQLiteAdapter) -> dict[str, pd.DataFrame]:
    """All database sheets for the maintenance export (one xlsx, one sheet per entity)."""
    competitions = storage.get_competitions()
    export_custom_fields = [field for field in storage.get_custom_fields() if field.show_in_export]
    return {
        'Записи': build_index_dataframe(competitions, export_custom_fields),
        'Пользователи': build_users_dataframe(storage.list_users()),
        'Уровни': build_levels_dataframe(storage.list_levels()),
        'Виды спорта': build_sports_dataframe(storage.get_sport_names()),
    }


def write_database_workbook(frames: dict[str, pd.DataFrame], buffer: BytesIO) -> None:
    with pd.ExcelWriter(buffer) as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)


@app.get('/admin/maintenance/export')
async def export_database(request: Request):
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    frames = await asyncio.to_thread(build_database_export_frames, storage)
    buffer = BytesIO()
    await asyncio.to_thread(write_database_workbook, frames, buffer)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    filename = f'База_данных_{now_str}.xlsx'
    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="{filename}"',
        },
    )


def remove_attachment_files() -> None:
    """Remove uploaded attachment files only (data/files), nothing else inside data/."""
    root = files_dir()
    if root.is_dir():
        shutil.rmtree(root, ignore_errors=True)


def remove_attachment_record_dirs(record_ids: Sequence[int]) -> None:
    """Удалить каталоги вложений указанных записей (очистка по дате)."""
    root = files_dir()
    for record_id in record_ids:
        shutil.rmtree(root / str(record_id), ignore_errors=True)


def attachment_source_path(attachment: dict) -> Path:
    return files_dir() / str(attachment['record_id']) / attachment['stored_name']


def attachments_files_bytes(attachments: Sequence[dict]) -> int:
    """Сумма размеров файлов вложений на диске (что реально освободится, №18)."""
    return sum(file_size(attachment_source_path(attachment)) for attachment in attachments)


def parse_wipe_request(request: Request) -> tuple[dict | None, str | None]:
    """Валидация параметров очистки: режим (полная/по дате), объём, дата, галочка вложений."""
    mode = get_form_value(request, 'mode').strip() or 'scope'
    if mode not in WIPE_MODES:
        return None, 'Некорректный режим очистки'
    params: dict = {'mode': mode, 'scope': None, 'date_before': None, 'with_attachments': True}
    if mode == 'scope':
        scope = get_form_value(request, 'scope').strip()
        if scope not in WIPE_SCOPES:
            return None, 'Некорректный объём очистки'
        params['scope'] = scope
    else:
        raw_date = get_form_value(request, 'date_before').strip()
        try:
            params['date_before'] = datetime.strptime(raw_date, settings.date_format)
        except ValueError:
            return None, 'Некорректная дата очистки (ожидается дд.мм.гггг)'
        # Галочка «вместе с вложениями» включена по умолчанию: нет поля —
        # считаем включённой; скрытый маркер 0 + чекбокс 1 различают снятие.
        raw_flags = request.form.getlist('with_attachments')
        params['with_attachments'] = '1' in raw_flags or not raw_flags
    return params, None


def wipe_scope_label(params: dict) -> str:
    return params['scope'] if params['mode'] == 'scope' else 'date'


def collect_wipe_targets(storage: SQLiteAdapter, params: dict) -> tuple[list, list[dict]]:
    """Записи и вложения, попадающие под очистку (для превью и архива)."""
    if params['mode'] == 'scope':
        records = list(storage.get_competitions()) if params['scope'] != 'attachments' else []
        attachments = storage.get_attachments() if params['scope'] != 'records' else []
    else:
        records = storage.get_competitions_before(params['date_before'])
        record_ids = [int(comp.record_id) for comp in records]
        attachments = storage.get_attachments_for_records(record_ids) if params['with_attachments'] else []
    return records, attachments


def perform_wipe(storage: SQLiteAdapter, params: dict, records: Sequence) -> tuple[int, int]:
    """Исполнить очистку; возвращает (удалено записей, удалено вложений)."""
    records_deleted = 0
    attachments_deleted = 0
    if params['mode'] == 'scope':
        if params['scope'] in ('records', 'all'):
            records_deleted = storage.delete_all_competitions()
        if params['scope'] in ('attachments', 'all'):
            attachments_deleted = storage.delete_all_attachments()
            remove_attachment_files()
    else:
        record_ids = [int(comp.record_id) for comp in records]
        records_deleted = storage.delete_competitions_before(params['date_before'])
        if params['with_attachments']:
            attachments_deleted = storage.delete_attachments_for_records(record_ids)
            remove_attachment_record_dirs(record_ids)
    return records_deleted, attachments_deleted


def create_pre_wipe_archive(
    storage: SQLiteAdapter,
    *,
    scope_label: str,
    records: Sequence,
    attachments: Sequence[dict],
) -> dict:
    """Архив удаляемого перед очисткой (№12): xlsx записей + zip их вложений.

    В data/backups: pre-wipe-<дата>-<scope>.xlsx (когда удаляются записи)
    и pre-wipe-<дата>-<scope>.zip (когда удаляются вложения с файлами).
    """
    root = backups_dir()
    root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime('%Y%m%d-%H%M%S')
    result = {
        'scope': scope_label,
        'xlsx': None,
        'zip': None,
        'records': len(records),
        'attachments': len(attachments),
    }

    if records:
        export_custom_fields = [field for field in storage.get_custom_fields() if field.show_in_export]
        df = build_index_dataframe(records, export_custom_fields)
        xlsx_path = root / f'pre-wipe-{stamp}-{scope_label}.xlsx'
        df.to_excel(xlsx_path, index=False)
        result['xlsx'] = xlsx_path.name

    if attachments:
        zip_path = root / f'pre-wipe-{stamp}-{scope_label}.zip'
        written = False
        seen_names: set[str] = set()
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for attachment in attachments:
                source = attachment_source_path(attachment)
                if not source.is_file():
                    continue
                # Внутри архива — понятное имя файла; при коллизии (одноимённые
                # вложения у одной записи) откатываемся на уникальное stored_name.
                arc_name = f'{attachment["record_id"]}/{attachment["filename"]}'
                if arc_name in seen_names:
                    arc_name = f'{attachment["record_id"]}/{attachment["stored_name"]}'
                seen_names.add(arc_name)
                archive.write(source, arc_name)
                written = True
        if written:
            result['zip'] = zip_path.name
        else:
            zip_path.unlink(missing_ok=True)
    return result


def build_wipe_preview(storage: SQLiteAdapter, params: dict) -> dict:
    """Превью очистки: сколько записей/вложений уйдёт и сколько места освободится (№18)."""
    records, attachments = collect_wipe_targets(storage, params)
    attachments_bytes = attachments_files_bytes(attachments)
    return {
        'records': len(records),
        'attachments': len(attachments),
        'attachments_size_bytes': attachments_bytes,
        'attachments_size': format_size(attachments_bytes),
    }


def log_wipe_execution(
    request: Request,
    params: dict,
    archive: dict,
    records_deleted: int,
    attachments_deleted: int,
    freed_total: int,
) -> str:
    """Аудит выполненной очистки (№16а) и итоговое flash-сообщение."""
    log_audit_event(
        request,
        'pre_wipe_archive_created',
        {
            'scope': archive['scope'],
            'archive_xlsx': archive['xlsx'],
            'archive_zip': archive['zip'],
            'records': archive['records'],
            'attachments': archive['attachments'],
        },
    )
    wipe_details = {
        'scope': archive['scope'],
        'records_deleted': records_deleted,
        'attachments_deleted': attachments_deleted,
        'freed_bytes': freed_total,
        'archive_xlsx': archive['xlsx'],
        'archive_zip': archive['zip'],
    }
    if params['mode'] == 'date':
        wipe_details['date_before'] = params['date_before'].strftime(settings.date_format)
        wipe_details['with_attachments'] = params['with_attachments']
    log_audit_event(request, 'db_wiped', wipe_details)

    summary = (
        f'Очистка выполнена: удалено записей {records_deleted}, вложений {attachments_deleted}. '
        f'Освобождено: {format_size(freed_total)}'
    )
    archive_names = ', '.join(name for name in (archive['xlsx'], archive['zip']) if name)
    if archive_names:
        summary += f'. Архив: {archive_names}'
    return summary


@app.post('/admin/maintenance/wipe')
async def wipe_database(request: Request):
    # Очистка в два шага: превью (POST без confirm=1 → JSON с числами) и
    # выполнение с фразой «УДАЛИТЬ» (docs/data-model-decisions.md).
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    storage = get_storage(request.app)
    params, error = parse_wipe_request(request)
    if error is not None:
        return text(body=error, status=400)

    if get_form_value(request, 'confirm') != '1':
        return json_response(await asyncio.to_thread(build_wipe_preview, storage, params))

    confirm_phrase = get_form_value(request, 'confirm_phrase').strip()
    if confirm_phrase != WIPE_CONFIRM_PHRASE:
        return text(body=f'Для подтверждения введите слово «{WIPE_CONFIRM_PHRASE}»', status=400)

    db_size_before = file_size(Path(settings.database_path))
    records, attachments = await asyncio.to_thread(collect_wipe_targets, storage, params)
    freed_files_bytes = attachments_files_bytes(attachments)

    # Сначала архив — страховка от «очистил и пожалел»; не создался — не чистим.
    try:
        archive = await asyncio.to_thread(
            create_pre_wipe_archive,
            storage,
            scope_label=wipe_scope_label(params),
            records=records,
            attachments=attachments,
        )
    except Exception:
        logger.exception('Failed to create pre-wipe archive')
        return build_redirect_with_message(
            error='Не удалось создать архив очистки — очистка отменена',
            url='/admin/maintenance',
        )

    records_deleted, attachments_deleted = await asyncio.to_thread(perform_wipe, storage, params, records)
    try:
        # DELETE не сжимает файл SQLite: перестраиваем базу, чтобы место
        # реально освободилось (№18). Только здесь — редкая админская операция.
        await asyncio.to_thread(storage.vacuum)
    except Exception:
        logger.exception('VACUUM after wipe failed')

    freed_db_bytes = max(0, db_size_before - file_size(Path(settings.database_path)))
    freed_total = freed_db_bytes + freed_files_bytes

    summary = log_wipe_execution(request, params, archive, records_deleted, attachments_deleted, freed_total)
    return build_redirect_with_message(message=summary, url='/admin/maintenance')


@app.post('/admin/maintenance/backup')
async def backup_now(request: Request):
    """Ручной бэкап (№10): та же логика, что у таймерного scripts/backup_sqlite.py."""
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    result = await asyncio.to_thread(
        run_backup,
        Path(settings.database_path),
        backups_dir(),
        14,
        files_dir(),
    )
    log_audit_event(
        request,
        'backup_created',
        {
            'source': 'manual',
            'archive': result['archive'].name,
            'archive_size': result['archive'].stat().st_size,
            'files_archive': result['files_archive'].name if result['files_archive'] else None,
        },
    )
    message = f'Бэкап создан: {result["archive"].name}'
    if result['files_archive'] is not None:
        message += f', вложения: {result["files_archive"].name}'
    return build_redirect_with_message(message=message, url='/admin/maintenance')


BACKUP_DOWNLOAD_NAME = re.compile(r'\A[A-Za-z0-9][A-Za-z0-9._-]*\Z')
BACKUP_CONTENT_TYPES = {
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    '.zip': 'application/zip',
    '.gz': 'application/gzip',
}


@app.get('/admin/maintenance/backups/<filename>')
async def download_backup(request: Request, filename: str):
    """Скачать файл из data/backups (архивы очисток и бэкапы).

    Имя строго валидируется (без путей и спецсимволов) + контроль после
    resolve() — защита от path-traversal.
    """
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    if not BACKUP_DOWNLOAD_NAME.fullmatch(filename):
        return text(body='Invalid backup filename', status=400)

    root = backups_dir().resolve()
    target = (root / filename).resolve()
    if target.parent != root or not target.is_file():
        return text(body='Not Found', status=404)

    content_type = BACKUP_CONTENT_TYPES.get(target.suffix.lower(), 'application/octet-stream')
    return raw(
        await asyncio.to_thread(target.read_bytes),
        headers={
            'content-type': content_type,
            'content-disposition': f'attachment; filename="{filename}"',
        },
    )


def parse_audit_filters(request: Request) -> tuple[dict, dict, str | None]:
    """Фильтры журнала из GET-параметров.

    Возвращает (storage-фильтры с ISO-датами, сырые значения для формы/ссылок,
    текст ошибки или None).
    """
    args = dict(request.args)
    user_filter = (get_param(args, 'user') or '').strip()
    action_filter = (get_param(args, 'action') or '').strip()
    date_from_raw = (get_param(args, 'date_from') or '').strip()
    date_to_raw = (get_param(args, 'date_to') or '').strip()
    raw_values = {
        'user': user_filter,
        'action': action_filter,
        'date_from': date_from_raw,
        'date_to': date_to_raw,
    }

    iso_dates = {}
    for key in ('date_from', 'date_to'):
        raw_value = raw_values[key]
        if not raw_value:
            iso_dates[key] = None
            continue
        try:
            iso_dates[key] = datetime.strptime(raw_value, settings.date_format).date().isoformat()
        except ValueError:
            return {}, raw_values, f'Некорректная дата фильтра: {raw_value}'

    filters = {
        'username': user_filter or None,
        'action': action_filter or None,
        'date_from': iso_dates['date_from'],
        'date_to': iso_dates['date_to'],
    }
    return filters, raw_values, None


@app.get('/admin/audit')
async def audit_page(request: Request):
    # Журнал с фильтрами (пользователь/действие/период) и пагинацией —
    # «стрелки + N из M» (замечание №14 из FEEDBACK.md).
    auth_error = require_admin(request)
    if auth_error is not None:
        return auth_error

    filters, raw_values, error = parse_audit_filters(request)
    if error is not None:
        return text(body=error, status=400)

    try:
        page = max(1, int(get_param(dict(request.args), 'page') or 1))
    except ValueError:
        page = 1

    events: list[dict] = []
    actions: list[str] = []
    audit_available = True
    total = 0
    try:
        storage = get_storage(request.app)
        total = storage.count_audit_events(**filters)
        pages = max(1, -(-total // AUDIT_PAGE_SIZE))
        page = min(page, pages)
        events = storage.get_audit_events(limit=AUDIT_PAGE_SIZE, offset=(page - 1) * AUDIT_PAGE_SIZE, **filters)
        actions = storage.list_audit_actions()
    except Exception:
        logger.exception('Failed to read audit events')
        audit_available = False
        pages = 1
        page = 1

    query = urlencode({key: value for key, value in raw_values.items() if value})

    def page_url(page_number: int) -> str:
        return f'/admin/audit?{query}&page={page_number}' if query else f'/admin/audit?page={page_number}'

    return await render(
        template_name=jinja_env.get_template('audit.html'),
        context={
            'request': request,
            'events': events,
            'audit_available': audit_available,
            'actions': actions,
            'filters': raw_values,
            'page': page,
            'pages': pages,
            'total': total,
            'page_size': AUDIT_PAGE_SIZE,
            'prev_url': page_url(page - 1) if page > 1 else None,
            'next_url': page_url(page + 1) if page < pages else None,
        },
    )


@app.get('/report')
async def get_report(request: Request):
    # Доступ как раньше: всем ролям, кроме athlete (личный кабинет вместо отчётов).
    if user_is_athlete(request):
        return forbidden(request)
    error = validate_report_filters(dict(request.args)) or collect_custom_filters(request)[1]
    if error:
        return text(body=error, status=400)

    slice_key = parse_report_slice(dict(request.args))
    report_rows = get_report_rows(request, slice_key)
    custom_fields = get_storage(request.app).get_custom_fields()
    rows_count = len(report_rows)
    total_participations = sum(int(row.count_participation) for row in report_rows)
    return await render(
        template_name=jinja_env.get_template('filtered.html'),
        context={
            'request': request,
            'report_rows': report_rows,
            'report_slice': slice_key,
            'report_slice_column': REPORT_SLICE_COLUMNS[slice_key][0],
            'report_slice_sort_type': 'number' if slice_key in REPORT_SLICE_NUMERIC_KEYS else 'text',
            # Композиция отчёта (прототип 06): чипы условий и счётчик строк.
            'report_chips': build_report_chips(dict(request.args), custom_fields),
            'rows_count': rows_count,
            'rows_count_label': ru_plural(rows_count, 'строка', 'строки', 'строк'),
            'total_participations': total_participations,
            'participations_label': ru_plural(total_participations, 'участие', 'участия', 'участий'),
        },
    )


def sanitize_spreadsheet_value(value):
    if isinstance(value, str) and value[:1] in {'=', '+', '-', '@'}:
        return f"'{value}"
    return value


# Чипы применённых условий отчёта (прототип 06): активный фильтр — чип,
# сброс одного условия — перезагрузка отчёта без его GET-параметра.
REPORT_CHIP_LABELS: Sequence[tuple[str, str]] = (
    ('name', 'ФИО'),
    ('date_from', 'Дата (от)'),
    ('date_to', 'Дата (до)'),
    ('position', 'Место'),
    ('level', 'Уровень'),
    ('institute', 'Институт'),
    ('group', 'Группа'),
    ('sport', 'Вид спорта'),
    ('group_by', 'Срез'),
)
REPORT_POSITION_CHIP_LABELS = {'<2': 'Победа', '<4': 'Призовое место', '>3': 'Не призовое место'}


def build_report_chips(args: dict, custom_fields: Sequence[CustomField]) -> list[dict]:
    """Чипы применённых условий отчёта для страницы /report.

    Каждый активный GET-параметр — отдельный чип c data-remove-key; метки
    позиций и среза — человекочитаемые. Кастомные поля — по ярлыку поля.
    """
    slice_labels = {item['key']: item['label'] for item in REPORT_SLICES}
    chips: list[dict] = []
    for key, label in REPORT_CHIP_LABELS:
        raw_value = (get_param(args, key) or '').strip()
        if not raw_value:
            continue
        if key == 'position':
            raw_value = REPORT_POSITION_CHIP_LABELS.get(raw_value, raw_value)
        if key == 'group_by':
            raw_value = slice_labels.get(raw_value, raw_value)
        chips.append({'label': label, 'value': raw_value, 'remove_key': key})
    for field in custom_fields:
        raw_value = (get_param(args, f'custom__{field.key}') or '').strip()
        if raw_value:
            chips.append({'label': field.label, 'value': raw_value, 'remove_key': f'custom__{field.key}'})
    return chips


def parse_report_export_columns(request: Request, slice_key: str) -> tuple[Sequence[str], str | None]:
    """Выбор колонок для выгрузки отчёта (GET-параметр ``columns``).

    Формат — повторённые параметры и/или список через запятую. Набор колонок
    зависит от среза (group_by): «название среза + Участий/Побед/Призовых»,
    для «студента» — как раньше (ФИО/пол/институт/группа/курс + метрики).
    Порядок в файле всегда канонический (как в HTML-отчёте), неизвестные
    имена игнорируются, старое имя «Количество участий» — алиас «Участий».
    Без параметра — полный набор (обратная совместимость; пустой ``?columns=``
    парсер запроса отбрасывает так же, как отсутствие параметра). Параметр
    есть, но валидных колонок не осталось — 400: минимум одна колонка
    обязательна.
    """
    available = report_export_columns(slice_key)
    raw_values = request.args.getlist('columns')
    if not raw_values:
        return available, None
    requested = {
        LEGACY_REPORT_COLUMN_ALIASES.get(part.strip(), part.strip())
        for value in raw_values
        for part in value.split(',')
        if part.strip()
    }
    columns = [column for column in available if column in requested]
    if not columns:
        return [], 'Не выбрано ни одной колонки для выгрузки'
    return columns, None


def build_report_dataframe(
    report_rows: Sequence,
    columns: Sequence[str],
    slice_key: str = DEFAULT_REPORT_SLICE,
) -> pd.DataFrame:
    if slice_key == DEFAULT_REPORT_SLICE:
        records = [info.model_dump(by_alias=True) for info in report_rows]
    else:
        slice_column = REPORT_SLICE_COLUMNS[slice_key][0]
        records = [
            {
                slice_column: row.slice_value,
                'Участий': row.count_participation,
                'Побед': row.count_wins,
                'Призовых': row.count_prizes,
            }
            for row in report_rows
        ]
    df = pd.DataFrame.from_records(records)
    df = df.reindex(columns=columns)
    for column in df.columns:
        if df[column].dtype == object:
            df[column] = df[column].map(sanitize_spreadsheet_value)
    return df


@app.get('/export/report')
async def export_report(request: Request):
    # Доступ как у просмотра отчёта: всем, кроме athlete. Наследует
    # срез/метрики/фильтры применённого отчёта (GET-параметры).
    if user_is_athlete(request):
        return forbidden(request)
    error = validate_report_filters(dict(request.args)) or collect_custom_filters(request)[1]
    if error:
        return text(body=error, status=400)

    slice_key = parse_report_slice(dict(request.args))
    columns, columns_error = parse_report_export_columns(request, slice_key)
    if columns_error:
        return text(body=columns_error, status=400)

    report_rows = get_report_rows(request, slice_key)
    df = await asyncio.to_thread(build_report_dataframe, report_rows, columns, slice_key)

    now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
    filename = f'Отчет_{now_str}.xlsx'
    buffer = BytesIO()
    await asyncio.to_thread(df.to_excel, buffer, index=False)
    buffer.seek(0)

    return raw(
        buffer.getvalue(),
        headers={
            'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'content-disposition': f'attachment; filename="{filename}"',
        },
    )


@app.exception(IsADirectoryError)
async def handle_directory_request(_, exception: IsADirectoryError):
    return text(body='Not Found', status=404)


@app.get('/healthcheck')
def healthcheck(request: Request):
    return text('OK')
