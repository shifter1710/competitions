"""GTO-03: synthetic award-set parity, identity, text UIN and safe navigation."""
import re
from html import unescape
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import parse_qs
from urllib.parse import urlencode
from urllib.parse import urlsplit

import pytest
from openpyxl import load_workbook

from src.auth import create_auth_cookie_value
from src.gto_files import store_certificate
from src.gto_report import build_gto_report
from src.main_test import csrf_for
from src.main_test import get_auth_headers
from src.routes.students import safe_student_return
from src.settings import settings
from src.storage.sqlite import SQLiteAdapter


@pytest.fixture
def report_case(http_client, tmp_path, monkeypatch):
    storage = SQLiteAdapter(str(tmp_path / 'synthetic-gto-report.sqlite3'))
    monkeypatch.setattr(settings, 'data_folder', str(tmp_path / 'synthetic-files'))
    http_client.app.ctx.storage = storage
    for name, role in (
        (settings.auth_admin_username, 'admin'),
        (settings.auth_editor_username, 'editor'),
        (settings.auth_viewer_username or 'viewer', 'viewer'),
        ('synthetic-gto-athlete', 'athlete'),
    ):
        storage.create_user(name, 'synthetic-unused-hash', role)
    first = storage.create_student('Синтетический Тёзка Образец', '', 'Демо А', 'СИН-26', '')
    second = storage.create_student('Синтетический Тёзка Образец', '', 'Демо Б', 'СИН-25', '')
    non_award = storage.create_student('Синтетический Участник Пример', '', '', '', '')
    assert storage.set_student_gto_uin(first, '000123') == (True, None)
    storage.set_student_active(second, False)
    ids = {}
    for status, day in [('bronze', '2026-01-10'), ('silver', '2026-02-10'), ('gold', '2026-03-10')]:
        ids[status], error = storage.create_student_gto_record(first, year=2026, stage=8, status=status, awarded_on=day)
        assert error is None
    ids['other'], _ = storage.create_student_gto_record(second, year=2025, stage=7, status='bronze')
    for status in ('registered', 'participated'):
        storage.create_student_gto_record(non_award, year=2026, stage=8, status=status)
    # Preserved corrupt/unknown legacy facts cannot enter a staff award set.
    storage.connection.execute(
        'INSERT INTO student_gto_records(student_id,year,stage,status,created_at,updated_at) '
        "VALUES (99999,2026,8,'gold','synthetic','synthetic'), (?,2026,8,'unknown','synthetic','synthetic')",
        (first,),
    )
    storage.connection.commit()
    admin = storage.get_user(settings.auth_admin_username)['id']
    for index, status in enumerate(('bronze', 'bronze', 'silver')):
        upload = SimpleNamespace(
            name=f'synthetic-certificate-{index}.pdf',
            body=f'%PDF-1.4\nSynthetic certificate {index}\n'.encode(),
            type='application/pdf',
        )
        fields = store_certificate(ids[status], upload)
        assert storage.create_gto_certificate(first, ids[status], uploaded_by=admin, **fields)
    athlete = storage.get_user('synthetic-gto-athlete')['id']
    storage.link_user(athlete, first)
    yield dict(
        client=http_client,
        storage=storage,
        first=first,
        second=second,
        non_award=non_award,
        ids=ids,
        admin=admin,
        athlete=athlete,
    )
    storage.connection.close()


def html_ids(response):
    return [int(value) for value in re.findall(r'data-gto-award-id="([0-9]+)"', response.text)]


def metric(response, key):
    return int(re.search(r'data-gto-metric="' + key + r'">([0-9]+)', response.text).group(1))


def workbook_rows(response):
    book = load_workbook(BytesIO(response.body), data_only=False)
    sheet = book.active
    header = next(row[0].row for row in sheet.iter_rows() if row[0].value == 'ID награждения')
    return book, [row for row in sheet.iter_rows(min_row=header + 1) if row[0].value is not None]


def outputs(case, query=''):
    suffix = '?' + query if query else ''
    headers = get_auth_headers('viewer')
    html = case['client'].get('/reports/gto' + suffix, headers=headers)[1]
    excel = case['client'].get('/reports/gto/export.xlsx' + suffix, headers=headers)[1]
    return html, excel


def test_same_name_ids_three_awards_and_multiple_certificates_never_multiply_counts(report_case):
    case = report_case
    html, excel = outputs(case)
    assert html.status == excel.status == 200
    ids = case['ids']
    assert html_ids(html) == [ids['bronze'], ids['silver'], ids['gold'], ids['other']]
    assert metric(html, 'students') == 2 and metric(html, 'awards') == 4
    assert [metric(html, key) for key in ('bronze', 'silver', 'gold')] == [2, 1, 1]
    assert f'data-gto-student-id="{case["first"]}"' in html.text
    assert f'data-gto-student-id="{case["second"]}"' in html.text
    assert 'Не прикреплён' in html.text and 'Не указан' in html.text and 'Не указана' in html.text
    assert 'Синтетический Участник Пример' not in html.text
    assert html.text.count('synthetic-certificate-') == 3
    assert 'stored_name' not in html.text and '/files/gto/' not in html.text
    book, rows = workbook_rows(excel)
    assert [row[0].value for row in rows] == html_ids(html)
    assert all(row[3].value == '000123' and row[3].data_type == 's' and row[3].number_format == '@' for row in rows[:3])
    assert rows[-1][3].value is None
    assert len(rows[0][10].value.splitlines()) == 2
    assert all('/gto/' in link for link in rows[0][11].value.splitlines())
    assert all(cell.data_type != 'f' for row in book.active for cell in row)
    selected, selected_excel = outputs(case, 'q=000123')
    assert metric(selected, 'students') == 1 and metric(selected, 'awards') == 3
    assert len(workbook_rows(selected_excel)[1]) == 3
    assert 'no-store' in html.headers['cache-control'] and 'no-store' in excel.headers['cache-control']


@pytest.mark.parametrize(
    'query,keys',
    [
        ('year=2026', ['bronze', 'silver', 'gold']),
        ('year=2025', ['other']),
        ('stage=7', ['other']),
        ('award=silver', ['silver']),
        ('award=bronze', ['bronze', 'other']),
        ('institute=' + urlencode({'x': 'Демо Б'})[2:], ['other']),
        ('group=' + urlencode({'x': 'СИН-26'})[2:], ['bronze', 'silver', 'gold']),
        ('q=000123', ['bronze', 'silver', 'gold']),
        ('q=' + urlencode({'x': 'СИНТЕТИЧЕСКИЙ ТЁЗКА'})[2:], ['bronze', 'silver', 'gold', 'other']),
        (
            urlencode(
                {'year': '2026', 'stage': '8', 'award': 'bronze', 'institute': 'Демо А', 'group': 'СИН-26', 'q': '000'}
            ),
            ['bronze'],
        ),
        ('year=2024', []),
        ('institute=SyntheticUnknown', []),
        ('group=SyntheticMissing', []),
        ('q=not-present', []),
        ('stage=8&stage=8', ['bronze', 'silver', 'gold']),
    ],
)
def test_every_applied_filter_has_exact_html_counter_xlsx_parity(report_case, query, keys):
    html, excel = outputs(report_case, query)
    assert html.status == excel.status == 200
    expected = [report_case['ids'][key] for key in keys]
    assert html_ids(html) == expected
    assert metric(html, 'awards') == len(expected)
    assert sum(metric(html, key) for key in ('bronze', 'silver', 'gold')) == len(expected)
    assert [row[0].value for row in workbook_rows(excel)[1]] == expected
    if not expected:
        assert metric(html, 'students') == 0
        assert 'По выбранным фильтрам награждений нет.' in html.text
    if 'SyntheticUnknown' in query:
        assert 'value="SyntheticUnknown" selected' in html.text
    if 'SyntheticMissing' in query:
        assert 'value="SyntheticMissing" selected' in html.text
    export_url = unescape(
        re.search(r'class="btn btn-outline-primary gto-report-export" href="([^"]+)"', html.text).group(1)
    )
    assert [
        row[0].value for row in workbook_rows(report_case['client'].get(export_url, headers=get_auth_headers())[1])[1]
    ] == expected


@pytest.mark.parametrize(
    'query',
    [
        'year=bad',
        'year=1999',
        'year=2026&year=2025',
        'stage=0',
        'stage=19',
        'stage=8&stage=7',
        'stage=' + '9' * 5000,
        'year=' + '9' * 5000,
        'award=registered',
        'award=diamond',
        'award=gold&award=silver',
        'group=one&group=two',
        'q=one&q=two',
        'q=bad%00value',
        'q=%0a',
        'year=%0a',
        'group_by=student',
    ],
)
def test_invalid_criteria_never_silently_broaden_or_export(report_case, query):
    html, excel = outputs(report_case, query)
    assert html.status == excel.status == 400
    assert 'role="alert"' in html.text
    assert html_ids(html) == []
    assert 'gto-report-export" href=' not in html.text
    assert excel.headers['content-type'].startswith('text/plain')


def test_repeated_awards_missing_date_inactive_history_and_current_group_change(report_case):
    case = report_case
    storage = case['storage']
    first = case['first']
    for year, status, date in [(2026, 'gold', '2026-03-10'), (2024, 'silver', None)]:
        assert storage.create_student_gto_record(first, year=year, stage=8, status=status, awarded_on=date)[1] is None
    before = [
        tuple(row[key] for key in ('id', 'student_id', 'year', 'stage', 'status', 'awarded_on'))
        for row in storage.list_gto_report_awards()
    ]
    html, excel = outputs(case)
    assert metric(html, 'awards') == 6 and metric(html, 'students') == 2
    assert len(workbook_rows(excel)[1]) == 6
    person = storage.get_student_by_id(first)
    storage.update_student(first, person['full_name'], '', 'Демо Новый', 'СИН-НОВАЯ', '')
    after = [
        tuple(row[key] for key in ('id', 'student_id', 'year', 'stage', 'status', 'awarded_on'))
        for row in storage.list_gto_report_awards()
    ]
    assert after == before
    old, old_excel = outputs(case, urlencode({'group': 'СИН-26'}))
    assert metric(old, 'awards') == 0 and workbook_rows(old_excel)[1] == []
    current, current_excel = outputs(case, urlencode({'institute': 'Демо Новый', 'group': 'СИН-НОВАЯ'}))
    assert metric(current, 'awards') == 5 and len(workbook_rows(current_excel)[1]) == 5
    assert storage.get_student_by_id(case['second'])['active'] == 0


@pytest.mark.parametrize('role', ['admin', 'editor', 'viewer'])
def test_staff_report_and_discovery_available_without_mutation_controls(report_case, role):
    case = report_case
    headers = get_auth_headers(role)
    for path in ('/reports/gto', '/reports/gto/export.xlsx'):
        assert case['client'].get(path, headers=headers)[1].status == 200
    discovery = case['client'].get('/reports', headers=headers)[1]
    assert 'href="/reports/gto"' in discovery.text
    page = case['client'].get('/reports/gto', headers=headers)[1]
    assert 'class="gto-report-form row g-3"' in page.text
    assert 'class="report-form"' not in page.text
    assert 'method="post"' not in page.text.split('<div class="gto-report-page">', 1)[1]


def test_athlete_and_current_revoked_staff_denied_report_and_export(report_case):
    case = report_case
    value = create_auth_cookie_value('synthetic-gto-athlete', 'athlete')
    headers = {'cookie': f'{settings.auth_cookie_name}={value}'}
    for path in ('/reports/gto', '/reports/gto/export.xlsx'):
        assert case['client'].get(path, headers=headers)[1].status == 403
    viewer = case['storage'].get_user(settings.auth_viewer_username or 'viewer')['id']
    case['storage'].connection.execute("UPDATE users SET role='athlete' WHERE id=?", (viewer,))
    case['storage'].connection.commit()
    for path in ('/reports/gto', '/reports/gto/export.xlsx'):
        assert case['client'].get(path, headers=get_auth_headers('viewer'))[1].status == 403


def test_literal_uin_formula_safety_for_identity_filename_and_filter_text(report_case):
    case = report_case
    storage = case['storage']
    storage.update_student(case['first'], '=SYNTHETIC()', '', '+Synthetic institute', '@Synthetic group', '')
    assert storage.set_student_gto_uin(case['first'], '=1+1') == (True, None)
    storage.connection.execute("UPDATE student_gto_certificates SET filename='=synthetic.pdf'")
    storage.connection.commit()
    html, excel = outputs(case, urlencode({'q': '=1+1'}))
    assert metric(html, 'awards') == 3
    book, rows = workbook_rows(excel)
    for row in rows:
        assert row[3].value == '=1+1' and row[3].data_type == 's' and row[3].number_format == '@'
        assert row[2].value == "'=SYNTHETIC()"
        assert row[4].value == "'+Synthetic institute" and row[5].value == "'@Synthetic group"
    assert all(cell.data_type != 'f' for row in book.active for cell in row)
    assert rows[0][10].value.startswith("'=synthetic.pdf")


def test_excel_cannot_silently_truncate_long_personal_uin(report_case):
    case = report_case
    value = '0' * 40000
    assert case['storage'].set_student_gto_uin(case['first'], value) == (True, None)
    response = case['client'].get('/reports/gto/export.xlsx', headers=get_auth_headers())[1]
    assert response.status == 400 and 'без потерь' in response.text
    assert case['storage'].get_student_by_id(case['first'])['gto_uin'] == value


def test_report_snapshot_uses_constant_queries_and_does_not_change_markers(report_case):
    storage = report_case['storage']
    statements = []
    storage.connection.set_trace_callback(statements.append)
    try:
        rows = storage.list_gto_report_awards()
    finally:
        storage.connection.set_trace_callback(None)
    assert len(rows) == 4
    assert len([sql for sql in statements if sql.lstrip().startswith('SELECT')]) == 2
    assert not storage.connection.in_transaction
    assert dict(storage.connection.execute('SELECT key,value FROM app_settings')) == {
        'identity_mode': 'dual',
        'gto_history_v1': '1',
    }
    assert build_gto_report(rows, {'q': ['000123']})['metrics'] == {
        'students': 1,
        'awards': 3,
        'bronze': 1,
        'silver': 1,
        'gold': 1,
    }


@pytest.mark.parametrize(
    'value',
    [
        '/reports/gto/extra',
        '/reports/gto/export.xlsx',
        '//synthetic.invalid/path',
        'https://synthetic.invalid/reports/gto',
        '/reports/gto?x=%0a',
        '/reports/gto?x=%255c',
        '/reports/gto\\escape',
        '/reports-other',
    ],
)
def test_exact_report_return_allowlist_rejects_unsafe_paths(value):
    assert safe_student_return(value) is None


def test_native_award_uin_and_certificate_forms_preserve_applied_report_return(report_case):
    case = report_case
    target = '/reports/gto?' + urlencode({'year': '2026', 'group': 'СИН-26', 'q': '000123'})
    assert safe_student_return(target) == target
    headers = get_auth_headers()
    card = case['client'].get(
        f'/admin/people/{case["first"]}?' + urlencode({'tab': 'gto', 'return': target}), headers=headers
    )[1]
    assert 'К отчёту' in card.text
    assert f'name="return" value="{target.replace("&", "&amp;")}"' in card.text
    base = f'/admin/people/{case["first"]}/gto'
    for path, values in [
        (base, {'year': '', 'stage': '8', 'status': 'gold'}),
        (base + '/uin', {'gto_uin': '000123'}),
        (f'{base}/{case["ids"]["gold"]}/edit', {'year': '2026', 'stage': '8', 'status': 'gold'}),
    ]:
        response = case['client'].post(
            path, headers=headers, data={**csrf_for(headers), 'return': target, **values}, allow_redirects=False
        )[1]
        location = urlsplit(response.headers['location'])
        assert not location.fragment and parse_qs(location.query)['tab'] == ['gto']
        assert parse_qs(location.query)['return'] == [target]
        assert 'К отчёту' in case['client'].get(response.headers['location'], headers=headers)[1].text
    response = case['client'].post(
        f'{base}/{case["ids"]["gold"]}/certificates',
        headers=headers,
        data={**csrf_for(headers), 'return': target},
        files={'file': ('synthetic-invalid.pdf', b'invalid', 'application/pdf')},
        allow_redirects=False,
    )[1]
    assert parse_qs(urlsplit(response.headers['location']).query)['return'] == [target]
    malicious = case['client'].post(
        base + '/uin',
        headers=headers,
        data={**csrf_for(headers), 'return': '//synthetic.invalid', 'gto_uin': '000123'},
        allow_redirects=False,
    )[1]
    assert 'return' not in parse_qs(urlsplit(malicious.headers['location']).query)


def test_export_rechecks_current_staff_authority_after_workbook_await(report_case, monkeypatch):
    import src.routes.gto_reports as routes

    original = routes.build_gto_xlsx
    storage = report_case['storage']
    viewer = storage.get_user(settings.auth_viewer_username or 'viewer')['id']

    def build_then_revoke(report):
        result = original(report)
        with storage._lock:
            storage.connection.execute('UPDATE users SET active=0 WHERE id=?', (viewer,))
            storage.connection.commit()
        return result

    monkeypatch.setattr(routes, 'build_gto_xlsx', build_then_revoke)
    response = report_case['client'].get('/reports/gto/export.xlsx', headers=get_auth_headers('viewer'))[1]
    assert response.status == 403
    assert not response.body.startswith(b'PK')
