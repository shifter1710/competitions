import asyncio
import hmac
import json
import re
import sqlite3
import time
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from unittest.mock import ANY
from unittest.mock import Mock
from urllib.parse import quote
from urllib.parse import unquote_plus
from urllib.parse import urlencode

import pandas as pd
import pytest
from sanic_testing.testing import SanicTestClient

from src.auth import hash_password
from src.conftest import FailingStudentsDeleteConnection
from src.conftest import make_report_fixture
from src.conftest import make_report_record
from src.conftest import make_report_unapproved_fixture
from src.main import app
from src.main import build_competition
from src.main import calendar_event_status
from src.main import competition_duplicate_key
from src.main import competition_to_export_row
from src.main import create_auth_cookie_value
from src.main import decorate_calendar_event
from src.main import EVENT_IMPORT_SESSION_TTL_SECONDS
from src.main import event_import_sessions
from src.main import format_date_range
from src.main import group_calendar_events_by_month
from src.main import normalize_position
from src.main import parse_date_value
from src.main import split_import_competitions
from src.main import student_import_sessions
from src.models.competition import Competition
from src.models.custom_field import CustomField
from src.models.http.student_info import StudentInfo
from src.settings import settings
from src.storage.sqlite import SQLiteAdapter


def get_xlsx_headers(content: bytes) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True)
    sheet = workbook.worksheets[0]
    row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(value) for value in row if value is not None]


def get_xlsx_rows(content: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True)
    sheet = workbook.worksheets[0]
    return list(sheet.iter_rows(values_only=True))


def get_auth_headers(role: str = 'admin') -> dict[str, str]:
    username = settings.auth_admin_username
    if role == 'editor':
        username = settings.auth_editor_username
    if role == 'viewer':
        username = settings.auth_viewer_username or 'viewer'
    cookie = create_auth_cookie_value(username=username, role=role)
    return {'cookie': f'{settings.auth_cookie_name}={cookie}'}


def csrf_for(headers: dict[str, str]) -> dict[str, str]:
    cookie = headers['cookie'].split('=', 1)[1]
    token = hmac.new(
        settings.auth_secret_key.encode(),
        b'csrf:' + cookie.encode(),
        sha256,
    ).hexdigest()
    return {'csrf_token': token}


def fake_get_user(username: str) -> dict | None:
    if username == 'sportik':
        return {
            'id': 1,
            'username': 'sportik',
            'password_hash': 'x',
            'role': 'athlete',
            'active': 1,
            'pwd_ver': 0,
        }
    accounts = [
        (settings.auth_admin_username, settings.auth_admin_password, 'admin'),
        (settings.auth_editor_username, settings.auth_editor_password, 'editor'),
        (settings.auth_viewer_username, settings.auth_viewer_password, 'viewer'),
    ]
    for account_username, password, role in accounts:
        if username == account_username and password:
            return {
                'id': 1,
                'username': account_username,
                'password_hash': hash_password(password),
                'role': role,
                'active': 1,
                'pwd_ver': 0,
            }
    return None


@pytest.fixture(scope='module')
def client() -> SanicTestClient:
    fake_storage = Mock()
    fake_storage.get_competitions.return_value = []
    # Серверные фильтры/пагинация главной (прототип 02)
    fake_storage.get_competitions_page.return_value = []
    fake_storage.count_competitions_filtered.return_value = 0
    fake_storage.get_filtered.return_value = []
    fake_storage.get_custom_fields.return_value = []
    fake_storage.save_competitions.return_value = None
    fake_storage.import_competitions.return_value = None
    fake_storage.update_competition.return_value = None
    fake_storage.delete_competition.return_value = None
    fake_storage.create_custom_field.return_value = None
    fake_storage.update_custom_field.return_value = None
    fake_storage.disable_custom_field.return_value = None
    fake_storage.get_competition_review.return_value = None
    fake_storage.get_profile.return_value = {}
    fake_storage.set_profile.return_value = None
    fake_storage.get_name_aliases.return_value = []
    fake_storage.add_name_alias.return_value = None
    fake_storage.count_records_by_student_hash.return_value = 2
    fake_storage.get_student_names.return_value = ['Абрамов Артём Артёмович', 'Иванов Иван Иванович']
    fake_storage.merge_students.return_value = 2
    fake_storage.carry_name_aliases.return_value = 1
    fake_storage.get_attachment.return_value = None
    fake_storage.get_attachments.return_value = []
    fake_storage.create_attachment.return_value = 1
    fake_storage.delete_attachment.return_value = None
    fake_storage.set_competition_review.return_value = None
    fake_storage.get_level_names.return_value = ['внутривузовские', 'межвузовские']
    fake_storage.list_levels.return_value = []
    fake_storage.create_level.return_value = None
    fake_storage.disable_level.return_value = None
    fake_storage.rename_catalog_value.return_value = 0
    fake_storage.list_catalog.side_effect = lambda category: {
        'sport': ['Бег', 'Лыжи'],
        'institute': ['ИСИ'],
    }.get(category, [])
    fake_storage.list_catalog_all.return_value = []
    fake_storage.list_catalog_tree.return_value = []
    # №1.1: канонизирующие подстановки по умолчанию выключены (None).
    fake_storage.find_catalog_row.return_value = None
    fake_storage.find_catalog_canonical.return_value = None
    fake_storage.find_level_canonical.return_value = None
    # №19а: автозаполнение института по группе по умолчанию выключено.
    fake_storage.find_unique_group_institute.return_value = None
    # №23доп: резолвер атлета по умолчанию ничего не знает.
    fake_storage.search_athletes.return_value = []
    # Карточки студентов в автодополнении (объединённый поиск) — по умолчанию пусты.
    fake_storage.search_student_suggestions.return_value = []
    fake_storage.find_athlete_fields.return_value = {}
    fake_storage.get_group_options_by_institute.return_value = {}
    fake_storage.ensure_catalog_pair.return_value = None
    fake_storage.count_child_groups.return_value = 0
    fake_storage.add_catalog_value.return_value = None
    fake_storage.get_catalog_value.return_value = None
    fake_storage.hide_catalog_value.return_value = None
    fake_storage.unhide_catalog_value.return_value = None
    fake_storage.delete_catalog_value.return_value = None
    fake_storage.count_records_using.return_value = 0
    fake_storage.touch_user_seen.return_value = False
    fake_storage.set_user_last_login.return_value = None
    fake_storage.get_user.side_effect = fake_get_user
    fake_storage.get_user_by_id.return_value = None
    fake_storage.list_users.return_value = []
    fake_storage.create_user.return_value = None
    fake_storage.set_user_password.return_value = None
    fake_storage.set_user_active.return_value = None
    fake_storage.count_records_by_owner.return_value = 0
    fake_storage.delete_user.return_value = 0
    fake_storage.count_competitions.return_value = 0
    fake_storage.count_attachments.return_value = 0
    fake_storage.delete_all_competitions.return_value = 0
    fake_storage.delete_all_attachments.return_value = 0
    fake_storage.get_competitions_before.return_value = []
    fake_storage.delete_competitions_before.return_value = 0
    fake_storage.get_attachments_for_records.return_value = []
    fake_storage.delete_attachments_for_records.return_value = 0
    fake_storage.vacuum.return_value = None
    fake_storage.get_sport_names.return_value = []
    fake_storage.add_audit_event.return_value = None
    fake_storage.get_audit_events.return_value = []
    fake_storage.count_audit_events.return_value = 0
    fake_storage.list_audit_actions.return_value = []
    # Лёгкий реестр полей + очередь конфликтов импорта (волна 3)
    fake_storage.get_field_settings.return_value = {}
    fake_storage.update_field_settings.return_value = None
    fake_storage.count_import_queue.return_value = 0
    fake_storage.list_import_queue.return_value = []
    fake_storage.get_import_queue_entry.return_value = None
    fake_storage.add_import_queue_entry.return_value = 1
    fake_storage.set_import_queue_status.return_value = None
    fake_storage.get_competition_by_id.return_value = None
    # Календарь соревнований (волна A, docs/feedback-live.md №23)
    fake_storage.list_calendar_events.return_value = []
    fake_storage.create_calendar_event.return_value = 1
    fake_storage.get_calendar_event.return_value = None
    fake_storage.update_calendar_event.return_value = None
    fake_storage.delete_calendar_event.return_value = None
    fake_storage.count_calendar_event_participants.return_value = 0
    fake_storage.list_calendar_event_participants.return_value = []
    app.ctx.storage = fake_storage
    return SanicTestClient(app)


def test_healthcheck(client: SanicTestClient):
    _, response = client.get('/healthcheck')
    assert response.status == 200


def test_get_report_empty(client: SanicTestClient):
    _, response = client.get('/report', headers=get_auth_headers())
    assert response.status == 200


def test_get_report_date(client: SanicTestClient):
    _, response = client.get('/report?date_from=03.02.2021', headers=get_auth_headers())
    assert response.status == 200

    _, response = client.get('/report?date_from=03.02.2021&date_to=04.02.2021', headers=get_auth_headers())
    assert response.status == 200


def test_get_report_position(client: SanicTestClient):
    _, response = client.get('/report?position=>3', headers=get_auth_headers())
    assert response.status == 200

    _, response = client.get('/report?position=<4', headers=get_auth_headers())
    assert response.status == 200


def test_get_report_level(client: SanicTestClient):
    _, response = client.get('/report?level=внутривузовские', headers=get_auth_headers())
    assert response.status == 200


def test_get_report_name(client: SanicTestClient):
    _, response = client.get('/report?name=Карлова', headers=get_auth_headers())
    assert response.status == 200


def test_clean_db_route_removed(client: SanicTestClient):
    # Старый эндпоинт очистки удалён: защищённая версия живёт на /admin/maintenance
    # (docs/data-model-decisions.md). Проверяем через GET: с валидной сессией он
    # доходит до роутера и возвращает 404. POST к несуществующему пути всегда
    # перехватывается глобальным CSRF-middleware (Sanic не парсит form-body для
    # unmatched-маршрутов → 403/401), поэтому 404 по POST недостижим в принципе.
    headers = get_auth_headers(role='admin')
    _, response = client.get('/clean_db', headers=headers, allow_redirects=False)
    assert response.status == 404

    _, response = client.post('/clean_db', headers=headers, data=csrf_for(headers))
    assert response.status == 403


def test_upload_rejects_missing_columns(client: SanicTestClient):
    df = pd.DataFrame([{'ФИО': 'Тест'}])
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers()
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'broken.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )
    assert response.status == 400


def test_upload_accepts_text_dates_in_app_format(client: SanicTestClient):
    app.ctx.storage.import_competitions.reset_mock()
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Тестов Тест Тестович',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 1,
                'Курс': 2,
            }
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    assert response.status == 200
    assert 'Импортировано записей: 1' in response.text
    assert 'дублей' not in response.text
    app.ctx.storage.import_competitions.assert_called_once()
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.date == datetime(2026, 3, 15)


def test_upload_skips_duplicates(client: SanicTestClient):
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='Тестов Тест Тестович',
            student_sex='М',
            institute='ИСИ',
            group='ПГС-101',
            course=2,
            sport='Бег',
            date=datetime(2026, 3, 15),
            level='внутривузовские',
            name='Кубок',
            position=1,
        )
    ]
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Тестов Тест Тестович',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 1,
                'Курс': 2,
            },
            {
                'ФИО': 'Новый Студент',
                'Пол': 'Ж',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-102',
                'Вид спорта': 'Бег',
                'Дата': '16.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 2,
                'Курс': 1,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )

    assert response.status == 200
    assert 'Импортировано записей: 1' in response.text
    assert 'Пропущено дублей: 1' in response.text
    saved_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert len(saved_rows) == 1
    assert saved_rows[0].student_name == 'Новый Студент'
    app.ctx.storage.get_competitions.return_value = []


def test_upload_error_mentions_row_number(client: SanicTestClient):
    app.ctx.storage.import_competitions.reset_mock()
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Нормальный Студент',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 1,
                'Курс': 2,
            },
            {
                'ФИО': 'Бедный Студент',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Бег',
                'Дата': 'не дата',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 1,
                'Курс': 2,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers()
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )

    assert response.status == 400
    assert 'строка 3' in response.text
    # Атомарность импорта: невалидная строка — ничего не сохраняется вовсе.
    app.ctx.storage.import_competitions.assert_not_called()


def test_editor_can_create_manual_competition(client: SanicTestClient):
    app.ctx.storage.save_competitions.reset_mock()
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Весенний кубок',
            'position': '1',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    assert response.headers['location'] == '/'
    app.ctx.storage.save_competitions.assert_called_once()


def test_manual_competition_create_rejects_invalid_date(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '2026-04-10',
            'level': 'межвузовские',
            'name': 'Весенний кубок',
            'position': '1',
        },
        allow_redirects=False,
    )

    assert response.status == 400


def test_editor_can_update_manual_competition(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/123',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Весенний кубок',
            'position': '1',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    assert response.headers['location'] == '/'
    app.ctx.storage.update_competition.assert_called_once()
    assert app.ctx.storage.update_competition.call_args[0][0] == 123
    # Правка модератора (editor входит в MODERATOR_ROLES) подтверждает запись.
    app.ctx.storage.set_competition_review.assert_called_once_with(123, 'approved')


def test_moderator_update_pending_competition_approves_it(client: SanicTestClient):
    app.ctx.storage.update_competition.reset_mock()
    app.ctx.storage.set_competition_review.reset_mock()
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/7',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Весенний кубок',
            'position': '1',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    app.ctx.storage.set_competition_review.assert_called_once_with(7, 'approved')


def test_athlete_update_approved_competition_returns_to_pending(client: SanicTestClient):
    app.ctx.storage.update_competition.reset_mock()
    app.ctx.storage.set_competition_review.reset_mock()
    app.ctx.storage.get_competition_review.return_value = {
        'id': 5,
        'review_status': 'approved',
        'owner_id': 1,  # sportik (athlete, id=1) владеет записью
    }
    headers = athlete_headers()
    _, response = client.post(
        '/competition/5',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )

    assert response.status == 302
    app.ctx.storage.update_competition.assert_called_once()
    assert app.ctx.storage.update_competition.call_args[0][0] == 5
    app.ctx.storage.set_competition_review.assert_called_once_with(5, 'pending')
    app.ctx.storage.get_competition_review.return_value = None


def test_admin_can_delete_competition(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    # M4: удаление записи удаляет и её вложения — строки в БД (метод
    # storage) и файлы на диске (каталог files/<id>).
    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    record_dir = tmp_path / 'files' / '123'
    record_dir.mkdir(parents=True)
    (record_dir / 'stored.png').write_bytes(b'fake-png-bytes')

    headers = get_auth_headers()
    _, response = client.post(
        '/competition/123/delete',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )

    assert response.status == 302
    assert response.headers['location'] == '/'
    app.ctx.storage.delete_competition_with_attachments.assert_called_once_with(123)
    assert not record_dir.exists()


def test_editor_cannot_delete_competition(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/abc123/delete',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403


def test_export_index_omits_dataframe_index_and_created_at(client: SanicTestClient):
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='Test 1',
            student_sex='M',
            institute='Inst',
            group='A',
            course=1,
            sport='Run',
            date=datetime(2024, 1, 1),
            level='межвузовские',
            name='Meet',
            position=1,
        )
    ]

    _, response = client.get('/export/index', headers=get_auth_headers())

    assert response.status == 200
    assert get_xlsx_headers(response.body) == [
        'ФИО',
        'Пол',
        'Институт',
        'Группа',
        'Вид спорта',
        'Дата',
        'Уровень соревнований',
        'Название соревнований',
        'Место',
        'Курс',
    ]


def test_export_index_includes_custom_fields(client: SanicTestClient):
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='Test 1',
            student_sex='M',
            institute='Inst',
            group='A',
            course=1,
            sport='Run',
            date=datetime(2024, 1, 1),
            level='межвузовские',
            name='Meet',
            position=1,
            extra_data={'trainer': 'Coach'},
        )
    ]
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(field_id=1, key='trainer', label='Тренер', show_in_export=True)
    ]

    _, response = client.get('/export/index', headers=get_auth_headers())

    assert response.status == 200
    assert 'Тренер' in get_xlsx_headers(response.body)


def test_export_report_omits_dataframe_index_and_student_id(client: SanicTestClient):
    app.ctx.storage.get_filtered.return_value = [
        StudentInfo(
            student_id='1',
            student_name='Test 1',
            student_sex='M',
            institute='Inst',
            group='A',
            course=1,
            count_participation=3,
            count_wins=1,
            count_prizes=2,
        )
    ]

    _, response = client.get('/export/report', headers=get_auth_headers())

    assert response.status == 200
    assert get_xlsx_headers(response.body) == [
        'ФИО',
        'Пол',
        'Институт',
        'Группа',
        'Курс',
        'Участий',
        'Побед',
        'Призовых',
    ]


def make_report_infos() -> list[StudentInfo]:
    return [
        StudentInfo(
            student_id='1',
            student_name='Иванов Иван',
            student_sex='М',
            institute='ИСИ',
            group='А-101',
            course=2,
            count_participation=5,
            count_wins=2,
            count_prizes=3,
        ),
        StudentInfo(
            student_id='2',
            student_name='Петров Пётр',
            student_sex='Ж',
            institute='ИМИ',
            group='Б-202',
            course=1,
            count_participation=1,
            count_wins=0,
            count_prizes=0,
        ),
    ]


def test_export_report_returns_rows_in_report_order(client: SanicTestClient):
    app.ctx.storage.get_filtered.return_value = make_report_infos()

    _, response = client.get('/export/report', headers=get_auth_headers())

    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('ФИО', 'Пол', 'Институт', 'Группа', 'Курс', 'Участий', 'Побед', 'Призовых'),
        ('Иванов Иван', 'М', 'ИСИ', 'А-101', 2, 5, 2, 3),
        ('Петров Пётр', 'Ж', 'ИМИ', 'Б-202', 1, 1, 0, 0),
    ]


def test_export_report_with_column_subset(client: SanicTestClient):
    app.ctx.storage.get_filtered.return_value = make_report_infos()

    _, response = client.get(
        '/export/report?' + urlencode({'columns': ['ФИО', 'Участий']}, doseq=True),
        headers=get_auth_headers(),
    )

    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('ФИО', 'Участий'),
        ('Иванов Иван', 5),
        ('Петров Пётр', 1),
    ]


def test_export_report_accepts_legacy_column_name(client: SanicTestClient):
    # Старые закладки с «Количество участий» продолжают работать: имя —
    # алиас «Участий» (замечание №19, обратная совместимость).
    app.ctx.storage.get_filtered.return_value = make_report_infos()

    _, response = client.get(
        '/export/report?' + urlencode({'columns': ['ФИО', 'Количество участий']}, doseq=True),
        headers=get_auth_headers(),
    )

    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('ФИО', 'Участий'),
        ('Иванов Иван', 5),
        ('Петров Пётр', 1),
    ]


def test_export_report_accepts_comma_separated_columns(client: SanicTestClient):
    app.ctx.storage.get_filtered.return_value = make_report_infos()

    _, response = client.get(
        '/export/report?columns=' + quote('Группа,ФИО'),
        headers=get_auth_headers(),
    )

    # порядок колонок всегда канонический, как в HTML-отчёте
    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('ФИО', 'Группа'),
        ('Иванов Иван', 'А-101'),
        ('Петров Пётр', 'Б-202'),
    ]


def test_export_report_requires_at_least_one_column(client: SanicTestClient):
    # значение из пробелов: параметр есть, валидных колонок нет — 400
    _, response = client.get('/export/report?columns=%20', headers=get_auth_headers())
    assert response.status == 400

    _, response = client.get('/export/report?columns=' + quote('Нет такой колонки'), headers=get_auth_headers())
    assert response.status == 400


def test_export_report_applies_report_filters(client: SanicTestClient):
    app.ctx.storage.get_filtered.reset_mock()

    _, response = client.get(
        '/export/report?date_from=01.02.2024&name=Иван&position=<4',
        headers=get_auth_headers(),
    )

    assert response.status == 200
    assert app.ctx.storage.get_filtered.call_args[1] == {
        'date_from': '01.02.2024',
        'date_to': None,
        'position': '<4',
        'level': None,
        'name': 'Иван',
        'institute': '',
        'group': '',
        'sport': '',
        'custom_filters': [],
    }


def test_export_report_available_for_viewer(client: SanicTestClient):
    _, response = client.get(
        '/export/report?' + urlencode({'columns': ['ФИО']}, doseq=True),
        headers=get_auth_headers(role='viewer'),
    )
    assert response.status == 200


def test_reports_page_shows_export_column_panel(client: SanicTestClient):
    _, response = client.get('/reports', headers=get_auth_headers())

    assert response.status == 200
    assert 'report-export-form' in response.text
    assert 'action="/export/report"' in response.text
    for column in ('ФИО', 'Пол', 'Институт', 'Группа', 'Курс', 'Участий', 'Побед', 'Призовых'):
        assert f'value="{column}"' in response.text


def test_reports_page_shows_slice_and_catalog_filters(client: SanicTestClient):
    # Срез и фильтры институт/группа/спорт — на странице фильтров отчёта;
    # группы знают иерархию (институт → его группы), срезы — фиксированный
    # набор из решения по замечанию №19.
    _, response = client.get('/reports', headers=get_auth_headers())

    assert response.status == 200
    assert 'Группировать по' in response.text
    assert 'name="group_by"' in response.text
    assert 'name="institute"' in response.text
    assert 'name="group"' in response.text
    assert 'name="sport"' in response.text
    for option in ('student', 'group', 'institute', 'course', 'sport', 'level', 'year'):
        assert f'value="{option}"' in response.text
    assert 'data-groups-by-institute' in response.text
    assert 'data-columns-by-slice' in response.text


# --- Расширение отчётов (замечание №19) на реальном SQLite-адаптере. ---
# Данные и ручной пересчёт метрик — src/conftest.py.


@pytest.fixture
def reports_client(client: SanicTestClient, tmp_path):
    storage = SQLiteAdapter(str(tmp_path / 'reports.sqlite3'))
    storage.save_competitions(make_report_fixture())
    for review_status, record in make_report_unapproved_fixture():
        storage.save_competitions([record], review_status=review_status, owner_id=9)
    storage.create_custom_field('trainer', 'Тренер', 'text', False, True, True, True, 0)
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


def get_report_cells(html: str) -> list[list[str]]:
    body = re.search(r'<tbody>(.*?)</tbody>', html, re.S)
    if body is None:
        return []
    rows = re.findall(r'<tr>(.*?)</tr>', body.group(1), re.S)
    return [
        [re.sub(r'<[^>]+>', '', cell).strip() for cell in re.findall(r'<td[^>]*>(.*?)</td>', row, re.S)] for row in rows
    ]


def test_report_page_student_slice_with_metrics(reports_client: SanicTestClient):
    # Историчность: Козлов в двух институтах — две отдельные строки; метрики
    # Побед/Призовых теперь есть и для среза «студент».
    _, response = reports_client.get('/report', headers=get_auth_headers())

    assert response.status == 200
    assert get_report_cells(response.text) == [
        ['1', 'Козлов Кирилл', 'М', 'ИМИ', 'ТД-303', '2', '1', '0', '1'],
        ['2', 'Козлов Кирилл', 'М', 'ИСИ', 'ПГС-101', '1', '1', '1', '1'],
        ['3', 'Сидоров Сидор', 'М', 'ИСИ', 'ПГС-102', '3', '1', '0', '0'],
        ['4', 'Петров Пётр', 'Ж', 'ИМИ', 'СБ-202', '2', '2', '1', '2'],
        ['5', 'Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '1', '3', '1', '2'],
    ]
    assert 'Участий' in response.text and 'Побед' in response.text and 'Призовых' in response.text


@pytest.mark.parametrize(
    'group_by,header,expected',
    [
        (
            'group',
            'Группа',
            [
                ['1', 'ПГС-102', '1', '0', '0'],
                ['2', 'ТД-303', '1', '0', '1'],
                ['3', 'СБ-202', '2', '1', '2'],
                ['4', 'ПГС-101', '4', '2', '3'],
            ],
        ),
        ('institute', 'Институт', [['1', 'ИМИ', '3', '1', '3'], ['2', 'ИСИ', '5', '2', '3']]),
        ('course', 'Курс', [['1', '3', '1', '0', '0'], ['2', '2', '3', '1', '3'], ['3', '1', '4', '2', '3']]),
        (
            'sport',
            'Вид спорта',
            [['1', 'Шахматы', '1', '0', '0'], ['2', 'Лыжи', '2', '1', '1'], ['3', 'Бег', '5', '2', '5']],
        ),
        (
            'level',
            'Уровень соревнований',
            [['1', 'межвузовские', '3', '0', '2'], ['2', 'внутривузовские', '5', '3', '4']],
        ),
        (
            'year',
            'Год',
            [['1', '2023', '2', '1', '1'], ['2', '2024', '3', '1', '3'], ['3', '2025', '3', '1', '2']],
        ),
    ],
)
def test_report_page_slice_rows(reports_client: SanicTestClient, group_by, header, expected):
    _, response = reports_client.get(f'/report?group_by={group_by}', headers=get_auth_headers())

    assert response.status == 200
    assert f'>{header}</th>' in response.text
    assert get_report_cells(response.text) == expected


def test_report_page_filters_institute_group_sport(reports_client: SanicTestClient):
    # Пара институт+группа по иерархии: остались Иванов и запись Козлова из
    # ИСИ/ПГС-101 (его ИМИ-запись не попадает).
    _, response = reports_client.get(
        '/report?' + urlencode({'institute': 'ИСИ', 'group': 'ПГС-101'}),
        headers=get_auth_headers(),
    )
    assert response.status == 200
    assert get_report_cells(response.text) == [
        ['1', 'Козлов Кирилл', 'М', 'ИСИ', 'ПГС-101', '1', '1', '1', '1'],
        ['2', 'Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '1', '3', '1', '2'],
    ]

    _, response = reports_client.get('/report?sport=' + quote('Шахматы'), headers=get_auth_headers())
    assert response.status == 200
    assert get_report_cells(response.text) == [
        ['1', 'Сидоров Сидор', 'М', 'ИСИ', 'ПГС-102', '3', '1', '0', '0'],
    ]

    # Новый фильтр работает и на другом срезе: группы института ИМИ.
    _, response = reports_client.get(
        '/report?' + urlencode({'group_by': 'group', 'institute': 'ИМИ'}),
        headers=get_auth_headers(),
    )
    assert response.status == 200
    assert get_report_cells(response.text) == [
        ['1', 'ТД-303', '1', '0', '1'],
        ['2', 'СБ-202', '2', '1', '2'],
    ]


def test_report_page_custom_filter_applies_to_slice(reports_client: SanicTestClient):
    _, response = reports_client.get(
        '/report?' + urlencode({'group_by': 'group', 'custom__trainer': 'Смит'}),
        headers=get_auth_headers(),
    )
    assert response.status == 200
    assert get_report_cells(response.text) == [['1', 'ПГС-101', '1', '1', '1']]


def test_report_url_is_reproducible(reports_client: SanicTestClient):
    # Шаримость: прямой GET с срезом и фильтрами даёт те же данные тем же
    # ответом, что и повторный заход по той же ссылке (закладка).
    url = '/report?' + urlencode(
        {'group_by': 'group', 'institute': 'ИСИ', 'date_from': '01.01.2023', 'date_to': '31.12.2025'}
    )
    _, first = reports_client.get(url, headers=get_auth_headers())
    _, second = reports_client.get(url, headers=get_auth_headers())

    assert first.status == second.status == 200
    assert first.text == second.text
    assert get_report_cells(first.text) == [
        ['1', 'ПГС-102', '1', '0', '0'],
        ['2', 'ПГС-101', '4', '2', '3'],
    ]


def test_report_rejects_unknown_slice(client: SanicTestClient):
    _, response = client.get('/report?group_by=bogus', headers=get_auth_headers())
    assert response.status == 400
    assert 'Неизвестный срез' in response.text


def test_report_available_for_viewer_with_slice(reports_client: SanicTestClient):
    _, response = reports_client.get('/report?group_by=year', headers=get_auth_headers(role='viewer'))
    assert response.status == 200
    assert get_report_cells(response.text) == [
        ['1', '2023', '2', '1', '1'],
        ['2', '2024', '3', '1', '3'],
        ['3', '2025', '3', '1', '2'],
    ]


def test_report_page_shows_row_counter(reports_client: SanicTestClient):
    # Композиция отчёта (прототип 06): счётчик «N строк · M участий».
    _, response = reports_client.get('/report', headers=get_auth_headers())

    assert response.status == 200
    # 5 строк-студентов, суммарно 8 участий в фикстуре.
    assert '5 строк · 8 участий' in response.text

    _, response = reports_client.get('/report?institute=' + quote('ИМИ'), headers=get_auth_headers())
    assert response.status == 200
    # Петров (2 участия) + Козлов в записи за ИМИ (1 участие) = 2 строки, 3 участия.
    assert '2 строки · 3 участия' in response.text


def test_report_page_shows_filter_chips(reports_client: SanicTestClient):
    # Каждый активный фильтр — чип с крестиком-сбросом (data-remove-key);
    # позиция и срез — человекочитаемые метки; кастомные поля — по ярлыку.
    params = urlencode(
        {
            'name': 'Иванов',
            'level': 'внутривузовские',
            'position': '<2',
            'group_by': 'sport',
            'custom__trainer': 'Смит',
        }
    )
    _, response = reports_client.get(f'/report?{params}', headers=get_auth_headers())

    assert response.status == 200
    assert 'ФИО: <strong>Иванов</strong>' in response.text
    assert 'Уровень: <strong>внутривузовские</strong>' in response.text
    assert 'Место: <strong>Победа</strong>' in response.text
    assert 'Срез: <strong>Вид спорта</strong>' in response.text
    assert 'Тренер: <strong>Смит</strong>' in response.text
    assert 'data-remove-key="name"' in response.text
    assert 'data-remove-key="custom__trainer"' in response.text


def test_report_page_without_filters_has_no_chips(reports_client: SanicTestClient):
    _, response = reports_client.get('/report', headers=get_auth_headers())
    assert response.status == 200
    assert 'filter-chip' not in response.text


# --- Серверная фильтрация, пагинация и счётчик главной (прототип 02) ---
# Реальный SQLite-адаптер: фильтры GET / — это SQL (storage), а не клиент.
# Набор подобран так, чтобы каждый фильтр и границы страниц считались руками:
# 12 подтверждённых + 2 «на проверке» + 1 «отклонена» = 15 для модератора.


def make_index_record(
    name: str,
    *,
    institute: str = 'ИСИ',
    sport: str = 'Бег',
    level: str = 'внутривузовские',
    date: datetime = datetime(2024, 3, 1),
) -> Competition:
    return Competition(
        student_id=f'id-{name}',
        student_name=name,
        student_sex='М',
        institute=institute,
        group='ГРП-101',
        course=2,
        sport=sport,
        date=date,
        level=level,
        name='Кубок',
        position=1,
    )


def seed_index_records(storage: SQLiteAdapter) -> None:
    approved = [
        make_index_record('Иванов Иван', date=datetime(2024, 3, 1)),
        make_index_record('Иванов Иван', date=datetime(2024, 4, 1)),
        make_index_record('Иванов Иван', date=datetime(2024, 5, 15)),
        make_index_record('Иванов Иван', date=datetime(2024, 9, 10)),
        make_index_record(
            'Петрова Анна', institute='ИМИ', sport='Лыжи', level='межвузовские', date=datetime(2024, 11, 1)
        ),
        make_index_record(
            'Петрова Анна', institute='ИМИ', sport='Лыжи', level='межвузовские', date=datetime(2024, 12, 1)
        ),
        make_index_record('Сидоров Сидор', date=datetime(2025, 1, 10)),
        make_index_record('Сидоров Сидор', date=datetime(2025, 2, 20)),
        make_index_record('Кузнецов Кирилл', institute='ИМИ', sport='Шахматы', date=datetime(2025, 3, 5)),
        make_index_record(
            'Иванов Иван', institute='ИМИ', sport='Лыжи', level='межвузовские', date=datetime(2025, 4, 1)
        ),
        make_index_record('Петрова Анна', sport='Бег', level='межвузовские', date=datetime(2025, 4, 15)),
        make_index_record('Петрова Анна', sport='Бег', level='межвузовские', date=datetime(2025, 4, 25)),
    ]
    storage.save_competitions(approved)
    storage.save_competitions(
        [
            make_index_record('Ждунов Ждун', date=datetime(2025, 5, 1)),
            make_index_record('Ждунов Ждун', date=datetime(2025, 5, 2)),
        ],
        review_status='pending',
        owner_id=9,
    )
    storage.save_competitions(
        [make_index_record('Отклонов Отклон', institute='ИМИ', sport='Лыжи', date=datetime(2025, 5, 3))],
        review_status='rejected',
        owner_id=9,
    )


@pytest.fixture
def index_client(client: SanicTestClient, tmp_path) -> SanicTestClient:
    storage = SQLiteAdapter(str(tmp_path / 'index.sqlite3'))
    seed_index_records(storage)
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


@pytest.fixture
def athlete_index_client(index_client: SanicTestClient) -> SanicTestClient:
    """Атлет sportik с двумя своими записями «на проверке» (owner_id — его id)."""
    storage = app.ctx.storage
    storage.create_user('sportik', hash_password('sportik-pass-123'), 'athlete')
    athlete_id = storage.get_user('sportik')['id']
    storage.save_competitions(
        [
            make_index_record('Атлетов Атлет', date=datetime(2025, 6, 1)),
            make_index_record('Атлетов Атлет', date=datetime(2025, 6, 2)),
        ],
        review_status='pending',
        owner_id=athlete_id,
    )
    return index_client


def get_tbody_rows(html: str) -> list[str]:
    body = re.search(r'<tbody>(.*?)</tbody>', html, re.S)
    if body is None:
        return []
    return re.findall(r'<tr>(.*?)</tr>', body.group(1), re.S)


def get_index_counter(html: str) -> tuple[int, int]:
    match = re.search(r'Показано <strong>(\d+)</strong> из <strong>(\d+)</strong> записей', html)
    assert match is not None, 'счётчик «Показано N из M» не найден'
    return int(match.group(1)), int(match.group(2))


def test_pager_items_edges_and_gaps():
    # Чистая функция пагинации: края (первая/последняя) и окно ±1 вокруг
    # текущей страницы, между непоследовательными номерами — разрыв «…».
    from src.main import pager_items

    assert pager_items(1, 0) == []
    assert pager_items(1, 1) == [{'page': 1}]
    assert pager_items(1, 2) == [{'page': 1}, {'page': 2}]
    assert pager_items(1, 4) == [{'page': 1}, {'page': 2}, {'gap': True}, {'page': 4}]
    assert pager_items(1, 9) == [{'page': 1}, {'page': 2}, {'gap': True}, {'page': 9}]
    assert pager_items(5, 9) == [
        {'page': 1},
        {'gap': True},
        {'page': 4},
        {'page': 5},
        {'page': 6},
        {'gap': True},
        {'page': 9},
    ]
    assert pager_items(9, 9) == [{'page': 1}, {'gap': True}, {'page': 8}, {'page': 9}]


def test_index_shows_counter_and_all_records(index_client: SanicTestClient):
    _, response = index_client.get('/', headers=get_auth_headers())

    assert response.status == 200
    shown, total = get_index_counter(response.text)
    assert (shown, total) == (15, 15)
    assert len(get_tbody_rows(response.text)) == 15
    # Композиция прототипа 02: заголовок с описанием, карточка фильтров, футер.
    assert 'Реестр соревнований' in response.text
    assert 'index-filter-card' in response.text
    assert 'Показывать по' in response.text


def test_index_linked_record_renders_event_link_and_edit_marker(index_client: SanicTestClient):
    """P2: связанная со событием запись — имя кликабельно (страница
    соревнования), кнопка правки несёт data-calendar-event-id (JS блокирует
    event-owned поля); NULL-запись — прежний вид без ссылки."""
    storage = app.ctx.storage
    event_id = storage.create_calendar_event(
        name='Кубок 2024', date='2024-03-01T00:00:00', date_to=None, level='', sport='', url=''
    )
    # Первая запись Иванова (name = «Кубок») — связываем raw SQL, как будто
    # это сделали будущие link/unlink.
    storage.connection.execute(
        'UPDATE competitions SET calendar_event_id = ? WHERE student_name = ? AND date = ?',
        (event_id, 'Иванов Иван', '2024-03-01T00:00:00'),
    )
    storage.connection.commit()

    _, response = index_client.get('/', headers=get_auth_headers())
    assert response.status == 200
    # Связанная строка: маркер для JS + ссылка на страницу соревнования.
    assert f'data-calendar-event-id="{event_id}"' in response.text
    assert (
        f'<a href="/calendar/{event_id}" class="link" title="Открыть страницу соревнования">Кубок</a>' in response.text
    )
    # NULL-строки: маркер пуст, ссылки события в колонке имени нет.
    assert 'data-calendar-event-id=""' in response.text
    null_rows = [row for row in get_tbody_rows(response.text) if 'data-calendar-event-id=""' in row]
    assert null_rows and 'Открыть страницу соревнования' not in null_rows[0]
    # Подсказка про нецензурируемые event-owned поля — на странице.
    assert 'они берутся из соревнования' in response.text


def test_index_athlete_sees_linked_record_name_as_text(athlete_index_client: SanicTestClient):
    """Атлету страница события недоступна (403) — имя связанной записи ему
    рендерится простым текстом, без ссылки."""
    storage = app.ctx.storage
    event_id = storage.create_calendar_event(
        name='Кубок атлета', date='2025-06-01T00:00:00', date_to=None, level='', sport='', url=''
    )
    storage.connection.execute(
        "UPDATE competitions SET calendar_event_id = ?, name = 'Кубок атлета' "
        "WHERE student_name = 'Атлетов Атлет' AND date = '2025-06-01T00:00:00'",
        (event_id,),
    )
    storage.connection.commit()

    _, response = athlete_index_client.get('/', headers=athlete_headers())
    assert response.status == 200
    assert f'<a href="/calendar/{event_id}"' not in response.text
    assert 'title="Открыть страницу соревнования"' not in response.text
    # Само название в строке осталось (простым текстом).
    assert 'Кубок атлета' in response.text


def test_index_filter_by_name_substring(index_client: SanicTestClient):
    _, response = index_client.get('/?' + urlencode({'name': 'Иванов'}), headers=get_auth_headers())

    assert response.status == 200
    assert get_index_counter(response.text) == (5, 5)
    assert len(get_tbody_rows(response.text)) == 5


def test_index_filter_by_catalog_exact_values(index_client: SanicTestClient):
    params = urlencode({'institute': 'ИСИ', 'sport': 'Бег', 'level': 'межвузовские'})
    _, response = index_client.get(f'/?{params}', headers=get_auth_headers())

    assert response.status == 200
    # Петрова Анна, ИСИ, Бег, межвузовские — две записи апреля 2025
    assert get_index_counter(response.text) == (2, 2)


def test_index_filter_by_date_range(index_client: SanicTestClient):
    params = urlencode({'date_from': '01.06.2024', 'date_to': '30.06.2024'})
    _, response = index_client.get(f'/?{params}', headers=get_auth_headers())

    assert response.status == 200
    assert get_index_counter(response.text) == (0, 0)
    assert 'Записей с текущим фильтром не найдено' in response.text

    params = urlencode({'date_from': '01.09.2024', 'date_to': '31.12.2024'})
    _, response = index_client.get(f'/?{params}', headers=get_auth_headers())
    # Иванов 10.09.2024 + Петрова 01.11.2024 и 01.12.2024
    assert get_index_counter(response.text) == (3, 3)


def test_index_filter_rejects_invalid_date(index_client: SanicTestClient):
    _, response = index_client.get('/?date_from=31.31.2024', headers=get_auth_headers())

    assert response.status == 400
    assert 'Некорректная дата' in response.text


def test_index_status_filter_moderator_only(index_client: SanicTestClient):
    _, response = index_client.get('/?status=pending', headers=get_auth_headers())
    assert response.status == 200
    assert get_index_counter(response.text) == (2, 2)

    _, response = index_client.get('/?status=rejected', headers=get_auth_headers())
    assert get_index_counter(response.text) == (1, 1)

    _, response = index_client.get('/?status=approved', headers=get_auth_headers())
    assert get_index_counter(response.text) == (12, 12)

    # Незнакомое значение — фильтр не применяется (все 15).
    _, response = index_client.get('/?status=unknown', headers=get_auth_headers())
    assert get_index_counter(response.text) == (15, 15)

    # viewer не модератор: параметр статуса игнорируется, селекта «Статус» нет.
    _, response = index_client.get('/?status=pending', headers=get_auth_headers(role='viewer'))
    assert response.status == 200
    assert get_index_counter(response.text) == (15, 15)
    assert 'id="index-filter-status"' not in response.text


def test_index_filters_kept_in_url_and_form(index_client: SanicTestClient):
    params = {'status': 'approved', 'per_page': 10, 'page': 2}
    _, response = index_client.get('/?' + urlencode(params), headers=get_auth_headers())

    assert response.status == 200
    # Поля формы переотображают применённые фильтры (ссылки шарятся).
    assert '<option value="approved" selected>' in response.text
    assert '<option value="10" selected>' in response.text
    # Селект «Показывать по» держит фильтры без per_page (ставит сам).
    assert f'data-query="{urlencode({"status": "approved"})}"' in response.text
    # Ссылки пагинации несут фильтры и per_page дальше (Jinja экранирует &).
    assert 'href="/?status=approved&amp;per_page=10&amp;page=1"' in response.text

    # Фильтр по ФИО тоже возвращается в поле ввода.
    _, response = index_client.get('/?' + urlencode({'name': 'Иванов'}), headers=get_auth_headers())
    assert 'value="Иванов"' in response.text


def test_index_per_page_default_shows_all(index_client: SanicTestClient):
    _, response = index_client.get('/', headers=get_auth_headers())
    assert get_index_counter(response.text) == (15, 15)
    # Дефолт — 100 (прежнее поведение показывало весь список): страница одна.
    assert '<option value="100" selected>' in response.text
    assert 'class="pager"' not in response.text


def test_index_per_page_splits_and_paginates(index_client: SanicTestClient):
    _, response = index_client.get('/?per_page=10', headers=get_auth_headers())
    assert get_index_counter(response.text) == (10, 15)
    assert len(get_tbody_rows(response.text)) == 10
    assert 'href="/?per_page=10&amp;page=2"' in response.text

    _, response = index_client.get('/?per_page=10&page=2', headers=get_auth_headers())
    assert get_index_counter(response.text) == (5, 15)
    assert len(get_tbody_rows(response.text)) == 5


def test_index_per_page_invalid_falls_back_to_default(index_client: SanicTestClient):
    _, response = index_client.get('/?per_page=7', headers=get_auth_headers())
    assert get_index_counter(response.text) == (15, 15)
    assert '<option value="100" selected>' in response.text


def test_index_page_out_of_range_clamps_to_last(index_client: SanicTestClient):
    _, response = index_client.get('/?per_page=10&page=999', headers=get_auth_headers())
    # Страница за пределами диапазона зажимается на последней (5 записей).
    assert get_index_counter(response.text) == (5, 15)


def test_index_page_invalid_number_falls_back_to_first(index_client: SanicTestClient):
    _, response = index_client.get('/?per_page=10&page=abc', headers=get_auth_headers())
    assert get_index_counter(response.text) == (10, 15)


def test_index_pager_two_pages_navigation(index_client: SanicTestClient):
    _, response = index_client.get('/?per_page=10', headers=get_auth_headers())
    # Первая из двух: «‹» неактивна, ссылка на вторую есть.
    assert 'pager__btn is-disabled' in response.text
    assert 'href="/?per_page=10&amp;page=2"' in response.text

    _, response = index_client.get('/?per_page=10&page=2', headers=get_auth_headers())
    # Вторая страница: 5 записей, активна кнопка «2», «›» неактивна.
    assert get_index_counter(response.text) == (5, 15)
    assert '<span class="pager__btn is-active" aria-current="page">2</span>' in response.text
    assert 'href="/?per_page=10&amp;page=1"' in response.text


def test_index_athlete_sees_own_records_and_ignores_status(athlete_index_client: SanicTestClient):
    _, response = athlete_index_client.get('/', headers=athlete_headers())
    assert response.status == 200
    assert get_index_counter(response.text) == (2, 2)
    # Селекта «Статус» у атлета нет; параметр статуса сервер игнорирует.
    assert 'id="index-filter-status"' not in response.text

    _, response = athlete_index_client.get('/?status=approved', headers=athlete_headers())
    assert get_index_counter(response.text) == (2, 2)


def test_index_athlete_filter_applies_to_own_records(athlete_index_client: SanicTestClient):
    storage = app.ctx.storage
    athlete_id = storage.get_user('sportik')['id']
    storage.save_competitions(
        [make_index_record('Атлетов Атлет', sport='Шахматы', date=datetime(2025, 7, 1))],
        review_status='pending',
        owner_id=athlete_id,
    )

    _, response = athlete_index_client.get('/?' + urlencode({'sport': 'Шахматы'}), headers=athlete_headers())
    assert get_index_counter(response.text) == (1, 1)


# --- Пустое состояние атлета: первая запись создаётся прямо с главной. ---


@pytest.fixture
def blank_index_client(client: SanicTestClient, tmp_path) -> SanicTestClient:
    """Реестр без единой записи (admin/editor/viewer сеются стартом приложения)."""
    storage = SQLiteAdapter(str(tmp_path / 'index-blank.sqlite3'))
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


@pytest.fixture
def athlete_empty_index_client(blank_index_client: SanicTestClient) -> SanicTestClient:
    """Атлет sportik без записей: раньше главная не рендерила таблицу вовсе."""
    app.ctx.storage.create_user('sportik', hash_password('sportik-pass-123'), 'athlete')
    return blank_index_client


def test_index_athlete_empty_renders_table_skeleton_and_cta(athlete_empty_index_client: SanicTestClient):
    _, response = athlete_empty_index_client.get('/', headers=athlete_headers())

    assert response.status == 200
    # Каркас таблицы рендерится: шапка с колонками на месте, tbody пуст.
    assert '<table class="table table-striped table-bordered interactive-table"' in response.text
    assert '<thead class="table-dark">' in response.text
    assert re.search(r'<tbody>\s*</tbody>', response.text)
    # Пустое состояние атлета и CTA на первую запись (та же инлайн-строка).
    assert 'У вас пока нет участий' in response.text
    assert (
        '<button type="button" class="btn btn-success empty-add-row-button">Добавить участие</button>' in response.text
    )
    # Тулбар «+ Пустая строка» и подсказка жестов сосуществуют с CTA.
    assert 'add-empty-row-button' in response.text
    assert 'table-gestures-hint' in response.text
    # Прежний маркер пустого реестра и футер «Показано 0 из 0» не рендерятся.
    assert 'Нет данных' not in response.text
    assert 'Показано <strong>' not in response.text


def test_index_athlete_empty_with_filter_keeps_plain_empty_state(athlete_empty_index_client: SanicTestClient):
    # Фильтр без совпадений у атлета — прежняя ветка, без каркаса и CTA.
    _, response = athlete_empty_index_client.get('/?' + urlencode({'name': 'Никого Нет'}), headers=athlete_headers())

    assert response.status == 200
    assert 'Записей с текущим фильтром не найдено' in response.text
    assert 'interactive-table' not in response.text
    assert 'empty-add-row-button' not in response.text
    assert 'У вас пока нет участий' not in response.text


@pytest.mark.parametrize('role', ['admin', 'editor'])
def test_index_moderator_empty_registry_keeps_plain_empty_state(blank_index_client: SanicTestClient, role: str):
    _, response = blank_index_client.get('/', headers=get_auth_headers(role=role))

    assert response.status == 200
    assert 'Нет данных' in response.text
    assert 'interactive-table' not in response.text
    assert 'empty-add-row-button' not in response.text
    assert 'У вас пока нет участий' not in response.text


def test_index_viewer_empty_registry_has_no_write_controls(blank_index_client: SanicTestClient):
    storage = app.ctx.storage
    # viewer сеётся из AUTH_VIEWER_*; в окружении без переменных — вручную.
    if storage.get_user('viewer') is None:
        storage.create_user('viewer', hash_password('viewer-pass-123'), 'viewer')
    _, response = blank_index_client.get('/', headers=get_auth_headers(role='viewer'))

    assert response.status == 200
    assert 'Нет данных' in response.text
    assert 'interactive-table' not in response.text
    assert 'empty-add-row-button' not in response.text
    assert 'add-empty-row-button' not in response.text
    assert 'competition-edit-button' not in response.text


def test_index_athlete_records_filter_no_match_keeps_plain_empty_state(
    athlete_index_client: SanicTestClient,
):
    # У атлета есть записи, но фильтр отсекает их все: без каркаса/CTA,
    # футер честно показывает «Показано 0 из 0».
    _, response = athlete_index_client.get('/?' + urlencode({'sport': 'Плавание'}), headers=athlete_headers())

    assert response.status == 200
    assert 'Записей с текущим фильтром не найдено' in response.text
    assert 'interactive-table' not in response.text
    assert 'empty-add-row-button' not in response.text
    assert get_index_counter(response.text) == (0, 0)


def test_index_athlete_rejected_record_shows_review_comment(index_client: SanicTestClient):
    storage = app.ctx.storage
    storage.create_user('sportik', hash_password('sportik-pass-123'), 'athlete')
    athlete_id = storage.get_user('sportik')['id']
    storage.save_competitions(
        [make_index_record('Атлетов Атлет', date=datetime(2025, 8, 1))],
        review_status='pending',
        owner_id=athlete_id,
    )
    record_id = int(storage.get_competitions_page(owner_id=athlete_id)[0].record_id)
    storage.set_competition_review(record_id, 'rejected', 'Уточните дату соревнования')

    _, response = index_client.get('/', headers=athlete_headers())

    assert response.status == 200
    # Комментарий модератора — в подсказке бейджа «отклонена» (регресс).
    assert 'Уточните дату соревнования' in response.text
    assert 'отклонена' in response.text


def test_export_report_slice_xlsx_rows(reports_client: SanicTestClient):
    _, response = reports_client.get('/export/report?group_by=group', headers=get_auth_headers())

    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('Группа', 'Участий', 'Побед', 'Призовых'),
        ('ПГС-102', 1, 0, 0),
        ('ТД-303', 1, 0, 1),
        ('СБ-202', 2, 1, 2),
        ('ПГС-101', 4, 2, 3),
    ]


def test_export_report_slice_with_column_subset(reports_client: SanicTestClient):
    _, response = reports_client.get(
        '/export/report?group_by=sport&columns=' + quote('Побед,Вид спорта'),
        headers=get_auth_headers(),
    )

    # порядок колонок канонический для среза, неизвестные игнорируются
    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('Вид спорта', 'Побед'),
        ('Шахматы', 0),
        ('Лыжи', 1),
        ('Бег', 2),
    ]


def test_export_report_slice_with_filters_and_year(reports_client: SanicTestClient):
    _, response = reports_client.get(
        '/export/report?' + urlencode({'group_by': 'year', 'institute': 'ИМИ', 'sport': 'Бег'}),
        headers=get_auth_headers(),
    )

    assert response.status == 200
    assert get_xlsx_rows(response.body) == [
        ('Год', 'Участий', 'Побед', 'Призовых'),
        (2024, 1, 0, 1),
        (2025, 1, 0, 1),
    ]


def test_empty_template_includes_custom_fields(client: SanicTestClient):
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(field_id=1, key='trainer', label='Тренер', show_in_template=True)
    ]

    _, response = client.get('/template/empty.xlsx', headers=get_auth_headers(role='editor'))

    assert response.status == 200
    assert 'Тренер' in get_xlsx_headers(response.body)


def test_admin_can_create_custom_field(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields',
        headers=headers,
        data={
            **csrf_for(headers),
            'label': 'Тренер',
            'field_type': 'text',
            'required': 'on',
            'show_in_table': 'on',
            'show_in_export': 'on',
            'show_in_template': 'on',
            'sort_order': '10',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    app.ctx.storage.create_custom_field.assert_called_once()


def test_viewer_cannot_manage_fields(client: SanicTestClient):
    headers = get_auth_headers(role='viewer')
    _, response = client.post(
        '/admin/fields',
        headers=headers,
        data={**csrf_for(headers), 'label': 'Тренер', 'field_type': 'text'},
        allow_redirects=False,
    )
    assert response.status == 403


def test_normalize_position_allows_multiple_digits():
    assert normalize_position('12') == 12


def test_requires_auth_for_main_page(client: SanicTestClient):
    _, response = client.get('/', allow_redirects=False)

    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_login_sets_auth_cookie(client: SanicTestClient):
    _, response = client.post(
        '/login',
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )

    assert response.status == 302
    assert response.headers['location'] == '/'
    assert settings.auth_cookie_name in response.headers.get('set-cookie', '')


# --- Присутствие пользователей: last_login_at / last_seen_at (№17) ---


def test_login_records_last_login(client: SanicTestClient):
    app.ctx.storage.set_user_last_login.reset_mock()
    _, response = client.post(
        '/login',
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )

    assert response.status == 302
    # фикстура отдаёт админу id = 1
    app.ctx.storage.set_user_last_login.assert_called_once_with(1)


def test_authenticated_requests_touch_last_seen_with_throttle(client: SanicTestClient):
    from src.main import reset_presence_tracking

    reset_presence_tracking()
    app.ctx.storage.touch_user_seen.reset_mock()
    headers = get_auth_headers()

    client.get('/', headers=headers)
    client.get('/', headers=headers)
    # два запроса подряд — один UPDATE благодаря троттлингу (окно 60 с)
    app.ctx.storage.touch_user_seen.assert_called_once_with(1)

    # без сессии активность не фиксируется
    app.ctx.storage.touch_user_seen.reset_mock()
    reset_presence_tracking()
    client.get('/login')
    app.ctx.storage.touch_user_seen.assert_not_called()


def test_admin_users_page_shows_presence(client: SanicTestClient):
    from datetime import timedelta
    from datetime import timezone

    now = datetime.now(timezone.utc)
    app.ctx.storage.list_users.return_value = [
        {
            'id': 1,
            'username': 'admin',
            'role': 'admin',
            'active': 1,
            'name_aliases': [],
            'last_login_at': now.isoformat(),
            'last_seen_at': now.isoformat(),
        },
        {
            'id': 2,
            'username': 'stale',
            'role': 'viewer',
            'active': 1,
            'name_aliases': [],
            'last_login_at': (now - timedelta(hours=2)).isoformat(),
            'last_seen_at': (now - timedelta(minutes=30)).isoformat(),
        },
        {
            'id': 3,
            'username': 'ghost',
            'role': 'viewer',
            'active': 0,
            'name_aliases': [],
            'last_login_at': None,
            'last_seen_at': None,
        },
    ]
    try:
        _, response = client.get('/admin/users', headers=get_auth_headers())
        assert response.status == 200
        # онлайн: зелёная точка
        assert 'presence-dot-online' in response.text
        assert 'онлайн' in response.text
        # был(а): 30 минут назад / —
        assert 'был(а): 30 минут назад' in response.text
        assert 'был(а): —' in response.text
        assert 'Последний вход:' in response.text
    finally:
        app.ctx.storage.list_users.return_value = []


def test_competition_created_at_default_factory():
    first = Competition(
        student_id='1',
        student_name='Test 1',
        student_sex='M',
        institute='Inst',
        group='A',
        course=1,
        sport='Run',
        date=datetime(2024, 1, 1),
        level='межвузовские',
        name='Meet',
        position=1,
    )
    second = Competition(
        student_id='2',
        student_name='Test 2',
        student_sex='F',
        institute='Inst',
        group='B',
        course=2,
        sport='Jump',
        date=datetime(2024, 1, 2),
        level='внутривузовские',
        name='Cup',
        position=2,
    )

    assert second.created_at >= first.created_at


def test_report_rejects_invalid_date(client: SanicTestClient):
    _, response = client.get('/report?date_from=garbage', headers=get_auth_headers())
    assert response.status == 400


def test_report_rejects_position_without_sign(client: SanicTestClient):
    _, response = client.get('/report?position=5', headers=get_auth_headers())
    assert response.status == 400


def test_report_rejects_position_with_bad_value(client: SanicTestClient):
    _, response = client.get('/report?position=%3Eabc', headers=get_auth_headers())
    assert response.status == 400


def test_update_competition_rejects_non_numeric_id(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/abc',
        headers=headers,
        data={**csrf_for(headers), 'student_name': 'X', 'course': '1', 'date': '10.04.2026'},
        allow_redirects=False,
    )
    assert response.status == 400


def test_delete_custom_field_rejects_non_numeric_id(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields/abc/delete',
        headers=headers,
        data=csrf_for(headers),
    )
    assert response.status == 400


def test_update_custom_field_rejects_invalid_sort_order(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields/1',
        headers=headers,
        data={**csrf_for(headers), 'label': 'Поле', 'field_type': 'text', 'sort_order': 'abc'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']


def test_update_custom_field_rejects_empty_label(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields/1',
        headers=headers,
        data={**csrf_for(headers), 'label': '', 'field_type': 'text'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']


def test_post_without_csrf_token_is_rejected(client: SanicTestClient):
    _, response = client.post(
        '/competition/123/delete',
        headers=get_auth_headers(),
        allow_redirects=False,
    )
    assert response.status == 403


def test_export_index_neutralizes_formula_values(client: SanicTestClient):
    from openpyxl import load_workbook

    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='=HYPERLINK("http://evil.example")',
            student_sex='M',
            institute='Inst',
            group='A',
            course=1,
            sport='Run',
            date=datetime(2024, 1, 1),
            level='межвузовские',
            name='Meet',
            position=1,
        )
    ]

    _, response = client.get('/export/index', headers=get_auth_headers())

    assert response.status == 200
    workbook = load_workbook(BytesIO(response.body), read_only=True)
    cell_value = next(workbook.worksheets[0].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert cell_value == "'=HYPERLINK(\"http://evil.example\")"


def test_expired_auth_cookie_is_rejected(client: SanicTestClient):
    from time import time as time_time

    expired_cookie = create_auth_cookie_value(
        username=settings.auth_admin_username,
        role='admin',
        issued_at=int(time_time()) - settings.auth_session_ttl_seconds - 10,
    )
    _, response = client.get(
        '/',
        headers={'cookie': f'{settings.auth_cookie_name}={expired_cookie}'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_login_rate_limit_locks_out_after_failures(client: SanicTestClient):
    for _ in range(5):
        client.post(
            '/login',
            data={'username': settings.auth_admin_username, 'password': 'wrong'},
            allow_redirects=False,
        )

    _, response = client.post(
        '/login',
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    assert response.status == 429


def test_login_rate_limit_counts_by_real_client_ip_behind_proxy(client: SanicTestClient):
    # M1: PROXIES_COUNT=1 — бакет лимита ведётся по X-Forwarded-For (реальный
    # клиент за nginx), а не по адресу прокси. client — module-scoped:
    # сбрасываем общий login_failures до и после.
    from src.main import login_failures

    login_failures.clear()
    for _ in range(5):
        _, response = client.post(
            '/login',
            headers={'X-Forwarded-For': '198.51.100.10'},
            data={'username': settings.auth_admin_username, 'password': 'wrong'},
            allow_redirects=False,
        )
        assert response.status == 401

    # Другой клиент — отдельный бакет: корректный вход не задет чужой блокировкой.
    _, response = client.post(
        '/login',
        headers={'X-Forwarded-For': '198.51.100.20'},
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    assert response.status == 302
    login_failures.clear()


def test_login_rate_limit_locks_forwarded_client_ip(client: SanicTestClient):
    from src.main import login_failures

    login_failures.clear()
    for _ in range(5):
        client.post(
            '/login',
            headers={'X-Forwarded-For': '198.51.100.10'},
            data={'username': settings.auth_admin_username, 'password': 'wrong'},
            allow_redirects=False,
        )

    _, response = client.post(
        '/login',
        headers={'X-Forwarded-For': '198.51.100.10'},
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    assert response.status == 429
    login_failures.clear()


def test_login_rate_limit_uses_last_forwarded_entry(client: SanicTestClient):
    # nginx дописывает реальный IP ПОСЛЕДНИМ ($proxy_add_x_forwarded_for):
    # бакет ключуется по последней записи, подделка первой записи ничего
    # не даёт.
    from src.main import login_failures

    login_failures.clear()
    for _ in range(5):
        client.post(
            '/login',
            headers={'X-Forwarded-For': '6.0.0.6, 198.51.100.10'},
            data={'username': settings.auth_admin_username, 'password': 'wrong'},
            allow_redirects=False,
        )

    # Заблокирован 198.51.100.10; «клиент» с другой последней записью
    # (пусть даже с той же подделанной первой) входит свободно.
    _, response = client.post(
        '/login',
        headers={'X-Forwarded-For': '6.0.0.6, 198.51.100.30'},
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    assert response.status == 302
    login_failures.clear()


@pytest.fixture
def bootstrap_settings(monkeypatch):
    """Синтетические настройки посева учёток: сильные уникальные пароли,
    viewer пустой (не создаётся). Тесты точечно подменяют нужное поле."""
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'auth_secret_key', 'unit-test-random-secret-key')
    monkeypatch.setattr(main_module.settings, 'auth_admin_username', 'admin')
    monkeypatch.setattr(main_module.settings, 'auth_admin_password', 'strong-admin-pass-123456')
    monkeypatch.setattr(main_module.settings, 'auth_editor_username', 'editor')
    monkeypatch.setattr(main_module.settings, 'auth_editor_password', 'strong-editor-pass-123456')
    monkeypatch.setattr(main_module.settings, 'auth_viewer_username', '')
    monkeypatch.setattr(main_module.settings, 'auth_viewer_password', '')
    return main_module.settings


def test_find_weak_bootstrap_accounts_flags_weak_passwords(bootstrap_settings, monkeypatch):
    from src.main import find_weak_bootstrap_accounts

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_password', 'change-me')
    storage = Mock()
    storage.get_user.return_value = None
    # учётки ещё нет → посев со слабым паролем → отказ; editor силён, viewer пуст
    assert find_weak_bootstrap_accounts(storage) == ['admin']


def test_find_weak_bootstrap_accounts_skips_existing_users(bootstrap_settings, monkeypatch):
    from src.main import find_weak_bootstrap_accounts

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_password', 'change-me')
    storage = Mock()
    storage.get_user.return_value = {'username': 'admin', 'role': 'admin'}
    # посев не перезаписывает существующие учётки — слабый .env не опасен
    assert find_weak_bootstrap_accounts(storage) == []


def test_find_weak_bootstrap_accounts_empty_for_strong_passwords(bootstrap_settings):
    from src.main import find_weak_bootstrap_accounts

    storage = Mock()
    storage.get_user.return_value = None
    assert find_weak_bootstrap_accounts(storage) == []


def test_find_weak_bootstrap_accounts_flags_default_viewer_placeholder(bootstrap_settings, monkeypatch):
    from src.main import find_weak_bootstrap_accounts

    monkeypatch.setattr(bootstrap_settings, 'auth_viewer_username', 'viewer')
    monkeypatch.setattr(bootstrap_settings, 'auth_viewer_password', 'change-me-viewer')
    storage = Mock()
    storage.get_user.return_value = None
    assert find_weak_bootstrap_accounts(storage) == ['viewer']


def make_bootstrap_app_holder():
    holder = Mock()
    holder.ctx.storage = None
    return holder


def test_init_storage_refuses_startup_with_weak_bootstrap_password(bootstrap_settings, monkeypatch, tmp_path):
    from src import main as main_module

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_password', 'change-me')
    monkeypatch.setattr(bootstrap_settings, 'database_path', str(tmp_path / 'guard.sqlite3'))
    monkeypatch.setattr(main_module.Sanic, 'test_mode', False)

    with pytest.raises(RuntimeError) as exc_info:
        asyncio.run(main_module.init_storage(make_bootstrap_app_holder(), None))

    message = str(exc_info.value)
    assert 'AUTH_ADMIN_PASSWORD' in message
    assert 'change-me' not in message  # значение пароля не раскрывается


def test_init_storage_starts_when_weak_env_accounts_already_exist(bootstrap_settings, monkeypatch, tmp_path):
    from src import main as main_module

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_password', 'change-me')
    monkeypatch.setattr(bootstrap_settings, 'database_path', str(tmp_path / 'seeded.sqlite3'))
    holder = make_bootstrap_app_holder()

    # «старая установка»: учётки посеяны до появления guard (в test-режиме)
    monkeypatch.setattr(main_module.Sanic, 'test_mode', True)
    asyncio.run(main_module.init_storage(holder, None))
    seeded = holder.ctx.storage

    monkeypatch.setattr(main_module.Sanic, 'test_mode', False)
    asyncio.run(main_module.init_storage(holder, None))  # не бросает
    assert seeded.get_user('admin') is not None


def test_init_storage_seeds_accounts_with_strong_passwords(bootstrap_settings, monkeypatch, tmp_path):
    from src import main as main_module

    monkeypatch.setattr(bootstrap_settings, 'database_path', str(tmp_path / 'fresh.sqlite3'))
    monkeypatch.setattr(main_module.Sanic, 'test_mode', False)
    holder = make_bootstrap_app_holder()

    asyncio.run(main_module.init_storage(holder, None))

    storage = holder.ctx.storage
    assert storage.get_user('admin') is not None
    assert storage.get_user('editor') is not None
    assert storage.get_user('viewer') is None  # пустые AUTH_VIEWER_* — учётки нет


def test_init_storage_refuses_bootstrap_username_with_colon(bootstrap_settings, monkeypatch, tmp_path):
    from src import main as main_module

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_username', 'bad:admin')
    monkeypatch.setattr(bootstrap_settings, 'database_path', str(tmp_path / 'colon.sqlite3'))
    monkeypatch.setattr(main_module.Sanic, 'test_mode', False)

    with pytest.raises(RuntimeError) as exc_info:
        asyncio.run(main_module.init_storage(make_bootstrap_app_holder(), None))

    assert 'AUTH_ADMIN_USERNAME' in str(exc_info.value)


def test_init_storage_starts_when_colon_user_already_exists(bootstrap_settings, monkeypatch, tmp_path):
    from src import main as main_module

    monkeypatch.setattr(bootstrap_settings, 'auth_admin_username', 'bad:admin')
    monkeypatch.setattr(bootstrap_settings, 'database_path', str(tmp_path / 'colon-seeded.sqlite3'))
    holder = make_bootstrap_app_holder()

    monkeypatch.setattr(main_module.Sanic, 'test_mode', True)
    asyncio.run(main_module.init_storage(holder, None))

    # существующая (пусть и некрасивая) учётка не пересеивается — не блокируем
    monkeypatch.setattr(main_module.Sanic, 'test_mode', False)
    asyncio.run(main_module.init_storage(holder, None))


def test_admin_can_create_user(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users',
        headers=headers,
        data={**csrf_for(headers), 'username': 'operator2', 'password': 'secret123', 'role': 'editor'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.create_user.assert_called_once()
    args = app.ctx.storage.create_user.call_args[0]
    assert args[0] == 'operator2'
    assert args[2] == 'editor'


def test_create_user_rejects_short_password(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users',
        headers=headers,
        data={**csrf_for(headers), 'username': 'operator3', 'password': '123', 'role': 'editor'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']


def test_create_user_rejects_colon_in_username(client: SanicTestClient):
    # L8: двоеточие ломает подписанную cookie «логин:роль:...».
    headers = get_auth_headers()
    app.ctx.storage.create_user.reset_mock()
    _, response = client.post(
        '/admin/users',
        headers=headers,
        data={**csrf_for(headers), 'username': 'bad:name', 'password': 'secret123', 'role': 'editor'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.create_user.assert_not_called()


def test_create_user_rejects_control_character_in_username(client: SanicTestClient):
    headers = get_auth_headers()
    app.ctx.storage.create_user.reset_mock()
    _, response = client.post(
        '/admin/users',
        headers=headers,
        data={**csrf_for(headers), 'username': 'na\x01me', 'password': 'secret123', 'role': 'editor'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.create_user.assert_not_called()


def test_create_user_accepts_regular_username(client: SanicTestClient):
    headers = get_auth_headers()
    app.ctx.storage.create_user.reset_mock()
    _, response = client.post(
        '/admin/users',
        headers=headers,
        data={**csrf_for(headers), 'username': 'operator-9', 'password': 'secret123', 'role': 'editor'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.create_user.assert_called_once()
    assert app.ctx.storage.create_user.call_args[0][0] == 'operator-9'


def test_admin_can_reset_user_password(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': 'someone',
        'password_hash': 'x',
        'role': 'editor',
        'active': 1,
    }
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/7/password',
        headers=headers,
        data={**csrf_for(headers), 'password': 'newsecret123'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.set_user_password.assert_called_once()
    app.ctx.storage.get_user_by_id.return_value = None


def test_admin_cannot_deactivate_self(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': settings.auth_admin_username,
        'password_hash': 'x',
        'role': 'admin',
        'active': 1,
    }
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/7/active',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.set_user_active.assert_not_called()
    app.ctx.storage.get_user_by_id.return_value = None


def delete_user_request(client, headers, user_id, confirm_username):
    return client.post(
        f'/admin/users/{user_id}/delete',
        headers=headers,
        data={**csrf_for(headers), 'confirm_username': confirm_username},
        allow_redirects=False,
    )


def test_admin_cannot_delete_self(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': settings.auth_admin_username,
        'password_hash': 'x',
        'role': 'admin',
        'active': 1,
    }
    app.ctx.storage.delete_user.reset_mock()
    headers = get_auth_headers()
    _, response = delete_user_request(client, headers, 7, settings.auth_admin_username)

    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.delete_user.assert_not_called()
    app.ctx.storage.get_user_by_id.return_value = None


def test_admin_cannot_delete_last_active_admin(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 9,
        'username': 'admin2',
        'password_hash': 'x',
        'role': 'admin',
        'active': 1,
    }
    # Единственный активный админ — сама цель удаления.
    app.ctx.storage.list_users.return_value = [{'id': 9, 'username': 'admin2', 'role': 'admin', 'active': 1}]
    app.ctx.storage.delete_user.reset_mock()
    headers = get_auth_headers()
    _, response = delete_user_request(client, headers, 9, 'admin2')

    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.delete_user.assert_not_called()

    # Неактивного админа удалить можно: активных админов останется больше нуля.
    app.ctx.storage.get_user_by_id.return_value['active'] = 0
    _, response = delete_user_request(client, headers, 9, 'admin2')
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.delete_user.assert_called_once()

    app.ctx.storage.get_user_by_id.return_value = None
    app.ctx.storage.list_users.return_value = []


def test_delete_user_rejects_wrong_confirm_username(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': 'operator9',
        'password_hash': 'x',
        'role': 'editor',
        'active': 1,
    }
    app.ctx.storage.list_users.return_value = [
        {'id': 1, 'username': 'admin', 'role': 'admin', 'active': 1},
        {'id': 7, 'username': 'operator9', 'role': 'editor', 'active': 1},
    ]
    app.ctx.storage.delete_user.reset_mock()
    headers = get_auth_headers()

    _, response = delete_user_request(client, headers, 7, 'wrong-login')
    assert response.status == 302
    assert 'admin_error' in response.headers['location']

    _, response = delete_user_request(client, headers, 7, '')
    assert response.status == 302
    assert 'admin_error' in response.headers['location']

    app.ctx.storage.delete_user.assert_not_called()
    app.ctx.storage.get_user_by_id.return_value = None
    app.ctx.storage.list_users.return_value = []


def test_admin_can_delete_user(client: SanicTestClient):
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': 'operator9',
        'password_hash': 'x',
        'role': 'editor',
        'active': 1,
    }
    app.ctx.storage.list_users.return_value = [
        {'id': 1, 'username': 'admin', 'role': 'admin', 'active': 1},
        {'id': 7, 'username': 'operator9', 'role': 'editor', 'active': 1},
    ]
    app.ctx.storage.delete_user.return_value = 3
    app.ctx.storage.delete_user.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    headers = get_auth_headers()

    _, response = delete_user_request(client, headers, 7, 'operator9')

    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    assert 'admin_message' in response.headers['location']
    app.ctx.storage.delete_user.assert_called_once_with(7)
    assert (
        'user_deleted',
        {'target_user_id': 7, 'target_username': 'operator9', 'records_detached': 3},
    ) in audit_calls()

    app.ctx.storage.get_user_by_id.return_value = None
    app.ctx.storage.list_users.return_value = []
    app.ctx.storage.delete_user.return_value = 0


def test_delete_user_rejects_non_numeric_id(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/abc/delete',
        headers=headers,
        data=csrf_for(headers),
    )
    assert response.status == 400


def test_delete_user_forbidden_for_non_admin(client: SanicTestClient):
    editor_headers = get_auth_headers(role='editor')
    _, response = delete_user_request(client, editor_headers, 7, 'operator9')
    assert response.status == 403

    viewer_headers = get_auth_headers(role='viewer')
    _, response = delete_user_request(client, viewer_headers, 7, 'operator9')
    assert response.status == 403


def test_editor_can_approve_pending_record(client: SanicTestClient):
    app.ctx.storage.get_competition_review.return_value = {
        'id': 5,
        'review_status': 'pending',
        'owner_id': 2,
    }
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/5/review/approve',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.set_competition_review.assert_called_with(5, 'approved', '')
    app.ctx.storage.get_competition_review.return_value = None


def test_editor_can_reject_record_with_comment(client: SanicTestClient):
    app.ctx.storage.get_competition_review.return_value = {
        'id': 5,
        'review_status': 'pending',
        'owner_id': 2,
    }
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition/5/review/reject',
        headers=headers,
        data={**csrf_for(headers), 'comment': 'укажите верное место'},
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.set_competition_review.assert_called_with(5, 'rejected', 'укажите верное место')
    app.ctx.storage.get_competition_review.return_value = None


def test_viewer_cannot_review(client: SanicTestClient):
    headers = get_auth_headers(role='viewer')
    _, response = client.post(
        '/competition/5/review/approve',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403


def test_review_rejects_unknown_decision(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/competition/5/review/delete',
        headers=headers,
        data=csrf_for(headers),
    )
    assert response.status == 400


def test_login_verifies_password_hash(client: SanicTestClient):
    from src.main import login_failures

    login_failures.clear()
    _, response = client.post(
        '/login',
        data={'username': settings.auth_editor_username, 'password': settings.auth_editor_password},
        allow_redirects=False,
    )
    assert response.status == 302

    _, response = client.post(
        '/login',
        data={'username': settings.auth_editor_username, 'password': 'definitely-wrong'},
        allow_redirects=False,
    )
    assert response.status == 401


def test_stale_pwd_ver_cookie_is_rejected(client: SanicTestClient):
    stale_cookie = create_auth_cookie_value(
        username=settings.auth_admin_username,
        role='admin',
        pwd_ver=1,
    )
    _, response = client.get(
        '/',
        headers={'cookie': f'{settings.auth_cookie_name}={stale_cookie}'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_template_date_column_is_typed(client: SanicTestClient):
    from openpyxl import load_workbook

    _, response = client.get('/template/empty.xlsx', headers=get_auth_headers(role='editor'))
    assert response.status == 200
    workbook = load_workbook(BytesIO(response.body))
    sheet = workbook.worksheets[0]
    date_column = get_xlsx_headers(response.body).index('Дата') + 1
    assert sheet.cell(row=1, column=date_column).number_format == 'DD.MM.YYYY'


def test_import_auto_adds_unknown_level(client: SanicTestClient):
    # Само автопополнение уровней — внутри одной транзакции импорта, значения
    # сеются из записей (storage_test: autofills levels and catalogs); здесь —
    # что роут отдаёт импорту вставляемые строки с новым уровнем как есть.
    app.ctx.storage.import_competitions.reset_mock()
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Лыжников Лыж Лыжович',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ЛС-101',
                'Вид спорта': 'Лыжи',
                'Дата': '01.02.2026',
                'Уровень соревнований': 'всероссийские',
                'Название соревнований': 'Чемпионат РФ',
                'Место': 5,
                'Курс': 1,
            }
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )

    assert response.status == 200
    inserted_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert [row.level for row in inserted_rows] == ['всероссийские']


def test_admin_can_create_level(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels',
        headers=headers,
        data={**csrf_for(headers), 'name': 'городские'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.create_level.assert_called_with('городские')


def test_admin_cannot_create_duplicate_level(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels',
        headers=headers,
        data={**csrf_for(headers), 'name': 'внутривузовские'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']


def level_hard_delete_mocks(level_id: int = 9, name: str = 'городские', records: int = 0):
    app.ctx.storage.list_levels.return_value = [{'id': level_id, 'name': name, 'sort_order': 0, 'active': 1}]
    app.ctx.storage.count_records_using.return_value = records
    app.ctx.storage.hard_delete_level.reset_mock()


def test_admin_hard_deletes_empty_level(client: SanicTestClient):
    level_hard_delete_mocks()
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels/9/hard-delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.hard_delete_level.assert_called_once_with(9)
    app.ctx.storage.list_levels.return_value = []


def test_admin_hard_delete_rejects_level_with_records(client: SanicTestClient):
    level_hard_delete_mocks(records=3)
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels/9/hard-delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.hard_delete_level.assert_not_called()
    app.ctx.storage.list_levels.return_value = []
    app.ctx.storage.count_records_using.return_value = 0


def test_level_hard_delete_forbidden_for_non_admin(client: SanicTestClient):
    level_hard_delete_mocks()
    editor_headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/admin/levels/9/hard-delete',
        headers=editor_headers,
        data=csrf_for(editor_headers),
    )
    assert response.status == 403
    app.ctx.storage.hard_delete_level.assert_not_called()
    app.ctx.storage.list_levels.return_value = []
    headers = get_auth_headers()
    _, response = client.get('/admin/catalogs', headers=headers)
    assert response.status == 200
    for title in ('Виды спорта', 'Институты', 'Уровни'):
        assert title in response.text
    assert 'action="/admin/catalogs/sport"' in response.text
    assert 'action="/admin/catalogs/institute"' in response.text
    assert 'action="/admin/levels"' in response.text  # секция уровней перенесена сюда

    _, response = client.get('/admin/catalogs', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = client.get('/admin/catalogs', headers=get_auth_headers(role='viewer'))
    assert response.status == 403

    _, response = client.get('/admin/catalogs', headers=athlete_headers())
    assert response.status == 403

    _, response = client.get('/admin/catalogs', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_old_levels_url_redirects_to_catalogs(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.get('/admin/levels', headers=headers, allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'] == '/admin/catalogs'

    # flash-параметры сохраняются в редиректе (query string как есть)
    _, response = client.get('/admin/levels?admin_message=готово', headers=headers, allow_redirects=False)
    expected_query = urlencode({'admin_message': 'готово'})
    assert response.headers['location'] == f'/admin/catalogs?{expected_query}'


def test_admin_can_add_catalog_value(client: SanicTestClient):
    app.ctx.storage.add_catalog_value.reset_mock()
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/catalogs/sport',
        headers=headers,
        data={**csrf_for(headers), 'value': 'Плавание'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.add_catalog_value.assert_called_once_with('sport', 'Плавание')

    _, response = client.post(
        '/admin/catalogs/sport',
        headers=headers,
        data={**csrf_for(headers), 'value': '   '},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']


def test_catalog_value_actions_forbidden_for_non_admin(client: SanicTestClient):
    editor_headers = get_auth_headers(role='editor')
    for path in (
        '/admin/catalogs/sport',
        '/admin/catalogs/sport/5/hide',
        '/admin/catalogs/sport/5/delete',
        '/admin/catalogs/group/5/hide',
        '/admin/catalogs/group/5/delete',
        '/admin/catalogs/institute/7/group',
    ):
        _, response = client.post(path, headers=editor_headers, data=csrf_for(editor_headers))
        assert response.status == 403, path


def test_catalog_hide_unhide_endpoints(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.return_value = {'id': 5, 'category': 'sport', 'value': 'Лыжи', 'active': 1}
    app.ctx.storage.hide_catalog_value.reset_mock()
    app.ctx.storage.unhide_catalog_value.reset_mock()
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/sport/5/hide', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.hide_catalog_value.assert_called_once_with(5)

    _, response = client.post(
        '/admin/catalogs/sport/5/unhide', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    app.ctx.storage.unhide_catalog_value.assert_called_once_with(5)

    app.ctx.storage.get_catalog_value.return_value = None


def test_catalog_delete_rejects_value_with_records(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.return_value = {'id': 5, 'category': 'sport', 'value': 'Бег', 'active': 1}
    app.ctx.storage.count_records_using.return_value = 3
    app.ctx.storage.delete_catalog_value.reset_mock()
    headers = get_auth_headers()

    _, response = client.post('/admin/catalogs/sport/5/delete', headers=headers, data=csrf_for(headers))
    assert response.status == 400
    app.ctx.storage.delete_catalog_value.assert_not_called()

    app.ctx.storage.count_records_using.return_value = 0
    _, response = client.post(
        '/admin/catalogs/sport/5/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.delete_catalog_value.assert_called_once_with(5)

    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.count_records_using.return_value = 0


# --- Переименование значений справочников (решение 2026-09-13) ---


def rename_sport_value_mock(row_id: int = 5, value: str = 'Бег'):
    app.ctx.storage.get_catalog_value.return_value = {
        'id': row_id,
        'category': 'sport',
        'value': value,
        'parent_id': None,
        'active': 1,
    }


def reset_rename_mocks():
    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.list_catalog_all.return_value = []
    app.ctx.storage.count_records_using.return_value = 0
    app.ctx.storage.list_levels.return_value = []


def test_catalog_rename_preview_returns_records_count(client: SanicTestClient):
    rename_sport_value_mock()
    app.ctx.storage.count_records_using.return_value = 4
    app.ctx.storage.rename_catalog_value.reset_mock()
    headers = get_auth_headers()

    # без confirm — только превью, ничего не выполняется
    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'Кросс'},
    )
    assert response.status == 200
    assert response.json['preview'] is True
    assert response.json['records'] == 4
    assert response.json['old'] == 'Бег'
    assert response.json['new'] == 'Кросс'
    app.ctx.storage.rename_catalog_value.assert_not_called()

    reset_rename_mocks()


def test_catalog_rename_executes_with_update_and_writes_audit(client: SanicTestClient):
    rename_sport_value_mock()
    app.ctx.storage.rename_catalog_value.reset_mock()
    app.ctx.storage.rename_catalog_value.return_value = 3
    app.ctx.storage.add_audit_event.reset_mock()
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'Кросс', 'confirm': 'on', 'update_records': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.rename_catalog_value.assert_called_once_with(
        'sport', 'Бег', 'Кросс', parent_value=None, update_records=True
    )

    audit_call = app.ctx.storage.add_audit_event.call_args
    assert audit_call[1]['action'] == 'catalog_value_renamed'
    details = json.loads(audit_call[1]['details'])
    assert details == {
        'category': 'sport',
        'old': 'Бег',
        'new': 'Кросс',
        'update_records': True,
        'records_updated': 3,
    }

    reset_rename_mocks()
    app.ctx.storage.rename_catalog_value.return_value = 0


def test_catalog_rename_execution_without_update_records_flag(client: SanicTestClient):
    rename_sport_value_mock()
    app.ctx.storage.rename_catalog_value.reset_mock()
    app.ctx.storage.rename_catalog_value.return_value = 0
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'Кросс', 'confirm': 'on'},  # галочка снята
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'записи не тронуты' in unquote_plus(response.headers['location'])
    app.ctx.storage.rename_catalog_value.assert_called_once_with(
        'sport', 'Бег', 'Кросс', parent_value=None, update_records=False
    )

    reset_rename_mocks()


def test_catalog_rename_level_uses_levels_table(client: SanicTestClient):
    app.ctx.storage.list_levels.return_value = [
        {'id': 9, 'name': 'городские', 'sort_order': 0, 'active': 1},
        {'id': 10, 'name': 'внутривузовские', 'sort_order': 0, 'active': 1},
    ]
    app.ctx.storage.rename_catalog_value.reset_mock()
    app.ctx.storage.rename_catalog_value.return_value = 2
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/level/9/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'муниципальные', 'confirm': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.rename_catalog_value.assert_called_once_with(
        'level', 'городские', 'муниципальные', parent_value=None, update_records=False
    )

    reset_rename_mocks()
    app.ctx.storage.rename_catalog_value.return_value = 0


def test_catalog_rename_confirmation_page(client: SanicTestClient):
    rename_sport_value_mock()
    app.ctx.storage.count_records_using.return_value = 7
    headers = get_auth_headers()

    _, response = client.get(f'/admin/catalogs/sport/5/rename?new_name={quote("Кросс")}', headers=headers)
    assert response.status == 200
    assert '«Бег» → «Кросс»' in response.text
    assert 'name="update_records"' in response.text  # галочка «обновить записи»
    assert 'name="confirm"' in response.text

    reset_rename_mocks()


def test_catalog_rename_validations(client: SanicTestClient):
    rename_sport_value_mock()
    app.ctx.storage.list_catalog_all.return_value = [
        {'id': 6, 'category': 'sport', 'value': 'Кросс', 'parent_id': None, 'active': 1}
    ]
    app.ctx.storage.rename_catalog_value.reset_mock()
    headers = get_auth_headers()

    # пустое имя
    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': '   ', 'confirm': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']

    # то же имя
    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'Бег', 'confirm': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']

    # конфликт с существующим значением той же категории
    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'Кросс', 'confirm': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.rename_catalog_value.assert_not_called()

    # неизвестная категория
    _, response = client.post(
        '/admin/catalogs/unknown/5/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'X', 'confirm': 'on'},
    )
    assert response.status == 404

    reset_rename_mocks()


def test_catalog_rename_group_uses_parent_institute(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.side_effect = lambda value_id: {
        7: {'id': 7, 'category': 'institute', 'value': 'ИСИ', 'parent_id': None, 'active': 1},
        8: {'id': 8, 'category': 'group', 'value': 'ПГС-101', 'parent_id': 7, 'active': 1},
    }.get(value_id)
    app.ctx.storage.list_catalog_all.return_value = []
    app.ctx.storage.rename_catalog_value.reset_mock()
    app.ctx.storage.rename_catalog_value.return_value = 2
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/group/8/rename',
        headers=headers,
        data={**csrf_for(headers), 'new_name': 'ПГС-101А', 'confirm': 'on', 'update_records': 'on'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.rename_catalog_value.assert_called_once_with(
        'group', 'ПГС-101', 'ПГС-101А', parent_value='ИСИ', update_records=True
    )
    audit_details = json.loads(app.ctx.storage.add_audit_event.call_args[1]['details'])
    assert audit_details['parent'] == 'ИСИ'

    app.ctx.storage.get_catalog_value.side_effect = None
    reset_rename_mocks()
    app.ctx.storage.rename_catalog_value.return_value = 0


def test_catalog_rename_forbidden_for_non_admin(client: SanicTestClient):
    editor_headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/admin/catalogs/sport/5/rename',
        headers=editor_headers,
        data={**csrf_for(editor_headers), 'new_name': 'Кросс', 'confirm': 'on'},
    )
    assert response.status == 403

    _, response = client.get('/admin/catalogs/sport/5/rename?new_name=Кросс', headers=editor_headers)
    assert response.status == 403


# ---- Перенос группы в другой институт (решение 2026-09-22) ----


def group_move_mock():
    """Иерархия для переноса: группа 11 «ПГС-101» в институте 1 «ИСИ»,
    целевой институт 2 «ФМА». find_catalog_row (конфликт в цели) — None."""
    catalog = {
        1: {'id': 1, 'category': 'institute', 'value': 'ИСИ', 'parent_id': None, 'active': 1},
        2: {'id': 2, 'category': 'institute', 'value': 'ФМА', 'parent_id': None, 'active': 1},
        11: {'id': 11, 'category': 'group', 'value': 'ПГС-101', 'parent_id': 1, 'active': 1},
    }
    app.ctx.storage.get_catalog_value.side_effect = catalog.get
    app.ctx.storage.find_catalog_row.return_value = None
    app.ctx.storage.count_records_using.return_value = 4
    app.ctx.storage.count_students_using_group_pair.return_value = 2
    app.ctx.storage.move_catalog_group.reset_mock()
    app.ctx.storage.move_catalog_group.return_value = {'students_updated': 2, 'competitions_updated': 4}


def reset_group_move_mock():
    app.ctx.storage.get_catalog_value.side_effect = None
    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.find_catalog_row.return_value = None
    app.ctx.storage.count_records_using.return_value = 0
    app.ctx.storage.count_students_using_group_pair.return_value = 0
    app.ctx.storage.move_catalog_group.reset_mock()


def test_catalog_group_move_confirmation_page(client: SanicTestClient):
    group_move_mock()
    headers = get_auth_headers()
    try:
        _, response = client.get('/admin/catalogs/group/11/move?target_institute_id=2', headers=headers)
        assert response.status == 200
        body = response.body.decode()
        assert 'Перенос группы: «ПГС-101» — «ИСИ» → «ФМА»' in body
        assert 'записей о соревнованиях — <strong>4</strong>' in body
        assert 'карточек студентов — <strong>2</strong>' in body
        assert 'Институт группы изменится в справочнике, в перечисленных записях о соревнованиях' in body
        assert 'Да, перенести группу «ПГС-101» в институт «ФМА»' in body
        assert 'name="confirm"' in body
        assert 'Перенести группу</button>' in body
        assert 'href="/admin/catalogs">Отмена' in body
        # Страница только читает — перенос не выполняется
        app.ctx.storage.move_catalog_group.assert_not_called()
    finally:
        reset_group_move_mock()


def test_catalog_group_move_post_without_confirm_changes_nothing(client: SanicTestClient):
    group_move_mock()
    headers = get_auth_headers()
    try:
        _, response = client.post(
            '/admin/catalogs/group/11/move',
            headers=headers,
            data={**csrf_for(headers), 'target_institute_id': '2'},  # без confirm
            allow_redirects=False,
        )
        assert response.status == 302
        assert response.headers['location'].startswith('/admin/catalogs/group/11/move')
        assert 'admin_error' not in response.headers['location']
        app.ctx.storage.move_catalog_group.assert_not_called()
    finally:
        reset_group_move_mock()


def test_catalog_group_move_executes_and_writes_audit(client: SanicTestClient):
    group_move_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    headers = get_auth_headers()
    try:
        _, response = client.post(
            '/admin/catalogs/group/11/move',
            headers=headers,
            data={**csrf_for(headers), 'target_institute_id': '2', 'confirm': 'on'},
            allow_redirects=False,
        )
        assert response.status == 302
        location = unquote_plus(response.headers['location'])
        assert 'admin_error' not in location
        assert 'Группа «ПГС-101» перенесена: «ИСИ» → «ФМА»; обновлено записей — 4, карточек студентов — 2' in location
        app.ctx.storage.move_catalog_group.assert_called_once_with(11, 2)

        audit_call = app.ctx.storage.add_audit_event.call_args
        assert audit_call[1]['action'] == 'catalog_group_moved'
        details = json.loads(audit_call[1]['details'])
        assert details == {
            'group': 'ПГС-101',
            'old_institute': 'ИСИ',
            'new_institute': 'ФМА',
            'students_updated': 2,
            'competitions_updated': 4,
            'group_id': 11,
        }
    finally:
        reset_group_move_mock()


def test_catalog_group_move_validations(client: SanicTestClient):
    group_move_mock()
    headers = get_auth_headers()
    try:
        # Тот же институт — ошибка
        _, response = client.post(
            '/admin/catalogs/group/11/move',
            headers=headers,
            data={**csrf_for(headers), 'target_institute_id': '1', 'confirm': 'on'},
            allow_redirects=False,
        )
        assert response.status == 302
        assert 'Группа уже относится к выбранному институту' in unquote_plus(response.headers['location'])

        # В целевом институте уже есть такая группа — ошибка с подсказкой
        app.ctx.storage.find_catalog_row.return_value = {
            'id': 12,
            'category': 'group',
            'value': 'ПГС-101',
            'parent_id': 2,
            'active': 1,
        }
        _, response = client.post(
            '/admin/catalogs/group/11/move',
            headers=headers,
            data={**csrf_for(headers), 'target_institute_id': '2', 'confirm': 'on'},
            allow_redirects=False,
        )
        assert response.status == 302
        assert (
            'В институте «ФМА» уже есть группа «ПГС-101» — переименуйте одну из групп перед переносом'
            in unquote_plus(response.headers['location'])
        )
        app.ctx.storage.find_catalog_row.return_value = None

        # Неизвестный целевой институт
        _, response = client.post(
            '/admin/catalogs/group/11/move',
            headers=headers,
            data={**csrf_for(headers), 'target_institute_id': '999', 'confirm': 'on'},
            allow_redirects=False,
        )
        assert response.status == 302
        assert 'Группа или институт не найдены' in unquote_plus(response.headers['location'])

        # GET без параметра цели
        _, response = client.get('/admin/catalogs/group/11/move', headers=headers, allow_redirects=False)
        assert response.status == 302
        assert 'Выберите институт для переноса' in unquote_plus(response.headers['location'])

        app.ctx.storage.move_catalog_group.assert_not_called()
    finally:
        reset_group_move_mock()


def test_catalog_group_move_forbidden_for_non_admin(client: SanicTestClient):
    editor_headers = get_auth_headers(role='editor')
    _, response = client.get('/admin/catalogs/group/11/move?target_institute_id=2', headers=editor_headers)
    assert response.status == 403
    _, response = client.post(
        '/admin/catalogs/group/11/move',
        headers=editor_headers,
        data={**csrf_for(editor_headers), 'target_institute_id': '2', 'confirm': 'on'},
    )
    assert response.status == 403


def test_admin_catalogs_page_shows_group_move_form(client: SanicTestClient):
    """Форма переноса в списке групп: активные институты кроме текущего."""
    app.ctx.storage.list_catalog_tree.return_value = [
        {
            'id': 1,
            'category': 'institute',
            'value': 'ИСИ',
            'parent_id': None,
            'active': 1,
            'records_count': 3,
            'groups': [
                {'id': 11, 'category': 'group', 'value': 'ПГС-101', 'parent_id': 1, 'active': 1, 'records_count': 2}
            ],
        },
        {
            'id': 2,
            'category': 'institute',
            'value': 'ФМА',
            'parent_id': None,
            'active': 1,
            'records_count': 0,
            'groups': [],
        },
        {
            'id': 3,
            'category': 'institute',
            'value': 'АРХ',
            'parent_id': None,
            'active': 0,  # скрытый институт не предлагается
            'records_count': 0,
            'groups': [],
        },
    ]
    app.ctx.storage.list_catalog_all.side_effect = lambda category: (
        [
            {'id': 1, 'category': 'institute', 'value': 'ИСИ', 'parent_id': None, 'active': 1},
            {'id': 2, 'category': 'institute', 'value': 'ФМА', 'parent_id': None, 'active': 1},
            {'id': 3, 'category': 'institute', 'value': 'АРХ', 'parent_id': None, 'active': 0},
        ]
        if category == 'institute'
        else []
    )
    try:
        _, response = client.get('/admin/catalogs', headers=get_auth_headers())
        assert response.status == 200
        body = response.body.decode()
        assert 'action="/admin/catalogs/group/11/move"' in body
        assert 'name="target_institute_id"' in body
        assert '<option value="2">ФМА</option>' in body
        assert '<option value="1">' not in body  # текущий институт не предлагается
        assert '<option value="3">' not in body  # скрытый тоже
        assert 'Перенести…</button>' in body
    finally:
        app.ctx.storage.list_catalog_tree.return_value = []
        app.ctx.storage.list_catalog_all.side_effect = None
        app.ctx.storage.list_catalog_all.return_value = []


def test_old_level_rename_route_removed(client: SanicTestClient):
    # Старый POST /admin/levels/<id> (rename без каскада) удалён: UI переехал
    # на /admin/catalogs/<category>/<id>/rename. GET несуществующего пути — 404;
    # POST перехватывается CSRF-middleware до роутера (Sanic не парсит form-body
    # для unmatched-маршрутов) — как у удалённого /clean_db.
    headers = get_auth_headers()
    _, response = client.get('/admin/levels/9', headers=headers, allow_redirects=False)
    assert response.status == 404

    _, response = client.post('/admin/levels/9', headers=headers, data=csrf_for(headers), allow_redirects=False)
    assert response.status == 403


def test_create_competition_auto_adds_catalog_values(client: SanicTestClient):
    app.ctx.storage.add_catalog_value.reset_mock()
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ФМА',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Плавание',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Кубок',
            'position': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    added = {(call[0][0], call[0][1]) for call in app.ctx.storage.add_catalog_value.call_args_list}
    assert ('sport', 'Плавание') in added
    assert ('institute', 'ФМА') in added


def test_import_auto_adds_catalog_values(client: SanicTestClient):
    # Справочники импорт пополняет из вставленных записей внутри
    # storage.import_competitions (одна транзакция с записями); само
    # пополнение проверено на реальном адаптере в storage_test (autofills
    # levels and catalogs). Здесь — что строки файла со значениями доходят
    # до импорта.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Шахматистов Шах Шахович',
                'Пол': 'М',
                'Институт': 'АДИ',
                'Группа': 'ША-101',
                'Вид спорта': 'Шахматы',
                'Дата': '01.02.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Турнир',
                'Место': 2,
                'Курс': 1,
            }
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )
    assert response.status == 200
    inserted_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert [(row.sport, row.institute) for row in inserted_rows] == [('Шахматы', 'АДИ')]


def test_index_datalists_show_active_catalog_values(client: SanicTestClient):
    # фикстура: list_catalog('sport') → ['Бег', 'Лыжи'], list_catalog('institute') → ['ИСИ']
    # отдельная форма «Добавить запись» убрана (ввод — инлайн-строкой), datalist
    # остаётся в разметке: script.js вешает его на поля инлайн-строки по id
    _, response = client.get('/', headers=get_auth_headers(role='editor'))
    assert response.status == 200
    assert '<datalist id="sport-options">' in response.text
    assert '<datalist id="institute-options">' in response.text
    assert '<option value="Бег"></option>' in response.text
    assert '<option value="Лыжи"></option>' in response.text
    assert '<option value="ИСИ"></option>' in response.text
    called_categories = {call[0][0] for call in app.ctx.storage.list_catalog.call_args_list}
    assert {'sport', 'institute'} <= called_categories


def test_hidden_catalog_value_not_in_datalist(client: SanicTestClient):
    # list_catalog отдаёт только активные значения — скрытое в подсказки не попадает
    original_side_effect = app.ctx.storage.list_catalog.side_effect
    app.ctx.storage.list_catalog.side_effect = lambda category: {'sport': ['Бег']}.get(category, [])
    try:
        _, response = client.get('/', headers=get_auth_headers(role='editor'))
        assert response.status == 200
        assert '<option value="Бег"></option>' in response.text
        assert '<option value="Лыжи"></option>' not in response.text
    finally:
        app.ctx.storage.list_catalog.side_effect = original_side_effect


# --- Иерархия справочника: институты содержат группы (№15/№4) ---


def test_index_table_carries_group_hints_by_institute(client: SanicTestClient):
    # карта институт → группы отдаётся таблице в data-атрибуте: комбобокс
    # «Группа» фильтрует список выбранным институтом
    record = Competition(
        record_id='1',
        student_id='hash-1',
        student_name='Иванов Иван Иванович',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 1, 1),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )
    app.ctx.storage.get_competitions_page.return_value = [record]
    app.ctx.storage.count_competitions_filtered.return_value = 1
    app.ctx.storage.get_group_options_by_institute.return_value = {'ИСИ': ['ПГС-101', 'ПГС-102']}
    try:
        _, response = client.get('/', headers=get_auth_headers(role='editor'))
        assert response.status == 200
        assert 'data-groups-by-institute' in response.text
        # tojson экранирует кириллицу в \u-последовательности — сверяем точный JSON
        import json

        assert json.dumps({'ИСИ': ['ПГС-101', 'ПГС-102']}) in response.text
    finally:
        app.ctx.storage.get_competitions_page.return_value = []
        app.ctx.storage.count_competitions_filtered.return_value = 0
        app.ctx.storage.get_group_options_by_institute.return_value = {}


def test_admin_catalogs_page_shows_institute_hierarchy(client: SanicTestClient):
    app.ctx.storage.list_catalog_tree.return_value = [
        {
            'id': 1,
            'category': 'institute',
            'value': 'ИСИ',
            'parent_id': None,
            'active': 1,
            'records_count': 3,
            'groups': [
                {'id': 11, 'category': 'group', 'value': 'ПГС-101', 'parent_id': 1, 'active': 1, 'records_count': 2},
                {'id': 12, 'category': 'group', 'value': 'ПГС-102', 'parent_id': 1, 'active': 0, 'records_count': 0},
            ],
        }
    ]
    try:
        _, response = client.get('/admin/catalogs', headers=get_auth_headers())
        assert response.status == 200
        assert 'Институты и группы' in response.text
        assert 'ПГС-101' in response.text
        assert 'action="/admin/catalogs/institute/1/group"' in response.text
        assert 'action="/admin/catalogs/group/11/hide"' in response.text
        assert 'action="/admin/catalogs/group/12/unhide"' in response.text
        # удаление — только у группы без записей
        assert 'action="/admin/catalogs/group/12/delete"' in response.text
        assert 'action="/admin/catalogs/group/11/delete"' not in response.text
    finally:
        app.ctx.storage.list_catalog_tree.return_value = []


def test_admin_can_add_group_to_institute(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.return_value = {
        'id': 7,
        'category': 'institute',
        'value': 'ИСИ',
        'parent_id': None,
        'active': 1,
    }
    app.ctx.storage.add_catalog_value.reset_mock()
    headers = get_auth_headers()

    _, response = client.post(
        '/admin/catalogs/institute/7/group',
        headers=headers,
        data={**csrf_for(headers), 'value': 'ПГС-101'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.add_catalog_value.assert_called_once_with('group', 'ПГС-101', parent_id=7)

    _, response = client.post(
        '/admin/catalogs/institute/7/group',
        headers=headers,
        data={**csrf_for(headers), 'value': '   '},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.get_catalog_value.return_value = None


def test_group_delete_uses_pair_record_count(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.side_effect = lambda value_id: {
        7: {'id': 7, 'category': 'institute', 'value': 'ИСИ', 'parent_id': None, 'active': 1},
        6: {'id': 6, 'category': 'group', 'value': 'ПГС-101', 'parent_id': 7, 'active': 1},
    }.get(value_id)
    app.ctx.storage.count_records_using.reset_mock()
    app.ctx.storage.count_records_using.return_value = 2
    app.ctx.storage.delete_catalog_value.reset_mock()
    headers = get_auth_headers()

    _, response = client.post('/admin/catalogs/group/6/delete', headers=headers, data=csrf_for(headers))
    assert response.status == 400
    app.ctx.storage.delete_catalog_value.assert_not_called()
    # счётчик — по паре институт+группа, а не по одному названию группы
    app.ctx.storage.count_records_using.assert_called_once_with('group', 'ПГС-101', parent_value='ИСИ')

    app.ctx.storage.count_records_using.return_value = 0
    _, response = client.post(
        '/admin/catalogs/group/6/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    app.ctx.storage.delete_catalog_value.assert_called_once_with(6)

    app.ctx.storage.get_catalog_value.side_effect = None
    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.count_records_using.return_value = 0


def test_institute_delete_blocked_until_groups_removed(client: SanicTestClient):
    app.ctx.storage.get_catalog_value.return_value = {
        'id': 7,
        'category': 'institute',
        'value': 'ИСИ',
        'parent_id': None,
        'active': 1,
    }
    app.ctx.storage.count_records_using.return_value = 0
    app.ctx.storage.count_child_groups.return_value = 2
    app.ctx.storage.delete_catalog_value.reset_mock()
    headers = get_auth_headers()

    _, response = client.post('/admin/catalogs/institute/7/delete', headers=headers, data=csrf_for(headers))
    assert response.status == 400
    assert 'сначала удалите' in response.text
    app.ctx.storage.delete_catalog_value.assert_not_called()

    app.ctx.storage.count_child_groups.return_value = 0
    _, response = client.post(
        '/admin/catalogs/institute/7/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    app.ctx.storage.delete_catalog_value.assert_called_once_with(7)

    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.count_child_groups.return_value = 0


def test_create_competition_auto_adds_catalog_pair(client: SanicTestClient):
    app.ctx.storage.ensure_catalog_pair.reset_mock()
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ФМА',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Плавание',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Кубок',
            'position': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.ensure_catalog_pair.assert_called_once_with('ФМА', 'ПГС-101')


def test_import_auto_adds_catalog_pair(client: SanicTestClient):
    # Пара институт→группа из импорта попадает в иерархию справочника тем же
    # батчем (проверено на реальном адаптере в storage_test); здесь — что
    # строки с парой доходят до импорта целиком.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Шахматистов Шах Шахович',
                'Пол': 'М',
                'Институт': 'АДИ',
                'Группа': 'ША-101',
                'Вид спорта': 'Шахматы',
                'Дата': '01.02.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Турнир',
                'Место': 2,
                'Курс': 1,
            }
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )
    assert response.status == 200
    inserted_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert [(row.institute, row.group) for row in inserted_rows] == [('АДИ', 'ША-101')]


def test_url_custom_field_validated(client: SanicTestClient):
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(field_id=1, key='link', label='Ссылка', field_type='url')
    ]
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Кубок',
            'position': '1',
            'custom__link': 'ftp://bad.example',
        },
        allow_redirects=False,
    )
    assert response.status == 400

    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '10.04.2026',
            'level': 'межвузовские',
            'name': 'Кубок',
            'position': '1',
            'custom__link': 'https://example.com/results',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.get_custom_fields.return_value = []


PNG_BYTES = b'\x89PNG\r\n\x1a\n' + b'0' * 32


def upload_attachment(client, headers, filename, payload):
    return client.post(
        '/competition/3/attachments',
        headers=headers,
        data=csrf_for(headers),
        files={'file': (filename, payload, 'application/octet-stream')},
    )


def test_attachment_upload_and_download(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    app.ctx.storage.get_competition_review.return_value = {'id': 3, 'review_status': 'approved', 'owner_id': 1}
    app.ctx.storage.create_attachment.return_value = 42
    app.ctx.storage.get_attachment.return_value = {
        'id': 42,
        'record_id': 3,
        'filename': 'diploma.png',
        'stored_name': 'stored.png',
        'content_type': 'image/png',
        'size': len(PNG_BYTES),
    }
    headers = get_auth_headers(role='editor')

    _, response = upload_attachment(client, headers, 'diploma.png', PNG_BYTES)
    assert response.status == 200
    assert (tmp_path / 'files' / '3' / 'stored-name-unused').exists() is False  # имя генерится
    app.ctx.storage.create_attachment.assert_called_once()
    stored_name = app.ctx.storage.create_attachment.call_args[1]['stored_name']
    assert (tmp_path / 'files' / '3' / stored_name).is_file()

    (tmp_path / 'files' / '3' / 'stored.png').write_bytes(PNG_BYTES)
    _, response = client.get('/attachment/42', headers=get_auth_headers())
    assert response.status == 200
    assert response.body == PNG_BYTES

    app.ctx.storage.get_competition_review.return_value = None
    app.ctx.storage.get_attachment.return_value = None


def test_attachment_rejects_wrong_extension(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    app.ctx.storage.get_competition_review.return_value = {'id': 3, 'review_status': 'approved', 'owner_id': 1}
    headers = get_auth_headers(role='editor')

    _, response = upload_attachment(client, headers, 'notes.txt', b'hello')
    assert response.status == 400
    app.ctx.storage.get_competition_review.return_value = None


def test_attachment_rejects_fake_signature(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    app.ctx.storage.get_competition_review.return_value = {'id': 3, 'review_status': 'approved', 'owner_id': 1}
    headers = get_auth_headers(role='editor')

    _, response = upload_attachment(client, headers, 'fake.png', b'%PDF-1.4 not a png')
    assert response.status == 400
    app.ctx.storage.get_competition_review.return_value = None


def test_attachment_delete_removes_file(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    stored = tmp_path / 'files' / '3' / 'stored.png'
    stored.parent.mkdir(parents=True)
    stored.write_bytes(PNG_BYTES)
    app.ctx.storage.get_attachment.return_value = {
        'id': 42,
        'record_id': 3,
        'filename': 'diploma.png',
        'stored_name': 'stored.png',
        'content_type': 'image/png',
        'size': len(PNG_BYTES),
    }
    headers = get_auth_headers()
    _, response = client.post('/attachment/42/delete', headers=headers, data=csrf_for(headers), allow_redirects=False)
    assert response.status == 302
    assert not stored.exists()
    app.ctx.storage.get_attachment.return_value = None


def athlete_headers() -> dict[str, str]:
    cookie = create_auth_cookie_value(username='sportik', role='athlete')
    return {'cookie': f'{settings.auth_cookie_name}={cookie}'}


def athlete_headers_for(username: str) -> dict[str, str]:
    """Cookie произвольного аккаунта атлета (фикстуры создают их в БД явно)."""
    cookie = create_auth_cookie_value(username=username, role='athlete')
    return {'cookie': f'{settings.auth_cookie_name}={cookie}'}


ATHLETE_RECORD_DATA = {
    'student_name': 'Спортсменов Спорт Спортович',
    'student_sex': 'М',
    'institute': 'ИСИ',
    'group': 'СБ-101',
    'course': '1',
    'sport': 'Бег',
    'date': '10.04.2026',
    'level': 'внутривузовские',
    'name': 'Кубок',
    'position': '2',
}


def test_athlete_create_goes_to_moderation(client: SanicTestClient):
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    headers = athlete_headers()
    _, response = client.post(
        '/competition',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )
    assert response.status == 302
    kwargs = app.ctx.storage.save_competitions.call_args[1]
    assert kwargs['review_status'] == 'pending'
    app.ctx.storage.get_competitions.return_value = []


def test_athlete_sees_only_own_records(client: SanicTestClient):
    headers = athlete_headers()
    _, response = client.get('/', headers=headers)
    assert response.status == 200
    # Реестр читается через страницу с серверными фильтрами/пагинацией:
    # видимость атлета (только свои записи) задаётся owner_id в обоих
    # запросах — счётчик и выборку страницы.
    page_kwargs = app.ctx.storage.get_competitions_page.call_args[1]
    assert page_kwargs['owner_id'] == 1
    assert page_kwargs['student_id_hashes'] == []
    count_kwargs = app.ctx.storage.count_competitions_filtered.call_args[1]
    assert count_kwargs['owner_id'] == 1
    assert count_kwargs['student_id_hashes'] == []


def test_athlete_cannot_update_foreign_record(client: SanicTestClient):
    app.ctx.storage.get_competition_review.return_value = {
        'id': 9,
        'review_status': 'approved',
        'owner_id': 42,
    }
    headers = athlete_headers()
    _, response = client.post(
        '/competition/9',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )
    assert response.status == 403
    app.ctx.storage.get_competition_review.return_value = None


def test_athlete_cannot_import_review_or_report(client: SanicTestClient):
    headers = athlete_headers()
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403

    _, response = client.post(
        '/competition/5/review/approve',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403

    _, response = client.get('/report', headers=headers)
    assert response.status == 403

    _, response = client.get('/report?group_by=group', headers=headers)
    assert response.status == 403

    _, response = client.get('/export/index', headers=headers)
    assert response.status == 403

    _, response = client.get('/export/report', headers=headers)
    assert response.status == 403

    _, response = client.get('/export/report?group_by=institute', headers=headers)
    assert response.status == 403


def test_profile_save_and_fetch(client: SanicTestClient):
    headers = athlete_headers()
    _, response = client.post(
        '/api/profile',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Спортсменов Спорт Спортович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'СБ-101',
            'course': '2',
        },
    )
    assert response.status == 200
    app.ctx.storage.set_profile.assert_called_once()
    saved = app.ctx.storage.set_profile.call_args[0][1]
    assert saved['group'] == 'СБ-101'

    _, response = client.get('/api/profile', headers=headers)
    assert response.status == 200


def test_profile_rejects_bad_sex_and_course(client: SanicTestClient):
    headers = athlete_headers()
    _, response = client.post(
        '/api/profile',
        headers=headers,
        data={**csrf_for(headers), 'student_sex': 'другое'},
    )
    assert response.status == 400

    _, response = client.post(
        '/api/profile',
        headers=headers,
        data={**csrf_for(headers), 'course': 'второй'},
    )
    assert response.status == 400


def test_athlete_record_uses_profile_defaults(client: SanicTestClient):
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.get_profile.return_value = {
        'student_name': 'Спортсменов Спорт Спортович',
        'student_sex': 'М',
        'institute': 'ИСИ',
        'group': 'СБ-101',
        'course': '2',
    }
    headers = athlete_headers()
    _, response = client.post(
        '/competition',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.group == 'СБ-101'

    # а если в форме данные есть — профиль не подменяет их
    app.ctx.storage.get_profile.return_value = {
        'student_name': 'Другое ФИО',
    }
    _, response = client.post(
        '/competition',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.student_name == ATHLETE_RECORD_DATA['student_name']
    app.ctx.storage.get_profile.return_value = {}


def test_admin_merge_endpoint_preview_and_apply(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/students/merge',
        headers=headers,
        data={**csrf_for(headers), 'from_name': 'Иванова Анна', 'to_name': 'Петрова Анна'},
    )
    assert response.status == 200
    assert response.json['preview'] is True
    assert 'records_to_merge' in response.json

    _, response = client.post(
        '/admin/students/merge',
        headers=headers,
        data={
            **csrf_for(headers),
            'from_name': 'Иванова Анна',
            'to_name': 'Петрова Анна',
            'confirm': 'on',
            'rewrite_names': 'on',
        },
    )
    assert response.status == 200
    app.ctx.storage.merge_students.assert_called_once()
    call = app.ctx.storage.merge_students.call_args
    assert call[1]['new_name'] == 'Петрова Анна'
    app.ctx.storage.carry_name_aliases.assert_called_once_with('Иванова Анна', 'Петрова Анна')
    assert 'aliases_updated' in response.json


def test_admin_can_add_alias_to_user(client: SanicTestClient):
    app.ctx.storage.add_name_alias.reset_mock()
    app.ctx.storage.get_user_by_id.return_value = {'id': 7, 'username': 'anna', 'role': 'athlete'}
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/7/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванова Анна Петровна'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.add_name_alias.assert_called_once_with(7, 'Иванова Анна Петровна')
    app.ctx.storage.get_user_by_id.return_value = None


def test_profile_save_appends_alias(client: SanicTestClient):
    app.ctx.storage.add_name_alias.reset_mock()
    app.ctx.storage.get_name_aliases.return_value = []
    headers = athlete_headers()
    _, response = client.post(
        '/api/profile',
        headers=headers,
        data={**csrf_for(headers), 'student_name': 'Спортсменов Спорт Спортович'},
    )
    assert response.status == 200
    app.ctx.storage.add_name_alias.assert_called_once_with(1, 'Спортсменов Спорт Спортович')
    app.ctx.storage.get_name_aliases.return_value = []


STUDENT_NAMES = ['Абрамов Артём Артёмович', 'Иванов Иван Иванович']


def test_students_list_for_admin_and_editor(client: SanicTestClient):
    app.ctx.storage.get_student_names.return_value = STUDENT_NAMES
    for role in ('admin', 'editor'):
        _, response = client.get('/api/students', headers=get_auth_headers(role))
        assert response.status == 200
        assert response.json == STUDENT_NAMES


def test_students_list_forbidden_for_viewer_and_athlete(client: SanicTestClient):
    _, response = client.get('/api/students', headers=get_auth_headers('viewer'))
    assert response.status == 403

    _, response = client.get('/api/students', headers=athlete_headers())
    assert response.status == 403


def test_students_lookup_requires_auth_and_role(client: SanicTestClient):
    _, response = client.get(
        '/api/students/lookup?name=Иванов Иван Иванович',
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'] == '/login'

    _, response = client.get(
        '/api/students/lookup',
        headers=get_auth_headers('viewer'),
        params={'name': 'Иванов Иван Иванович'},
    )
    assert response.status == 403


def test_students_lookup_exact_match(client: SanicTestClient):
    known_name = 'Иванов Иван Иванович'
    known_hash = sha256(known_name.encode()).hexdigest()
    app.ctx.storage.count_records_by_student_hash.side_effect = lambda value: 1 if value == known_hash else 0
    headers = athlete_headers()
    _, response = client.get(
        '/api/students/lookup',
        headers=headers,
        params={'name': known_name},
    )
    assert response.status == 200
    assert response.json == {'found': True}

    # опечатка или другой регистр — уже другой sha256, совпадения нет
    _, response = client.get(
        '/api/students/lookup',
        headers=headers,
        params={'name': 'Иванов Иван Ивановеч'},
    )
    assert response.status == 200
    assert response.json == {'found': False}

    _, response = client.get(
        '/api/students/lookup',
        headers=headers,
        params={'name': 'иванов иван иванович'},
    )
    assert response.status == 200
    assert response.json == {'found': False}
    app.ctx.storage.count_records_by_student_hash.side_effect = None


def test_students_lookup_requires_name_param(client: SanicTestClient):
    _, response = client.get('/api/students/lookup', headers=get_auth_headers())
    assert response.status == 400


def audit_calls():
    import json as json_module

    return [
        (call[1]['action'], json_module.loads(call[1]['details']))
        for call in app.ctx.storage.add_audit_event.call_args_list
    ]


def test_login_success_writes_audit_event(client: SanicTestClient):
    from src.main import login_failures

    login_failures.clear()
    app.ctx.storage.add_audit_event.reset_mock()
    _, response = client.post(
        '/login',
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.add_audit_event.assert_called_once()
    kwargs = app.ctx.storage.add_audit_event.call_args[1]
    assert kwargs['action'] == 'login_success'
    assert kwargs['username'] == settings.auth_admin_username
    assert kwargs['user_id'] == 1


def test_login_failure_writes_audit_event_without_user_id(client: SanicTestClient):
    from src.main import login_failures

    login_failures.clear()
    app.ctx.storage.add_audit_event.reset_mock()
    _, response = client.post(
        '/login',
        data={'username': settings.auth_admin_username, 'password': 'wrong'},
        allow_redirects=False,
    )
    assert response.status == 401
    app.ctx.storage.add_audit_event.assert_called_once()
    kwargs = app.ctx.storage.add_audit_event.call_args[1]
    assert kwargs['action'] == 'login_failed'
    assert kwargs['username'] == settings.auth_admin_username
    assert kwargs['user_id'] is None


def test_audit_failure_does_not_break_login(client: SanicTestClient):
    from src.main import login_failures

    login_failures.clear()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.add_audit_event.side_effect = RuntimeError('audit db is down')
    _, response = client.post(
        '/login',
        data={
            'username': settings.auth_admin_username,
            'password': settings.auth_admin_password,
        },
        allow_redirects=False,
    )
    app.ctx.storage.add_audit_event.side_effect = None
    assert response.status == 302


def test_merge_writes_audit_event(client: SanicTestClient):
    app.ctx.storage.add_audit_event.reset_mock()
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/students/merge',
        headers=headers,
        data={
            **csrf_for(headers),
            'from_name': 'Иванова Анна',
            'to_name': 'Петрова Анна',
            'confirm': 'on',
            'rewrite_names': 'on',
        },
    )
    assert response.status == 200
    assert (
        'students_merged',
        {
            'from_name': 'Иванова Анна',
            'to_name': 'Петрова Анна',
            'merged': 2,
            'rewrite_names': True,
            'aliases_updated': 1,
        },
    ) in audit_calls()


def test_review_decisions_write_audit_events(client: SanicTestClient):
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.get_competition_review.return_value = {
        'id': 5,
        'review_status': 'pending',
        'owner_id': 2,
    }
    headers = get_auth_headers(role='editor')

    _, response = client.post(
        '/competition/5/review/approve',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    assert ('record_approved', {'record_id': 5, 'comment': ''}) in audit_calls()

    app.ctx.storage.add_audit_event.reset_mock()
    _, response = client.post(
        '/competition/5/review/reject',
        headers=headers,
        data={**csrf_for(headers), 'comment': 'укажите верное место'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert ('record_rejected', {'record_id': 5, 'comment': 'укажите верное место'}) in audit_calls()
    app.ctx.storage.get_competition_review.return_value = None


def test_alias_added_writes_audit_event(client: SanicTestClient):
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.get_user_by_id.return_value = {'id': 7, 'username': 'anna', 'role': 'athlete'}
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/7/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванова Анна Петровна'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert (
        'alias_added',
        {
            'target_user_id': 7,
            'target_username': 'anna',
            'name': 'Иванова Анна Петровна',
        },
    ) in audit_calls()
    app.ctx.storage.get_user_by_id.return_value = None


def test_password_change_writes_audit_event(client: SanicTestClient):
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.get_user_by_id.return_value = {
        'id': 7,
        'username': 'someone',
        'password_hash': 'x',
        'role': 'editor',
        'active': 1,
    }
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/users/7/password',
        headers=headers,
        data={**csrf_for(headers), 'password': 'newsecret123'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert ('password_changed', {'target_user_id': 7, 'target_username': 'someone'}) in audit_calls()
    app.ctx.storage.get_user_by_id.return_value = None


def test_audit_page_available_for_admin_only(client: SanicTestClient):
    _, response = client.get('/admin/audit', headers=get_auth_headers())
    assert response.status == 200
    assert 'Журнал безопасности' in response.text

    _, response = client.get('/admin/audit', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = client.get('/admin/audit', headers=get_auth_headers(role='viewer'))
    assert response.status == 403

    _, response = client.get('/admin/audit', headers=athlete_headers())
    assert response.status == 403

    _, response = client.get('/admin/audit', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_audit_page_filters_and_pagination(client: SanicTestClient):
    app.ctx.storage.count_audit_events.return_value = 120
    app.ctx.storage.get_audit_events.reset_mock()
    app.ctx.storage.count_audit_events.reset_mock()
    app.ctx.storage.get_audit_events.return_value = []
    app.ctx.storage.list_audit_actions.return_value = ['db_wiped', 'login_success']

    _, response = client.get(
        '/admin/audit?page=2&user=admin&action=login_success&date_from=01.01.2026&date_to=31.01.2026',
        headers=get_auth_headers(),
    )

    assert response.status == 200
    assert 'Стр. 2 из 3' in response.text
    assert 'page=1' in response.text  # стрелка назад с сохранением фильтров
    assert 'page=3' in response.text  # стрелка вперёд
    app.ctx.storage.count_audit_events.assert_called_once_with(
        username='admin', action='login_success', date_from='2026-01-01', date_to='2026-01-31'
    )
    app.ctx.storage.get_audit_events.assert_called_once_with(
        limit=50,
        offset=50,
        username='admin',
        action='login_success',
        date_from='2026-01-01',
        date_to='2026-01-31',
    )
    assert 'login_success' in response.text  # выпадающий фильтр действий

    app.ctx.storage.count_audit_events.return_value = 0
    app.ctx.storage.list_audit_actions.return_value = []


def test_audit_page_rejects_invalid_filter_date(client: SanicTestClient):
    _, response = client.get('/admin/audit?date_from=31.31.2026', headers=get_auth_headers())
    assert response.status == 400


ADMIN_SECTION_PAGES = (
    '/admin/import',
    '/admin/fields',
    '/admin/catalogs',
    '/admin/users',
    '/admin/students',
    '/admin/maintenance',
    '/admin/maintenance/export',
)


def test_admin_hub_shows_all_sections_for_admin(client: SanicTestClient):
    _, response = client.get('/admin', headers=get_auth_headers())
    assert response.status == 200
    for path in ADMIN_SECTION_PAGES + ('/admin/audit',):
        if path in ('/admin/maintenance/export', '/admin/students'):
            # хаб ссылается на раздел, не на выгрузку; объединение студентов
            # (IA 2026-09) — внутри раздела «Студенты», а не плитка хаба
            continue
        assert path in response.text


def test_admin_hub_shows_only_import_for_editor(client: SanicTestClient):
    _, response = client.get('/admin', headers=get_auth_headers(role='editor'))
    assert response.status == 200
    assert '/admin/import' in response.text
    for path in ('/admin/fields', '/admin/catalogs', '/admin/users', '/admin/students', '/admin/maintenance'):
        assert path not in response.text


def test_admin_hub_access_for_viewer_athlete_and_anonymous(client: SanicTestClient):
    _, response = client.get('/admin', headers=get_auth_headers(role='viewer'))
    assert response.status == 403

    _, response = client.get('/admin', headers=athlete_headers())
    assert response.status == 403

    _, response = client.get('/admin', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'] == '/login'


# IA редизайна 2026-09: браузерной навигации (Accept: text/html) вместо
# голого text() отдаются страница «Доступ запрещён» (403) и редирект на
# вход (401); программным клиентам (fetch, Accept: */*) — как раньше.


def test_admin_forbidden_renders_html_page_for_browser(client: SanicTestClient):
    headers = {**get_auth_headers(role='editor'), 'accept': 'text/html'}
    _, response = client.get('/admin/users', headers=headers, allow_redirects=False)
    assert response.status == 403
    assert response.headers['content-type'].startswith('text/html')
    assert 'Доступ запрещён' in response.text
    assert 'У вашей роли нет доступа к этому разделу.' in response.text


def test_admin_forbidden_stays_plain_text_for_fetch_like_client(client: SanicTestClient):
    _, response = client.get('/admin/users', headers=get_auth_headers(role='editor'), allow_redirects=False)
    assert response.status == 403
    assert response.text == 'Forbidden'
    assert response.headers['content-type'].startswith('text/plain')


def test_admin_unauthorized_redirects_browser_to_login(client: SanicTestClient):
    _, response = client.get('/admin/users', headers={'accept': 'text/html'}, allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'] == '/login'

    _, response = client.post(
        '/admin/users',
        headers={'accept': 'text/html'},
        data={'anything': '1'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'] == '/login'


def test_admin_section_pages_available_for_admin(client: SanicTestClient):
    headers = get_auth_headers()
    for path in ADMIN_SECTION_PAGES:
        _, response = client.get(path, headers=headers)
        assert response.status == 200, path


def test_admin_import_page_available_for_editor(client: SanicTestClient):
    _, response = client.get('/admin/import', headers=get_auth_headers(role='editor'))
    assert response.status == 200
    assert 'import-form' in response.text
    assert '/template/empty.xlsx' in response.text


def test_admin_sections_forbidden_for_editor(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    for path in ADMIN_SECTION_PAGES:
        if path == '/admin/import':
            continue
        _, response = client.get(path, headers=headers)
        assert response.status == 403, path


def test_admin_sections_forbidden_for_viewer_and_athlete(client: SanicTestClient):
    for headers in (get_auth_headers(role='viewer'), athlete_headers()):
        for path in ADMIN_SECTION_PAGES:
            _, response = client.get(path, headers=headers)
            assert response.status == 403, path


def test_admin_pages_render_own_markup(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.get('/admin/students', headers=headers)
    assert response.status == 200
    assert 'merge-form' in response.text
    assert 'student-names-list' in response.text

    _, response = client.get('/admin/maintenance', headers=headers)
    assert response.status == 200
    assert '/admin/maintenance/export' in response.text
    assert 'УДАЛИТЬ' in response.text

    _, response = client.get('/admin/users', headers=headers)
    assert response.status == 200
    assert 'student-names-list' in response.text


def test_index_shows_import_link_for_moderators_only(client: SanicTestClient):
    _, response = client.get('/', headers=get_auth_headers(role='editor'))
    assert response.status == 200
    assert 'href="/admin/import"' in response.text

    _, response = client.get('/', headers=athlete_headers())
    assert response.status == 200
    assert 'href="/admin/import"' not in response.text


def wipe_request(
    client,
    headers,
    scope=None,
    mode='scope',
    date_before=None,
    with_attachments=None,
    phrase='УДАЛИТЬ',
    confirm='1',
):
    data = {**csrf_for(headers), 'mode': mode, 'confirm': confirm, 'confirm_phrase': phrase}
    if scope is not None:
        data['scope'] = scope
    if date_before is not None:
        data['date_before'] = date_before
    if with_attachments is not None:
        data['with_attachments'] = with_attachments
    return client.post(
        '/admin/maintenance/wipe',
        headers=headers,
        data=data,
        allow_redirects=False,
    )


def make_wipe_record(record_id, name, month):
    return Competition(
        record_id=str(record_id),
        student_id=f'hash-{record_id}',
        student_name=name,
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2025, month, 1),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )


def make_wipe_attachment(record_id, stored_name, filename='diploma.png'):
    return {
        'id': record_id * 10,
        'record_id': record_id,
        'filename': filename,
        'stored_name': stored_name,
        'content_type': 'image/png',
        'size': len(PNG_BYTES),
    }


def wipe_events(action):
    return [details for act, details in audit_calls() if act == action]


def test_maintenance_page_shows_status_panel(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    monkeypatch.setattr(main_module.settings, 'database_path', str(tmp_path / 'competitions.sqlite3'))
    (tmp_path / 'competitions.sqlite3').write_bytes(b'x' * 2048)
    (tmp_path / 'competitions.sqlite3-wal').write_bytes(b'y' * 1024)
    backups = tmp_path / 'backups'
    backups.mkdir()
    (backups / 'competitions-20260101-000000.sqlite3.gz').write_bytes(b'z' * 512)

    _, response = client.get('/admin/maintenance', headers=get_auth_headers())

    assert response.status == 200
    assert 'Диск:' in response.text  # полоса + подпись места на диске (№16б)
    assert '2 КБ' in response.text  # файл БД
    assert '1 КБ' in response.text  # WAL
    assert '512 Б' in response.text  # вложения/бэкапы или сам бэкап
    assert 'competitions-20260101-000000.sqlite3.gz' in response.text  # последний бэкап
    assert 'Сделать бэкап сейчас' in response.text
    assert '/admin/maintenance/backup' in response.text
    assert 'Архивы очисток' in response.text


def test_wipe_records_scope(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    app.ctx.storage.get_competitions.return_value = [make_wipe_record(1, 'Первый', 1)]
    app.ctx.storage.delete_all_competitions.reset_mock()
    app.ctx.storage.delete_all_attachments.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.delete_all_competitions.return_value = 5

    _, response = wipe_request(client, get_auth_headers(), 'records')

    assert response.status == 302
    assert response.headers['location'].startswith('/admin/maintenance?')
    app.ctx.storage.delete_all_competitions.assert_called_once()
    app.ctx.storage.delete_all_attachments.assert_not_called()
    events = wipe_events('db_wiped')
    assert events[0]['scope'] == 'records'
    assert events[0]['records_deleted'] == 5
    assert events[0]['attachments_deleted'] == 0
    # архив удалённых записей: только xlsx, без zip (вложения не трогаются)
    assert len(list((tmp_path / 'backups').glob('pre-wipe-*-records.xlsx'))) == 1
    assert not list((tmp_path / 'backups').glob('pre-wipe-*.zip'))
    archive_events = wipe_events('pre_wipe_archive_created')
    assert archive_events[0]['scope'] == 'records'
    assert archive_events[0]['archive_xlsx']
    assert archive_events[0]['archive_zip'] is None
    app.ctx.storage.get_competitions.return_value = []


def test_wipe_attachments_scope_removes_files_only(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    files_dir = tmp_path / 'files' / '3'
    files_dir.mkdir(parents=True)
    (files_dir / 'stored.png').write_bytes(PNG_BYTES)
    keep_me = tmp_path / 'competitions.sqlite3'
    keep_me.write_bytes(b'db')
    app.ctx.storage.get_attachments.return_value = [make_wipe_attachment(3, 'stored.png')]
    app.ctx.storage.delete_all_competitions.reset_mock()
    app.ctx.storage.delete_all_attachments.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.delete_all_attachments.return_value = 2

    _, response = wipe_request(client, get_auth_headers(), 'attachments')

    assert response.status == 302
    app.ctx.storage.delete_all_attachments.assert_called_once()
    app.ctx.storage.delete_all_competitions.assert_not_called()
    assert not (tmp_path / 'files').exists()
    assert keep_me.exists()
    events = wipe_events('db_wiped')
    assert events[0]['scope'] == 'attachments'
    assert events[0]['records_deleted'] == 0
    assert events[0]['attachments_deleted'] == 2
    assert events[0]['freed_bytes'] >= len(PNG_BYTES)
    # архив: только zip вложений, без xlsx (записи не удаляются)
    assert not list((tmp_path / 'backups').glob('pre-wipe-*.xlsx'))
    zip_archives = list((tmp_path / 'backups').glob('pre-wipe-*-attachments.zip'))
    assert len(zip_archives) == 1
    import zipfile as zipfile_module

    with zipfile_module.ZipFile(zip_archives[0]) as archive:
        assert archive.namelist() == ['3/diploma.png']
    app.ctx.storage.get_attachments.return_value = []


def test_wipe_all_scope_creates_xlsx_and_zip_archives(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    record_dir = tmp_path / 'files' / '1'
    record_dir.mkdir(parents=True)
    (record_dir / 'stored.png').write_bytes(PNG_BYTES)
    app.ctx.storage.get_competitions.return_value = [make_wipe_record(1, 'Первый', 1)]
    app.ctx.storage.get_attachments.return_value = [make_wipe_attachment(1, 'stored.png')]
    app.ctx.storage.delete_all_competitions.reset_mock()
    app.ctx.storage.delete_all_attachments.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.delete_all_competitions.return_value = 4
    app.ctx.storage.delete_all_attachments.return_value = 3

    _, response = wipe_request(client, get_auth_headers(), 'all')

    assert response.status == 302
    app.ctx.storage.delete_all_competitions.assert_called_once()
    app.ctx.storage.delete_all_attachments.assert_called_once()
    assert not (tmp_path / 'files').exists()
    events = wipe_events('db_wiped')
    assert events[0]['scope'] == 'all'
    assert events[0]['records_deleted'] == 4
    assert events[0]['attachments_deleted'] == 3
    backups = tmp_path / 'backups'
    assert len(list(backups.glob('pre-wipe-*-all.xlsx'))) == 1
    assert len(list(backups.glob('pre-wipe-*-all.zip'))) == 1
    archive_events = wipe_events('pre_wipe_archive_created')
    assert archive_events[0]['scope'] == 'all'
    assert archive_events[0]['records'] == 1
    assert archive_events[0]['attachments'] == 1
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.get_attachments.return_value = []


def test_wipe_preview_returns_counts_without_confirm(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    record_dir = tmp_path / 'files' / '3'
    record_dir.mkdir(parents=True)
    (record_dir / 'stored.png').write_bytes(PNG_BYTES)
    app.ctx.storage.get_attachments.return_value = [make_wipe_attachment(3, 'stored.png')]
    app.ctx.storage.delete_all_attachments.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()

    _, response = wipe_request(client, get_auth_headers(), 'attachments', confirm='')

    assert response.status == 200
    payload = response.json
    assert payload['records'] == 0
    assert payload['attachments'] == 1
    assert payload['attachments_size_bytes'] == len(PNG_BYTES)
    assert payload['attachments_size']
    app.ctx.storage.delete_all_attachments.assert_not_called()
    assert wipe_events('db_wiped') == []
    app.ctx.storage.get_attachments.return_value = []


def test_wipe_by_date_deletes_older_records_and_their_attachments(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    monkeypatch.setattr(main_module.settings, 'database_path', str(tmp_path / 'competitions.sqlite3'))
    (tmp_path / 'competitions.sqlite3').write_bytes(b'db')
    for record_id, stored in ((1, 'a.png'), (1, 'b.png'), (2, 'c.png')):
        record_dir = tmp_path / 'files' / str(record_id)
        record_dir.mkdir(parents=True, exist_ok=True)
        (record_dir / stored).write_bytes(PNG_BYTES)
    keep_dir = tmp_path / 'files' / '9'
    keep_dir.mkdir(parents=True)
    (keep_dir / 'keep.png').write_bytes(PNG_BYTES)

    old_records = [make_wipe_record(1, 'Старый', 1), make_wipe_record(2, 'Старая', 2)]
    app.ctx.storage.get_competitions_before.return_value = old_records
    app.ctx.storage.get_attachments_for_records.return_value = [
        make_wipe_attachment(1, 'a.png'),
        make_wipe_attachment(1, 'b.png'),
        make_wipe_attachment(2, 'c.png'),
    ]
    app.ctx.storage.delete_competitions_before.reset_mock()
    app.ctx.storage.delete_attachments_for_records.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.delete_competitions_before.return_value = 2
    app.ctx.storage.delete_attachments_for_records.return_value = 3

    _, response = wipe_request(client, get_auth_headers(), mode='date', date_before='01.06.2025')

    assert response.status == 302
    app.ctx.storage.get_competitions_before.assert_called_once_with(datetime(2025, 6, 1))
    app.ctx.storage.delete_competitions_before.assert_called_once_with(datetime(2025, 6, 1))
    app.ctx.storage.delete_attachments_for_records.assert_called_once_with([1, 2])
    # файлы вложений удалённых записей ушли, чужая запись осталась
    assert not (tmp_path / 'files' / '1').exists()
    assert not (tmp_path / 'files' / '2').exists()
    assert (tmp_path / 'files' / '9' / 'keep.png').exists()
    events = wipe_events('db_wiped')
    assert events[0]['scope'] == 'date'
    assert events[0]['records_deleted'] == 2
    assert events[0]['attachments_deleted'] == 3
    assert events[0]['date_before'] == '01.06.2025'
    assert events[0]['with_attachments'] is True
    assert events[0]['freed_bytes'] >= 3 * len(PNG_BYTES)
    backups = tmp_path / 'backups'
    assert len(list(backups.glob('pre-wipe-*-date.xlsx'))) == 1
    assert len(list(backups.glob('pre-wipe-*-date.zip'))) == 1
    app.ctx.storage.get_competitions_before.return_value = []
    app.ctx.storage.get_attachments_for_records.return_value = []


def test_wipe_by_date_keeps_attachments_when_checkbox_off(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    record_dir = tmp_path / 'files' / '1'
    record_dir.mkdir(parents=True)
    (record_dir / 'a.png').write_bytes(PNG_BYTES)
    app.ctx.storage.get_competitions_before.return_value = [make_wipe_record(1, 'Старый', 1)]
    app.ctx.storage.get_attachments_for_records.reset_mock()
    app.ctx.storage.delete_attachments_for_records.reset_mock()
    app.ctx.storage.delete_competitions_before.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.delete_competitions_before.return_value = 1

    # скрытый маркер 0 без чекбокса 1 — галочка снята
    _, response = wipe_request(client, get_auth_headers(), mode='date', date_before='01.06.2025', with_attachments='0')

    assert response.status == 302
    app.ctx.storage.get_attachments_for_records.assert_not_called()
    app.ctx.storage.delete_attachments_for_records.assert_not_called()
    assert (tmp_path / 'files' / '1' / 'a.png').exists()  # вложения не тронуты
    assert len(list((tmp_path / 'backups').glob('pre-wipe-*-date.xlsx'))) == 1
    assert not list((tmp_path / 'backups').glob('pre-wipe-*.zip'))
    events = wipe_events('db_wiped')
    assert events[0]['with_attachments'] is False
    assert events[0]['attachments_deleted'] == 0
    app.ctx.storage.get_competitions_before.return_value = []


def test_wipe_by_date_preview_counts_affected(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    record_dir = tmp_path / 'files' / '2'
    record_dir.mkdir(parents=True)
    (record_dir / 'c.png').write_bytes(PNG_BYTES)
    app.ctx.storage.get_competitions_before.return_value = [
        make_wipe_record(1, 'Старый', 1),
        make_wipe_record(2, 'Старая', 2),
    ]
    app.ctx.storage.get_attachments_for_records.return_value = [make_wipe_attachment(2, 'c.png')]
    app.ctx.storage.delete_competitions_before.reset_mock()

    _, response = wipe_request(client, get_auth_headers(), mode='date', date_before='01.06.2025', confirm='')

    assert response.status == 200
    payload = response.json
    assert payload['records'] == 2
    assert payload['attachments'] == 1
    assert payload['attachments_size_bytes'] == len(PNG_BYTES)
    app.ctx.storage.delete_competitions_before.assert_not_called()
    app.ctx.storage.get_competitions_before.return_value = []
    app.ctx.storage.get_attachments_for_records.return_value = []


def test_wipe_rejects_wrong_confirm_phrase(client: SanicTestClient):
    app.ctx.storage.delete_all_competitions.reset_mock()
    _, response = wipe_request(client, get_auth_headers(), 'records', phrase='удалить')
    assert response.status == 400
    app.ctx.storage.delete_all_competitions.assert_not_called()


def test_wipe_rejects_unknown_scope_and_mode(client: SanicTestClient):
    _, response = wipe_request(client, get_auth_headers(), 'everything')
    assert response.status == 400

    _, response = wipe_request(client, get_auth_headers(), mode='date', date_before='июнь')
    assert response.status == 400

    _, response = wipe_request(client, get_auth_headers(), mode='yesterday')
    assert response.status == 400


def test_wipe_forbidden_for_non_admin(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = wipe_request(client, headers, 'records')
    assert response.status == 403

    viewer_headers = get_auth_headers(role='viewer')
    _, response = client.post(
        '/admin/maintenance/wipe',
        headers=viewer_headers,
        data={**csrf_for(viewer_headers), 'mode': 'scope', 'scope': 'records', 'confirm_phrase': 'УДАЛИТЬ'},
    )
    assert response.status == 403


def make_tmp_database(db_path):
    import sqlite3 as sqlite3_module

    connection = sqlite3_module.connect(db_path)
    connection.execute(
        'CREATE TABLE audit_log (id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT NOT NULL, '
        'user_id INTEGER, username TEXT NOT NULL, action TEXT NOT NULL, details TEXT NOT NULL DEFAULT \'\')'
    )
    connection.commit()
    connection.close()


def test_manual_backup_creates_archive_and_audit_event(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    db_path = tmp_path / 'competitions.sqlite3'
    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    monkeypatch.setattr(main_module.settings, 'database_path', str(db_path))
    make_tmp_database(db_path)
    record_dir = tmp_path / 'files' / '5'
    record_dir.mkdir(parents=True)
    (record_dir / 'stored.png').write_bytes(PNG_BYTES)
    app.ctx.storage.add_audit_event.reset_mock()

    headers = get_auth_headers()
    _, response = client.post(
        '/admin/maintenance/backup',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )

    assert response.status == 302
    assert response.headers['location'].startswith('/admin/maintenance?')
    backups = tmp_path / 'backups'
    db_archives = list(backups.glob('competitions-*.sqlite3.gz'))
    assert len(db_archives) == 1
    assert db_archives[0].stat().st_size > 0
    files_archives = list(backups.glob('competitions-*-files.zip'))
    assert len(files_archives) == 1  # вложения заархивированы тем же кодом, что и скриптом
    events = [details for action, details in audit_calls() if action == 'backup_created']
    assert len(events) == 1
    assert events[0]['source'] == 'manual'
    assert events[0]['archive'] == db_archives[0].name
    assert events[0]['files_archive'] == files_archives[0].name


def test_manual_backup_forbidden_for_editor(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.post('/admin/maintenance/backup', headers=headers, data=csrf_for(headers))
    assert response.status == 403


def test_download_backup_file(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    backups = tmp_path / 'backups'
    backups.mkdir()
    (backups / 'pre-wipe-20260101-000000-date.xlsx').write_bytes(b'xlsx-content')

    _, response = client.get(
        '/admin/maintenance/backups/pre-wipe-20260101-000000-date.xlsx', headers=get_auth_headers()
    )
    assert response.status == 200
    assert response.body == b'xlsx-content'
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


def test_download_backup_rejects_traversal_and_missing(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))

    _, response = client.get('/admin/maintenance/backups/..%2F..%2Fcompetitions.sqlite3', headers=get_auth_headers())
    assert response.status in (400, 404)

    _, response = client.get('/admin/maintenance/backups/no-such-file.zip', headers=get_auth_headers())
    assert response.status == 404

    _, response = client.get('/admin/maintenance/backups/.env', headers=get_auth_headers())
    assert response.status in (400, 404)

    headers = get_auth_headers(role='editor')
    _, response = client.get('/admin/maintenance/backups/whatever.zip', headers=headers)
    assert response.status == 403


def test_maintenance_export_for_admin(client: SanicTestClient):
    from openpyxl import load_workbook

    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.get_custom_fields.return_value = []
    app.ctx.storage.list_users.return_value = [
        {'id': 1, 'username': 'admin', 'role': 'admin', 'active': 1, 'name_aliases': ['Иванов Иван']},
    ]
    app.ctx.storage.list_levels.return_value = [{'id': 1, 'name': 'внутривузовские', 'sort_order': 0, 'active': 1}]
    app.ctx.storage.get_sport_names.return_value = ['Бег', 'Лыжи']

    _, response = client.get('/admin/maintenance/export', headers=get_auth_headers())

    assert response.status == 200
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    workbook = load_workbook(BytesIO(response.body), read_only=True)
    assert workbook.sheetnames == ['Записи', 'Пользователи', 'Уровни', 'Виды спорта']
    users_sheet = workbook['Пользователи']
    rows = list(users_sheet.iter_rows(values_only=True))
    assert rows[0] == ('Логин', 'Роль', 'Активен', 'Псевдонимы ФИО')
    assert rows[1][0] == 'admin'
    assert rows[1][3] == 'Иванов Иван'
    sports_sheet = workbook['Виды спорта']
    assert [row[0] for row in sports_sheet.iter_rows(min_row=2, values_only=True)] == ['Бег', 'Лыжи']

    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.get_custom_fields.return_value = []
    app.ctx.storage.list_users.return_value = []
    app.ctx.storage.list_levels.return_value = []
    app.ctx.storage.get_sport_names.return_value = []


def test_maintenance_export_forbidden_for_editor(client: SanicTestClient):
    _, response = client.get('/admin/maintenance/export', headers=get_auth_headers(role='editor'))
    assert response.status == 403


def test_audit_page_shows_message_when_log_unavailable(client: SanicTestClient):
    app.ctx.storage.get_audit_events.side_effect = RuntimeError('audit db is down')
    _, response = client.get('/admin/audit', headers=get_auth_headers())
    app.ctx.storage.get_audit_events.side_effect = None

    assert response.status == 200
    assert 'Журнал недоступен' in response.text


def test_old_admin_post_still_works_and_returns_to_section(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields',
        headers=headers,
        data={**csrf_for(headers), 'label': 'Тренер', 'field_type': 'text'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'].startswith('/admin/fields')


# --- Волна 0 живого фидбека: баг «nan», №1.1 регистронезависимость, №19б аккордеон ---


def make_import_row(**overrides) -> dict:
    row = {
        'ФИО': 'Тестов Тест Тестович',
        'Пол': 'М',
        'Институт': 'ИСИ',
        'Группа': 'ПГС-101',
        'Вид спорта': 'Бег',
        'Дата': '15.03.2026',
        'Уровень соревнований': 'внутривузовские',
        'Название соревнований': 'Кубок',
        'Место': 1,
        'Курс': 2,
    }
    row.update(overrides)
    return row


def upload_xlsx(client: SanicTestClient, rows: list[dict]):
    df = pd.DataFrame(rows)
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    return response


def test_upload_empty_cells_become_empty_not_nan(client: SanicTestClient):
    # Баг с живого импорта (docs/feedback-live.md «nan»): пустые ячейки Excel
    # (pandas NaN) не должны превращаться в текст «nan» в значениях записи.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    response = upload_xlsx(client, [make_import_row(**{'Пол': None, 'Институт': None, 'Группа': None})])

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.student_sex == ''
    assert saved.institute == ''
    assert saved.group == ''
    assert saved.sport == 'Бег'
    app.ctx.storage.get_competitions.return_value = []


def test_upload_empty_level_becomes_empty_not_nan(client: SanicTestClient):
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    response = upload_xlsx(client, [make_import_row(**{'Уровень соревнований': None})])

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.level == ''
    app.ctx.storage.get_competitions.return_value = []


def test_upload_case_variants_get_canonical_catalog_values(client: SanicTestClient):
    # №1.1: значение, отличающееся от справочника только регистром,
    # заменяется каноническим написанием ещё на входе в запись.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []

    def fake_canonical(category, value, parent_id=None):
        by_category = {
            'institute': {'иси': 'ИСИ'},
            'sport': {'бег': 'Бег'},
            'group': {'пгс-101': 'ПГС-101'},
        }
        return by_category.get(category, {}).get(value.strip().lower())

    app.ctx.storage.find_catalog_canonical.side_effect = fake_canonical
    app.ctx.storage.find_catalog_row.return_value = {
        'id': 1,
        'category': 'institute',
        'value': 'ИСИ',
        'parent_id': None,
        'active': 1,
    }
    app.ctx.storage.find_level_canonical.return_value = 'внутривузовские'

    response = upload_xlsx(
        client,
        [
            make_import_row(
                **{
                    'Институт': 'иси',
                    'Группа': 'ПГС-101',
                    'Вид спорта': 'БЕГ',
                    'Уровень соревнований': 'Внутривузовские',
                }
            )
        ],
    )

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.institute == 'ИСИ'
    assert saved.sport == 'Бег'
    assert saved.group == 'ПГС-101'
    assert saved.level == 'внутривузовские'
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.find_catalog_canonical.side_effect = None
    app.ctx.storage.find_catalog_canonical.return_value = None
    app.ctx.storage.find_catalog_row.return_value = None
    app.ctx.storage.find_level_canonical.return_value = None


def test_manual_competition_uses_canonical_catalog_values(client: SanicTestClient):
    app.ctx.storage.save_competitions.reset_mock()

    def fake_canonical(category, value, parent_id=None):
        if category == 'institute' and value.lower() == 'иси':
            return 'ИСИ'
        return None

    app.ctx.storage.find_catalog_canonical.side_effect = fake_canonical

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Новиков Новел Новикович',
            'student_sex': 'М',
            'institute': 'иси',
            'group': 'ПГС-101',
            'sport': 'Бег',
            'date': '20.03.2026',
            'level': 'внутривузовские',
            'name': 'Кубок',
            'position': '2',
            'course': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.institute == 'ИСИ'
    app.ctx.storage.find_catalog_canonical.side_effect = None
    app.ctx.storage.find_catalog_canonical.return_value = None


def test_import_autofills_institute_from_unique_group(client: SanicTestClient):
    # №19a: группа известна, институт пуст, группа принадлежит ровно одному
    # институту — институт подставляется каноническим значением.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.find_unique_group_institute.return_value = 'ИСИ'
    app.ctx.storage.find_catalog_row.return_value = {
        'id': 1,
        'category': 'institute',
        'value': 'ИСИ',
        'parent_id': None,
        'active': 1,
    }
    app.ctx.storage.find_catalog_canonical.return_value = 'ПГС-101-и'

    response = upload_xlsx(
        client,
        [make_import_row(**{'Институт': '', 'Группа': 'пгс-101-и'})],
    )

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.institute == 'ИСИ'
    assert saved.group == 'ПГС-101-и'
    app.ctx.storage.find_unique_group_institute.return_value = None
    app.ctx.storage.find_catalog_row.return_value = None
    app.ctx.storage.find_catalog_canonical.return_value = None


def test_import_ambiguous_group_leaves_institute_empty(client: SanicTestClient):
    # №19a: одно имя группы в разных институтах — институт не подставляется.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.find_unique_group_institute.return_value = None

    response = upload_xlsx(
        client,
        [make_import_row(**{'Институт': '', 'Группа': 'ПГС-101'})],
    )

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.institute == ''
    assert saved.group == 'ПГС-101'
    app.ctx.storage.find_unique_group_institute.return_value = None


def test_import_empty_group_keeps_institute_empty(client: SanicTestClient):
    # №19a: группа не указана — автозаполнение не запускается.
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []

    response = upload_xlsx(
        client,
        [make_import_row(**{'Институт': '', 'Группа': ''})],
    )

    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.institute == ''
    app.ctx.storage.find_unique_group_institute.return_value = None


def test_manual_competition_autofills_institute_from_unique_group(client: SanicTestClient):
    # №19a: ручной ввод — тот же путь build_competition.
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.find_unique_group_institute.return_value = 'ИСИ'

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/competition',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Абрамов Артём Абрамович',
            'student_sex': 'М',
            'institute': '',
            'group': 'ПГС-101',
            'sport': 'Бег',
            'date': '20.03.2026',
            'level': 'внутривузовские',
            'name': 'Кубок',
            'position': '2',
            'course': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.institute == 'ИСИ'
    app.ctx.storage.find_unique_group_institute.return_value = None


def test_index_renders_url_custom_field_as_link(client: SanicTestClient):
    # №3: url-кастомное поле в таблице — гиперссылка target=_blank;
    # значение с посторонней схемой — обычный текст.
    record = Competition(
        student_id='2',
        student_name='Тестов Тест Тестович',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 3, 15),
        level='внутривузовские',
        name='Кубок',
        position=1,
        extra_data={
            'link': 'https://sport.example.com/race',
            'bad': 'javascript:alert(1)',
        },
    )
    app.ctx.storage.get_competitions_page.return_value = [record]
    app.ctx.storage.count_competitions_filtered.return_value = 1
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(field_id=1, key='link', label='Ссылка на соревнование', field_type='url'),
        CustomField(field_id=2, key='bad', label='Плохая ссылка', field_type='url'),
    ]

    headers = get_auth_headers(role='editor')
    _, response = client.get('/', headers=headers)
    assert response.status == 200
    body = response.text
    expected = (
        '<a href="https://sport.example.com/race" class="link" '
        'rel="noopener noreferrer" target="_blank">https://sport.example.com/race</a>'
    )
    assert expected in body
    assert 'javascript:alert(1)' in body
    assert '<a href="javascript:' not in body
    app.ctx.storage.get_custom_fields.return_value = []
    app.ctx.storage.get_competitions_page.return_value = []
    app.ctx.storage.count_competitions_filtered.return_value = 0


def test_admin_catalog_add_rejects_case_duplicate(client: SanicTestClient):
    # №1.1: ручное добавление регистрового дубля — отказ с подсказкой.
    app.ctx.storage.add_catalog_value.reset_mock()
    app.ctx.storage.find_catalog_canonical.return_value = 'Бег'

    headers = get_auth_headers()
    _, response = client.post(
        '/admin/catalogs/sport',
        headers=headers,
        data={**csrf_for(headers), 'value': 'бег'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    assert 'Бег' in unquote_plus(response.headers['location'])
    app.ctx.storage.add_catalog_value.assert_not_called()
    app.ctx.storage.find_catalog_canonical.return_value = None


def test_admin_catalog_add_allows_exact_and_new_values(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/catalogs/sport',
        headers=headers,
        data={**csrf_for(headers), 'value': 'Плавание'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.add_catalog_value.assert_called_with('sport', 'Плавание')


def test_admin_group_add_rejects_case_duplicate_within_institute(client: SanicTestClient):
    app.ctx.storage.add_catalog_value.reset_mock()
    app.ctx.storage.get_catalog_value.return_value = {
        'id': 5,
        'category': 'institute',
        'value': 'ИСИ',
        'parent_id': None,
        'active': 1,
    }
    app.ctx.storage.find_catalog_canonical.return_value = 'ПГС-101'

    headers = get_auth_headers()
    _, response = client.post(
        '/admin/catalogs/institute/5/group',
        headers=headers,
        data={**csrf_for(headers), 'value': 'пгс-101'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    app.ctx.storage.add_catalog_value.assert_not_called()
    app.ctx.storage.get_catalog_value.return_value = None
    app.ctx.storage.find_catalog_canonical.return_value = None


def test_admin_level_add_rejects_case_duplicate(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Внутривузовские'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' in response.headers['location']
    # Имя приведено к lower до проверки, поэтому «Внутривузовские» ловится
    # как дубль «внутривузовские» (источник регистровых дублей устранён).
    assert unquote_plus(response.headers['location']).startswith('/admin/catalogs?admin_error=')


def test_admin_level_add_normalizes_to_lower(client: SanicTestClient):
    # Уровень всегда хранится в lower (как из импорта); ручное добавление
    # приведено к тому же виду — источник регистровых дублей устранён.
    app.ctx.storage.create_level.reset_mock()
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/levels',
        headers=headers,
        data={**csrf_for(headers), 'name': 'ГОРОДСКИЕ'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'admin_error' not in response.headers['location']
    app.ctx.storage.create_level.assert_called_with('городские')


def test_admin_catalogs_institutes_are_collapsed_accordion(client: SanicTestClient):
    # №19б: институты — <details>/<summary> без JS, по умолчанию свёрнуто.
    app.ctx.storage.list_catalog_tree.return_value = [
        {
            'id': 1,
            'category': 'institute',
            'value': 'ИСИ',
            'parent_id': None,
            'active': 1,
            'records_count': 3,
            'groups': [
                {
                    'id': 2,
                    'category': 'group',
                    'value': 'ПГС-101',
                    'parent_id': 1,
                    'active': 1,
                    'records_count': 2,
                }
            ],
        }
    ]
    headers = get_auth_headers()
    _, response = client.get('/admin/catalogs', headers=headers)
    assert response.status == 200
    assert '<details class="catalog-institute' in response.text
    assert '<summary>' in response.text
    assert not re.search(r'<details[^>]*\bopen\b', response.text), 'институты должны быть свёрнуты по умолчанию'
    assert 'ПГС-101' in response.text
    app.ctx.storage.list_catalog_tree.return_value = []


# Даты-диапазоны (решение 2026-09-13, docs/data-model-decisions.md
# «Даты-диапазоны: визуально одно, под капотом два»).


def test_parse_date_value_single_day():
    date, date_to = parse_date_value('25.06.2026')
    assert date == datetime(2026, 6, 25)
    assert date_to is None


def test_parse_date_value_one_month_range():
    date, date_to = parse_date_value('25-27.06.2026')
    assert date == datetime(2026, 6, 25)
    assert date_to == datetime(2026, 6, 27)


def test_parse_date_value_one_year_range():
    date, date_to = parse_date_value('30.01-01.02.2026')
    assert date == datetime(2026, 1, 30)
    assert date_to == datetime(2026, 2, 1)


def test_parse_date_value_full_range():
    date, date_to = parse_date_value('28.12.2025-02.01.2026')
    assert date == datetime(2025, 12, 28)
    assert date_to == datetime(2026, 1, 2)


def test_parse_date_value_rejects_end_before_start():
    with pytest.raises(ValueError):
        parse_date_value('27-25.06.2026')


def test_parse_date_value_rejects_invalid_date():
    with pytest.raises(ValueError):
        parse_date_value('32.06.2026')
    with pytest.raises(ValueError):
        parse_date_value('25-26.13.2026')


def test_format_date_range_variants():
    assert format_date_range(datetime(2026, 6, 25), None) == '25.06.2026'
    assert format_date_range(datetime(2026, 6, 25), datetime(2026, 6, 27)) == '25-27.06.2026'
    assert format_date_range(datetime(2026, 1, 30), datetime(2026, 2, 1)) == '30.01-01.02.2026'
    assert format_date_range(datetime(2025, 12, 28), datetime(2026, 1, 2)) == '28.12.2025-02.01.2026'


def test_manual_competition_accepts_range(client: SanicTestClient):
    _, response = client.post(
        '/competition',
        headers=get_auth_headers(),
        data={
            **csrf_for(get_auth_headers()),
            'student_name': 'Диапазон Дарья',
            'student_sex': 'Ж',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '25-27.06.2026',
            'level': 'межвузовские',
            'name': 'Летний кубок',
            'position': '2',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    competition = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert competition.date == datetime(2026, 6, 25)
    assert competition.date_to == datetime(2026, 6, 27)


def test_manual_competition_rejects_range_end_before_start(client: SanicTestClient):
    _, response = client.post(
        '/competition',
        headers=get_auth_headers(),
        data={
            **csrf_for(get_auth_headers()),
            'student_name': 'Диапазон Дарья',
            'student_sex': 'Ж',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
            'sport': 'Легкая атлетика',
            'date': '27-25.06.2026',
            'level': 'межвузовские',
            'name': 'Летний кубок',
            'position': '2',
        },
        allow_redirects=False,
    )

    assert response.status == 400


def test_import_accepts_range_in_date_column(client: SanicTestClient):
    storage = app.ctx.storage
    storage.find_catalog_canonical.return_value = None
    storage.find_level_canonical.return_value = None
    storage.find_catalog_row.return_value = None
    storage.find_unique_group_institute.return_value = None
    competition = build_competition(
        {
            'ФИО': 'Импорт Игорь',
            'Пол': 'М',
            'Институт': 'ИСИ',
            'Группа': 'ПГС-101',
            'Вид спорта': 'Бег',
            'Дата': '28.12.2025-02.01.2026',
            'Уровень соревнований': 'межвузовские',
            'Название соревнований': 'Новогодний кубок',
            'Место': '1',
            'Курс': 2,
        },
        custom_fields=[],
        storage=None,
    )
    assert competition.date == datetime(2025, 12, 28)
    assert competition.date_to == datetime(2026, 1, 2)


def test_export_range_roundtrips_through_import():
    competition = Competition(
        student_id='id-1',
        student_name='Экспорт Елена',
        student_sex='Ж',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 1, 30),
        date_to=datetime(2026, 2, 1),
        level='внутривузовские',
        name='Зимний кубок',
        position=1,
    )
    row = competition_to_export_row(competition, [])
    # Компактная строка в существующей колонке «Дата» — без новой колонки.
    assert 'Дата по' not in row
    assert row['Дата'] == '30.01-01.02.2026'

    buffer = BytesIO()
    pd.DataFrame.from_records([row]).to_excel(buffer, index=False)
    buffer.seek(0)
    imported = build_competition(pd.read_excel(buffer).iloc[0].to_dict(), custom_fields=[])

    assert competition_duplicate_key(imported) == competition_duplicate_key(competition)
    assert imported.date == datetime(2026, 1, 30)
    assert imported.date_to == datetime(2026, 2, 1)


def test_duplicate_key_uses_start_date_only():
    existing = Competition(
        student_id='id-1',
        student_name='Дубль Дмитрий',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 6, 25),
        date_to=datetime(2026, 6, 27),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )
    incoming = Competition(
        student_id='id-2',
        student_name='Дубль Дмитрий',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 6, 25),
        level='внутривузовские',
        name='Кубок',
        position=3,
    )

    new_competitions, skipped, conflicts = split_import_competitions([incoming], [existing])
    assert new_competitions == []
    assert skipped == 1
    assert conflicts == []


# --- Волна 3: лёгкий реестр полей (№2/№7) ---


class FakeSettingsStorage:
    """Минимальный storage для build_competition: только настройки полей
    и «пустые» справочники канонизации."""

    def __init__(self, field_settings=None):
        self._field_settings = field_settings or {}

    def get_field_settings(self):
        return self._field_settings

    def find_catalog_row(self, *args, **kwargs):
        return None

    def find_catalog_canonical(self, *args, **kwargs):
        return None

    def find_level_canonical(self, *args, **kwargs):
        return None

    def find_unique_group_institute(self, *args, **kwargs):
        return None


def make_import_record(**overrides):
    record = {
        'ФИО': 'Реестров Реестр Реестрович',
        'Пол': 'М',
        'Институт': 'ИСИ',
        'Группа': 'ПГС-101',
        'Вид спорта': 'Бег',
        'Дата': '15.03.2026',
        'Уровень соревнований': 'внутривузовские',
        'Название соревнований': 'Кубок',
        'Место': 1,
        'Курс': 2,
    }
    record.update(overrides)
    return record


def test_field_settings_defaults_lock_name_and_date():
    from src.main import get_base_field_settings

    merged = get_base_field_settings(None)
    assert merged['student_name'] == {'value_type': 'text', 'required': True}
    assert merged['date'] == {'value_type': 'text', 'required': True}
    assert merged['course'] == {'value_type': 'number', 'required': True}


def test_field_settings_locked_fields_survive_bad_storage_rows():
    from src.main import get_base_field_settings

    storage = FakeSettingsStorage(
        {
            'student_name': {'value_type': 'number', 'required': False},
            'date': {'value_type': 'number', 'required': False},
        }
    )
    merged = get_base_field_settings(storage)
    assert merged['student_name'] == {'value_type': 'text', 'required': True}
    assert merged['date'] == {'value_type': 'text', 'required': True}


def test_build_competition_text_course_graduate_case():
    from src.main import BASE_FIELD_SETTING_DEFAULTS

    settings = {
        key: {'value_type': value_type, 'required': required}
        for key, (value_type, required) in BASE_FIELD_SETTING_DEFAULTS.items()
    }
    settings['course'] = {'value_type': 'text', 'required': False}
    storage = FakeSettingsStorage(settings)
    competition = build_competition(
        make_import_record(Курс='Выпускник 2025/26'),
        custom_fields=[],
        storage=storage,
    )
    assert competition.course == 'Выпускник 2025/26'


def test_build_competition_number_course_rejects_graduate_case():
    from src.main import BASE_FIELD_SETTING_DEFAULTS

    settings = {
        key: {'value_type': value_type, 'required': required}
        for key, (value_type, required) in BASE_FIELD_SETTING_DEFAULTS.items()
    }
    settings['course'] = {'value_type': 'number', 'required': True}
    storage = FakeSettingsStorage(settings)
    with pytest.raises(ValueError):
        build_competition(
            make_import_record(Курс='Выпускник 2025/26'),
            custom_fields=[],
            storage=storage,
        )


def test_build_competition_required_text_field_rejects_empty():
    from src.main import BASE_FIELD_SETTING_DEFAULTS

    settings = {
        key: {'value_type': value_type, 'required': required}
        for key, (value_type, required) in BASE_FIELD_SETTING_DEFAULTS.items()
    }
    settings['group'] = {'value_type': 'text', 'required': True}
    storage = FakeSettingsStorage(settings)
    with pytest.raises(ValueError, match='Группа'):
        build_competition(
            make_import_record(Группа=''),
            custom_fields=[],
            storage=storage,
        )


def test_build_competition_name_and_date_always_required():
    storage = FakeSettingsStorage(
        {
            'student_name': {'value_type': 'text', 'required': False},
            'date': {'value_type': 'text', 'required': False},
        }
    )
    with pytest.raises(ValueError, match='ФИО'):
        build_competition(make_import_record(ФИО=''), custom_fields=[], storage=storage)
    with pytest.raises(ValueError, match='Дата'):
        build_competition(make_import_record(Дата=''), custom_fields=[], storage=storage)


def test_save_base_field_settings_updates_and_audits(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.update_field_settings.reset_mock()
    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/fields/base',
        headers=headers,
        data={
            **csrf_for(headers),
            'type_course': 'text',
            'required_course': 'on',
            'type_position': 'number',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.update_field_settings.assert_called_once()
    saved = app.ctx.storage.update_field_settings.call_args[0][0]
    assert saved['course'] == ('text', True)
    assert saved['position'] == ('number', False)
    # Заблокированные поля игнорируют форму: обязательность не отключается.
    assert saved['student_name'][1] is True
    assert saved['date'][1] is True


def test_base_field_settings_editor_forbidden(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.get('/admin/fields', headers=headers)
    assert response.status == 403
    _, response = client.post(
        '/admin/fields/base',
        headers=headers,
        data=csrf_for(headers),
    )
    assert response.status == 403


# --- Волна 3: очередь конфликтов импорта (№6/№8) ---


def test_upload_similar_row_goes_to_queue(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.add_import_queue_entry.reset_mock()
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            record_id='7',
            student_id='1',
            student_name='Похожий Павел',
            student_sex='М',
            institute='ИСИ',
            group='ПГС-101',
            course=2,
            sport='Бег',
            date=datetime(2026, 3, 15),
            level='внутривузовские',
            name='Кубок',
            position=1,
        )
    ]
    df = pd.DataFrame(
        [
            # Похожая строка: ФИО+дата совпадают, вид спорта другой → очередь.
            {
                'ФИО': 'Похожий Павел',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Лыжи',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 2,
                'Курс': 2,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    assert response.status == 200
    assert 'Импортировано записей: 0' in response.text
    assert 'На подтверждение: 1 — передано администратору.' in response.text
    assert 'import-queue' not in response.text
    assert app.ctx.storage.import_competitions.call_args[0][0] == []
    app.ctx.storage.add_import_queue_entry.assert_called_once()
    payload = app.ctx.storage.add_import_queue_entry.call_args[0][0]
    assert payload['sport'] == 'Лыжи'
    assert app.ctx.storage.add_import_queue_entry.call_args[1]['matched_record_id'] == 7


def make_similar_row_file() -> bytes:
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Похожий Павел',
                'Пол': 'М',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Лыжи',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 2,
                'Курс': 2,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    return file_obj.getvalue()


def upload_registry_file(client, role: str):
    headers = get_auth_headers(role=role)
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                make_similar_row_file(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    return response


def test_upload_similar_row_summary_admin_vs_editor(client: SanicTestClient):
    # IA-редизайн 2026-09: admin видит адресата разбора очереди, editor —
    # только факт передачи администратору.
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.add_import_queue_entry.reset_mock()
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            record_id='7',
            student_id='1',
            student_name='Похожий Павел',
            student_sex='М',
            institute='ИСИ',
            group='ПГС-101',
            course=2,
            sport='Бег',
            date=datetime(2026, 3, 15),
            level='внутривузовские',
            name='Кубок',
            position=1,
        )
    ]
    try:
        response = upload_registry_file(client, 'admin')
        assert response.status == 200
        assert 'На подтверждение: 1.' in response.text
        assert 'передано администратору' not in response.text

        response = upload_registry_file(client, 'editor')
        assert response.status == 200
        assert 'На подтверждение: 1 — передано администратору.' in response.text
        assert 'import-queue' not in response.text
    finally:
        app.ctx.storage.get_competitions.return_value = []


def test_upload_exact_duplicate_still_skipped_as_before(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.add_import_queue_entry.reset_mock()
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='Дубль Дарья',
            student_sex='Ж',
            institute='ИСИ',
            group='ПГС-101',
            course=1,
            sport='Бег',
            date=datetime(2026, 3, 15),
            level='внутривузовские',
            name='Кубок',
            position=1,
        )
    ]
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Дубль Дарья',
                'Пол': 'Ж',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 1,
                'Курс': 1,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    assert response.status == 200
    assert 'Пропущено дублей: 1' in response.text
    assert 'На подтверждение' not in response.text
    assert app.ctx.storage.import_competitions.call_args[0][0] == []
    app.ctx.storage.add_import_queue_entry.assert_not_called()


def test_import_queue_page_admin_only_and_renders(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    entry_payload = {
        'student_id': 'hash-1',
        'student_name': 'Похожий Павел',
        'student_sex': 'М',
        'institute': 'ИСИ',
        'group': 'ПГС-101',
        'course': 2,
        'sport': 'Лыжи',
        'date': '2026-03-15T00:00:00',
        'date_to': None,
        'level': 'внутривузовские',
        'name': 'Кубок',
        'position': 2,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-03-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.list_import_queue.return_value = [
        {
            'id': 5,
            'created_at': '2026-03-01T12:00:00',
            'payload_json': json.dumps(entry_payload),
            'payload': entry_payload,
            'status': 'pending',
            'matched_record_id': 1,
            'created_by': 1,
        }
    ]
    matched = Competition(
        student_id='1',
        student_name='Похожий Павел',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 3, 15),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )
    app.ctx.storage.get_competition_by_id.return_value = matched

    headers = get_auth_headers(role='admin')
    _, response = client.get('/admin/import-queue', headers=headers)
    assert response.status == 200
    assert 'Похожий Павел' in response.text
    assert 'Лыжи' in response.text

    headers = get_auth_headers(role='editor')
    _, response = client.get('/admin/import-queue', headers=headers)
    assert response.status == 403
    _, response = client.post(
        '/admin/import-queue/5/accept',
        headers=headers,
        data=csrf_for(headers),
    )
    assert response.status == 403


def test_import_queue_accept_inserts_record_and_audits(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    entry_payload = {
        'student_id': 'hash-2',
        'student_name': 'Принятый Пётр',
        'student_sex': 'М',
        'institute': 'ИСИ',
        'group': 'ПГС-101',
        'course': 2,
        'sport': 'Лыжи',
        'date': '2026-03-15T00:00:00',
        'date_to': None,
        'level': 'внутривузовские',
        'name': 'Кубок',
        'position': 2,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-03-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 9,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': json.dumps(entry_payload),
        'payload': entry_payload,
        'status': 'pending',
        'matched_record_id': 3,
        'created_by': 2,
    }
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.set_import_queue_status.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()

    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/9/accept',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.save_competitions.assert_called_once()
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.student_name == 'Принятый Пётр'
    assert saved.sport == 'Лыжи'
    app.ctx.storage.set_import_queue_status.assert_called_once_with(9, 'accepted')
    app.ctx.storage.add_audit_event.assert_called_once()
    audit_kwargs = app.ctx.storage.add_audit_event.call_args[1]
    assert audit_kwargs['action'] == 'import_conflict_resolved'
    assert json.loads(audit_kwargs['details'])['decision'] == 'accepted'


def test_import_queue_skip_marks_skipped_and_audits(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 11,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': '{}',
        'payload': {},
        'status': 'pending',
        'matched_record_id': 4,
        'created_by': 1,
    }
    app.ctx.storage.set_import_queue_status.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()

    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/11/skip',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.set_import_queue_status.assert_called_once_with(11, 'skipped')
    assert app.ctx.storage.add_audit_event.call_args[1]['action'] == 'import_conflict_resolved'
    assert json.loads(app.ctx.storage.add_audit_event.call_args[1]['details'])['decision'] == 'skipped'


def test_import_queue_replace_updates_existing_record_and_audits(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    entry_payload = {
        'student_id': 'hash-r',
        'student_name': 'Заменяемый Захар',
        'student_sex': 'М',
        'institute': 'ИСИ',
        'group': 'ПГС-101',
        'course': 2,
        'sport': 'Лыжи',
        'date': '2026-03-15T00:00:00',
        'date_to': None,
        'level': 'внутривузовские',
        'name': 'Кубок',
        'position': 2,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-03-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 13,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': json.dumps(entry_payload),
        'payload': entry_payload,
        'status': 'pending',
        'matched_record_id': 3,
        'created_by': 2,
    }
    existing = Competition(
        student_id='hash-r',
        student_name='Заменяемый Захар',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 3, 15),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )
    app.ctx.storage.get_competition_by_id.return_value = existing
    app.ctx.storage.update_competition.reset_mock()
    app.ctx.storage.set_import_queue_status.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()
    app.ctx.storage.save_competitions.reset_mock()

    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/13/replace',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    # Существующая запись обновляется по месту: record_id сохраняется, вставки нет.
    app.ctx.storage.update_competition.assert_called_once_with(3, ANY)
    replaced = app.ctx.storage.update_competition.call_args[0][1]
    assert replaced.sport == 'Лыжи'
    assert replaced.position == 2
    app.ctx.storage.save_competitions.assert_not_called()
    app.ctx.storage.set_import_queue_status.assert_called_once_with(13, 'replaced')
    audit_kwargs = app.ctx.storage.add_audit_event.call_args[1]
    assert audit_kwargs['action'] == 'import_conflict_resolved'
    details = json.loads(audit_kwargs['details'])
    assert details['decision'] == 'replaced'
    assert details['matched_record_id'] == 3
    assert details['changes']['sport'] == {'old': 'Бег', 'new': 'Лыжи'}
    assert details['changes']['position'] == {'old': '1', 'new': '2'}


def test_import_queue_replace_without_matched_record_rejected(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 14,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': '{}',
        'payload': {},
        'status': 'pending',
        'matched_record_id': None,
        'created_by': 1,
    }
    app.ctx.storage.update_competition.reset_mock()
    app.ctx.storage.set_import_queue_status.reset_mock()
    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/14/replace',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 400
    app.ctx.storage.update_competition.assert_not_called()
    app.ctx.storage.set_import_queue_status.assert_not_called()


def test_import_queue_edit_accepts_with_edited_values_and_audits(client: SanicTestClient):
    app.ctx.storage.get_field_settings.return_value = {}
    entry_payload = {
        'student_id': 'hash-e',
        'student_name': 'Ручной Роман',
        'student_sex': 'М',
        'institute': 'ИСИ',
        'group': 'ПГС-101',
        'course': 2,
        'sport': 'Лыжи',
        'date': '2026-03-15T00:00:00',
        'date_to': None,
        'level': 'внутривузовские',
        'name': 'Кубок',
        'position': 2,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-03-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 15,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': json.dumps(entry_payload),
        'payload': entry_payload,
        'status': 'pending',
        'matched_record_id': 3,
        'created_by': 2,
    }
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.set_import_queue_status.reset_mock()
    app.ctx.storage.add_audit_event.reset_mock()

    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/15/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            # №27в: подмена ФИО и даты формой игнорируется — берутся из payload.
            'student_name': 'Подмена ФИО',
            'student_sex': 'Ж',
            'institute': 'ИГДН',
            'group': 'ТД-202',
            'course': '3',
            'sport': 'Шахматы',
            'date': '20.03.2026',
            'level': 'городские',
            'name': 'Кубок edited',
            'position': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    app.ctx.storage.save_competitions.assert_called_once()
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.student_name == 'Ручной Роман'
    assert saved.student_sex == 'Ж'
    assert saved.institute == 'ИГДН'
    assert saved.group == 'ТД-202'
    assert saved.sport == 'Шахматы'
    assert saved.date == datetime(2026, 3, 15)
    assert saved.date_to is None
    assert saved.level == 'городские'
    assert saved.name == 'Кубок edited'
    assert saved.position == 1
    assert saved.course == 3
    app.ctx.storage.set_import_queue_status.assert_called_once_with(15, 'accepted')
    audit_kwargs = app.ctx.storage.add_audit_event.call_args[1]
    assert audit_kwargs['action'] == 'import_conflict_resolved'
    details = json.loads(audit_kwargs['details'])
    assert details['decision'] == 'accepted'
    assert details['edited'] is True


def test_import_queue_edit_uses_payload_date_range_from_payload(client: SanicTestClient):
    """№27в: диапазон дат кандидата сохраняется из payload, форма дату не меняет."""
    app.ctx.storage.get_field_settings.return_value = {}
    entry_payload = {
        'student_name': 'Диапазон Дарья',
        'student_sex': 'Ж',
        'institute': 'ИСИ',
        'group': 'ПГС-102',
        'course': 1,
        'sport': 'Биатлон',
        'date': '2026-06-25T00:00:00',
        'date_to': '2026-06-27T00:00:00',
        'level': 'областные',
        'name': 'Спартакиада',
        'position': 3,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-06-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 17,
        'created_at': '2026-06-01T12:00:00',
        'payload_json': json.dumps(entry_payload),
        'payload': entry_payload,
        'status': 'pending',
        'matched_record_id': None,
        'created_by': 2,
    }
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.save_competitions.reset_mock()
    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/17/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_sex': 'Ж',
            'institute': 'ИСИ',
            'group': 'ПГС-103',
            'course': '2',
            'sport': 'Лёгкая атлетика',
            'date': '01.01.2027',
            'level': 'городские',
            'name': 'Спартакиада edited',
            'position': '5',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.student_name == 'Диапазон Дарья'
    assert saved.date == datetime(2026, 6, 25)
    assert saved.date_to == datetime(2026, 6, 27)
    assert saved.group == 'ПГС-103'
    assert saved.sport == 'Лёгкая атлетика'
    assert saved.level == 'городские'
    assert saved.name == 'Спартакиада edited'
    assert saved.position == 5
    assert saved.course == 2


def test_import_queue_edit_course_text_allowed_by_field_settings(client: SanicTestClient):
    """№27в: валидация по field_settings — course type=text принимает «Выпускник 2025/26»."""
    app.ctx.storage.get_field_settings.return_value = {'course': {'value_type': 'text', 'required': True}}
    entry_payload = {
        'student_name': 'Выпускник Влада',
        'student_sex': 'Ж',
        'institute': 'ИСИ',
        'group': 'ПГС-103',
        'course': 'Выпускник 2025/26',
        'sport': 'Волейбол',
        'date': '2026-03-15T00:00:00',
        'date_to': None,
        'level': 'внутривузовские',
        'name': 'Кубок',
        'position': 2,
        'extra_data': {},
        'record_id': None,
        'created_at': '2026-03-01T00:00:00',
        'review_status': 'approved',
        'review_comment': '',
    }
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 18,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': json.dumps(entry_payload),
        'payload': entry_payload,
        'status': 'pending',
        'matched_record_id': None,
        'created_by': 2,
    }
    app.ctx.storage.get_competitions.return_value = []
    app.ctx.storage.save_competitions.reset_mock()
    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/18/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_sex': 'Ж',
            'institute': 'ИСИ',
            'group': 'ПГС-103',
            'course': 'Выпускник 2025/26',
            'sport': 'Волейбол',
            'level': 'внутривузовские',
            'name': 'Кубок',
            'position': '2',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    saved = app.ctx.storage.save_competitions.call_args[0][0][0]
    assert saved.course == 'Выпускник 2025/26'


def test_import_queue_edit_rejects_bad_date(client: SanicTestClient):
    """№27в: дата берётся из payload — битая дата payload даёт 400."""
    app.ctx.storage.get_field_settings.return_value = {}
    app.ctx.storage.get_import_queue_entry.return_value = {
        'id': 16,
        'created_at': '2026-03-01T12:00:00',
        'payload_json': '{}',
        'payload': {},
        'status': 'pending',
        'matched_record_id': 3,
        'created_by': 1,
    }
    app.ctx.storage.save_competitions.reset_mock()
    app.ctx.storage.set_import_queue_status.reset_mock()
    headers = get_auth_headers(role='admin')
    _, response = client.post(
        '/admin/import-queue/16/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Ручной Роман',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '3',
            'sport': 'Шахматы',
            'date': '2026-03-20',
            'level': 'внутривузовские',
            'name': 'Кубок',
            'position': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 400
    app.ctx.storage.save_competitions.assert_not_called()
    app.ctx.storage.set_import_queue_status.assert_not_called()


def test_import_queue_editor_forbidden_on_replace_and_edit(client: SanicTestClient):
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/admin/import-queue/5/replace',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403
    _, response = client.post(
        '/admin/import-queue/5/edit',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 403


def test_split_import_similar_rows_within_file(client: SanicTestClient):
    first = Competition(
        student_id='1',
        student_name='Файл Фёдор',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 3, 15),
        level='внутривузовские',
        name='Кубок',
        position=1,
    )
    similar = first.model_copy(update={'sport': 'Лыжи', 'student_id': '2'})
    new_competitions, skipped, conflicts = split_import_competitions([first, similar], [])
    assert [comp.sport for comp in new_competitions] == ['Бег']
    assert skipped == 0
    assert len(conflicts) == 1
    assert conflicts[0][0].sport == 'Лыжи'
    assert conflicts[0][1] is None


# --- №23доп (docs/feedback-live.md): резолвер атлета ---


def get_athlete_headers() -> dict[str, str]:
    cookie = create_auth_cookie_value(username='sportik', role='athlete')
    return {'cookie': f'{settings.auth_cookie_name}={cookie}'}


def test_athletes_search_returns_variants_for_moderator(client: SanicTestClient):
    app.ctx.storage.search_athletes.return_value = [
        {'name': 'Иванов Иван', 'sex': 'М', 'institute': 'ИСИ', 'group': 'ПГС-101', 'course': '1'}
    ]
    headers = get_auth_headers(role='editor')
    _, response = client.get('/api/athletes/search?q=Ива', headers=headers)
    assert response.status == 200
    # Каждый вариант помечен kind и student_id (объединённый поиск:
    # карточки Students + легаси-подсказки) — см. search_athletes.
    assert response.json == [
        {
            'kind': 'legacy',
            'student_id': None,
            'sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '1',
            'name': 'Иванов Иван',
        }
    ]
    args = app.ctx.storage.search_athletes.call_args[0]
    assert args[0] == 'Ива'
    app.ctx.storage.search_athletes.return_value = []


def test_athletes_search_empty_query_returns_empty_list(client: SanicTestClient):
    _, response = client.get('/api/athletes/search?q=', headers=get_auth_headers(role='admin'))
    assert response.status == 200
    assert response.json == []


def test_athletes_search_forbidden_for_athlete(client: SanicTestClient):
    # Полный список ФИО атлету не раскрывается (docs/data-model-decisions.md):
    # 403 и никаких обращений к хранилищу.
    app.ctx.storage.search_athletes.reset_mock()
    _, response = client.get('/api/athletes/search?q=Ива', headers=get_athlete_headers())
    assert response.status == 403
    app.ctx.storage.search_athletes.assert_not_called()


def test_athletes_search_forbidden_for_viewer(client: SanicTestClient):
    _, response = client.get('/api/athletes/search?q=Ива', headers=get_auth_headers(role='viewer'))
    assert response.status == 403


def test_upload_autofills_empty_fields_from_known_athlete(client: SanicTestClient):
    # Строка импорта: ФИО известно, Пол/Институт/Группа/Курс пусты —
    # резолвер подставляет непустые значения (№23доп).
    app.ctx.storage.find_athlete_fields.side_effect = lambda name: (
        {
            'sex': 'Ж',
            'institute': 'ИМИ',
            'group': 'СБ-202',
            'course': '2',
        }
        if name.strip().lower() == 'петрова петра'
        else {}
    )
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Петрова Петра',
                'Пол': '',
                'Институт': '',
                'Группа': '',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 2,
                'Курс': '',
            }
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )

    assert response.status == 200
    assert 'Импортировано записей: 1' in response.text
    saved_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert len(saved_rows) == 1
    assert saved_rows[0].student_sex == 'Ж'
    assert saved_rows[0].institute == 'ИМИ'
    assert saved_rows[0].group == 'СБ-202'
    assert saved_rows[0].course == 2
    app.ctx.storage.find_athlete_fields.side_effect = None
    app.ctx.storage.find_athlete_fields.return_value = {}


def test_upload_keeps_filled_fields_and_counts_duplicates_after_autofill(client: SanicTestClient):
    # Заполненные поля строки НЕ перезаписываются; дубль-ключ считается
    # ПОСЛЕ автозаполнения: строка с пустыми полями, совпавшая по ключу
    # с существующей записью после подстановки, пропускается как дубль.
    app.ctx.storage.find_athlete_fields.side_effect = lambda name: (
        {'sex': 'М', 'institute': 'ИСИ', 'group': 'ПГС-101', 'course': '2'}
        if name.strip().lower() == 'иванов иван'
        else {}
    )
    app.ctx.storage.get_competitions.return_value = [
        Competition(
            student_id='1',
            student_name='Иванов Иван',
            student_sex='М',
            institute='ИСИ',
            group='ПГС-101',
            course=2,
            sport='Бег',
            date=datetime(2026, 3, 15),
            level='внутривузовские',
            name='Кубок',
            position=1,
        )
    ]
    df = pd.DataFrame(
        [
            {
                'ФИО': 'Иванов Иван',
                'Пол': '',
                'Институт': '',
                'Группа': '',
                'Вид спорта': 'Бег',
                'Дата': '15.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 3,
                'Курс': '',
            },
            {
                'ФИО': 'Петров Пётр',
                'Пол': 'М',
                'Институт': '',
                'Группа': '',
                'Вид спорта': 'Бег',
                'Дата': '16.03.2026',
                'Уровень соревнований': 'внутривузовские',
                'Название соревнований': 'Кубок',
                'Место': 2,
                'Курс': 1,
            },
        ]
    )
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)

    app.ctx.storage.find_athlete_fields.side_effect = lambda name: (
        {'sex': 'М', 'institute': 'ИСИ', 'group': 'ПГС-101', 'course': '2'}
        if name.strip().lower() == 'иванов иван'
        else {'sex': 'Ж', 'institute': 'ИМИ', 'group': 'СБ-202', 'course': '9'}
    )
    headers = get_auth_headers(role='editor')
    _, response = client.post(
        '/',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                'import.xlsx',
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
    )

    assert response.status == 200
    assert 'Импортировано записей: 1' in response.text
    assert 'Пропущено дублей: 1' in response.text
    saved_rows = app.ctx.storage.import_competitions.call_args[0][0]
    assert len(saved_rows) == 1
    # Заполненные поля (Пол=М, Курс=1) резолвером с другим ответом (Ж, 9)
    # не перезаписаны; пустые — заполнены.
    row = saved_rows[0]
    assert row.student_name == 'Петров Пётр'
    assert row.student_sex == 'М'
    assert row.course == 1
    assert row.institute == 'ИМИ'
    assert row.group == 'СБ-202'
    app.ctx.storage.find_athlete_fields.side_effect = None
    app.ctx.storage.find_athlete_fields.return_value = {}
    app.ctx.storage.get_competitions.return_value = []


# №24 (docs/feedback-live.md): привязка link-поля к другой колонке.
def test_index_renders_bound_link_in_target_column(client: SanicTestClient):
    """Целевая колонка — гиперссылка на link-поле ЭТОЙ записи; без ссылки —
    текст; сама link-колонка скрыта (d-none)."""
    with_link = Competition(
        student_id='2',
        student_name='Тестов Тест Тестович',
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 3, 15),
        level='внутривузовские',
        name='Кубок «Весна»',
        position=1,
        extra_data={'comp_link': 'https://sport.example.com/race'},
    )
    without_link = Competition(
        student_id='3',
        student_name='Без Ссылки',
        student_sex='Ж',
        institute='ИСИ',
        group='ПГС-102',
        course=3,
        sport='Бег',
        date=datetime(2026, 3, 16),
        level='внутривузовские',
        name='Кубок «Осень»',
        position=2,
        extra_data={},
    )
    app.ctx.storage.get_competitions_page.return_value = [with_link, without_link]
    app.ctx.storage.count_competitions_filtered.return_value = 2
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(
            field_id=1,
            key='comp_link',
            label='Ссылка на соревнование',
            field_type='url',
            link_target='name',
        ),
    ]

    headers = get_auth_headers(role='editor')
    _, response = client.get('/', headers=headers)
    assert response.status == 200
    body = response.text
    expected = (
        '<a href="https://sport.example.com/race" class="link" '
        'rel="noopener noreferrer" target="_blank">Кубок «Весна»</a>'
    )
    assert expected in body
    # Запись без ссылки — обычный текст целевой колонки.
    assert '>Кубок «Осень»<' in body
    # Скрытая link-колонка: остаётся в DOM (инлайн-строка читает типы), но d-none.
    assert 'data-column-key="comp_link" data-field-type="url" data-sort-type="url" class="d-none"' in body
    # Не дублируется: у записи со ссылкой отдельная колонка пуста и скрыта.
    assert '<td data-column-key="comp_link" class="d-none">' in body

    app.ctx.storage.get_custom_fields.return_value = []
    app.ctx.storage.get_competitions_page.return_value = []
    app.ctx.storage.count_competitions_filtered.return_value = 0


def test_inline_new_row_hides_bound_link_column():
    """Хотфикс: инлайн-строка новой записи строится по всем th, включая
    скрытую d-none колонку привязанного link-поля. JS обязан переносить
    d-none с th на td, иначе в строке появляется лишнее видимое поле
    (placeholder https://) и вёрстка съезжает."""
    script = (Path(__file__).parent / 'static' / 'scripts' / 'script.js').read_text(encoding='utf-8')
    start = script.index('startNewRowEdit()')
    end = script.index('saveNewRowEdit()', start)
    body = script[start:end]
    assert 'header.classList.contains("d-none")' in body
    assert 'cell.classList.add("d-none")' in body


def test_update_custom_field_saves_link_target(client: SanicTestClient):
    # Старое поле нужно маршруту для сравнения привязки (audit-событие).
    app.ctx.storage.get_custom_fields.return_value = [
        CustomField(field_id=1, key='comp_link', label='Ссылка', field_type='url'),
    ]
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields/1',
        headers=headers,
        data={
            **csrf_for(headers),
            'label': 'Ссылка на соревнование',
            'field_type': 'url',
            'sort_order': '0',
            'link_target': 'name',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    _, kwargs = app.ctx.storage.update_custom_field.call_args
    assert kwargs['link_target'] == 'name'
    # Изменение привязки — в audit (field_settings_changed).
    audit_calls = [
        call
        for call in app.ctx.storage.add_audit_event.call_args_list
        if call.kwargs.get('action') == 'field_settings_changed'
    ]
    assert audit_calls, 'ожидалось audit-событие field_settings_changed'
    app.ctx.storage.get_custom_fields.return_value = []


def test_update_custom_field_rejects_unknown_link_target(client: SanicTestClient):
    headers = get_auth_headers()
    _, response = client.post(
        '/admin/fields/1',
        headers=headers,
        data={
            **csrf_for(headers),
            'label': 'Ссылка на соревнование',
            'field_type': 'url',
            'sort_order': '0',
            'link_target': '../../evil',
        },
        allow_redirects=False,
    )

    assert response.status == 302
    _, kwargs = app.ctx.storage.update_custom_field.call_args
    assert kwargs['link_target'] is None


# ---- Календарь соревнований (волна A, docs/feedback-live.md №23) ----


def test_calendar_event_status_auto_by_end_date():
    today = datetime(2026, 9, 14).date()
    # Будущее → запланировано
    assert calendar_event_status({'date': '2026-09-20', 'date_to': None}, today) == 'planned'
    # Многодневное, сегодня внутри диапазона → ещё запланировано
    assert calendar_event_status({'date': '2026-09-10', 'date_to': '2026-09-15'}, today) == 'planned'
    # Дата окончания = сегодня → ещё запланировано
    assert calendar_event_status({'date': '2026-09-14', 'date_to': None}, today) == 'planned'
    # Однодневное вчера → прошло
    assert calendar_event_status({'date': '2026-09-13', 'date_to': None}, today) == 'past'
    # Диапазон закончился вчера → прошло
    assert calendar_event_status({'date': '2026-09-01', 'date_to': '2026-09-13'}, today) == 'past'


def test_calendar_events_grouped_by_month():
    events = [
        decorate_calendar_event({'date': '2026-09-12', 'date_to': None, 'name': 'A'}, datetime(2026, 9, 14).date()),
        decorate_calendar_event(
            {'date': '2026-09-26', 'date_to': '2026-09-27', 'name': 'B'}, datetime(2026, 9, 14).date()
        ),
        decorate_calendar_event({'date': '2026-10-17', 'date_to': None, 'name': 'C'}, datetime(2026, 9, 14).date()),
    ]
    groups = group_calendar_events_by_month(events)
    assert [group['label'] for group in groups] == ['Сентябрь 2026', 'Октябрь 2026']
    assert [group['count'] for group in groups] == [2, 1]
    assert groups[0]['events'][0]['date_label'] == '12.09.2026'
    assert groups[0]['events'][1]['date_label'] == '26-27.09.2026'


def test_calendar_page_available_for_viewer(client: SanicTestClient):
    headers = get_auth_headers('viewer')
    _, response = client.get('/calendar', headers=headers)
    assert response.status == 200
    # Viewer — просмотр без кнопок управления.
    assert 'Запланировать'.encode() not in response.body


def test_calendar_page_forbidden_for_athlete(client: SanicTestClient):
    # У атлета нет дев-аккаунта в настройках — собираем cookie вручную.
    cookie = create_auth_cookie_value(username='sportik', role='athlete')
    headers = {'cookie': f'{settings.auth_cookie_name}={cookie}'}
    _, response = client.get('/calendar', headers=headers)
    assert response.status == 403


def test_editor_can_create_calendar_event(client: SanicTestClient):
    headers = get_auth_headers('editor')
    _, response = client.post(
        '/calendar/new',
        headers=headers,
        data={
            **csrf_for(headers),
            'name': 'Осенний кросс СибАДИ',
            'date': '25-27.06.2026',
            'level': 'внутривузовские',
            'sport': 'Бег',
            'url': 'https://example.com/reglement',
        },
    )
    assert response.status == 200
    _, kwargs = app.ctx.storage.create_calendar_event.call_args
    assert kwargs['name'] == 'Осенний кросс СибАДИ'
    assert kwargs['date'] == '2026-06-25T00:00:00'
    assert kwargs['date_to'] == '2026-06-27T00:00:00'
    assert kwargs['url'] == 'https://example.com/reglement'


def test_calendar_create_requires_name(client: SanicTestClient):
    # client — module-scoped Mock: сбрасываем вызовы прошлых тестов.
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('editor')
    _, response = client.post(
        '/calendar/new',
        headers=headers,
        data={**csrf_for(headers), 'name': '  ', 'date': '25.06.2026'},
        allow_redirects=False,
    )
    assert response.status == 400
    app.ctx.storage.create_calendar_event.assert_not_called()


def test_calendar_create_rejects_bad_date_and_reversed_range(client: SanicTestClient):
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('editor')
    for bad_date in ('завтра', '27-25.06.2026'):
        _, response = client.post(
            '/calendar/new',
            headers=headers,
            data={**csrf_for(headers), 'name': 'Кросс', 'date': bad_date},
            allow_redirects=False,
        )
        assert response.status == 400
    app.ctx.storage.create_calendar_event.assert_not_called()


def test_calendar_create_rejects_non_http_url(client: SanicTestClient):
    # M2: ссылка календаря рендерится как href — принимаем только http/https.
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('editor')
    for bad_url in ('javascript:alert(1)', 'data:text/html,<b>', 'vbscript:msgbox'):
        _, response = client.post(
            '/calendar/new',
            headers=headers,
            data={**csrf_for(headers), 'name': 'Кросс', 'date': '25.06.2026', 'url': bad_url},
            allow_redirects=False,
        )
        assert response.status == 400
    app.ctx.storage.create_calendar_event.assert_not_called()


def test_calendar_create_accepts_http_https_and_empty_url(client: SanicTestClient):
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('editor')
    for good_url in ('http://example.com/a', 'https://example.com', ''):
        _, response = client.post(
            '/calendar/new',
            headers=headers,
            data={**csrf_for(headers), 'name': 'Кросс', 'date': '25.06.2026', 'url': good_url},
            allow_redirects=False,
        )
        assert response.status == 302
        assert app.ctx.storage.create_calendar_event.call_args[1]['url'] == good_url


def test_viewer_cannot_create_calendar_event(client: SanicTestClient):
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('viewer')
    _, response = client.post(
        '/calendar/new',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Кросс', 'date': '25.06.2026'},
        allow_redirects=False,
    )
    assert response.status == 403


def test_editor_can_edit_calendar_event(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = {
        'id': 7,
        'name': 'Кросс',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': '',
        'created_at': '2026-01-01T00:00:00',
    }
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/7/edit',
            headers=headers,
            data={
                **csrf_for(headers),
                'name': 'Осенний кросс СибАДИ',
                'date': '30.01-01.02.2026',
                'level': 'региональные',
                'sport': 'Лыжи',
                'url': '',
            },
        )
        assert response.status == 200
        _, kwargs = app.ctx.storage.update_calendar_event.call_args
        assert kwargs['event_id'] == 7
        assert kwargs['name'] == 'Осенний кросс СибАДИ'
        assert kwargs['date'] == '2026-01-30T00:00:00'
        assert kwargs['date_to'] == '2026-02-01T00:00:00'
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_edit_rejects_non_http_url(client: SanicTestClient):
    # Тот же парсер формы, что у создания: правка не обходит проверку схемы.
    app.ctx.storage.get_calendar_event.return_value = {
        'id': 7,
        'name': 'Кросс',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': '',
        'created_at': '2026-01-01T00:00:00',
    }
    app.ctx.storage.update_calendar_event.reset_mock()
    try:
        headers = get_auth_headers('editor')
        for bad_url in ('javascript:alert(1)', 'data:text/html,<b>', 'vbscript:msgbox'):
            _, response = client.post(
                '/calendar/7/edit',
                headers=headers,
                data={**csrf_for(headers), 'name': 'Кросс', 'date': '25.06.2026', 'url': bad_url},
                allow_redirects=False,
            )
            assert response.status == 400
        app.ctx.storage.update_calendar_event.assert_not_called()
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_pages_render_unsafe_stored_url_as_text(client: SanicTestClient):
    # Эшелонированная защита: значения, сохранённые до валидации схемы,
    # рендерятся как обычный текст, а не как href.
    unsafe_event = {
        'id': 9,
        'name': 'Кросс',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': 'javascript:alert(1)',
        'created_at': '2026-01-01T00:00:00',
        'participant_count': 0,
        'no_result_count': 0,
    }
    app.ctx.storage.list_calendar_events.return_value = [unsafe_event]
    app.ctx.storage.get_calendar_event.return_value = unsafe_event
    try:
        headers = get_auth_headers('viewer')

        _, response = client.get('/calendar', headers=headers)
        assert response.status == 200
        assert 'href="javascript:' not in response.text

        _, response = client.get('/calendar/9', headers=headers)
        assert response.status == 200
        assert 'href="javascript:' not in response.text
        # значение не теряется — показывается обычным текстом
        assert 'javascript:alert(1)' in response.text
    finally:
        app.ctx.storage.list_calendar_events.return_value = []
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_delete_refuses_with_participants(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = {
        'id': 5,
        'name': 'Кросс',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': '',
        'created_at': '2026-01-01T00:00:00',
    }
    app.ctx.storage.count_calendar_event_participants.return_value = 14
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/5/delete',
            headers=headers,
            data=csrf_for(headers),
            allow_redirects=False,
        )
        assert response.status == 302
        assert 'calendar' in response.headers['location']
        assert '14' in response.headers['location']
        app.ctx.storage.delete_calendar_event.assert_not_called()
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.count_calendar_event_participants.return_value = 0


def test_calendar_delete_without_participants_writes_audit(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = {
        'id': 5,
        'name': 'Пустой турнир',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': '',
        'created_at': '2026-01-01T00:00:00',
    }
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/5/delete',
            headers=headers,
            data=csrf_for(headers),
            allow_redirects=False,
        )
        assert response.status == 302
        app.ctx.storage.delete_calendar_event.assert_called_once_with(5)
        audit_calls = [
            call
            for call in app.ctx.storage.add_audit_event.call_args_list
            if call.kwargs.get('action') == 'calendar_event_deleted'
        ]
        assert audit_calls, 'ожидалось audit-событие calendar_event_deleted'
        app.ctx.storage.delete_calendar_event.reset_mock()
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_delete_flashes_message(client: SanicTestClient):
    # IA-редизайн 2026-09: тихие редиректы календаря несут flash-сообщение.
    app.ctx.storage.get_calendar_event.return_value = {
        'id': 5,
        'name': 'Пустой турнир',
        'date': '2026-06-25',
        'date_to': None,
        'level': '',
        'sport': '',
        'url': '',
        'created_at': '2026-01-01T00:00:00',
    }
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/5/delete',
            headers=headers,
            data=csrf_for(headers),
            allow_redirects=False,
        )
        assert response.status == 302
        assert response.headers['location'].startswith('/calendar?admin_message=')
        assert 'Соревнование удалено' in unquote_plus(response.headers['location'])
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_create_flashes_message(client: SanicTestClient):
    app.ctx.storage.create_calendar_event.reset_mock()
    headers = get_auth_headers('editor')
    _, response = client.post(
        '/calendar/new',
        headers=headers,
        data={
            **csrf_for(headers),
            'name': 'Осенний кросс СибАДИ',
            'date': '25.06.2026',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'].startswith('/calendar?admin_message=')
    assert 'Соревнование «Осенний кросс СибАДИ» запланировано' in unquote_plus(response.headers['location'])


def test_calendar_filters_pass_sport_to_storage(client: SanicTestClient):
    headers = get_auth_headers('viewer')
    _, response = client.get('/calendar?status=past&sport=Бег', headers=headers)
    assert response.status == 200
    app.ctx.storage.list_calendar_events.assert_called_with(sport='Бег')


# ---- Страница соревнования: участники (волна B, прототип 16) ----


def _event_for_page():
    return {
        'id': 7,
        'name': 'Кросс СибАДИ',
        'date': '2026-06-25',
        'date_to': None,
        'level': 'внутривузовские',
        'sport': 'Бег',
        'url': 'https://example.com/reglement',
        'created_at': '2026-01-01T00:00:00',
    }


def _participants_sample():
    return [
        {
            'record_id': 11,
            'student_name': 'Иванов Дмитрий Сергеевич',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group_name': 'ПГСб-41',
            'course': 4,
            'position': 1,
        },
        {
            'record_id': 12,
            'student_name': 'Волков Артём Игоревич',
            'student_sex': 'М',
            'institute': 'ИТМ',
            'group_name': 'ТМб-12',
            'course': 1,
            'position': 0,
        },
    ]


def test_calendar_event_page_shows_participants_and_counts(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.list_calendar_event_participants.return_value = _participants_sample()
    try:
        headers = get_auth_headers('viewer')
        _, response = client.get('/calendar/7', headers=headers)
        assert response.status == 200
        body = response.body.decode()
        # Заголовок события и пресет-метаданные
        assert 'Кросс СибАДИ' in body
        assert '25.06.2026' in body
        assert 'внутривузовские' in body
        assert 'Бег' in body
        # Счётчики: участников 2, без результата 1
        assert 'Участников: <strong>2</strong>' in body
        assert 'Без результата: <strong>1</strong>' in body
        # Бейдж «ждёт результата» у записи без места
        assert 'ждёт результата' in body
        # Viewer — просмотр без кнопок управления
        assert 'participant-add-form' not in body
        assert 'participant-edit-button' not in body
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.list_calendar_event_participants.return_value = []


def test_calendar_event_page_result_filter(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.list_calendar_event_participants.return_value = _participants_sample()
    try:
        headers = get_auth_headers('viewer')
        _, response = client.get('/calendar/7?result=without', headers=headers)
        assert response.status == 200
        body = response.body.decode()
        # Фильтр клиентский на списке: показан только «без результата»
        assert 'Волков Артём Игоревич' in body
        assert 'Иванов Дмитрий Сергеевич' not in body
        # Показано 1 из 2 (между «из» и числом в шаблоне перенос строки)
        assert 'Показано <strong>1</strong>' in body
        assert '<strong>2</strong> участников' in body
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.list_calendar_event_participants.return_value = []


def test_calendar_event_page_forbidden_for_athlete(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    try:
        cookie = create_auth_cookie_value(username='sportik', role='athlete')
        headers = {'cookie': f'{settings.auth_cookie_name}={cookie}'}
        _, response = client.get('/calendar/7', headers=headers)
        assert response.status == 403
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


def test_calendar_event_page_not_found(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = None
    headers = get_auth_headers('editor')
    _, response = client.get('/calendar/999', headers=headers)
    assert response.status == 404


def test_calendar_event_page_invalid_id(client: SanicTestClient):
    headers = get_auth_headers('editor')
    _, response = client.get('/calendar/abc', headers=headers)
    assert response.status == 400


def test_calendar_event_page_participant_delete_button_roles(client: SanicTestClient):
    """Кнопка удаления участника на странице события — только admin
    (эндпоинт /competition/<id>/delete админский)."""
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.list_calendar_event_participants.return_value = _participants_sample()
    try:
        _, response = client.get('/calendar/7', headers=get_auth_headers('admin'))
        assert response.status == 200
        assert 'participant-delete-button' in response.body.decode()

        # Editor: правка есть, удаления нет
        _, response = client.get('/calendar/7', headers=get_auth_headers('editor'))
        body = response.body.decode()
        assert 'participant-edit-button' in body
        assert 'participant-delete-button' not in body

        # Viewer: колонки действий нет вовсе
        _, response = client.get('/calendar/7', headers=get_auth_headers('viewer'))
        assert 'participant-edit-button' not in response.body.decode()
        assert 'participant-delete-button' not in response.body.decode()

        # Athlete: страница события недоступна
        _, response = client.get('/calendar/7', headers=athlete_headers())
        assert response.status == 403
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.list_calendar_event_participants.return_value = []


# ---- Файл положения события календаря (решение 2026-09-22) и аудит
# удаления записи. Реальный SQLite-адаптер: колонки regulation_*, файлы в
# data/files/calendar/<event_id>/ и записи аудита.


@pytest.fixture
def calendar_client(client: SanicTestClient, tmp_path, monkeypatch):
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    storage = SQLiteAdapter(str(tmp_path / 'calendar.sqlite3'))
    storage.create_user(settings.auth_admin_username, 'hash', 'admin')
    storage.create_user(settings.auth_editor_username, 'hash', 'editor')
    if settings.auth_viewer_username:
        storage.create_user(settings.auth_viewer_username, 'hash', 'viewer')
    storage.create_user('sportik', 'hash', 'athlete')
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


def make_calendar_participant(name: str, event_name: str = 'Кросс СибАДИ', day: int = 25) -> Competition:
    return Competition(
        student_id=f'id-{name}',
        student_name=name,
        student_sex='М',
        institute='ИСИ',
        group='ПГС-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 6, day),
        level='внутривузовские',
        name=event_name,
        position=1,
        extra_data={},
    )


def upload_regulation(client, headers, event_id, filename, payload):
    _, response = client.post(
        f'/calendar/{event_id}/regulation',
        headers=headers,
        data=csrf_for(headers),
        files={'regulation': (filename, payload, 'application/octet-stream')},
        allow_redirects=False,
    )
    return response


def test_calendar_regulation_upload_download_replace_delete(calendar_client: SanicTestClient, tmp_path):
    storage = app.ctx.storage
    event_id = storage.create_calendar_event('Кросс СибАДИ', '2026-06-25', None, 'внутривузовские', 'Бег', '')
    admin = get_auth_headers('admin')

    # Страница без файла модератору: заглушка + форма прикрепления
    _, page = calendar_client.get(f'/calendar/{event_id}', headers=admin)
    assert 'Положение о соревновании' in page.body.decode()
    assert 'Положение не прикреплено.' in page.body.decode()
    assert 'Прикрепить положение' in page.body.decode()
    # viewer без файла секции не видит
    _, page = calendar_client.get(f'/calendar/{event_id}', headers=get_auth_headers('viewer'))
    assert 'Положение о соревновании' not in page.body.decode()

    # Загрузка валидного PDF
    pdf_bytes = b'%PDF-1.4 ' + b'0' * 32
    response = upload_regulation(calendar_client, admin, event_id, 'polozhenie.pdf', pdf_bytes)
    assert response.status == 302
    assert 'Положение прикреплено' in unquote_plus(response.headers['location'])
    event = storage.get_calendar_event(event_id)
    assert event['regulation_filename'] == 'polozhenie.pdf'
    stored_files = list((tmp_path / 'files' / 'calendar' / str(event_id)).iterdir())
    assert len(stored_files) == 1 and stored_files[0].name == event['regulation_stored_name']
    assert audit_details(storage, 'calendar_regulation_uploaded') == [
        {
            'event_id': event_id,
            'event_name': 'Кросс СибАДИ',
            'filename': 'polozhenie.pdf',
            'replaced': False,
        }
    ]

    # Скачивание viewer'ом: оригинальное имя в content-disposition
    _, response = calendar_client.get(
        f'/calendar/{event_id}/regulation', headers=get_auth_headers('viewer'), allow_redirects=False
    )
    assert response.status == 200
    assert response.body == pdf_bytes
    assert response.headers['content-disposition'] == 'attachment; filename="polozhenie.pdf"'
    assert response.headers['content-type'] == 'application/pdf'

    # Замена: колонки обновлены, старый файл удалён
    old_stored = event['regulation_stored_name']
    png_bytes = b'\x89PNG\r\n\x1a\n' + b'1' * 16
    response = upload_regulation(calendar_client, admin, event_id, 'новое.png', png_bytes)
    assert response.status == 302
    assert 'Положение заменено' in unquote_plus(response.headers['location'])
    event = storage.get_calendar_event(event_id)
    assert event['regulation_filename'] == 'новое.png'
    assert not (tmp_path / 'files' / 'calendar' / str(event_id) / old_stored).exists()
    # Журнал — свежие сверху: замена несёт replaced=True
    replaced_events = audit_details(storage, 'calendar_regulation_uploaded')
    assert len(replaced_events) == 2
    assert replaced_events[0] == {
        'event_id': event_id,
        'event_name': 'Кросс СибАДИ',
        'filename': 'новое.png',
        'replaced': True,
    }

    # Удаление: колонки NULL, каталог файла удалён
    _, response = calendar_client.post(
        f'/calendar/{event_id}/regulation/delete',
        headers=admin,
        data=csrf_for(admin),
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Положение удалено' in unquote_plus(response.headers['location'])
    event = storage.get_calendar_event(event_id)
    assert event['regulation_filename'] is None and event['regulation_stored_name'] is None
    assert not (tmp_path / 'files' / 'calendar' / str(event_id)).exists()
    assert audit_details(storage, 'calendar_regulation_deleted') == [
        {'event_id': event_id, 'event_name': 'Кросс СибАДИ', 'filename': 'новое.png'}
    ]

    # Повторное удаление без файла — ошибка, не 500
    _, response = calendar_client.post(
        f'/calendar/{event_id}/regulation/delete',
        headers=admin,
        data=csrf_for(admin),
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Положение не прикреплено.' in unquote_plus(response.headers['location'])


def test_calendar_regulation_rejects_invalid_files(calendar_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = storage.create_calendar_event('Кросс', '2026-06-25', None, '', '', '')
    admin = get_auth_headers('admin')

    # Пустая отправка
    _, response = calendar_client.post(
        f'/calendar/{event_id}/regulation', headers=admin, data=csrf_for(admin), allow_redirects=False
    )
    assert response.status == 302
    assert 'Выберите файл.' in unquote_plus(response.headers['location'])

    # Неизвестное расширение
    response = upload_regulation(calendar_client, admin, event_id, 'notes.txt', b'hello')
    assert response.status == 302
    assert 'Допустимы только PDF, JPEG и PNG' in unquote_plus(response.headers['location'])

    # Поддельная сигнатура
    response = upload_regulation(calendar_client, admin, event_id, 'fake.pdf', b'not a pdf at all')
    assert response.status == 302
    assert 'Допустимы только PDF, JPEG и PNG' in unquote_plus(response.headers['location'])

    # Больше 5 МБ
    response = upload_regulation(calendar_client, admin, event_id, 'big.pdf', b'%PDF- ' + b'0' * (5 * 1024 * 1024))
    assert response.status == 302
    assert 'Файл больше 5 МБ' in unquote_plus(response.headers['location'])

    event = storage.get_calendar_event(event_id)
    assert event['regulation_filename'] is None and event['regulation_stored_name'] is None
    assert not (Path(settings.data_folder) / 'files' / 'calendar').exists()


def test_calendar_regulation_access_rights(calendar_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = storage.create_calendar_event('Кросс', '2026-06-25', None, '', '', '')
    admin = get_auth_headers('admin')
    pdf_bytes = b'%PDF-1.4 ok'
    assert upload_regulation(calendar_client, admin, event_id, 'p.pdf', pdf_bytes).status == 302

    # Скачивание: viewer может, атлет — 403, аноним — редирект на вход
    _, response = calendar_client.get(
        f'/calendar/{event_id}/regulation', headers=get_auth_headers('viewer'), allow_redirects=False
    )
    assert response.status == 200
    _, response = calendar_client.get(f'/calendar/{event_id}/regulation', headers=athlete_headers())
    assert response.status == 403
    _, response = calendar_client.get(f'/calendar/{event_id}/regulation', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')

    # Upload/delete — только модераторы: viewer POST → 403
    viewer = get_auth_headers('viewer')
    _, response = calendar_client.post(
        f'/calendar/{event_id}/regulation',
        headers=viewer,
        data=csrf_for(viewer),
        files={'regulation': ('x.pdf', b'%PDF-1.4', 'application/pdf')},
    )
    assert response.status == 403
    _, response = calendar_client.post(f'/calendar/{event_id}/regulation/delete', headers=viewer, data=csrf_for(viewer))
    assert response.status == 403

    # Без CSRF-токена POST отклоняется middleware
    _, response = calendar_client.post(
        f'/calendar/{event_id}/regulation', headers=admin, data={}, allow_redirects=False
    )
    assert response.status == 403

    # Нет файла положения — скачивание 404
    other = storage.create_calendar_event('Вторая', '2026-07-01', None, '', '', '')
    _, response = calendar_client.get(
        f'/calendar/{other}/regulation', headers=get_auth_headers('viewer'), allow_redirects=False
    )
    assert response.status == 404


def test_calendar_event_deletion_removes_regulation_file(calendar_client: SanicTestClient, tmp_path):
    storage = app.ctx.storage
    event_id = storage.create_calendar_event('Кросс', '2026-06-25', None, '', '', '')
    admin = get_auth_headers('admin')
    assert upload_regulation(calendar_client, admin, event_id, 'p.pdf', b'%PDF-1.4 ok').status == 302
    regulation_dir = tmp_path / 'files' / 'calendar' / str(event_id)
    assert regulation_dir.is_dir()

    # У события нет участников — удаление доступно, каталог положения уходит тоже
    _, response = calendar_client.post(f'/calendar/{event_id}/delete', headers=admin, data=csrf_for(admin))
    assert response.status == 200
    assert storage.get_calendar_event(event_id) is None
    assert not regulation_dir.exists()


def test_delete_competition_writes_record_deleted_audit(calendar_client: SanicTestClient):
    storage = app.ctx.storage
    storage.save_competitions([make_calendar_participant('Иванов Дмитрий Сергеевич')])
    record_id = storage.connection.execute('SELECT id FROM competitions ORDER BY id').fetchone()['id']

    admin = get_auth_headers('admin')
    _, response = calendar_client.post(
        f'/competition/{record_id}/delete', headers=admin, data=csrf_for(admin), allow_redirects=False
    )
    assert response.status == 302
    assert response.headers['location'] == '/'
    assert storage.get_competition_review(record_id) is None
    assert audit_details(storage, 'record_deleted') == [
        {'record_id': record_id, 'student_name': 'Иванов Дмитрий Сергеевич'}
    ]


def test_editor_adds_participant_as_registry_record(client: SanicTestClient):
    # Участник — ОБЫЧНАЯ запись реестра: пресет события (name/date/date_to,
    # уровень, спорт) копируется в запись; никаких FK (историчность —
    # docs/data-model-decisions.md).
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.save_competitions.reset_mock()
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/7/participants',
            headers=headers,
            data={
                **csrf_for(headers),
                'student_name': 'Соколова Екатерина Дмитриевна',
                'student_sex': 'Ж',
                'institute': 'ИСИ',
                'group': 'ГТб-42',
                'course': '4',
                'position': '',
            },
            allow_redirects=False,
        )
        assert response.status == 302
        assert response.headers['location'].startswith('/calendar/7?admin_message=')
        assert 'Соколова Екатерина Дмитриевна' in unquote_plus(response.headers['location'])
        app.ctx.storage.save_competitions.assert_called_once()
        args, kwargs = app.ctx.storage.save_competitions.call_args
        competition = args[0][0]
        assert competition.student_name == 'Соколова Екатерина Дмитриевна'
        assert competition.name == 'Кросс СибАДИ'
        assert competition.date.isoformat() == '2026-06-25T00:00:00'
        assert competition.date_to is None
        assert competition.level == 'внутривузовские'
        assert competition.sport == 'Бег'
        # Место опционально («дописать позже») → position = 0.
        assert competition.position == 0
        assert kwargs['review_status'] == 'approved'
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.save_competitions.reset_mock()


def test_editor_adds_participant_with_place(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.save_competitions.reset_mock()
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/7/participants',
            headers=headers,
            data={
                **csrf_for(headers),
                'student_name': 'Новый Атлет',
                'student_sex': 'М',
                'institute': '',
                'group': '',
                'course': '2',
                'position': '3',
            },
        )
        assert response.status == 200
        args, _ = app.ctx.storage.save_competitions.call_args
        assert args[0][0].position == 3
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.save_competitions.reset_mock()


def test_add_participant_requires_name(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.save_competitions.reset_mock()
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/7/participants',
            headers=headers,
            data={**csrf_for(headers), 'student_name': '   ', 'course': '1'},
            allow_redirects=False,
        )
        assert response.status == 302
        assert 'admin_error' in response.headers['location']
        app.ctx.storage.save_competitions.assert_not_called()
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.save_competitions.reset_mock()


def test_viewer_cannot_add_participant(client: SanicTestClient):
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    app.ctx.storage.save_competitions.reset_mock()
    try:
        headers = get_auth_headers('viewer')
        _, response = client.post(
            '/calendar/7/participants',
            headers=headers,
            data={**csrf_for(headers), 'student_name': 'Кто-то', 'course': '1'},
            allow_redirects=False,
        )
        assert response.status == 403
        app.ctx.storage.save_competitions.assert_not_called()
    finally:
        app.ctx.storage.get_calendar_event.return_value = None
        app.ctx.storage.save_competitions.reset_mock()


def test_calendar_event_edit_redirects_back_to_event_page(client: SanicTestClient):
    # Правка со страницы соревнования (next=/calendar/7) возвращает внутрь
    # события; открытый редирект исключён — только пути /calendar/.
    app.ctx.storage.get_calendar_event.return_value = _event_for_page()
    try:
        headers = get_auth_headers('editor')
        _, response = client.post(
            '/calendar/7/edit',
            headers=headers,
            data={
                **csrf_for(headers),
                'name': 'Кросс СибАДИ',
                'date': '25.06.2026',
                'level': 'внутривузовские',
                'sport': 'Бег',
                'url': '',
                'next': '/calendar/7',
            },
            allow_redirects=False,
        )
        assert response.status == 302
        assert response.headers['location'].startswith('/calendar/7?admin_message=')
        assert 'Соревнование обновлено' in unquote_plus(response.headers['location'])

        _, response = client.post(
            '/calendar/7/edit',
            headers=headers,
            data={
                **csrf_for(headers),
                'name': 'Кросс СибАДИ',
                'date': '25.06.2026',
                'next': 'https://evil.example',
            },
            allow_redirects=False,
        )
        assert response.status == 302
        assert response.headers['location'].startswith('/calendar?admin_message=')
        assert 'Соревнование обновлено' in unquote_plus(response.headers['location'])
    finally:
        app.ctx.storage.get_calendar_event.return_value = None


# --- Карточки студентов (Student Identity v1, Phase 1 — фундамент). ---
# Полные сценарии — на реальном SQLite-адаптере (паттерн reports_client):
# валидация, редиректы, flash-сообщения и аудит проверяются вместе с данными.


@pytest.fixture
def people_client(client: SanicTestClient, tmp_path):
    storage = SQLiteAdapter(str(tmp_path / 'people.sqlite3'))
    storage.create_user(settings.auth_admin_username, 'hash', 'admin')
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


def create_person(client: SanicTestClient, **overrides) -> object:
    headers = get_auth_headers()
    data = {
        **csrf_for(headers),
        'full_name': 'Иванов Иван Иванович',
        'sex': 'М',
        'institute': 'ИСИ',
        'group': 'ПГС-101',
        'course': '2',
        **overrides,
    }
    _, response = client.post('/admin/people', headers=headers, data=data, allow_redirects=False)
    return response


def test_admin_people_requires_admin(client: SanicTestClient):
    _, response = client.get('/admin/people', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = client.get('/admin/people', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')

    _, response = client.post('/admin/people', data={'full_name': 'Кто-то'})
    assert response.status == 401


def test_admin_people_create_and_card_flow(people_client: SanicTestClient):
    _, response = people_client.get('/admin/people', headers=get_auth_headers())
    assert response.status == 200
    assert 'Студентов пока нет.' in response.text

    response = create_person(people_client)
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/1?')
    assert 'Студент «Иванов Иван Иванович» добавлен.' in location

    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert response.status == 200
    assert 'Иванов Иван Иванович' in response.text
    assert 'ИСИ' in response.text
    assert 'ПГС-101' in response.text
    assert 'Псевдонимов пока нет.' in response.text

    # Поиск: найденная карточка и счётчик «Найдено: N из M»
    _, response = people_client.get('/admin/people?q=Иванов', headers=get_auth_headers())
    assert response.status == 200
    assert 'Найдено: 1 из 1' in response.text
    _, response = people_client.get('/admin/people?q=Неттакова', headers=get_auth_headers())
    assert 'По запросу «Неттакова» ничего не найдено.' in response.text


def test_admin_people_create_requires_full_name(people_client: SanicTestClient):
    response = create_person(people_client, full_name='   ')
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Укажите ФИО студента.' in location

    response = create_person(people_client, sex='Оно')
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'admin_error' in location

    # Ни одна из попыток не создала карточку
    _, response = people_client.get('/admin/people', headers=get_auth_headers())
    assert 'Студентов пока нет.' in response.text


def test_admin_person_edit_updates_fields(people_client: SanicTestClient):
    create_person(people_client)
    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people/1/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'full_name': 'Иванов Иван Ильич',
            'sex': '',
            'institute': 'ИМИ',
            'group': 'СБ-202',
            'course': '3',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'Данные студента «Иванов Иван Ильич» сохранены.' in location

    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert 'Иванов Иван Ильич' in response.text
    assert 'СБ-202' in response.text
    assert 'ПГС-101' not in response.text


def test_admin_person_toggle_active(people_client: SanicTestClient):
    create_person(people_client)
    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'помечен неактивным' in unquote_plus(response.headers['location'])

    _, response = people_client.get('/admin/people', headers=get_auth_headers())
    assert 'неактивен' in response.text

    _, response = people_client.post(
        '/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'снова активен' in unquote_plus(response.headers['location'])


def test_admin_person_alias_add_and_remove(people_client: SanicTestClient):
    create_person(people_client)
    headers = get_auth_headers()

    _, response = people_client.post(
        '/admin/people/1/alias', headers=headers, data={**csrf_for(headers), 'name': ''}, allow_redirects=False
    )
    assert 'admin_error' in unquote_plus(response.headers['location'])

    _, response = people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Псевдоним «Иванов И.И.» добавлен.' in unquote_plus(response.headers['location'])

    _, response = people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'Псевдоним «Иванов И.И.» уже есть у этого студента.' in location

    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert 'Иванов И.И.' in response.text

    _, response = people_client.post(
        '/admin/people/1/alias/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'Псевдоним «Иванов И.И.» удалён.' in unquote_plus(response.headers['location'])

    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert 'Псевдонимов пока нет.' in response.text


def test_student_admin_actions_write_audit(people_client: SanicTestClient):
    storage = app.ctx.storage
    headers = get_auth_headers()

    create_person(people_client)
    people_client.post(
        '/admin/people/1/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'full_name': 'Иванов Иван Иванович',
            'sex': 'Ж',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
        },
        allow_redirects=False,
    )
    people_client.post('/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False)
    people_client.post('/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False)
    people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )
    people_client.post('/admin/people/1/alias/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False)

    events = [(event['action'], json.loads(event['details'])) for event in storage.get_audit_events(limit=20)]
    actions = [action for action, _ in events]
    assert actions == [
        'student_alias_removed',
        'student_alias_added',
        'student_activated',
        'student_deactivated',
        'student_updated',
        'student_created',
    ]

    by_action = dict(reversed(events))
    assert by_action['student_created'] == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}
    # Diff — только изменённые поля, old→new
    assert by_action['student_updated']['changed'] == {'sex': {'old': 'М', 'new': 'Ж'}}
    assert by_action['student_deactivated'] == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}
    assert by_action['student_activated'] == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}
    assert by_action['student_alias_added'] == {'student_id': 1, 'name': 'Иванов И.И.'}
    assert by_action['student_alias_removed'] == {'student_id': 1, 'alias_id': 1, 'name': 'Иванов И.И.'}


def test_admin_person_endpoints_reject_missing_csrf(people_client: SanicTestClient):
    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people', headers=headers, data={'full_name': 'Без токена'}, allow_redirects=False
    )
    assert response.status == 403
    assert 'CSRF' in response.text

    _, response = people_client.post('/admin/people/1/delete', headers=headers, allow_redirects=False)
    assert response.status == 403
    assert 'CSRF' in response.text


def test_admin_person_unknown_id_redirects(people_client: SanicTestClient):
    _, response = people_client.get('/admin/people/999999', headers=get_auth_headers(), allow_redirects=False)
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Студент не найден.' in location

    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people/999999/edit',
        headers=headers,
        data={**csrf_for(headers), 'full_name': 'Никто'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Студент не найден.' in unquote_plus(response.headers['location'])


def test_admin_person_non_numeric_id_returns_400(people_client: SanicTestClient):
    _, response = people_client.get('/admin/people/abc', headers=get_auth_headers())
    assert response.status == 400

    headers = get_auth_headers()
    _, response = people_client.post('/admin/people/abc/active', headers=headers, data=csrf_for(headers))
    assert response.status == 400

    _, response = people_client.post('/admin/people/abc/delete', headers=headers, data=csrf_for(headers))
    assert response.status == 400


# --- Полное удаление карточки студента (hard delete, admin-only). ---


def save_person_record(storage, student_name: str) -> int:
    """Синтетическая запись реестра для связи с карточкой; возвращает id."""
    storage.save_competitions(
        [make_report_record(student_name, 'М', 'ИСИ', 'ПГС-101', 1, 'Бег', datetime(2024, 3, 1), 'внутривузовские', 1)]
    )
    row = storage.connection.execute('SELECT id FROM competitions ORDER BY id DESC LIMIT 1').fetchone()
    return row['id']


def test_admin_person_delete_requires_admin_role(people_client: SanicTestClient):
    create_person(people_client)
    # sportik должен существовать в реальном storage, иначе cookie не распарсится
    app.ctx.storage.create_user('sportik', 'hash', 'athlete')
    for role in ('editor', 'viewer'):
        headers = get_auth_headers(role)
        _, response = people_client.post(
            '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
        )
        assert response.status == 403

    headers = athlete_headers()
    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 403

    _, response = people_client.post('/admin/people/1/delete', allow_redirects=False)
    assert response.status == 401

    # Ни одна попытка ничего не удалила
    assert app.ctx.storage.get_student_by_id(1) is not None


def test_admin_person_delete_removes_card_aliases_and_hides_everywhere(people_client: SanicTestClient):
    storage = app.ctx.storage
    create_person(people_client)
    headers = get_auth_headers()
    people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )
    # Псевдоним с тем же именем у ДРУГОЙ карточки — не должен пострадать
    create_person(people_client, full_name='Сидор Сидор Сидорович')
    people_client.post(
        '/admin/people/2/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )

    # Блок удаления на карточке без связей: форма с confirm и кнопкой
    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert 'Удаление карточки' in response.text
    assert '/admin/people/1/delete' in response.text
    assert 'Удалить карточку «Иванов Иван Иванович» (#1) безвозвратно?' in response.text
    assert 'Удаление недоступно' not in response.text

    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Карточка «Иванов Иван Иванович» (#1) удалена.' in location

    # Карточка и её псевдонимы исчезли из БД; чужой псевдоним на месте
    assert storage.get_student_by_id(1) is None
    assert storage.list_student_aliases(1) == []
    assert [alias['name'] for alias in storage.list_student_aliases(2)] == ['Иванов И.И.']

    # Нет в списке и в поиске страницы
    _, response = people_client.get('/admin/people', headers=get_auth_headers())
    assert response.status == 200
    assert 'Иванов Иван Иванович' not in response.text
    _, response = people_client.get('/admin/people?q=Иванов', headers=get_auth_headers())
    assert 'По запросу «Иванов» ничего не найдено.' in response.text

    # Нет в кандидатах сопоставления и подсказках (storage-level);
    # одноимённый псевдоним ДРУГОЙ карточки по-прежнему находится
    assert storage.find_student_candidates('Иванов Иван Иванович') == []
    assert [candidate['student_id'] for candidate in storage.find_student_candidates('Иванов И.И.')] == [2]
    assert [candidate['student_id'] for candidate in storage.search_student_candidates('Иванов')] == [2]

    # Аудит успешного удаления
    events = storage.get_audit_events(limit=5)
    deleted = next(event for event in events if event['action'] == 'student_deleted')
    assert json.loads(deleted['details']) == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}


def test_admin_person_delete_blocked_by_linked_records(people_client: SanicTestClient):
    storage = app.ctx.storage
    create_person(people_client)
    headers = get_auth_headers()
    people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )
    record_id = save_person_record(storage, 'Старое ФИО')
    storage.link_competitions([record_id], 1)
    before = storage.connection.execute('SELECT * FROM competitions WHERE id = ?', (record_id,)).fetchone()

    # Блок удаления на карточке со связями: вместо формы — счётчик блокера
    _, response = people_client.get('/admin/people/1', headers=get_auth_headers())
    assert 'Удаление недоступно' in response.text
    assert 'записей о соревнованиях — <strong>1</strong>' in response.text
    assert '/admin/people/1/delete' not in response.text

    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/1?')
    assert 'Удалить карточку нельзя' in location
    assert 'записей — 1' in location
    assert 'аккаунтов' not in location

    # Записи неизменны: состав и количество, байт-в-байт
    after = storage.connection.execute('SELECT * FROM competitions WHERE id = ?', (record_id,)).fetchone()
    assert tuple(after) == tuple(before)
    assert storage.linked_records_count(1) == 1

    # Карточка осталась со всеми полями и псевдонимами
    student = storage.get_student_by_id(1)
    assert student is not None
    assert student['full_name'] == 'Иванов Иван Иванович'
    assert student['sex'] == 'М'
    assert student['institute'] == 'ИСИ'
    assert student['group_name'] == 'ПГС-101'
    assert student['course'] == '2'
    assert student['active'] == 1
    assert [alias['name'] for alias in storage.list_student_aliases(1)] == ['Иванов И.И.']


def test_admin_person_delete_blocked_by_linked_account(people_client: SanicTestClient):
    storage = app.ctx.storage
    create_person(people_client)
    storage.create_user('anna', 'hash', 'athlete')
    anna_id = storage.get_user('anna')['id']
    storage.link_user(anna_id, 1)

    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'Удалить карточку нельзя' in location
    assert 'аккаунтов атлета — 1' in location
    assert 'записей' not in location

    # Связь аккаунта не тронута, карточка на месте
    anna_ref = storage.connection.execute('SELECT student_ref_id FROM users WHERE id = ?', (anna_id,)).fetchone()
    assert anna_ref['student_ref_id'] == 1
    assert storage.get_student_by_id(1) is not None


def test_admin_person_delete_blocked_writes_no_audit_and_toggle_audit_unchanged(
    people_client: SanicTestClient,
):
    """Заблокированная попытка в аудит не пишется; deactivate/activate
    аудируются как раньше (student_deactivated/student_activated)."""
    storage = app.ctx.storage
    create_person(people_client)
    headers = get_auth_headers()
    record_id = save_person_record(storage, 'Старое ФИО')
    storage.link_competitions([record_id], 1)

    people_client.post('/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False)
    people_client.post('/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False)
    people_client.post('/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False)

    events = [(event['action'], json.loads(event['details'])) for event in storage.get_audit_events(limit=10)]
    actions = [action for action, _ in events]
    assert actions == [
        'student_activated',
        'student_deactivated',
        'student_created',
    ]
    by_action = dict(events)
    assert by_action['student_deactivated'] == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}
    assert by_action['student_activated'] == {'student_id': 1, 'full_name': 'Иванов Иван Иванович'}
    assert 'student_deleted' not in by_action


def test_admin_person_delete_inactive_card_without_links(people_client: SanicTestClient):
    storage = app.ctx.storage
    create_person(people_client)
    headers = get_auth_headers()
    people_client.post('/admin/people/1/active', headers=headers, data=csrf_for(headers), allow_redirects=False)

    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'Карточка «Иванов Иван Иванович» (#1) удалена.' in unquote_plus(response.headers['location'])
    assert storage.get_student_by_id(1) is None


def test_admin_person_delete_get_deleted_id_redirects_not_found(people_client: SanicTestClient):
    create_person(people_client)
    headers = get_auth_headers()
    people_client.post('/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False)

    _, response = people_client.get('/admin/people/1', headers=get_auth_headers(), allow_redirects=False)
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Студент не найден.' in location


def test_admin_person_delete_unknown_id_redirects(people_client: SanicTestClient):
    headers = get_auth_headers()
    _, response = people_client.post(
        '/admin/people/999999/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Студент не найден.' in location


def test_admin_person_delete_rolls_back_on_mid_delete_failure(people_client: SanicTestClient, monkeypatch):
    """Сбой между DELETE псевдонимов и DELETE карточки: полный откат,
    понятная ошибка сервера, ни карточки, ни псевдонима не потеряны."""
    storage = app.ctx.storage
    create_person(people_client)
    headers = get_auth_headers()
    people_client.post(
        '/admin/people/1/alias',
        headers=headers,
        data={**csrf_for(headers), 'name': 'Иванов И.И.'},
        allow_redirects=False,
    )

    monkeypatch.setattr(storage, 'connection', FailingStudentsDeleteConnection(storage.connection))
    _, response = people_client.post(
        '/admin/people/1/delete', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 500
    monkeypatch.undo()

    student = storage.get_student_by_id(1)
    assert student is not None and student['full_name'] == 'Иванов Иван Иванович'
    assert [alias['name'] for alias in storage.list_student_aliases(1)] == ['Иванов И.И.']
    # Аудита удаления нет — его и не было
    assert all(event['action'] != 'student_deleted' for event in storage.get_audit_events(limit=10))

    _, response = people_client.post('/admin/people/1/alias/abc/delete', headers=headers, data=csrf_for(headers))
    assert response.status == 400


# --- Сопоставление данных (Student Identity v1, Phase 2). ---
# Полные сценарии на реальном SQLite-адаптере (паттерн people_client):
# привязки записей/аккаунтов к карточкам, аудит и главная гарантия —
# runtime (кабинет атлета, отчёты) не замечает student_ref_id.


@pytest.fixture
def reconcile_client(client: SanicTestClient, tmp_path):
    storage = SQLiteAdapter(str(tmp_path / 'reconcile.sqlite3'))
    storage.create_user(settings.auth_admin_username, 'hash', 'admin')
    # Атлет для проверок кабинета (athlete_headers использует логин sportik).
    storage.create_user('sportik', 'hash', 'athlete')
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    try:
        yield client
    finally:
        app.ctx.storage = previous


def save_reconcile_record(storage, name: str, date: datetime, position: int = 1, owner_id=None) -> int:
    """Одна одобренная запись реестра; возвращает её id.

    Легаси-ключ личности — sha256(ФИО), как в реальных данных. owner_id —
    владелец записи (P3: ветка owner в scope кабинета атлета).
    """
    storage.save_competitions(
        [
            Competition(
                student_id=sha256(name.encode()).hexdigest(),
                student_name=name,
                student_sex='М',
                institute='ИСИ',
                group='ПГС-101',
                course=2,
                sport='Бег',
                date=date,
                level='внутривузовские',
                name='Кубок',
                position=position,
                extra_data={},
            )
        ],
        owner_id=owner_id,
    )
    row = storage.connection.execute('SELECT MAX(id) AS id FROM competitions').fetchone()
    return row['id']


def record_ref(storage, record_id: int):
    row = storage.connection.execute('SELECT student_ref_id FROM competitions WHERE id = ?', (record_id,)).fetchone()
    return row['student_ref_id'] if row else 'missing'


def user_ref(storage, user_id: int):
    row = storage.connection.execute('SELECT student_ref_id FROM users WHERE id = ?', (user_id,)).fetchone()
    return row['student_ref_id'] if row else 'missing'


def audit_details(storage, action: str) -> list[dict]:
    events = storage.get_audit_events(limit=100)
    return [json.loads(event['details']) for event in events if event['action'] == action]


def test_reconcile_pages_require_admin(reconcile_client: SanicTestClient):
    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = reconcile_client.get('/admin/people/reconcile/users', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = reconcile_client.get('/admin/people/reconcile', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')

    _, response = reconcile_client.post('/admin/people/reconcile/link', data={'record_ids[]': ['1']})
    assert response.status == 401

    _, response = reconcile_client.post('/admin/people/reconcile/users/link', data={'user_id': '1'})
    assert response.status == 401


def test_reconcile_route_priority_over_student_id(reconcile_client: SanicTestClient):
    """/admin/people/reconcile — статический маршрут, а не id карточки."""
    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert response.status == 200
    assert 'Записи без студента' in response.text

    _, response = reconcile_client.get('/admin/people/reconcile/users', headers=get_auth_headers())
    assert response.status == 200
    assert 'Аккаунты атлетов без студента' in response.text


def test_reconcile_records_tab_renders_and_paginates(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    ivanov_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    save_reconcile_record(storage, 'Петров Пётр', datetime(2026, 2, 10))

    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert response.status == 200
    # Сводка счётчиков: всего/без студента по записям и аккаунтам
    assert 'Записей о соревнованиях: 2' in response.text
    assert 'Аккаунтов атлетов: 1' in response.text
    assert '<strong>2</strong>' in response.text
    assert '<strong>1</strong>' in response.text
    # Запись и кандидаты: точное совпадение ФИО — карточка Иванова
    assert 'Иванов Иван' in response.text
    assert f'<a href="/admin/people/{student_id}">Иванов Иван</a>' in response.text
    assert 'Подходящих студентов не найдено.' in response.text
    assert 'Создать студента из записи' in response.text
    # Возврат для форм — текущая страница без поиска
    assert 'value="/admin/people/reconcile?page=1"' in response.text

    # Пагинация в стиле журнала аудита: 55 записей «Бегун» — две страницы
    for index in range(55):
        save_reconcile_record(storage, f'Бегун Бежит {index:02d}', datetime(2026, 3, 1))
    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert 'страница 1 из 2' in response.text
    assert 'Всего по фильтру: 57' in response.text
    assert '/admin/people/reconcile?page=2' in response.text

    # Страница 2: хвост «Бегунов» + Иванов и Петров (алфавит), первой страницы тут нет
    _, response = reconcile_client.get('/admin/people/reconcile?page=2', headers=get_auth_headers())
    assert 'страница 2 из 2' in response.text
    assert 'Бегун Бежит 00' not in response.text
    assert 'Иванов Иван' in response.text
    assert f'#{ivanov_id}' in response.text

    # Поиск сохраняется в ссылках пагинации
    _, response = reconcile_client.get(
        '/admin/people/reconcile?q=%D0%91%D0%B5%D0%B3%D1%83%D0%BD', headers=get_auth_headers()
    )
    assert 'страница 1 из 2' in response.text
    assert 'q=%D0%91%D0%B5%D0%B3%D1%83%D0%BD&amp;page=2' in response.text

    # Пустые состояния: поиск без результата и полностью пустой раздел
    _, response = reconcile_client.get('/admin/people/reconcile?q=%D0%9D%D0%B5%D1%82', headers=get_auth_headers())
    assert 'непривязанных записей не найдено' in response.text

    storage.connection.execute('UPDATE competitions SET student_ref_id = 1')
    storage.connection.commit()
    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert 'Всё сопоставлено: записей без студента нет.' in response.text


def test_reconcile_records_tab_does_not_autolink(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))

    _, response = reconcile_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert response.status == 200

    # Кандидаты — только предложения: страница ничего не связывает сама
    assert record_ref(storage, record_id) is None
    user = storage.get_user('sportik')
    assert user_ref(storage, user['id']) is None


def test_reconcile_single_link_via_candidate(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': [str(record_id)],
            'student_id': str(student_id),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/reconcile?')
    assert 'Запись №1 привязана к студенту «Иванов Иван».' in location

    assert record_ref(storage, record_id) == student_id
    # Снимок записи и легаси-ключ не тронуты
    row = storage.connection.execute('SELECT * FROM competitions WHERE id = ?', (record_id,)).fetchone()
    assert row['student_name'] == 'Иванов Иван'
    assert row['student_id'] == sha256('Иванов Иван'.encode()).hexdigest()
    assert audit_details(storage, 'competition_linked_to_student') == [
        {'record_id': record_id, 'old_ref': None, 'new_ref': student_id, 'student_id': student_id}
    ]


def test_reconcile_bulk_link_single_audit_event(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    first = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    second = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 2, 10))

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': [str(first), str(second)],
            'student_id': str(student_id),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'Привязано записей: 2 — студент «Иванов Иван».' in location

    assert record_ref(storage, first) == student_id
    assert record_ref(storage, second) == student_id
    # Ровно одно событие на всю массовую привязку
    assert audit_details(storage, 'student_records_bulk_linked') == [
        {'student_id': student_id, 'record_ids': [first, second], 'count': 2}
    ]
    assert audit_details(storage, 'competition_linked_to_student') == []


def test_reconcile_bulk_rejects_partially_linked(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    first_student = storage.create_student('Первый', 'М', '', '', '')
    second_student = storage.create_student('Второй', 'Ж', '', '', '')
    first = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    second = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 2, 10))
    storage.link_competitions([first], first_student)

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': [str(first), str(second)],
            'student_id': str(second_student),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert 'Ничего не привязано: часть выбранных записей уже привязана. Обновите список и повторите.' in location
    # Всё или ничего: вторая запись не связана, первая осталась у прежнего студента
    assert record_ref(storage, first) == first_student
    assert record_ref(storage, second) is None
    assert audit_details(storage, 'student_records_bulk_linked') == []

    # Одиночная привязка занятой записи — свой текст
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': [str(first)],
            'student_id': str(second_student),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Запись №1 уже привязана к студенту.' in location
    assert record_ref(storage, first) == first_student


def test_reconcile_link_validation_errors(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    storage.create_student('Иванов Иван', 'М', '', '', '')
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    headers = get_auth_headers()

    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={**csrf_for(headers), 'student_id': '1', 'back': '/admin/people/reconcile'},
        allow_redirects=False,
    )
    assert 'Не выбрано ни одной записи.' in unquote_plus(response.headers['location'])

    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={**csrf_for(headers), 'record_ids[]': ['1'], 'back': '/admin/people/reconcile'},
        allow_redirects=False,
    )
    assert 'Выберите студента.' in unquote_plus(response.headers['location'])

    # Неактивная карточка — привязка запрещена с понятным текстом
    inactive = storage.create_student('Спящий Студент', 'М', '', '', '')
    storage.set_student_active(inactive, False)
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': ['1'],
            'student_id': str(inactive),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Студент «Спящий Студент» неактивен: привязка возможна только к активным студентам.' in location
    assert record_ref(storage, 1) is None

    # Неизвестная карточка
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={**csrf_for(headers), 'record_ids[]': ['1'], 'student_id': '999999', 'back': '/x'},
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Студент не найден.' in location
    # Недоверенный back заменён вкладкой сопоставления
    assert location.startswith('/admin/people/reconcile?')


def test_reconcile_create_from_record(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    headers = get_auth_headers()

    _, response = reconcile_client.get(
        f'/admin/people/reconcile/create?record_id={record_id}&back=/admin/people/reconcile',
        headers=headers,
    )
    assert response.status == 200
    assert 'Источник — запись №1' in response.text
    assert 'value="Иванов Иван"' in response.text
    assert 'создан, запись №1 привязана' in response.text or 'сразу привяжется' in response.text
    assert 'Сама запись' not in response.text or 'не изменится' in response.text

    _, response = reconcile_client.post(
        '/admin/people/reconcile/create',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_id': str(record_id),
            'back': '/admin/people/reconcile',
            'full_name': 'Иванов Иван',
            'sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-101',
            'course': '2',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/1?')
    assert 'Студент «Иванов Иван» создан, запись №1 привязана к нему.' in location

    student = storage.get_student_by_id(1)
    assert student['full_name'] == 'Иванов Иван'
    assert record_ref(storage, record_id) == 1
    # Снимок записи не переписан
    row = storage.connection.execute('SELECT * FROM competitions WHERE id = ?', (record_id,)).fetchone()
    assert row['student_name'] == 'Иванов Иван'
    assert row['institute'] == 'ИСИ'
    assert audit_details(storage, 'student_created') == [
        {'student_id': 1, 'full_name': 'Иванов Иван', 'source': 'reconciliation-record'}
    ]
    assert audit_details(storage, 'competition_linked_to_student')[0]['record_id'] == record_id

    # Гонка: запись уже связали — студент НЕ создаётся, запись не трогаем
    another = storage.create_student('Другой', 'Ж', '', '', '')
    storage.unlink_competition(record_id)
    assert storage.link_competitions([record_id], another) == (1, None)
    _, response = reconcile_client.post(
        '/admin/people/reconcile/create',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_id': str(record_id),
            'back': '/admin/people/reconcile',
            'full_name': 'Ещё Студент',
            'sex': '',
            'institute': '',
            'group': '',
            'course': '',
        },
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Запись №1 уже привязана к студенту.' in location
    assert record_ref(storage, record_id) == another
    # Новая карточка не создана: в базе по-прежнему «Другой» и исходный студент
    assert [student['full_name'] for student in storage.list_students()] == ['Другой', 'Иванов Иван']

    # Окно гонки между проверкой и привязкой: карточка уже создана — она
    # остаётся, запись не меняется, flash объясняет (симуляция подменой ответа
    # хранилища на already_linked). Запись на момент запроса свободна.
    storage.unlink_competition(record_id)
    original_link = storage.link_competitions
    storage.link_competitions = lambda record_ids, student_id: (0, 'already_linked')
    try:
        _, response = reconcile_client.post(
            '/admin/people/reconcile/create',
            headers=headers,
            data={
                **csrf_for(headers),
                'record_id': str(record_id),
                'back': '/admin/people/reconcile',
                'full_name': 'Гонка Гонщиков',
                'sex': '',
                'institute': '',
                'group': '',
                'course': '',
            },
            allow_redirects=False,
        )
    finally:
        storage.link_competitions = original_link
    location = unquote_plus(response.headers['location'])
    assert 'Студент «Гонка Гонщиков» создан, но запись №1 уже была привязана другим действием.' in location
    assert location.startswith('/admin/people/3?')
    # Запись в окне гонки не была тронута этим запросом
    assert record_ref(storage, record_id) is None
    assert storage.get_student_by_id(3)['full_name'] == 'Гонка Гонщиков'

    # GET-ошибки: неизвестная запись и уже привязанная
    _, response = reconcile_client.get(
        '/admin/people/reconcile/create?record_id=999999&back=/admin/people/reconcile',
        headers=headers,
        allow_redirects=False,
    )
    assert 'Запись не найдена.' in unquote_plus(response.headers['location'])

    storage.link_competitions([record_id], another)
    _, response = reconcile_client.get(
        f'/admin/people/reconcile/create?record_id={record_id}&back=/admin/people/reconcile',
        headers=headers,
        allow_redirects=False,
    )
    assert 'Запись №1 уже привязана к студенту.' in unquote_plus(response.headers['location'])


def test_reconcile_create_from_user_profile(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    user = storage.get_user('sportik')
    storage.set_profile(
        user['id'],
        {'student_name': 'Спортсменов Спорт', 'student_sex': 'М', 'institute': 'ИСИ', 'group': 'СБ-101', 'course': '1'},
    )
    storage.add_name_alias(user['id'], 'Спортсменов С.С.')
    headers = get_auth_headers()

    # Вкладка аккаунтов: карточка с профилем и предупреждением о псевдонимах
    _, response = reconcile_client.get('/admin/people/reconcile/users', headers=get_auth_headers())
    assert response.status == 200
    assert 'sportik' in response.text
    assert 'Спортсменов Спорт' in response.text
    assert 'Псевдонимы ФИО аккаунта' in response.text
    assert 'Спортсменов С.С.' in response.text
    assert 'не копирует псевдонимы аккаунта студенту' in response.text

    _, response = reconcile_client.get(
        f'/admin/people/reconcile/users/create?user_id={user["id"]}&back=/admin/people/reconcile/users',
        headers=headers,
    )
    assert response.status == 200
    assert 'Источник — аккаунт sportik (профиль атлета)' in response.text
    assert 'value="Спортсменов Спорт"' in response.text

    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/create',
        headers=headers,
        data={
            **csrf_for(headers),
            'user_id': str(user['id']),
            'back': '/admin/people/reconcile/users',
            'full_name': 'Спортсменов Спорт',
            'sex': 'М',
            'institute': 'ИСИ',
            'group': 'СБ-101',
            'course': '1',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/1?')
    assert 'Студент «Спортсменов Спорт» создан, аккаунт sportik привязан к нему.' in location

    assert user_ref(storage, user['id']) == 1
    # Псевдонимы аккаунта НЕ перенесены в student_aliases и остались у аккаунта
    assert storage.list_student_aliases(1) == []
    assert storage.get_name_aliases(user['id']) == ['Спортсменов С.С.']
    assert audit_details(storage, 'student_created') == [
        {'student_id': 1, 'full_name': 'Спортсменов Спорт', 'source': 'reconciliation-user'}
    ]
    assert audit_details(storage, 'user_linked_to_student') == [
        {'user_id': user['id'], 'username': 'sportik', 'old_ref': None, 'new_ref': 1}
    ]

    # Повторное создание из уже привязанного аккаунта — ошибка без новой карточки
    _, response = reconcile_client.get(
        f'/admin/people/reconcile/users/create?user_id={user["id"]}&back=/admin/people/reconcile/users',
        headers=headers,
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Аккаунт sportik уже привязан к студенту.' in location
    assert len(storage.list_students()) == 1


def test_reconcile_user_link_unlink_relink(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    user = storage.get_user('sportik')
    admin_user = storage.get_user(settings.auth_admin_username)
    first = storage.create_student('Иванов Иван', 'М', '', '', '')
    second = storage.create_student('Новый ФИО', 'Ж', '', '', '')
    headers = get_auth_headers()

    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'user_id': str(user['id']),
            'student_id': str(first),
            'back': '/admin/people/reconcile/users',
        },
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert 'Аккаунт sportik привязан к студенту «Иванов Иван».' in location

    # Повторная привязка и не-атлет
    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'user_id': str(user['id']),
            'student_id': str(second),
            'back': '/admin/people/reconcile/users',
        },
        allow_redirects=False,
    )
    assert 'Аккаунт sportik уже привязан к студенту.' in unquote_plus(response.headers['location'])

    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'user_id': str(admin_user['id']),
            'student_id': str(first),
            'back': '/admin/people/reconcile/users',
        },
        allow_redirects=False,
    )
    assert 'Пользователь не является атлетом.' in unquote_plus(response.headers['location'])
    assert user_ref(storage, admin_user['id']) is None

    # Перепривязка и отвязка
    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/relink',
        headers=headers,
        data={**csrf_for(headers), 'user_id': str(user['id']), 'student_id': str(second)},
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert location.startswith(f'/admin/people/{second}?')
    assert 'Аккаунт sportik перепривязан на студента «Новый ФИО».' in location
    assert user_ref(storage, user['id']) == second

    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/unlink',
        headers=headers,
        data={**csrf_for(headers), 'user_id': str(user['id'])},
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert location.startswith(f'/admin/people/{second}?')
    assert 'Аккаунт sportik отвязан от студента «Новый ФИО».' in location
    assert user_ref(storage, user['id']) is None

    assert audit_details(storage, 'user_linked_to_student') == [
        {'user_id': user['id'], 'username': 'sportik', 'old_ref': None, 'new_ref': first}
    ]
    assert audit_details(storage, 'user_relinked_to_student') == [
        {'user_id': user['id'], 'username': 'sportik', 'old_ref': first, 'new_ref': second}
    ]
    assert audit_details(storage, 'user_unlinked_from_student') == [
        {'user_id': user['id'], 'username': 'sportik', 'old_ref': second, 'new_ref': None}
    ]


def test_reconcile_record_unlink_and_relink(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    first = storage.create_student('Иванов Иван', 'М', '', '', '')
    second = storage.create_student('Новый ФИО', 'Ж', '', '', '')
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    storage.link_competitions([record_id], first)
    headers = get_auth_headers()

    _, response = reconcile_client.post(
        '/admin/people/reconcile/unlink',
        headers=headers,
        data={**csrf_for(headers), 'record_id': str(record_id)},
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert location.startswith(f'/admin/people/{first}?')
    assert 'Запись №1 отвязана от студента «Иванов Иван».' in location
    assert record_ref(storage, record_id) is None
    assert audit_details(storage, 'competition_unlinked_from_student') == [{'record_id': record_id, 'old_ref': first}]

    _, response = reconcile_client.post(
        '/admin/people/reconcile/relink',
        headers=headers,
        data={**csrf_for(headers), 'record_id': str(record_id), 'student_id': str(second)},
        allow_redirects=False,
    )
    location = unquote_plus(response.headers['location'])
    assert location.startswith(f'/admin/people/{second}?')
    assert 'Запись №1 перепривязана на студента «Новый ФИО».' in location
    assert record_ref(storage, record_id) == second
    assert audit_details(storage, 'competition_relinked') == [
        {'record_id': record_id, 'old_ref': None, 'new_ref': second, 'student_id': second}
    ]

    # Неизвестные цели
    _, response = reconcile_client.post(
        '/admin/people/reconcile/unlink',
        headers=headers,
        data={**csrf_for(headers), 'record_id': '999999'},
        allow_redirects=False,
    )
    assert 'Запись не найдена.' in unquote_plus(response.headers['location'])
    _, response = reconcile_client.post(
        '/admin/people/reconcile/relink',
        headers=headers,
        data={**csrf_for(headers), 'record_id': str(record_id), 'student_id': '999999'},
        allow_redirects=False,
    )
    assert 'Студент не найден.' in unquote_plus(response.headers['location'])


def test_person_card_shows_linked_data(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    other = storage.create_student('Другой Студент', 'Ж', '', '', '')

    _, response = reconcile_client.get(f'/admin/people/{student_id}', headers=get_auth_headers())
    assert response.status == 200
    assert 'Связанные данные' in response.text
    assert 'Привязанных записей пока нет.' in response.text
    assert 'Аккаунтов атлетов нет.' in response.text

    first = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    second = save_reconcile_record(storage, 'Старое ФИО', datetime(2026, 2, 10))
    storage.link_competitions([first, second], student_id)
    sportik = storage.get_user('sportik')
    storage.create_user('backuper', 'hash', 'athlete')
    backuper = storage.get_user('backuper')
    storage.link_user(sportik['id'], student_id)
    storage.link_user(backuper['id'], student_id)

    _, response = reconcile_client.get(f'/admin/people/{student_id}', headers=get_auth_headers())
    assert 'Записи о соревнованиях' in response.text
    # ФИО в записи — снимок записи, а не актуальное ФИО карточки
    assert 'Старое ФИО' in response.text
    assert 'Отвязать' in response.text
    assert 'Перепривязать' in response.text
    # Перепривязка аккаунта атлета (IA 2026-09): details-паттерн как у записей
    assert 'action="/admin/people/reconcile/users/relink"' in response.text
    assert 'Перепривязать аккаунт sportik на выбранного студента?' in response.text
    # Предупреждение о двух аккаунтах
    assert 'привязано 2 аккаунта атлета' in response.text
    assert 'sportik' in response.text
    assert 'backuper' in response.text
    # Перепривязка предлагает всех активных, кроме текущего студента
    assert f'/admin/people/{other}' not in response.text or 'Другой Студент' not in response.text

    # Больше 50 записей — показываются первые 50 с припиской
    for index in range(51):
        save_reconcile_record(storage, f'Бегун Бежит {index:02d}', datetime(2026, 3, 1))
    ids = [row['id'] for row in storage.connection.execute('SELECT id FROM competitions WHERE student_ref_id IS NULL')]
    storage.link_competitions(ids, student_id)
    _, response = reconcile_client.get(f'/admin/people/{student_id}', headers=get_auth_headers())
    assert '…и ещё 3 записей (показаны первые 50).' in response.text


def test_reconcile_link_extends_athlete_cabinet_by_ref(reconcile_client: SanicTestClient):
    """P3 dual-инвариант (замена test_reconcile_link_keeps_athlete_cabinet_
    unchanged): привязка по-прежнему меняет только student_ref_id, но
    кабинет атлета теперь читает связь (dual-read). Своя по ФИО запись
    видима до и после привязки; чужая запись ПОЯВЛЯЕТСЯ после привязки
    её карточки к аккаунту — плановое поведение dual-read."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    own_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    alien_id = save_reconcile_record(storage, 'Петров Пётр', datetime(2026, 2, 10))

    # Один и тот же cookie на оба запроса: иначе отличается csrf-мета страницы
    athlete = athlete_headers()
    _, before = reconcile_client.get('/', headers=athlete)
    assert before.status == 200
    assert 'Иванов Иван' in before.text
    assert 'Петров Пётр' not in before.text

    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    assert storage.link_competitions([own_id, alien_id], student_id) == (2, None)
    assert storage.link_user(sportik['id'], student_id) == (1, None)

    _, after = reconcile_client.get('/', headers=athlete)
    assert after.status == 200
    # Своя (по ФИО) запись не пропала...
    assert 'Иванов Иван' in after.text
    # ...а чужая появилась: кабинет читает стабильную связь с карточкой
    assert 'Петров Пётр' in after.text

    # Отчёт администратора от привязок не изменился (снимок по записям)
    _, report_before = reconcile_client.get('/report?slice=student', headers=get_auth_headers())
    storage.unlink_competition(own_id)
    storage.unlink_competition(alien_id)
    _, report_after = reconcile_client.get('/report?slice=student', headers=get_auth_headers())
    assert report_before.text == report_after.text


def test_reconcile_endpoints_reject_missing_csrf(reconcile_client: SanicTestClient):
    headers = get_auth_headers()
    for url, data in (
        ('/admin/people/reconcile/link', {'record_ids[]': ['1'], 'student_id': '1'}),
        ('/admin/people/reconcile/unlink', {'record_id': '1'}),
        ('/admin/people/reconcile/users/link', {'user_id': '1', 'student_id': '1'}),
        ('/admin/people/reconcile/users/unlink', {'user_id': '1'}),
    ):
        _, response = reconcile_client.post(url, headers=headers, data=data, allow_redirects=False)
        assert response.status == 403, url
        assert 'CSRF' in response.text


def test_reconcile_non_numeric_ids_return_400(reconcile_client: SanicTestClient):
    _, response = reconcile_client.get('/admin/people/reconcile/create?record_id=abc', headers=get_auth_headers())
    assert response.status == 400

    _, response = reconcile_client.get('/admin/people/reconcile/users/create?user_id=abc', headers=get_auth_headers())
    assert response.status == 400

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={**csrf_for(headers), 'record_ids[]': ['abc'], 'student_id': '1'},
    )
    assert response.status == 400

    _, response = reconcile_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={**csrf_for(headers), 'record_ids[]': ['1'], 'student_id': 'abc'},
    )
    assert response.status == 400

    _, response = reconcile_client.post(
        '/admin/people/reconcile/unlink', headers=headers, data={**csrf_for(headers), 'record_id': 'abc'}
    )
    assert response.status == 400

    _, response = reconcile_client.post(
        '/admin/people/reconcile/users/link',
        headers=headers,
        data={**csrf_for(headers), 'user_id': 'abc', 'student_id': '1'},
    )
    assert response.status == 400

    _, response = reconcile_client.post(
        '/admin/people/reconcile/relink',
        headers=headers,
        data={**csrf_for(headers), 'record_id': '1', 'student_id': 'abc'},
    )
    assert response.status == 400


def test_reconcile_links_in_admin_hub_and_people(reconcile_client: SanicTestClient):
    storage = app.ctx.storage
    # IA редизайна 2026-09: сопоставление живёт в разделе «Студенты», хаб
    # «Админ» на него больше не ссылается (пункт «Студенты» в меню — норма).
    _, response = reconcile_client.get('/admin', headers=get_auth_headers())
    assert response.status == 200
    assert 'href="/admin/people/reconcile"' not in response.text

    # Бейдж непривязанных записей — только когда они есть
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    _, response = reconcile_client.get('/admin/people', headers=get_auth_headers())
    assert 'Сопоставление данных' in response.text
    assert 'href="/admin/people/reconcile"' in response.text
    assert '<span class="badge text-bg-light border">1</span>' in response.text

    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    storage.link_competitions([record_id], student_id)
    _, response = reconcile_client.get('/admin/people', headers=get_auth_headers())
    assert 'badge text-bg-light border' not in response.text


# --- P3 Runtime Identity: dual-read кабинета атлета и режим dual/ref. ---
#
# HTTP-сценарии на реальном SQLite-адаптере (паттерн reconcile_client):
# ветки видимости главной (owner / student_ref_id / легаси-хеш), права на
# запись по связи, сужение кабинета в ref, отчёт проверки и guard-переключение
# режима. Синтетические данные.


def test_dual_mode_visibility_branches_and_owns_by_ref(reconcile_client: SanicTestClient, tmp_path, monkeypatch):
    """Dual: главная атлета = owner OR ref OR легаси-хеш; связь даёт и право
    на запись (правка, вложения)."""
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    storage.link_user(sportik['id'], student_id)

    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))  # хеш
    ref_id = save_reconcile_record(storage, 'Петров Пётр', datetime(2026, 2, 10))  # связь
    storage.link_competitions([ref_id], student_id)
    save_reconcile_record(storage, 'Чужое Имя', datetime(2026, 3, 10), owner_id=sportik['id'])  # owner
    save_reconcile_record(storage, 'Невидимый Ник', datetime(2026, 4, 10))  # мимо всех веток

    headers = athlete_headers()
    _, response = reconcile_client.get('/', headers=headers)
    assert response.status == 200
    assert 'Иванов Иван' in response.text
    assert 'Петров Пётр' in response.text
    assert 'Чужое Имя' in response.text
    assert 'Невидимый Ник' not in response.text

    # Правка записи, видимой только по связи, разрешена атлету
    _, response = reconcile_client.post(
        f'/competition/{ref_id}',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA, 'student_name': 'Петров Пётр'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'] == '/'

    # Вложение на запись по связи тоже разрешено
    _, response = reconcile_client.post(
        f'/competition/{ref_id}/attachments',
        headers=headers,
        data=csrf_for(headers),
        files={'file': ('diploma.png', PNG_BYTES, 'application/octet-stream')},
    )
    assert response.status == 200
    assert response.text == 'Файл загружен'
    # Связь пережила правку записи
    assert record_ref(storage, ref_id) == student_id


def test_ref_mode_narrows_athlete_cabinet(reconcile_client: SanicTestClient):
    """Ref: кабинет = owner OR ref; хеши и псевдонимы ФИО не расширяют;
    модератор видит весь реестр независимо от режима."""
    storage = app.ctx.storage
    storage.set_identity_mode('ref')
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    storage.add_name_alias(sportik['id'], 'Иванов И.И.')
    student_id = storage.create_student('Иванов Иван', 'М', '', '', '')
    storage.link_user(sportik['id'], student_id)

    ref_id = save_reconcile_record(storage, 'Петров Пётр', datetime(2026, 2, 10))
    storage.link_competitions([ref_id], student_id)
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))  # только хеш
    save_reconcile_record(storage, 'Иванов И.И.', datetime(2026, 3, 10))  # псевдоним
    save_reconcile_record(storage, 'Чужое Имя', datetime(2026, 4, 10), owner_id=sportik['id'])

    headers = athlete_headers()
    _, response = reconcile_client.get('/', headers=headers)
    assert response.status == 200
    assert 'Петров Пётр' in response.text
    assert 'Чужое Имя' in response.text
    assert 'Иванов Иван' not in response.text
    assert 'Иванов И.И.' not in response.text

    _, response = reconcile_client.get('/', headers=get_auth_headers())
    for name in ('Петров Пётр', 'Чужое Имя', 'Иванов Иван', 'Иванов И.И.'):
        assert name in response.text


def test_namesakes_isolated_in_ref_mode(reconcile_client: SanicTestClient):
    """Два аккаунта с одинаковым ФИО на разных карточках. Dual (легаси):
    хеш один — оба видят обе записи (фиксируем поведение). Ref: каждый
    видит только запись своей карточки."""
    storage = app.ctx.storage
    storage.create_user('twin1', 'hash', 'athlete')
    storage.create_user('twin2', 'hash', 'athlete')
    twin1 = storage.get_user('twin1')
    twin2 = storage.get_user('twin2')
    for user in (twin1, twin2):
        storage.set_profile(user['id'], {'student_name': 'Одинаков Фамилия', 'student_sex': 'М'})
    card1 = storage.create_student('Одинаков Фамилия', 'М', '', '', '')
    card2 = storage.create_student('Одинаков Фамилия', 'М', '', '', '')
    storage.link_user(twin1['id'], card1)
    storage.link_user(twin2['id'], card2)

    storage.link_competitions([save_reconcile_record(storage, 'Одинаков Фамилия', datetime(2026, 1, 10))], card1)
    storage.link_competitions([save_reconcile_record(storage, 'Одинаков Фамилия', datetime(2026, 2, 10))], card2)

    # Dual: обе записи в кабинете (даты различают записи с одинаковым ФИО)
    _, response = reconcile_client.get('/', headers=athlete_headers_for('twin1'))
    assert '10.01.2026' in response.text
    assert '10.02.2026' in response.text

    storage.set_identity_mode('ref')
    _, response = reconcile_client.get('/', headers=athlete_headers_for('twin1'))
    assert '10.01.2026' in response.text
    assert '10.02.2026' not in response.text


def test_ref_mode_namesake_hash_cannot_edit_or_touch_attachments(
    reconcile_client: SanicTestClient, tmp_path, monkeypatch
):
    """QA D2: в ref право на запись = owner OR ref ТОЛЬКО. Атлет-тёзка
    (то же ФИО в профиле, другая карточка) при хеш-совпадении НЕ может
    править чужую запись, загружать вложения и скачивать чужие вложения.
    В dual легаси-хеш-ветка прав сохраняется (инвариант зафиксирован)."""
    from src import main as main_module

    monkeypatch.setattr(main_module.settings, 'data_folder', str(tmp_path))
    storage = app.ctx.storage
    storage.create_user('twin1', 'hash', 'athlete')
    storage.create_user('twin2', 'hash', 'athlete')
    twin1 = storage.get_user('twin1')
    twin2 = storage.get_user('twin2')
    for user in (twin1, twin2):
        storage.set_profile(user['id'], {'student_name': 'Одинаков Фамилия', 'student_sex': 'М'})
    card1 = storage.create_student('Одинаков Фамилия', 'М', '', '', '')
    card2 = storage.create_student('Одинаков Фамилия', 'М', '', '', '')
    storage.link_user(twin1['id'], card1)
    storage.link_user(twin2['id'], card2)

    # Запись twin1: связана с его карточкой, owner нет — для twin2 она
    # «чужая по связи», но совпадает по легаси-хешу ФИО.
    record_id = save_reconcile_record(storage, 'Одинаков Фамилия', datetime(2026, 1, 10))
    storage.link_competitions([record_id], card1)

    # Dual (легаси-инвариант): хеш-ветка прав разрешает правку и вложения
    assert storage.get_identity_mode() == 'dual'
    twin2_headers = athlete_headers_for('twin2')
    _, response = reconcile_client.post(
        f'/competition/{record_id}',
        headers=twin2_headers,
        data={**csrf_for(twin2_headers), **ATHLETE_RECORD_DATA, 'student_name': 'Одинаков Фамилия'},
        allow_redirects=False,
    )
    assert response.status == 302
    _, response = reconcile_client.post(
        f'/competition/{record_id}/attachments',
        headers=twin2_headers,
        data=csrf_for(twin2_headers),
        files={'file': ('note.png', PNG_BYTES, 'application/octet-stream')},
    )
    assert response.status == 200
    assert response.text == 'Файл загружен'
    attachment_id = storage.get_attachments(record_id)[0]['id']
    _, response = reconcile_client.get(f'/attachment/{attachment_id}', headers=twin2_headers)
    assert response.status == 200

    # Ref: те же действия тёзкой — запрещены (owner OR ref только)
    storage.set_identity_mode('ref')
    _, response = reconcile_client.post(
        f'/competition/{record_id}',
        headers=twin2_headers,
        data={**csrf_for(twin2_headers), **ATHLETE_RECORD_DATA, 'student_name': 'Одинаков Фамилия'},
        allow_redirects=False,
    )
    assert response.status == 403
    _, response = reconcile_client.post(
        f'/competition/{record_id}/attachments',
        headers=twin2_headers,
        data=csrf_for(twin2_headers),
        files={'file': ('note2.png', PNG_BYTES, 'application/octet-stream')},
    )
    assert response.status == 403
    _, response = reconcile_client.get(f'/attachment/{attachment_id}', headers=twin2_headers)
    assert response.status == 403

    # Владелец по связи (twin1) работает в обоих режимах: правка и вложения
    twin1_headers = athlete_headers_for('twin1')
    _, response = reconcile_client.post(
        f'/competition/{record_id}',
        headers=twin1_headers,
        data={**csrf_for(twin1_headers), **ATHLETE_RECORD_DATA, 'student_name': 'Одинаков Фамилия'},
        allow_redirects=False,
    )
    assert response.status == 302
    _, response = reconcile_client.get(f'/attachment/{attachment_id}', headers=twin1_headers)
    assert response.status == 200


def test_account_without_student_card(reconcile_client: SanicTestClient):
    """Аккаунт без карточки: dual — owner и хеши; ref — только owner."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Сидоров Сидор', 'student_sex': 'М'})

    save_reconcile_record(storage, 'Сидоров Сидор', datetime(2026, 1, 10))  # хеш
    save_reconcile_record(storage, 'Чужое Имя', datetime(2026, 2, 10), owner_id=sportik['id'])

    headers = athlete_headers()
    _, response = reconcile_client.get('/', headers=headers)
    assert 'Сидоров Сидор' in response.text
    assert 'Чужое Имя' in response.text

    storage.set_identity_mode('ref')
    _, response = reconcile_client.get('/', headers=headers)
    assert 'Сидоров Сидор' not in response.text
    assert 'Чужое Имя' in response.text


def test_unlinked_participation_hash_only_disappears_in_ref(reconcile_client: SanicTestClient):
    """Запись без ref, созданная модератором: владелец видит её в обоих
    режимах; атлет, совпадающий только по ФИО, — только в dual."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    admin = storage.get_user(settings.auth_admin_username)
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10), owner_id=admin['id'])

    athlete = athlete_headers()
    _, response = reconcile_client.get('/', headers=athlete)
    assert 'Иванов Иван' in response.text

    storage.set_identity_mode('ref')
    _, response = reconcile_client.get('/', headers=athlete)
    assert 'Иванов Иван' not in response.text
    # Владелец-модератор видит свою запись в обоих режимах
    _, response = reconcile_client.get('/', headers=get_auth_headers())
    assert 'Иванов Иван' in response.text
    assert record_ref(storage, record_id) is None


def test_person_card_relink_preview_counts(reconcile_client: SanicTestClient):
    """Предпросмотр «Видимо атлету» следует за перепривязкой аккаунта:
    счётчики на карточке-цели учитывают её записи."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    first = storage.create_student('Иванов Иван', 'М', '', '', '')
    second = storage.create_student('Петров Пётр', 'Ж', '', '', '')
    storage.link_user(sportik['id'], first)

    own = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10), owner_id=sportik['id'])
    storage.link_competitions([own], first)
    alien = save_reconcile_record(storage, 'Петров Пётр', datetime(2026, 2, 10))
    storage.link_competitions([alien], second)

    _, response = reconcile_client.get(f'/admin/people/{first}', headers=get_auth_headers())
    assert 'сейчас 1 · после отвязки 1' in response.text

    # Перепривязка аккаунта на вторую карточку: кабинет читает уже её
    assert storage.relink_user(sportik['id'], second) == (first, None)
    _, response = reconcile_client.get(f'/admin/people/{second}', headers=get_auth_headers())
    assert 'сейчас 2 · после отвязки 1' in response.text


def test_person_card_unlink_preview_counts_and_confirm_variants(reconcile_client: SanicTestClient):
    """Три варианта подтверждения отвязки: ничего не меняется / частично /
    пропадает всё; счётчики соответствуют scope атлета."""
    storage = app.ctx.storage
    student = storage.create_student('Иванов Иван', 'М', '', '', '')

    # Все записи карточки названы «Иванов Иван»: stable покрывает их хешем
    # (отвязка ничего не меняет), partial/doomed теряют их целиком.
    storage.create_user('stable', 'hash', 'athlete')
    stable = storage.get_user('stable')
    storage.set_profile(stable['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    storage.link_user(stable['id'], student)
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10), owner_id=stable['id'])
    storage.link_competitions(
        [
            save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 2, 10)),
            save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 4, 10)),
            save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 5, 10)),
        ],
        student,
    )

    storage.create_user('partial', 'hash', 'athlete')
    partial = storage.get_user('partial')
    storage.set_profile(partial['id'], {'student_name': 'Чужое Имя', 'student_sex': 'М'})
    storage.link_user(partial['id'], student)
    save_reconcile_record(storage, 'Чужое Имя', datetime(2026, 3, 10), owner_id=partial['id'])

    storage.create_user('doomed', 'hash', 'athlete')
    storage.link_user(storage.get_user('doomed')['id'], student)

    _, response = reconcile_client.get(f'/admin/people/{student}', headers=get_auth_headers())
    assert response.status == 200
    assert 'сейчас 4 · после отвязки 4' in response.text
    assert 'сейчас 4 · после отвязки 1' in response.text
    assert 'сейчас 3 · после отвязки 0' in response.text
    # Подтверждения различают три случая
    assert 'Видимость кабинета атлета не изменится (4 записей).' in response.text
    assert 'Атлет будет видеть 1 из 4 записей.' in response.text
    assert 'Атлет перестанет видеть записи этого студента (сейчас 3, останется 0).' in response.text


def test_identity_report_counts_and_tab(reconcile_client: SanicTestClient):
    """Отчёт проверки: сводка аккаунтов, строка H, риск тёзки, пропадающие
    записи и вкладка с бейджем hash_only_total."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    storage.create_user('free', 'hash', 'athlete')
    free = storage.get_user('free')
    storage.set_profile(free['id'], {'student_name': 'Безкарточкин Бес', 'student_sex': 'М'})

    student = storage.create_student('Иванов Иван', 'М', '', '', '')
    namesake_card = storage.create_student('Иванов Иван', 'Ж', '', '', '')
    storage.link_user(sportik['id'], student)

    ref_only = save_reconcile_record(storage, 'По связи', datetime(2026, 1, 10))
    storage.link_competitions([ref_only], student)
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 2, 10))
    namesake = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 3, 10))
    storage.link_competitions([namesake], namesake_card)
    save_reconcile_record(storage, 'Безкарточкин Бес', datetime(2026, 4, 10))

    # H: hash_only (спорт) + namesake (спорт, чужая карточка) + free_hash (free)
    assert storage.count_hash_only_visible() == 3

    _, response = reconcile_client.get('/admin/people/reconcile/identity', headers=get_auth_headers())
    assert response.status == 200
    assert 'Режим идентификации' in response.text
    assert 'двойной (dual)' in response.text
    assert 'Аккаунтов атлетов: 2, из них привязано к карточкам: 1, без карточки: 1.' in response.text
    assert 'пропадут из кабинетов при переходе на «только связи»:' in response.text
    assert '<strong>3</strong>' in response.text
    assert 'риск тёзки: 1' in response.text
    assert 'Записи, которые пропадут из кабинетов' in response.text
    # Бейдж вкладки — на всех трёх страницах раздела
    for url in ('/admin/people/reconcile', '/admin/people/reconcile/users', '/admin/people/reconcile/identity'):
        _, page = reconcile_client.get(url, headers=get_auth_headers())
        assert page.status == 200
        assert 'Режим идентификации' in page.text


def test_identity_flip_blocked_with_hash_only(reconcile_client: SanicTestClient):
    """Guard: при H≥1 формы нет, POST-подделка отклоняется — flash с числом,
    режим прежний, аудита нет."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    assert storage.count_hash_only_visible() == 1

    _, response = reconcile_client.get('/admin/people/reconcile/identity', headers=get_auth_headers())
    assert 'заблокирован' in response.text
    assert 'action="/admin/people/reconcile/identity/flip"' not in response.text

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/identity/flip',
        headers=headers,
        data={**csrf_for(headers), 'mode': 'ref'},
        allow_redirects=False,
    )
    assert response.status == 302
    location = unquote_plus(response.headers['location'])
    assert location.startswith('/admin/people/reconcile/identity?')
    assert '1 записей видны атлетам только по ФИО' in location
    assert storage.get_identity_mode() == 'dual'
    assert audit_details(storage, 'identity_mode_changed') == []


def test_identity_flip_dual_to_ref_with_zero_hash_only(reconcile_client: SanicTestClient):
    """Dual→ref при H=0: режим меняется, аудит пишется, кабинет атлета
    меняется ровно на H=0 записей — то есть не меняется."""
    storage = app.ctx.storage
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    student = storage.create_student('Иванов Иван', 'М', '', '', '')
    storage.link_user(sportik['id'], student)
    # Всё, что видит атлет, связано с его карточкой или принадлежит ему
    linked = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10), owner_id=sportik['id'])
    storage.link_competitions([linked], student)
    assert storage.count_hash_only_visible() == 0

    athlete = athlete_headers()
    _, before = reconcile_client.get('/', headers=athlete)
    assert before.status == 200

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/identity/flip',
        headers=headers,
        data={**csrf_for(headers), 'mode': 'ref'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert storage.get_identity_mode() == 'ref'
    assert audit_details(storage, 'identity_mode_changed') == [{'old': 'dual', 'new': 'ref', 'hash_only_visible': 0}]

    _, after = reconcile_client.get('/', headers=athlete)
    assert after.status == 200
    assert after.text == before.text


def test_identity_flip_ref_back_to_dual_always_allowed(reconcile_client: SanicTestClient):
    """Ref→dual — всегда разрешён (даже при H≥1) и пишется в аудит."""
    storage = app.ctx.storage
    storage.set_identity_mode('ref')
    sportik = storage.get_user('sportik')
    storage.set_profile(sportik['id'], {'student_name': 'Иванов Иван', 'student_sex': 'М'})
    # Появилась hash-only запись — возврат всё равно разрешён
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    assert storage.count_hash_only_visible() == 1

    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/identity/flip',
        headers=headers,
        data={**csrf_for(headers), 'mode': 'dual'},
        allow_redirects=False,
    )
    assert response.status == 302
    assert storage.get_identity_mode() == 'dual'
    assert audit_details(storage, 'identity_mode_changed') == [{'old': 'ref', 'new': 'dual', 'hash_only_visible': 1}]
    # Атлет снова видит запись по ФИО
    _, response = reconcile_client.get('/', headers=athlete_headers())
    assert 'Иванов Иван' in response.text


def test_athlete_self_submit_visible_by_owner_in_both_modes(reconcile_client: SanicTestClient):
    """Своя заявка атлета видна по owner в dual и в ref (модерация —
    как раньше: новая запись уходит в pending)."""
    storage = app.ctx.storage
    headers = athlete_headers()
    _, response = reconcile_client.post(
        '/competition',
        headers=headers,
        data={**csrf_for(headers), **ATHLETE_RECORD_DATA},
        allow_redirects=False,
    )
    assert response.status == 302
    row = storage.connection.execute(
        'SELECT id, review_status, owner_id FROM competitions WHERE student_name = ?',
        (ATHLETE_RECORD_DATA['student_name'],),
    ).fetchone()
    assert row['review_status'] == 'pending'
    assert row['owner_id'] == storage.get_user('sportik')['id']

    _, response = reconcile_client.get('/', headers=headers)
    assert ATHLETE_RECORD_DATA['student_name'] in response.text

    storage.set_identity_mode('ref')
    _, response = reconcile_client.get('/', headers=headers)
    assert ATHLETE_RECORD_DATA['student_name'] in response.text


def test_admin_editor_semantics_unchanged_in_ref_mode(reconcile_client: SanicTestClient):
    """Режим касается только кабинета атлета: admin/editor видят весь
    реестр, editor правит любую запись, модерация подтверждает."""
    storage = app.ctx.storage
    storage.set_identity_mode('ref')
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))

    _, response = reconcile_client.get('/', headers=get_auth_headers(role='editor'))
    assert 'Иванов Иван' in response.text

    editor = get_auth_headers(role='editor')
    _, response = reconcile_client.post(
        f'/competition/{record_id}',
        headers=editor,
        data={**csrf_for(editor), **ATHLETE_RECORD_DATA, 'student_name': 'Иванов Иван'},
        allow_redirects=False,
    )
    assert response.status == 302
    # Правка модератора подтверждает запись
    assert storage.get_competition_review(record_id)['review_status'] == 'approved'


def test_identity_endpoints_require_admin_and_csrf(reconcile_client: SanicTestClient):
    _, response = reconcile_client.get('/admin/people/reconcile/identity', headers=get_auth_headers(role='editor'))
    assert response.status == 403

    _, response = reconcile_client.get('/admin/people/reconcile/identity', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')

    _, response = reconcile_client.post('/admin/people/reconcile/identity/flip', data={'mode': 'ref'})
    assert response.status == 401

    # CSRF проверяет общий middleware POST-запросов
    _, response = reconcile_client.post(
        '/admin/people/reconcile/identity/flip', headers=get_auth_headers(), data={'mode': 'ref'}
    )
    assert response.status == 403
    assert 'CSRF' in response.text

    # Неизвестный режим — 400, ничего не меняется
    headers = get_auth_headers()
    _, response = reconcile_client.post(
        '/admin/people/reconcile/identity/flip',
        headers=headers,
        data={**csrf_for(headers), 'mode': 'single'},
    )
    assert response.status == 400
    assert app.ctx.storage.get_identity_mode() == 'dual'


# --- Импорт студентов из Excel (Student Identity v1, Phase 2.5). ---
# Полные сценарии на реальном SQLite-адаптере (паттерн reconcile_client):
# предпросмотр ничего не пишет в БД, создание — только явным подтверждением,
# импорт трогает только students + audit_log, справочники только читаются.


@pytest.fixture
def student_import_client(client: SanicTestClient, tmp_path):
    storage = SQLiteAdapter(str(tmp_path / 'student_import.sqlite3'))
    storage.create_user(settings.auth_admin_username, 'hash', 'admin')
    storage.create_user('sportik', 'hash', 'athlete')
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    student_import_sessions.clear()
    try:
        yield client
    finally:
        app.ctx.storage = previous
        student_import_sessions.clear()


def upload_student_xlsx(
    client: SanicTestClient,
    rows: list[dict],
    *,
    role: str = 'admin',
    filename: str = 'students.xlsx',
    columns: list[str] | None = None,
):
    """Загрузить файл импорта студентов; ответ — редирект (без перехода)."""
    df = pd.DataFrame(rows, columns=columns) if columns else pd.DataFrame(rows)
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)
    headers = get_auth_headers(role=role)
    _, response = client.post(
        '/admin/people/import',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                filename,
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    return response


def import_session_token(response) -> str:
    """Токен staging-сессии из редиректа загрузки."""
    return response.headers['location'].rsplit('/', 1)[-1]


def post_import_action(client: SanicTestClient, token: str, suffix: str, data: dict | None = None):
    headers = get_auth_headers()
    _, response = client.post(
        f'/admin/people/import/preview/{token}/{suffix}',
        headers=headers,
        data={**csrf_for(headers), **(data or {})},
        allow_redirects=False,
    )
    return response


def get_preview(client: SanicTestClient, token: str):
    _, response = client.get(f'/admin/people/import/preview/{token}', headers=get_auth_headers(), allow_redirects=False)
    return response


def students_snapshot(storage) -> list[tuple]:
    rows = storage.connection.execute(
        'SELECT id, full_name, sex, institute, group_name, course FROM students ORDER BY id'
    ).fetchall()
    return [tuple(row) for row in rows]


def catalog_snapshot(storage) -> list[tuple]:
    rows = storage.connection.execute(
        'SELECT id, category, value, parent_id, active FROM catalog_values ORDER BY id'
    ).fetchall()
    return [tuple(row) for row in rows]


def seed_catalog(storage, institute: str, group: str) -> None:
    storage.add_catalog_value('institute', institute)
    storage.ensure_catalog_pair(institute, group)


def test_student_import_pages_require_admin(student_import_client: SanicTestClient):
    _, response = student_import_client.get('/admin/people/import', headers=get_auth_headers(role='editor'))
    assert response.status == 403
    _, response = student_import_client.get('/admin/people/import/template', headers=get_auth_headers(role='editor'))
    assert response.status == 403
    _, response = student_import_client.get('/admin/people/import', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')
    _, response = student_import_client.post('/admin/people/import', data={'full_name': 'X'})
    assert response.status == 401


def test_student_import_template_headers_only(student_import_client: SanicTestClient):
    _, response = student_import_client.get('/admin/people/import/template', headers=get_auth_headers())
    assert response.status == 200
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    assert response.headers['content-disposition'].startswith('attachment; filename="Шаблон_студенты_')
    assert get_xlsx_headers(response.body) == ['ФИО', 'Пол', 'Институт', 'Группа', 'Курс']
    # Только заголовки: ни одной строки данных
    assert len(get_xlsx_rows(response.body)) == 1


def test_student_import_instruction_page_renders(student_import_client: SanicTestClient):
    _, response = student_import_client.get('/admin/people/import', headers=get_auth_headers())
    assert response.status == 200
    assert 'Как заполнить файл' in response.text
    assert 'Скачать шаблон (xlsx)' in response.text
    assert 'В шаблоне — только заголовки колонок.' in response.text
    assert 'Обязательное поле — только ФИО' in response.text
    assert 'Если институт не указан, он определится автоматически по группе из справочника.' in response.text
    assert 'Лишние колонки в файле игнорируются.' in response.text
    assert 'Ничего не сохраняется, пока вы не подтвердите строки' in response.text
    assert 'action="/admin/people/import"' in response.text
    # Обычная форма, а не AJAX-форма импорта записей
    assert 'import-form' not in response.text
    # Кнопка входа в раздел — на странице списка студентов
    _, response = student_import_client.get('/admin/people', headers=get_auth_headers())
    assert 'href="/admin/people/import">Импорт из Excel' in response.text


def test_student_import_upload_rejects_bad_files(student_import_client: SanicTestClient, monkeypatch):
    headers = get_auth_headers()

    _, response = student_import_client.post(
        '/admin/people/import', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'Выберите файл.' in unquote_plus(response.headers['location'])

    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Тестов'}], filename='students.csv')
    assert response.status == 302
    assert 'Файл должен быть в формате .xlsx.' in unquote_plus(response.headers['location'])

    file_obj = BytesIO(b'this is not an excel file at all')
    _, response = student_import_client.post(
        '/admin/people/import',
        headers=headers,
        data=csrf_for(headers),
        files={'file': ('broken.xlsx', file_obj.getvalue(), 'application/octet-stream')},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Не удалось прочитать файл. Проверьте, что это не повреждённый .xlsx.' in unquote_plus(
        response.headers['location']
    )

    response = upload_student_xlsx(student_import_client, [{'Имя': 'Тестов', 'Пол': 'М'}])
    assert response.status == 302
    assert 'В файле нет обязательной колонки «ФИО». Скачайте шаблон и заполните его.' in unquote_plus(
        response.headers['location']
    )

    response = upload_student_xlsx(student_import_client, [], columns=['ФИО', 'Пол', 'Институт', 'Группа', 'Курс'])
    assert response.status == 302
    assert 'В файле нет строк с данными.' in unquote_plus(response.headers['location'])

    monkeypatch.setattr('src.main.STUDENT_IMPORT_MAX_ROWS', 3)
    rows = [{'ФИО': f'Студентов Студент {index:02d}'} for index in range(4)]
    response = upload_student_xlsx(student_import_client, rows)
    assert response.status == 302
    assert 'В файле больше 3 строк. Разбейте список на части.' in unquote_plus(response.headers['location'])

    assert student_import_sessions == {}


def test_student_import_upload_success_unknown_columns_and_single_session(
    student_import_client: SanicTestClient,
):
    storage = app.ctx.storage
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван', 'Комментарий': 'прим.'}])
    assert response.status == 302
    assert response.headers['location'].startswith('/admin/people/import/preview/')
    token = import_session_token(response)

    _, page = student_import_client.get(f'/admin/people/import/preview/{token}', headers=get_auth_headers())
    assert page.status == 200
    assert 'Неизвестные колонки файла игнорируются: Комментарий.' in page.text
    assert 'Иванов Иван' in page.text
    # Загрузка ничего не создала и в БД не писала
    assert students_snapshot(storage) == []
    assert audit_details(storage, 'student_created') == []

    # Одна сессия на админа: новая загрузка заменяет прежнюю
    second = upload_student_xlsx(student_import_client, [{'ФИО': 'Петров Пётр'}])
    second_token = import_session_token(second)
    assert second_token != token
    response = get_preview(student_import_client, token)
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(response.headers['location'])
    page = get_preview(student_import_client, second_token)
    assert page.status == 200
    assert 'Петров Пётр' in page.text


def test_student_import_preview_row_normalization(student_import_client: SanicTestClient):
    response = upload_student_xlsx(
        student_import_client,
        [
            {
                'ФИО': '  Иванов Иван  ',
                'Пол': None,
                'Институт': None,
                'Группа': None,
                'Курс': 2.0,
            },
            {'ФИО': 'Петров Пётр', 'Пол': 'Муж', 'Институт': 'ИСИ', 'Группа': 'ПГС-101', 'Курс': '3'},
            {'ФИО': '   ', 'Пол': '', 'Институт': '', 'Группа': '', 'Курс': ''},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    # Счётчики категорий: одна новая, две ошибочные
    assert 'Всего строк: 3 · новые: 1 · требуют решения: 0 · ошибочные: 2' in page.text
    assert '>Новые' in page.text and '>Ошибочные' in page.text
    assert 'Пол должен быть «М» или «Ж».' in page.text
    assert 'Пустое ФИО.' in page.text
    # strip значений, NaN → пустые, курс 2.0 → «2»
    assert '<strong>Иванов Иван</strong>' in page.text
    assert '>2</td>' in page.text
    # Кнопка завершения задизейблена, пока есть неразобранные строки
    assert 'Кнопка «Завершить импорт» станет активной' in page.text
    assert 'disabled>Завершить импорт</button>' in page.text
    # Правка — bootstrap-collapse по data-атрибутам + мини-формы действий
    assert f'action="/admin/people/import/preview/{token}/row/2/create"' in page.text
    assert f'action="/admin/people/import/preview/{token}/row/2/skip"' in page.text
    assert 'data-bs-target="#import-row-2-edit"' in page.text
    assert 'data-bs-toggle="collapse"' in page.text
    assert 'Сохранить правку' in page.text
    assert 'placeholder="например, 2"' in page.text


def test_student_import_infile_duplicates_flagged_and_both_creatable(
    student_import_client: SanicTestClient,
):
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Пол': 'М', 'Институт': 'ИСИ', 'Группа': 'ПГС-101', 'Курс': '2'},
            {'ФИО': 'иванов иван', 'Пол': 'Ж', 'Институт': 'ИМИ', 'Группа': 'СБ-202', 'Курс': '1'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    # Обе строки — «Требуют решения» с перечнем строк группы дублей
    assert 'Всего строк: 2 · новые: 0 · требуют решения: 2' in page.text
    assert 'Одинаковые ФИО в файле: строки 2, 3' in page.text

    # Полные тёзки — разные люди: обе строки создаются явно и раздельно
    first = post_import_action(student_import_client, token, 'row/2/create')
    assert first.status == 302
    assert 'Строка 2: студент «Иванов Иван» создан.' in unquote_plus(first.headers['location'])
    second = post_import_action(student_import_client, token, 'row/3/create')
    assert second.status == 302
    snapshot = students_snapshot(storage)
    assert [(row[1], row[2], row[4]) for row in snapshot] == [
        ('Иванов Иван', 'М', 'ПГС-101'),
        ('иванов иван', 'Ж', 'СБ-202'),
    ]


def test_student_import_catalog_autofill_by_group(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    seed_catalog(storage, 'ИСИ', 'ПГС-101')

    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван', 'Группа': 'пгс-101'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    assert 'ИСИ' in page.text
    assert 'определён по группе' in page.text
    assert 'title="Институт определён по группе из справочника"' in page.text
    assert 'ПГС-101' in page.text
    assert 'Институт не определён' not in page.text

    # Создаётся ровно то, что показано: институт автозаполнен, группа канонична
    created = post_import_action(student_import_client, token, 'row/2/create')
    assert created.status == 302
    snapshot = students_snapshot(storage)
    assert snapshot[0][3] == 'ИСИ'
    assert snapshot[0][4] == 'ПГС-101'


def test_student_import_catalog_warnings(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    seed_catalog(storage, 'ИСИ', 'ПГС-101')
    seed_catalog(storage, 'ИМИ', 'ПГС-101')  # то же имя группы в двух институтах
    seed_catalog(storage, 'ИМИ', 'ТД-303')

    # Неизвестная группа: предупреждение, значения не меняются
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван', 'Группа': 'ХЗ-999'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Институт не определён: группа отсутствует в справочнике или относится к нескольким институтам' in page.text
    assert 'определён по группе' not in page.text
    assert 'ХЗ-999' in page.text

    # Неоднозначная группа (два института): то же предупреждение
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Петров Пётр', 'Группа': 'ПГС-101'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Институт не определён: группа отсутствует в справочнике или относится к нескольким институтам' in page.text

    # Группа из другого института: предупреждение, значения НЕ переписываются
    response = upload_student_xlsx(
        student_import_client,
        [{'ФИО': 'Сидоров Сидор', 'Институт': 'ИСИ', 'Группа': 'ТД-303'}],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Требует внимания: группа относится к другому институту' in page.text
    assert 'ИСИ' in page.text and 'ТД-303' in page.text
    assert 'ИМИ' not in page.text  # институт строки не заменён владельцем группы

    # Только институт: канонизация по справочнику (без бейджа автозаполнения)
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Козлов Кирилл', 'Институт': 'иси'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'ИСИ' in page.text
    assert 'определён по группе' not in page.text


def test_student_import_sex_normalized_to_canonical(student_import_client: SanicTestClient):
    """Пол из файла приводится к «М»/«Ж» независимо от регистра (Задача 1):
    «м»/«ж» — валидны и канонизируются, «муж» и латинское «m» — ошибка строки."""
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Пол': 'М'},
            {'ФИО': 'Петров Пётр', 'Пол': 'м'},
            {'ФИО': 'Сидорова Сидора', 'Пол': 'Ж'},
            {'ФИО': 'Козлова Кира', 'Пол': 'ж'},
            {'ФИО': 'Волков Вольдемар', 'Пол': ' м '},
            {'ФИО': 'Павлов Павел', 'Пол': 'муж'},
            {'ФИО': 'Латинов Латин', 'Пол': 'm'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    assert 'Всего строк: 7 · новые: 5 · требуют решения: 0 · ошибочные: 2' in page.text
    assert page.text.count('Пол должен быть «М» или «Ж».') == 2

    # Предпросмотр показывает каноническое значение, создание сохраняет его
    created = post_import_action(student_import_client, token, 'row/2/create')
    assert created.status == 302
    created = post_import_action(student_import_client, token, 'row/3/create')
    assert created.status == 302
    snapshot = {row[1]: row[2] for row in students_snapshot(storage)}
    assert snapshot['Иванов Иван'] == 'М'
    assert snapshot['Петров Пётр'] == 'М'


def test_student_import_infile_group_institute_autofill(student_import_client: SanicTestClient):
    """Группа→институт из ТЕКУЩЕГО файла (Задача 2): строки без института
    получают институт строк-«доноров» того же файла + бейдж «определён по
    файлу»; создаётся ровно то, что показано (в т.ч. через bulk-create)."""
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Доноров Донор', 'Институт': 'ИСИ', 'Группа': 'ПГС-101'},
            {'ФИО': 'Пустов Пуст', 'Группа': 'ПГС-101'},
            {'ФИО': 'Второв Втор', 'Группа': 'пгс-101'},  # другой регистр — та же группа
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    assert page.text.count('определён по файлу') == 2
    assert 'title="Институт определён по группе из этого файла"' in page.text
    assert 'определён по группе</span>' not in page.text  # источник — файл, не справочник
    assert 'Институт не определён' not in page.text

    # Bulk-create создаёт ровно значения предпросмотра: институт донора файла
    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    assert 'Создано студентов: 3.' in unquote_plus(bulk.headers['location'])
    by_name = {row[1]: (row[3], row[4]) for row in students_snapshot(storage)}
    assert by_name['Доноров Донор'] == ('ИСИ', 'ПГС-101')
    assert by_name['Пустов Пуст'] == ('ИСИ', 'ПГС-101')
    # Группы справочник не знает: написание строки сохраняется (канонизация
    # группы — только через справочник, как и у прежнего пути подсказок)
    assert by_name['Второв Втор'] == ('ИСИ', 'пгс-101')


def test_student_import_infile_conflicting_institutes(student_import_client: SanicTestClient):
    """Группа в файле у разных институтов: автозаполнения НЕТ, строка без
    института получает предупреждение; явные строки (институт указан)
    предупреждения о конфликте файла не получают."""
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Институт': 'ИСИ', 'Группа': 'ПГС-101'},
            {'ФИО': 'Петров Пётр', 'Институт': 'ИМИ', 'Группа': 'ПГС-101'},
            {'ФИО': 'Сидоров Сидор', 'Группа': 'ПГС-101'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Институт не определён: в файле группа относится к разным институтам' in page.text
    assert page.text.count('Институт не определён: в файле группа относится к разным институтам') == 1
    assert 'определён по файлу' not in page.text
    assert 'определён по группе' not in page.text


def test_student_import_file_conflicts_with_catalog(student_import_client: SanicTestClient):
    """Файл против справочника (правило 5): файл уникален, справочник
    однозначен, значения расходятся — берётся файловое, предупреждение;
    значения справочника не меняются (снимок ниже в ..._never_written)."""
    storage = app.ctx.storage
    seed_catalog(storage, 'ИСИ', 'ПГС-101')  # справочник: ПГС-101 → ИСИ
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Доноров Донор', 'Институт': 'ИМИ', 'Группа': 'ПГС-101'},
            {'ФИО': 'Пустов Пуст', 'Группа': 'ПГС-101'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    # Строка без института: файловое значение + предупреждение о справочнике
    assert page.text.count('Требует внимания: справочник относит группу к другому институту') == 1
    assert 'определён по файлу' in page.text

    created = post_import_action(student_import_client, token, 'row/3/create')
    assert created.status == 302
    by_name = {row[1]: (row[3], row[4]) for row in students_snapshot(storage)}
    assert by_name['Пустов Пуст'] == ('ИМИ', 'ПГС-101')


def test_student_import_infile_institute_case_variants_one_entry(student_import_client: SanicTestClient):
    """Регистровые варианты института в файле — ОДНА запись карты: строка без
    института автозаполняется первым написанием (или каноном справочника)."""
    storage = app.ctx.storage
    seed_catalog(storage, 'ИСИ', 'ТД-303')
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Доноров Донор', 'Институт': 'ИСИ', 'Группа': 'ТД-303'},
            {'ФИО': 'Регистров Регистр', 'Институт': 'иси', 'Группа': 'ТД-303'},
            {'ФИО': 'Пустов Пуст', 'Группа': 'тд-303'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'определён по файлу' in page.text
    assert 'Институт не определён: в файле группа относится к разным институтам' not in page.text

    created = post_import_action(student_import_client, token, 'row/4/create')
    assert created.status == 302
    by_name = {row[1]: (row[3], row[4]) for row in students_snapshot(storage)}
    # каноническое написание института и группы из справочника
    assert by_name['Пустов Пуст'] == ('ИСИ', 'ТД-303')


def test_student_import_donor_edit_changes_autofill(student_import_client: SanicTestClient):
    """Правка строки-«донора» меняет карту файла: на следующем рендере
    автозаполнение строк без института следует новым данным сессии."""
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Доноров Донор', 'Институт': 'ИСИ', 'Группа': 'ПГС-101'},
            {'ФИО': 'Пустов Пуст', 'Группа': 'ПГС-101'},
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'определён по файлу' in page.text

    edited = post_import_action(
        student_import_client,
        token,
        'row/2/edit',
        data={
            'full_name': 'Доноров Донор',
            'sex': '',
            'institute': 'ИМИ',
            'group': 'ПГС-101',
            'course': '',
        },
    )
    assert edited.status == 302
    page = get_preview(student_import_client, token)
    assert page.text.count('определён по файлу') == 1  # только строка без института
    body_around = page.text
    # Создаётся строка без института — институт теперь ИМИ из правки донора
    created = post_import_action(student_import_client, token, 'row/3/create')
    assert created.status == 302
    by_name = {row[1]: (row[3], row[4]) for row in students_snapshot(app.ctx.storage)}
    assert by_name['Пустов Пуст'] == ('ИМИ', 'ПГС-101')
    assert body_around.count('ИСИ') == 0


def test_student_import_catalog_tables_never_written(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    seed_catalog(storage, 'ИСИ', 'ПГС-101')
    seed_catalog(storage, 'ИМИ', 'ТД-303')
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Институт': 'НОВЫЙ ИНСТИТУТ', 'Группа': 'НОВАЯ-ГРУППА'},
            {'ФИО': 'Петров Пётр', 'Группа': 'НЕИЗВЕСТНАЯ-ГРУППА'},
            {'ФИО': 'Сидоров Сидор', 'Институт': 'ИСИ', 'Группа': 'ТД-303'},
            # Автозаполнение института из файла (Задача 2): донор + строка
            # без института — тоже не должно сеять ничего в справочники.
            {'ФИО': 'Доноров Донор', 'Институт': 'ФАЙЛ-ИНСТИТУТ', 'Группа': 'ФГ-1'},
            {'ФИО': 'Пустов Пуст', 'Группа': 'ФГ-1'},
        ],
    )
    token = import_session_token(response)
    # Снимки после загрузки (посев уровней при старте приложения не должен
    # попадать в сравнение): дальше предпросмотр/создание/завершение.
    before = catalog_snapshot(storage)
    levels_before = [
        tuple(row) for row in storage.connection.execute('SELECT id, name, active FROM levels ORDER BY id').fetchall()
    ]

    assert get_preview(student_import_client, token).status == 200
    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    assert 'Создано студентов: 5.' in unquote_plus(bulk.headers['location'])
    finish = post_import_action(student_import_client, token, 'finish')
    assert finish.status == 302

    # Инвариант Phase 2.5: импорт студентов не пишет в справочники НИЧЕГО
    assert catalog_snapshot(storage) == before
    levels_after = [
        tuple(row) for row in storage.connection.execute('SELECT id, name, active FROM levels ORDER BY id').fetchall()
    ]
    assert levels_after == levels_before


def test_student_import_candidates_matching_rules(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    ivanov = storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    twin = storage.create_student('Иванов Иван', 'Ж', 'ИМИ', 'СБ-202', '1')
    storage.add_student_alias(ivanov, 'Иванов И.И.')
    # Неактивная карточка кандидатом быть не должна
    inactive = storage.create_student('Петров Пётр', 'М', '', '', '')
    storage.set_student_active(inactive, False)

    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'ИВАНОВ ИВАН'},  # точное совпадение без учёта регистра — тёзки обе
            {'ФИО': 'Иванов И.И.'},  # совпадение по псевдониму
            {'ФИО': 'петров пётр'},  # неактивная карточка не кандидат
            {'ФИО': 'Иваноф Иван'},  # опечатка — не кандидат
        ],
    )
    token = import_session_token(response)
    before = students_snapshot(storage)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    assert 'Всего строк: 4 · новые: 2 · требуют решения: 2 · ошибочные: 0' in page.text
    assert f'<a href="/admin/people/{ivanov}">Иванов Иван</a>' in page.text
    assert f'<a href="/admin/people/{twin}">Иванов Иван</a>' in page.text
    assert '>псевдоним</span>' in page.text
    assert 'Совпадений с существующими студентами нет.' in page.text
    assert 'Петров Пётр</a>' not in page.text
    # Мини-форма кандидата: использовать существующего по строке предпросмотра
    assert (
        f'action="/admin/people/import/preview/{token}/row/2/use-existing"' in page.text
        and f'name="student_id" value="{ivanov}"' in page.text
        and f'Использовать существующего #{ivanov}</button>' in page.text
    )

    # Предпросмотр не создаёт и не меняет ничего
    assert students_snapshot(storage) == before


def test_student_import_bulk_create_atomic(student_import_client: SanicTestClient, monkeypatch):
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Пол': 'М'},
            {'ФИО': 'Петров Пётр', 'Пол': 'Ж'},
        ],
    )
    token = import_session_token(response)

    # Сбой в середине батча (NOT NULL) откатывает всё: ничего не вставлено
    real_create_students = storage.create_students

    def failing_create_students(students):
        broken = list(students)
        broken.insert(1, (None, '', '', '', ''))
        return real_create_students(broken)

    monkeypatch.setattr(storage, 'create_students', failing_create_students)
    _, response = student_import_client.post(
        f'/admin/people/import/preview/{token}/bulk-create',
        headers=get_auth_headers(),
        data=csrf_for(get_auth_headers()),
        allow_redirects=False,
    )
    assert response.status == 500
    assert students_snapshot(storage) == []
    monkeypatch.undo()

    # После сбоя строки остались неразобранными — повтор проходит
    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    assert 'Создано студентов: 2.' in unquote_plus(bulk.headers['location'])
    assert len(students_snapshot(storage)) == 2


def test_student_import_bulk_cancelled_when_candidate_appeared(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван', 'Пол': 'М'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 1 · новые: 1' in page.text

    # Между рендером и кликом появился студент с таким же ФИО
    storage.create_student('Иванов Иван', 'Ж', '', '', '')

    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    location = unquote_plus(bulk.headers['location'])
    assert 'Ничего не создано: у части строк появились совпадения с существующими студентами.' in location
    assert 'Проверьте раздел «Требуют решение».' in location
    # Ничего не создано (остался только вручную созданный), строка переквалифицировалась
    assert len(students_snapshot(storage)) == 1
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 1 · новые: 0 · требуют решения: 1' in page.text


def test_student_import_bulk_creates_only_clean_rows_with_unresolved_candidates_left(
    student_import_client: SanicTestClient,
):
    """D1 (QA): батч массового создания — ТОЛЬКО чистые A-строки; неразобранные
    строки с кандидатами в батч не входят и отмены НЕ вызывают. Полная отмена —
    только когда у ранее чистой строки кандидат появился между рендером и кликом
    (второй сценарий)."""
    storage = app.ctx.storage
    existing = storage.create_student('Иванов Иван', 'М', '', '', '')
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван'},  # B: кандидат с самого рендера, НЕ разобрана
            {'ФИО': 'Петров Пётр', 'Пол': 'М'},  # A: чистая
            {'ФИО': 'Сидоров Сидор', 'Пол': 'Ж'},  # A: чистая
        ],
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 3 · новые: 2 · требуют решения: 1' in page.text
    assert 'Создать 2 новых' in page.text

    # Клик «Создать 2 новых»: чистые созданы, строка-кандидат не отменяет батч
    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    location = unquote_plus(bulk.headers['location'])
    assert 'Создано студентов: 2.' in location
    assert 'Ничего не создано' not in location
    snapshot = students_snapshot(storage)
    assert [row[1] for row in snapshot] == ['Иванов Иван', 'Петров Пётр', 'Сидоров Сидор']
    page = get_preview(student_import_client, token)
    assert 'Решено: создано 2 · использовано существующих 0 · пропущено 0' in page.text
    assert 'Всего строк: 3 · новые: 0 · требуют решения: 1 · ошибочные: 0' in page.text
    assert f'<a href="/admin/people/{existing}">Иванов Иван</a>' in page.text

    # Второй сценарий: чистая строка ПОЛУЧИЛА кандидата после рендера — отмена всего
    response = upload_student_xlsx(
        student_import_client, [{'ФИО': 'Козлов Кирилл', 'Пол': 'М'}, {'ФИО': 'Орлов Олег', 'Пол': 'Ж'}]
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert 'Создать 2 новых' in page.text
    storage.create_student('Орлов Олег', 'М', '', '', '')

    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert bulk.status == 302
    location = unquote_plus(bulk.headers['location'])
    assert 'Ничего не создано: у части строк появились совпадения с существующими студентами.' in location
    # Ничего из этого файла не создано (в БД — только прежние 4 карточки);
    # строка «Орлов» переквалифицировалась в «требуют решения», «Козлов»
    # остался чистой (батч отменён целиком — повторный клик создаст её одну)
    assert len(students_snapshot(storage)) == 4
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 2 · новые: 1 · требуют решения: 1' in page.text


def test_student_import_actions_on_error_row_rejected(student_import_client: SanicTestClient):
    """N2: create/use-existing на ошибочной строке — точное сообщение об
    ошибке строки (пропуск ошибочных остаётся доступен — см. skip-тесты)."""
    response = upload_student_xlsx(student_import_client, [{'ФИО': '   ', 'Пол': 'Мужской'}])
    token = import_session_token(response)

    created = post_import_action(student_import_client, token, 'row/2/create')
    assert created.status == 302
    assert 'Строка 2 содержит ошибку — исправьте её.' in unquote_plus(created.headers['location'])

    used = post_import_action(student_import_client, token, 'row/2/use-existing', {'student_id': '1'})
    assert used.status == 302
    assert 'Строка 2 содержит ошибку — исправьте её.' in unquote_plus(used.headers['location'])

    assert len(students_snapshot(app.ctx.storage)) == 0


def test_student_import_row_create_on_full_match(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    existing = storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван', 'Пол': 'Ж'}])
    token = import_session_token(response)

    created = post_import_action(student_import_client, token, 'row/2/create')
    assert created.status == 302
    assert 'Строка 2: студент «Иванов Иван» создан.' in unquote_plus(created.headers['location'])
    snapshot = students_snapshot(storage)
    assert len(snapshot) == 2
    assert snapshot[-1][1] == 'Иванов Иван' and snapshot[-1][2] == 'Ж'
    assert snapshot[-1][0] != existing
    # Аудит создания — с источником импорта
    created_events = audit_details(storage, 'student_created')
    assert created_events == [{'student_id': snapshot[-1][0], 'full_name': 'Иванов Иван', 'source': 'excel-import'}]


def test_student_import_use_existing_no_writes_no_links(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    existing = storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'ИВАНОВ ИВАН'}])
    token = import_session_token(response)

    # Некорректный student_id — 400; несуществующий — flash-ошибка
    assert post_import_action(student_import_client, token, 'row/2/use-existing', {'student_id': 'abc'}).status == 400
    unknown = post_import_action(student_import_client, token, 'row/2/use-existing', {'student_id': '999'})
    assert unknown.status == 302
    assert 'Студент не найден.' in unquote_plus(unknown.headers['location'])

    used = post_import_action(student_import_client, token, 'row/2/use-existing', {'student_id': str(existing)})
    assert used.status == 302
    assert f'Строка 2: использован существующий студент #{existing}.' in unquote_plus(used.headers['location'])

    # Только пометка в сессии: ничего не создано, не связано, не записано в аудит
    assert students_snapshot(storage) == [(existing, 'Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')]
    assert audit_details(storage, 'student_created') == []
    refs = storage.connection.execute('SELECT student_ref_id FROM competitions').fetchall()
    assert [row['student_ref_id'] for row in refs] == [None]
    page = get_preview(student_import_client, token)
    assert f'<span class="badge text-bg-primary">использован #{existing}</span>' in page.text


def test_student_import_skip_rows(student_import_client: SanicTestClient):
    response = upload_student_xlsx(
        student_import_client,
        [{'ФИО': 'Иванов Иван', 'Пол': 'М'}, {'ФИО': '   ', 'Пол': 'Мужской'}],
    )
    token = import_session_token(response)

    skipped = post_import_action(student_import_client, token, 'row/2/skip')
    assert skipped.status == 302
    assert 'Строка 2 пропущена.' in unquote_plus(skipped.headers['location'])
    # Ошибочную строку тоже можно пропустить (не блокирует завершение)
    skipped_error = post_import_action(student_import_client, token, 'row/3/skip')
    assert skipped_error.status == 302
    assert 'Строка 3 пропущена.' in unquote_plus(skipped_error.headers['location'])

    page = get_preview(student_import_client, token)
    assert 'Решено: создано 0 · использовано существующих 0 · пропущено 2' in page.text
    assert '<span class="badge text-bg-secondary">пропущен</span>' in page.text


def test_student_import_edit_and_candidate_rerun(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    storage.create_student('Иванов Иван', 'М', 'ИСИ', 'ПГС-101', '2')
    response = upload_student_xlsx(
        student_import_client,
        [{'ФИО': 'Иванов Иван', 'Пол': 'Муж'}, {'ФИО': 'Сидоров Сидор', 'Пол': 'М'}],
    )
    token = import_session_token(response)

    # Некорректная правка: строка обновлена, но осталась ошибочной
    edited = post_import_action(
        student_import_client,
        token,
        'row/2/edit',
        {'full_name': 'Иванов Иван', 'sex': 'Мужской', 'institute': '', 'group': '', 'course': ''},
    )
    assert edited.status == 302
    assert 'Строка 2 обновлена, но данные некорректны: Пол должен быть «М» или «Ж».' in unquote_plus(
        edited.headers['location']
    )
    page = get_preview(student_import_client, token)
    assert 'ошибочные: 1' in page.text

    # Исправление: строка снова валидна и попадает в «Требуют решения» (кандидат по ФИО)
    fixed = post_import_action(
        student_import_client,
        token,
        'row/2/edit',
        {'full_name': 'Иванов Иван', 'sex': 'М', 'institute': '', 'group': '', 'course': '2'},
    )
    assert fixed.status == 302
    assert 'Строка 2 обновлена.' in unquote_plus(fixed.headers['location'])
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 2 · новые: 1 · требуют решения: 1 · ошибочные: 0' in page.text

    # Смена ФИО на написание существующего студента переводит строку A→B
    renamed = post_import_action(
        student_import_client,
        token,
        'row/3/edit',
        {'full_name': 'Иванов Иван', 'sex': 'М', 'institute': '', 'group': '', 'course': ''},
    )
    assert renamed.status == 302
    page = get_preview(student_import_client, token)
    assert 'Всего строк: 2 · новые: 0 · требуют решения: 2 · ошибочные: 0' in page.text


def test_student_import_finish_flow_and_audit(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client,
        [
            {'ФИО': 'Иванов Иван', 'Пол': 'М', 'Институт': 'ИСИ', 'Группа': 'ПГС-101', 'Курс': '2'},
            {'ФИО': 'Петров Пётр', 'Пол': 'Ж'},
            {'ФИО': '   ', 'Пол': ''},
        ],
    )
    token = import_session_token(response)

    # Завершение недоступно, пока есть неразобранные строки
    blocked = post_import_action(student_import_client, token, 'finish')
    assert blocked.status == 302
    assert 'Завершение недоступно: не разобрано строк: 2.' in unquote_plus(blocked.headers['location'])
    # Сессия жива
    assert get_preview(student_import_client, token).status == 200

    bulk = post_import_action(student_import_client, token, 'bulk-create')
    assert 'Создано студентов: 2.' in unquote_plus(bulk.headers['location'])
    post_import_action(student_import_client, token, 'row/4/skip')

    finished = post_import_action(student_import_client, token, 'finish')
    assert finished.status == 302
    location = unquote_plus(finished.headers['location'])
    assert location.startswith('/admin/people?')
    assert 'Импорт завершён: создано 2, использовано существующих 0, пропущено 1.' in location

    # Аудит: по событию student_created на строку + один итог импорта
    created = audit_details(storage, 'student_created')
    assert len(created) == 2
    assert {event['source'] for event in created} == {'excel-import'}
    completed = audit_details(storage, 'student_import_completed')
    assert completed == [{'created': 2, 'reused_existing': 0, 'skipped': 1, 'errors': 0, 'total': 3}]

    # Сессия удалена; карточки остались
    assert token not in student_import_sessions
    assert len(students_snapshot(storage)) == 2


def test_student_import_discard(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    response = upload_student_xlsx(
        student_import_client, [{'ФИО': 'Иванов Иван', 'Пол': 'М'}, {'ФИО': 'Петров Пётр', 'Пол': 'Ж'}]
    )
    token = import_session_token(response)
    post_import_action(student_import_client, token, 'row/2/create')

    discarded = post_import_action(student_import_client, token, 'discard')
    assert discarded.status == 302
    assert 'Импорт отменён. Созданные карточки (1) сохранены.' in unquote_plus(discarded.headers['location'])
    assert token not in student_import_sessions
    # Созданное подтверждением остаётся, итогового события нет
    assert len(students_snapshot(storage)) == 1
    assert audit_details(storage, 'student_import_completed') == []

    # Отмена без созданного — короткое сообщение
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Сидоров Сидор'}])
    token = import_session_token(response)
    discarded = post_import_action(student_import_client, token, 'discard')
    assert 'Импорт отменён.' in unquote_plus(discarded.headers['location'])
    assert 'Созданные карточки' not in unquote_plus(discarded.headers['location'])
    # Состояние БД не изменилось: остаётся только подтверждённая ранее карточка
    snapshot = students_snapshot(storage)
    assert len(snapshot) == 1
    assert snapshot[0][1] == 'Иванов Иван'


def test_student_import_sessions_access_control(student_import_client: SanicTestClient):
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван'}])
    token = import_session_token(response)

    # Неизвестный токен
    missing = get_preview(student_import_client, 'no-such-token')
    assert missing.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(missing.headers['location'])

    # Просроченная сессия
    student_import_sessions[token]['created_at'] = time.time() - (2 * 60 * 60 + 60)
    expired = get_preview(student_import_client, token)
    assert expired.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(expired.headers['location'])
    assert token not in student_import_sessions

    # Чужая сессия (другой user_id) для текущего админа недоступна
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Петров Пётр'}])
    token = import_session_token(response)
    student_import_sessions[token]['user_id'] = 999
    foreign = get_preview(student_import_client, token)
    assert foreign.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(foreign.headers['location'])


def test_student_import_post_routes_reject_missing_csrf(student_import_client: SanicTestClient):
    response = upload_student_xlsx(
        student_import_client,
        [{'ФИО': 'Иванов Иван', 'Пол': 'М'}, {'ФИО': '   ', 'Пол': ''}],
    )
    token = import_session_token(response)
    suffixes = [
        'bulk-create',
        'row/2/create',
        'row/2/use-existing',
        'row/2/skip',
        'row/2/edit',
        'finish',
        'discard',
    ]
    for suffix in suffixes:
        _, response = student_import_client.post(
            f'/admin/people/import/preview/{token}/{suffix}',
            headers=get_auth_headers(),
            data={'full_name': 'X'},
            allow_redirects=False,
        )
        assert response.status == 403, suffix
        assert 'CSRF' in response.text

    # Загрузка файла — тоже POST с CSRF
    _, response = student_import_client.post('/admin/people/import', headers=get_auth_headers(), allow_redirects=False)
    assert response.status == 403
    assert 'CSRF' in response.text


def test_student_import_action_roles(student_import_client: SanicTestClient):
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван'}])
    token = import_session_token(response)
    editor_headers = get_auth_headers(role='editor')

    _, response = student_import_client.get(f'/admin/people/import/preview/{token}', headers=editor_headers)
    assert response.status == 403
    for suffix in ('bulk-create', 'row/2/create', 'row/2/skip', 'finish', 'discard'):
        _, response = student_import_client.post(
            f'/admin/people/import/preview/{token}/{suffix}',
            headers=editor_headers,
            data={**csrf_for(editor_headers), 'student_id': '1'},
            allow_redirects=False,
        )
        assert response.status == 403, suffix

    # Без аутентификации — 401 (кроме GET — редирект в логин)
    _, response = student_import_client.post(f'/admin/people/import/preview/{token}/finish', data={'csrf_token': 'x'})
    assert response.status == 401
    _, response = student_import_client.get(f'/admin/people/import/preview/{token}', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')


def test_student_import_invalid_row_numbers(student_import_client: SanicTestClient):
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван'}])
    token = import_session_token(response)
    for suffix in ('row/abc/create', 'row/2.5/create', 'row/999/create', 'row/-1/skip'):
        _, response = student_import_client.post(
            f'/admin/people/import/preview/{token}/{suffix}',
            headers=get_auth_headers(),
            data=csrf_for(get_auth_headers()),
            allow_redirects=False,
        )
        assert response.status == 400, suffix


def test_student_import_already_resolved_row_rejected(student_import_client: SanicTestClient):
    response = upload_student_xlsx(student_import_client, [{'ФИО': 'Иванов Иван'}])
    token = import_session_token(response)
    assert post_import_action(student_import_client, token, 'row/2/skip').status == 302
    repeated = post_import_action(student_import_client, token, 'row/2/create')
    assert repeated.status == 302
    assert 'Строка 2 уже разобрана.' in unquote_plus(repeated.headers['location'])
    assert len(students_snapshot(app.ctx.storage)) == 0


def test_student_import_xss_rendered_as_text(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    payload = '<script>alert("xss")</script>'
    response = upload_student_xlsx(
        student_import_client, [{'ФИО': payload, 'Институт': '<img src=x onerror=alert(1)>'}]
    )
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    assert page.status == 200
    # Значения файла рендерятся как текст: сырой payload в разметке отсутствует
    assert payload not in page.text
    assert '<img src=x onerror=' not in page.text
    assert '&lt;script&gt;alert(' in page.text
    assert '&lt;/script&gt;' in page.text
    assert '&lt;img src=x onerror=alert(1)&gt;' in page.text

    created = post_import_action(student_import_client, token, 'row/2/create')
    assert created.status == 302
    # Значение файла хранится как есть, но рендерится экранированным:
    # flash с ФИО из файла отображается как текст (Jinja-экранирование)
    stored = students_snapshot(storage)
    assert stored[0][1] == payload
    _, flashed = student_import_client.get(created.headers['location'], headers=get_auth_headers())
    assert flashed.status == 200
    assert payload not in flashed.text
    assert '&lt;script&gt;alert(' in flashed.text
    page = get_preview(student_import_client, token)
    assert payload not in page.text


def test_student_import_confirm_strings_numbers_only(student_import_client: SanicTestClient):
    # Источник шаблона: confirm-строки содержат только счётчики, не значения файла
    template_source = (Path(__file__).parent / 'templates' / 'admin_people_import_preview.html').read_text()
    assert "confirm('Создать {{ new_count }} новых студентов?')" in template_source
    assert (
        "confirm('Отменить импорт? Неразобранные строки будут отброшены, уже созданные карточки останутся.')"
        in template_source
    )

    response = upload_student_xlsx(student_import_client, [{'ФИО': 'ИМПОРТ-НЕ-В-КОНФИРМЕ Иванов', 'Пол': 'М'}])
    token = import_session_token(response)
    page = get_preview(student_import_client, token)
    confirms = re.findall(r'onsubmit="return confirm\((.*?)\)"', page.text)
    assert confirms
    for confirm_text in confirms:
        assert 'ИМПОРТ-НЕ-В-КОНФИРМЕ' not in confirm_text


def test_student_import_route_priority_over_student_id(student_import_client: SanicTestClient):
    _, response = student_import_client.get('/admin/people/import', headers=get_auth_headers())
    assert response.status == 200
    assert 'Загрузка файла' in response.text
    assert 'Студент не найден.' not in response.text

    _, response = student_import_client.get('/admin/people/import/template', headers=get_auth_headers())
    assert response.status == 200
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


def test_student_import_isolation_and_reconcile_compat(student_import_client: SanicTestClient):
    storage = app.ctx.storage
    record_id = save_reconcile_record(storage, 'Иванов Иван', datetime(2026, 1, 10))

    response = upload_student_xlsx(
        student_import_client,
        [{'ФИО': 'Иванов Иван', 'Пол': 'М', 'Институт': 'ИСИ', 'Группа': 'ПГС-101', 'Курс': '2'}],
    )
    token = import_session_token(response)
    # Снимок аккаунтов после загрузки: предпросмотр сам ничего не пишет;
    # поля last_seen/посев учёток в сравнении не участвуют.
    users_before = [
        tuple(row)
        for row in storage.connection.execute(
            'SELECT id, username, role, active, student_ref_id FROM users ORDER BY id'
        ).fetchall()
    ]
    post_import_action(student_import_client, token, 'row/2/create')
    post_import_action(student_import_client, token, 'finish')

    # Изоляция: записи и аккаунты не тронуты, student_ref_id остался NULL
    record = storage.connection.execute('SELECT * FROM competitions WHERE id = ?', (record_id,)).fetchone()
    assert record['student_ref_id'] is None
    assert record['student_name'] == 'Иванов Иван'
    users_after = [
        tuple(row)
        for row in storage.connection.execute(
            'SELECT id, username, role, active, student_ref_id FROM users ORDER BY id'
        ).fetchall()
    ]
    assert users_after == users_before

    # Совместимость: импортированный студент — кандидат в сопоставлении
    student_id = students_snapshot(storage)[0][0]
    _, page = student_import_client.get('/admin/people/reconcile', headers=get_auth_headers())
    assert page.status == 200
    assert f'<a href="/admin/people/{student_id}">Иванов Иван</a>' in page.text


# --- Student-backed participant entry + импорт участников события. ---
# Полные сценарии на реальном SQLite-адаптере (паттерн student_import_client):
# автодополнение из карточек Students, запись student_ref_id при явном
# выборе, Excel-импорт участников события с предпросмотром и явными
# решениями (Excel > Student > пусто; карточка студента не мутируется).


@pytest.fixture
def event_import_client(client: SanicTestClient, tmp_path):
    storage = SQLiteAdapter(str(tmp_path / 'event_import.sqlite3'))
    storage.create_user(settings.auth_admin_username, 'hash', 'admin')
    storage.create_user(settings.auth_editor_username, 'hash', 'editor')
    storage.create_user(settings.auth_viewer_username or 'viewer', 'hash', 'viewer')
    storage.create_user('sportik', 'hash', 'athlete')
    previous = getattr(app.ctx, 'storage', None)
    app.ctx.storage = storage
    student_import_sessions.clear()
    event_import_sessions.clear()
    try:
        yield client
    finally:
        app.ctx.storage = previous
        student_import_sessions.clear()
        event_import_sessions.clear()


def make_calendar_event(storage, **kwargs) -> int:
    """Событие календаря с дефолтами синтетического «Забега 2026»."""
    values = {
        'name': 'Забег 2026',
        'date': '2026-05-10T00:00:00',
        'date_to': '2026-05-11T00:00:00',
        'level': 'внутривузовские',
        'sport': 'Бег',
        'url': '',
    }
    values.update(kwargs)
    return storage.create_calendar_event(**values)


def upload_event_xlsx(
    client: SanicTestClient,
    event_id: int,
    rows: list[dict],
    *,
    role: str = 'admin',
    filename: str = 'participants.xlsx',
    columns: list[str] | None = None,
):
    """Загрузить файл импорта участников события; ответ — редирект."""
    df = pd.DataFrame(rows, columns=columns) if columns else pd.DataFrame(rows)
    file_obj = BytesIO()
    df.to_excel(file_obj, index=False)
    file_obj.seek(0)
    headers = get_auth_headers(role=role)
    _, response = client.post(
        f'/calendar/{event_id}/participants/import',
        headers=headers,
        data=csrf_for(headers),
        files={
            'file': (
                filename,
                file_obj.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        },
        allow_redirects=False,
    )
    return response


def post_event_import_action(
    client: SanicTestClient,
    event_id: int,
    token: str,
    suffix: str,
    data: dict | None = None,
    *,
    role: str = 'admin',
):
    headers = get_auth_headers(role=role)
    _, response = client.post(
        f'/calendar/{event_id}/participants/import/preview/{token}/{suffix}',
        headers=headers,
        data={**csrf_for(headers), **(data or {})},
        allow_redirects=False,
    )
    return response


def get_event_preview(client: SanicTestClient, event_id: int, token: str, *, role: str = 'admin'):
    _, response = client.get(
        f'/calendar/{event_id}/participants/import/preview/{token}',
        headers=get_auth_headers(role=role),
        allow_redirects=False,
    )
    return response


def event_records_snapshot(storage) -> list[tuple]:
    rows = storage.connection.execute(
        'SELECT student_name, student_sex, institute, "group", course, sport, date, date_to, level, name, '
        'position, student_ref_id FROM competitions ORDER BY id'
    ).fetchall()
    return [tuple(row) for row in rows]


# --- Автодополнение из Students (Task A). ---


def test_athletes_search_finds_student_without_history(event_import_client: SanicTestClient):
    # Студент только импортирован в Students и никогда не участвовал —
    # раньше источник (_known_athletes) его не знал вовсе.
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    _, response = event_import_client.get('/api/athletes/search?q=Иван', headers=get_auth_headers())
    assert response.status == 200
    assert response.json == [
        {
            'kind': 'student',
            'student_id': student_id,
            'name': 'Иванов Иван Иванович',
            'sex': 'М',
            'institute': 'ИСЭиУ',
            'group': 'ЭБ-241',
            'course': '2',
        }
    ]


def test_athletes_search_dedup_legacy_and_student_by_exact_name(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    save_reconcile_record(storage, 'Иванов Иван Иванович', datetime(2026, 1, 10))
    _, response = event_import_client.get('/api/athletes/search?q=Иван', headers=get_auth_headers())
    assert response.status == 200
    # Точное совпадение ФИО (strip + casefold): карточка Student приоритетна,
    # легаси-дубль не показывается. Слияния/авто-связи нет — только вид списка.
    assert [item['kind'] for item in response.json] == ['student']

    # Разные написания остаются обоими вариантами.
    save_reconcile_record(storage, 'Иванов Иван Иванович (Ивано-Ивановск)', datetime(2026, 2, 1))
    _, response = event_import_client.get('/api/athletes/search?q=Иванов Иван', headers=get_auth_headers())
    kinds = [item['kind'] for item in response.json]
    assert kinds.count('student') == 1
    assert kinds.count('legacy') == 1


def test_athletes_search_students_first_and_namesakes_distinct(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    first = storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИСИ', 'ПГС-201', '1')
    second = storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИМИ', 'СБ-202', '2')
    save_reconcile_record(storage, 'Абрамов Артём Артёмович', datetime(2026, 1, 10))
    _, response = event_import_client.get('/api/athletes/search?q=ид', headers=get_auth_headers())
    assert response.status == 200
    # Оба тёзки различимы (id карточек), Students — раньше легаси-подсказок.
    assert [item['student_id'] for item in response.json if item['kind'] == 'student'] == [first, second]
    assert response.json[0]['kind'] == 'student'


def test_athletes_search_inactive_student_not_offered(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    storage.set_student_active(student_id, False)
    _, response = event_import_client.get('/api/athletes/search?q=Иван', headers=get_auth_headers())
    assert response.status == 200
    assert response.json == []


def test_calendar_participant_add_writes_student_ref(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_ref_id': str(student_id),
            'course': '2',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'].startswith(f'/calendar/{event_id}?admin_message=')
    record = storage.connection.execute(
        'SELECT student_name, student_sex, institute, "group", student_ref_id, calendar_event_id FROM competitions'
    ).fetchone()
    # Снимок = значения формы (не карточки), связь — по явному выбору;
    # P2: участие явно связано с событием (участники читаются по ссылке).
    assert tuple(record) == ('Иванов Иван Иванович', '', '', '', student_id, event_id)
    # Участник виден на странице события (id-first).
    assert [p['student_name'] for p in storage.list_calendar_event_participants(event_id)] == ['Иванов Иван Иванович']

    # Инлайн-правка записи (POST /competition/<id>) связь не сбрасывает.
    record_id = storage.connection.execute('SELECT id FROM competitions').fetchone()['id']
    _, response = event_import_client.post(
        f'/competition/{record_id}',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСЭиУ',
            'group': 'ЭБ-241',
            'sport': 'Бег',
            'date': '10.05.2026-11.05.2026',
            'level': 'внутривузовские',
            'name': 'Забег 2026',
            'position': '1',
            'course': '2',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    ref_after = storage.connection.execute(
        'SELECT student_ref_id FROM competitions WHERE id = ?', (record_id,)
    ).fetchone()['student_ref_id']
    assert ref_after == student_id


def test_calendar_participant_add_rejects_invalid_student_ref(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    storage.set_student_active(student_id, False)
    headers = get_auth_headers()
    for raw_ref in ('99999', str(student_id)):
        _, response = event_import_client.post(
            f'/calendar/{event_id}/participants',
            headers=headers,
            data={
                **csrf_for(headers),
                'student_name': 'Иванов Иван Иванович',
                'student_ref_id': raw_ref,
                'course': '2',
            },
            allow_redirects=False,
        )
        assert response.status == 302
        assert 'Выбранная карточка студента не найдена или неактивна.' in unquote_plus(response.headers['location'])
    assert storage.connection.execute('SELECT COUNT(*) FROM competitions').fetchone()[0] == 0


# --- Event Model, Wave 1 P2: стабильная связь «запись → событие». ---


def record_identity_snapshot(storage, record_id: int) -> tuple:
    """Сырой снимок event-owned + внутренних полей записи (byte-equivalent)."""
    row = storage.connection.execute(
        'SELECT name, sport, date, date_to, level, discipline, result, calendar_event_id, student_ref_id '
        'FROM competitions WHERE id = ?',
        (record_id,),
    ).fetchone()
    return tuple(row)


def test_linked_record_edit_ignores_forged_event_fields_and_keeps_identity(event_import_client: SanicTestClient):
    """P2 + QA O1: запись со ссылкой на событие не владеет полями
    название/вид спорта/дата/уровень — даже подделанная форма их не меняет;
    discipline/result (форма их не присылает) не затираются модельными None;
    не-event-owned поля (институт/группа/место) редактируются."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_ref_id': str(student_id),
            'course': '2',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    record_id = storage.connection.execute('SELECT id FROM competitions').fetchone()['id']
    storage.connection.execute(
        "UPDATE competitions SET discipline = 'Бег 100 м', result = '11.2' WHERE id = ?",
        (record_id,),
    )
    storage.connection.commit()
    before = record_identity_snapshot(storage, record_id)
    assert before[:5] == ('Забег 2026', 'Бег', '2026-05-10T00:00:00', '2026-05-11T00:00:00', 'внутривузовские')

    # Подделанная форма: чужие название/вид спорта/дата/уровень + правка
    # обычных полей.
    _, response = event_import_client.post(
        f'/competition/{record_id}',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Иванов Иван Иванович',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ПГС-201',
            'sport': 'Подделанный спорт',
            'date': '01.01.2030',
            'level': 'олимпийские',
            'name': 'Подделанное название',
            'position': '3',
            'course': '3',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert record_identity_snapshot(storage, record_id) == before
    row = storage.connection.execute(
        'SELECT institute, "group", position, course FROM competitions WHERE id = ?',
        (record_id,),
    ).fetchone()
    assert tuple(row) == ('ИСИ', 'ПГС-201', 3, 3)


def test_null_record_edit_updates_event_owned_fields(event_import_client: SanicTestClient):
    """QA O1 (NULL-путь): у записи без ссылки на событие event-owned поля
    по-прежнему редактируются формой; discipline/result сохраняются."""
    storage = app.ctx.storage
    record = Competition(
        student_id='id-null-record',
        student_name='Осипов Осип',
        student_sex='М',
        institute='ИСИ',
        group='ГРП-101',
        course=2,
        sport='Бег',
        date=datetime(2026, 1, 10),
        level='внутривузовские',
        name='Старый кубок',
        position=1,
        discipline='Эстафета',
        result='3:21',
    )
    storage.save_competitions([record])
    record_id = int(storage.connection.execute('SELECT id FROM competitions').fetchone()['id'])

    headers = get_auth_headers(role='editor')
    _, response = event_import_client.post(
        f'/competition/{record_id}',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Осипов Осип',
            'student_sex': 'М',
            'institute': 'ИСИ',
            'group': 'ГРП-101',
            'sport': 'Лыжи',
            'date': '20.02.2026',
            'level': 'региональные',
            'name': 'Новый кубок',
            'position': '2',
            'course': '3',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert record_identity_snapshot(storage, record_id) == (
        'Новый кубок',
        'Лыжи',
        '2026-02-20T00:00:00',
        None,
        'региональные',
        'Эстафета',
        '3:21',
        None,
        None,
    )


def test_calendar_edit_route_syncs_participations_and_audits(event_import_client: SanicTestClient):
    """Правка события синхронизирует связанные записи (5 полей) и пишет
    audit calendar_event_edited с old/new и synced_participations."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    headers = get_auth_headers()
    for name in ('Иванов Иван Иванович', 'Петров Пётр Петрович'):
        _, response = event_import_client.post(
            f'/calendar/{event_id}/participants',
            headers=headers,
            data={**csrf_for(headers), 'student_name': name, 'course': '1'},
            allow_redirects=False,
        )
        assert response.status == 302
    # NULL-сосед с тем же старым пресетом — синхронизация его не трогает.
    storage.save_competitions(
        [
            Competition(
                student_id='id-null-neighbor',
                student_name='Сидоров Сидор',
                student_sex='М',
                institute='ИСИ',
                group='ГРП-101',
                course=1,
                sport='Бег',
                date=datetime(2026, 5, 10),
                date_to=datetime(2026, 5, 11),
                level='внутривузовские',
                name='Забег 2026',
                position=1,
            )
        ]
    )

    _, response = event_import_client.post(
        f'/calendar/{event_id}/edit',
        headers=headers,
        data={
            **csrf_for(headers),
            'name': 'Забег 2026 — обновлённый',
            'date': '20-21.06.2026',
            'level': 'межвузовские',
            'sport': 'Лыжи',
            'url': '',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert response.headers['location'].startswith('/calendar?admin_message=')

    event = storage.get_calendar_event(event_id)
    assert event['name'] == 'Забег 2026 — обновлённый'
    assert event['date'] == '2026-06-20T00:00:00'
    assert event['date_to'] == '2026-06-21T00:00:00'

    rows = storage.connection.execute(
        'SELECT student_name, name, sport, date, date_to, level, calendar_event_id '
        'FROM competitions ORDER BY student_name'
    ).fetchall()
    assert len(rows) == 3
    for row in rows[:2]:  # Иванов/Петров — связанные, синхронизированы
        assert tuple(row)[1:] == (
            'Забег 2026 — обновлённый',
            'Лыжи',
            '2026-06-20T00:00:00',
            '2026-06-21T00:00:00',
            'межвузовские',
            event_id,
        )
    assert tuple(rows[2])[1:] == (
        'Забег 2026',
        'Бег',
        '2026-05-10T00:00:00',
        '2026-05-11T00:00:00',
        'внутривузовские',
        None,
    )

    assert audit_details(storage, 'calendar_event_edited') == [
        {
            'event_id': event_id,
            'name': {'old': 'Забег 2026', 'new': 'Забег 2026 — обновлённый'},
            'date': {'old': '2026-05-10T00:00:00', 'new': '2026-06-20T00:00:00'},
            'date_to': {'old': '2026-05-11T00:00:00', 'new': '2026-06-21T00:00:00'},
            'sport': {'old': 'Бег', 'new': 'Лыжи'},
            'level': {'old': 'внутривузовские', 'new': 'межвузовские'},
            'synced_participations': 2,
        }
    ]


def test_calendar_delete_blocked_by_linked_participant_real_storage(event_import_client: SanicTestClient):
    """Удаление события с участников: блокирует ССЫЛКА (не пресет — запись
    намеренно с другими полями); пустое событие удаляется как раньше."""
    storage = app.ctx.storage
    event_id = make_calendar_event(
        storage, name='Уникальный забег', date='2027-01-10T00:00:00', date_to=None, level='', sport='Плавание'
    )
    # Связанная запись с РАЗОШЕДШИМСЯ пресетом: блокирует только ссылка.
    storage.save_competitions(
        [
            Competition(
                student_id='id-linked-far',
                student_name='Далёков Далёк',
                student_sex='М',
                institute='ИСИ',
                group='ГРП-101',
                course=1,
                sport='Бег',
                date=datetime(2026, 2, 20),
                level='внутривузовские',
                name='Совсем другое название',
                position=1,
                calendar_event_id=event_id,
            )
        ]
    )
    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/delete',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'есть участники: 1.' in unquote_plus(response.headers['location'])
    assert storage.get_calendar_event(event_id) is not None

    # Пустое событие удаляется (записей по ссылке и по пресету нет).
    empty_id = make_calendar_event(
        storage, name='Пустой турнир', date='2027-03-01T00:00:00', date_to=None, level='', sport=''
    )
    _, response = event_import_client.post(
        f'/calendar/{empty_id}/delete',
        headers=headers,
        data=csrf_for(headers),
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Соревнование удалено' in unquote_plus(response.headers['location'])
    assert storage.get_calendar_event(empty_id) is None


def test_event_import_row_add_and_create_student_write_link(event_import_client: SanicTestClient):
    """Все явные пути импорта участников пишут calendar_event_id (row-add и
    create-student здесь; bulk — test_event_import_relations_event_preset)."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},
            {'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1},
        ],
    )
    token = import_session_token(response)

    response = post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    assert 'Строка 2: участник добавлен.' in unquote_plus(response.headers['location'])
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/3/create-student',
        {'full_name': 'Козлов Козьма Козьмич', 'sex': 'М', 'institute': 'ИСИ', 'group': 'Т-100', 'course': '1'},
    )
    assert 'студент «Козлов Козьма Козьмич» создан, участник добавлен.' in unquote_plus(response.headers['location'])

    links = [
        row['calendar_event_id']
        for row in storage.connection.execute('SELECT calendar_event_id FROM competitions ORDER BY id').fetchall()
    ]
    assert links == [event_id, event_id]
    assert [participant['student_name'] for participant in storage.list_calendar_event_participants(event_id)] == [
        'Иванов Иван Иванович',
        'Козлов Козьма Козьмич',
    ]


def test_calendar_event_page_shows_import_button_and_ref_field(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers())
    assert response.status == 200
    assert 'href="/calendar/%d/participants/import">Импортировать участников из Excel' % event_id in response.text
    assert 'name="student_ref_id" value="" id="participant-student-ref"' in response.text
    assert 'id="participant-student-badge"' in response.text
    # viewer — страница события доступна, кнопки импорта нет.
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers(role='viewer'))
    assert response.status == 200
    assert 'Импортировать участников из Excel' not in response.text


# --- Импорт участников события (Task B): страницы/шаблон/загрузка. ---


def test_event_import_pages_require_moderator(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(
        f'/calendar/{event_id}/participants/import', headers=get_auth_headers(role='viewer')
    )
    assert response.status == 403
    _, response = event_import_client.get(f'/calendar/{event_id}/participants/import', headers=get_athlete_headers())
    assert response.status == 403
    _, response = event_import_client.get(f'/calendar/{event_id}/participants/import', allow_redirects=False)
    assert response.status == 302
    assert response.headers['location'].startswith('/login')
    _, response = event_import_client.post(f'/calendar/{event_id}/participants/import', data={'x': '1'})
    assert response.status == 401
    # editor — доступен (импорт участников — модераторская операция).
    _, response = event_import_client.get(
        f'/calendar/{event_id}/participants/import', headers=get_auth_headers(role='editor')
    )
    assert response.status == 200


def test_event_import_upload_page_content(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}/participants/import', headers=get_auth_headers())
    assert response.status == 200
    assert '<h1 class="page-title mt-2">Импорт участников</h1>' in response.text
    assert 'Соревнование: Забег 2026 · 10-11.05.2026' in response.text
    assert 'href="/calendar/%d">← К соревнованию' % event_id in response.text
    assert 'Обязательное поле — только ФИО' in response.text
    assert 'пустые поля подтянутся из неё' in response.text
    assert 'Лишние колонки в файле игнорируются.' in response.text
    assert 'Ничего не добавляется, пока вы не подтвердите строки' in response.text
    assert 'Скачать пустой шаблон (xlsx)' in response.text
    assert 'action="/calendar/%d/participants/import"' % event_id in response.text
    assert 'import-form' not in response.text
    assert 'accept=".xlsx"' in response.text


def test_event_import_template_headers_only(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_custom_field('discipline', 'Дисциплина', 'text', False, True, True, True, 0)
    storage.create_custom_field('hidden_one', 'Скрытое', 'text', False, False, False, False, 1)
    _, response = event_import_client.get(
        f'/calendar/{event_id}/participants/import/template', headers=get_auth_headers()
    )
    assert response.status == 200
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    assert response.headers['content-disposition'].startswith('attachment; filename="Шаблон_участники_')
    # Базовые колонки + активное custom-поле с show_in_template; событие
    # (название/дата/уровень/спорт) в шаблоне НЕ нужно — оно берётся из события.
    assert get_xlsx_headers(response.body) == ['ФИО', 'Пол', 'Институт', 'Группа', 'Курс', 'Место', 'Дисциплина']
    assert len(get_xlsx_rows(response.body)) == 1


def test_event_import_upload_rejects_bad_files(event_import_client: SanicTestClient, monkeypatch):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    headers = get_auth_headers()

    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants/import', headers=headers, data=csrf_for(headers), allow_redirects=False
    )
    assert response.status == 302
    assert 'Выберите файл.' in unquote_plus(response.headers['location'])

    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Тестов'}], filename='participants.csv')
    assert response.status == 302
    assert 'Файл должен быть в формате .xlsx.' in unquote_plus(response.headers['location'])

    file_obj = BytesIO(b'this is not an excel file at all')
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants/import',
        headers=headers,
        data=csrf_for(headers),
        files={'file': ('broken.xlsx', file_obj.getvalue(), 'application/octet-stream')},
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Не удалось прочитать файл. Проверьте, что это не повреждённый .xlsx.' in unquote_plus(
        response.headers['location']
    )

    response = upload_event_xlsx(event_import_client, event_id, [{'Имя': 'Тестов', 'Пол': 'М'}])
    assert response.status == 302
    assert 'В файле нет обязательной колонки «ФИО». Скачайте шаблон и заполните его.' in unquote_plus(
        response.headers['location']
    )

    response = upload_event_xlsx(
        event_import_client, event_id, [], columns=['ФИО', 'Пол', 'Институт', 'Группа', 'Курс', 'Место']
    )
    assert response.status == 302
    assert 'В файле нет строк с данными.' in unquote_plus(response.headers['location'])

    monkeypatch.setattr('src.main.EVENT_PARTICIPANT_IMPORT_MAX_ROWS', 3)
    rows = [{'ФИО': f'Студентов Студент {index:02d}'} for index in range(4)]
    response = upload_event_xlsx(event_import_client, event_id, rows)
    assert response.status == 302
    assert 'В файле больше 3 строк. Разбейте список на части.' in unquote_plus(response.headers['location'])

    assert event_import_sessions == {}


def test_event_import_upload_stages_session_and_writes_nothing(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [{'ФИО': 'Иванов Иван Иванович', 'Комментарий': 'прим.'}],
    )
    assert response.status == 302
    assert response.headers['location'].startswith(f'/calendar/{event_id}/participants/import/preview/')
    token = import_session_token(response)

    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    assert 'Неизвестные колонки файла игнорируются: Комментарий.' in page.text
    assert 'Иванов Иван Иванович' in page.text
    # Загрузка ничего не создала: ни записей, ни аудита.
    assert event_records_snapshot(storage) == []
    assert audit_details(storage, 'event_participant_imported') == []

    # Одна сессия на пользователя: новая загрузка заменяет прежнюю.
    second = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Петров Пётр Петрович'}])
    second_token = import_session_token(second)
    assert second_token != token
    response = get_event_preview(event_import_client, event_id, token)
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(response.headers['location'])
    assert get_event_preview(event_import_client, event_id, second_token).status == 200

    # Сессия привязана к событию: чужое событие — то же поведение.
    other_event = make_calendar_event(storage, name='Другое')
    response = get_event_preview(event_import_client, other_event, second_token)
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло.' in unquote_plus(response.headers['location'])

    # Сессия привязана к владельцу: другой пользователь не видит её.
    response = get_event_preview(event_import_client, event_id, second_token, role='editor')
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло.' in unquote_plus(response.headers['location'])


# --- Разбор файла и матчинг кандидатов. ---


def test_event_import_parsing_full_set_lowercase_sex_and_place(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {
                'ФИО': 'Тестов Тест Тестович',
                'Пол': 'м',
                'Институт': 'ИСИ',
                'Группа': 'ПГС-101',
                'Курс': 2,
                'Место': 1,
            }
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    assert 'Тестов Тест Тестович' in page.text
    # Совпадений с карточками нет — но строка всё равно ГОТОВА: участие
    # добавится без связи, сопоставить можно позже.
    assert 'готовы к добавлению: 1' in page.text
    assert 'требуют решения: 0' in page.text
    assert 'Совпадений с активными карточками нет — можно сопоставить позже.' in page.text


def test_event_import_parsing_rejects_invalid_values(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Бесполый Гражданин', 'Пол': 'X', 'Курс': 1},
            {'ФИО': 'Некурсовый Гражданин', 'Курс': 'много'},
            {'ФИО': 'Неместный Гражданин', 'Курс': 1, 'Место': 'первое'},
            {'Курс': 1},
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'ошибочные: 4' in page.text
    assert 'Пол должен быть «М» или «Ж».' in page.text
    assert 'Поле &#34;Курс&#34; должно быть числом' in page.text
    assert 'Поле &#34;Место&#34; должно быть числом' in page.text
    assert 'Пустое ФИО.' in page.text


def test_event_import_matching_variants(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'ЭБ-241', '2')
    storage.create_student('Петров Пётр Петрович', 'М', 'ИСИ', 'ПГС-101', '1')
    storage.add_student_alias(2, 'Петя Петров')
    storage.create_student('Сидоров Сидор Сидорович', 'М', '', '', '')
    storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИМИ', 'СБ-202', '3')

    # Сопоставление с карточкой НЕ блокирует импорт: все шесть вариантов
    # (точное совпадение, псевдоним, тёзки ×2, 0 кандидатов ×2) — «готовы»;
    # «требуют решения» — только неразобранные дубли файла (см. отдельные
    # тесты дублей), поэтому строки-тёзки здесь различаются местом.
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович'},  # один кандидат → готова, авто-подбор
            {'ФИО': 'Петя Петров', 'Курс': 2},  # псевдоним → единственный, готова
            {'ФИО': 'Сидоров Сидор Сидорович', 'Курс': 1, 'Место': 1},  # тёзки → готова без подбора
            {'ФИО': 'сидоров сидор сидорович', 'Курс': 1, 'Место': 2},  # тот же casefold — те же тёзки
            {'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1},  # 0 кандидатов → готова без связи
            {'ФИО': 'Иванов Иван Иваныч', 'Курс': 1},  # опечатка — НЕ кандидат, готова без связи
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    # Все варианты совпадений — «готовы»: карточка необязательна.
    assert 'готовы к добавлению: 6' in page.text
    assert 'требуют решения: 0' in page.text
    # Единственный точный кандидат подбирается автоматически и виден.
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'найден автоматически · карточка #1' in collapsed  # Иванов
    assert 'найден автоматически · карточка #2' in collapsed  # Петя Петров (псевдоним)
    assert collapsed.count('найден автоматически') == 2
    assert 'выбрана вручную' not in collapsed
    # Тёзки: обе строки — по 2 карточки-кандидата с кнопкой (авто-подбора нет).
    assert 'Найдено нескольких подходящих студентов' in collapsed
    assert page.text.count('Выбрать карточку #') == 4
    # 0 кандидатов (Козлов и опечатка «Иваныч») — готовые строки без связи.
    assert page.text.count('Совпадений с активными карточками нет — можно сопоставить позже.') == 2


def test_event_import_autofill_priority_and_master_unchanged(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')

    # Excel: группа пуста (возьмётся G-101 из карточки), курс 2 (перекроет
    # карточку), место 1 (participation-specific, из Excel).
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [{'ФИО': 'Иванов Иван Иванович', 'Группа': '', 'Курс': 2, 'Место': 1}],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'из карточки #1' in page.text

    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert response.status == 302
    assert 'Добавлено участников из Excel: 1.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 1
    record = records[0]
    assert record[0] == 'Иванов Иван Иванович'
    assert record[1] == 'М'  # пол — из карточки (Excel пуст)
    assert record[2] == 'ИСЭиУ'  # институт — из карточки
    assert record[3] == 'G-101'  # группа — из карточки (Excel пуст)
    assert record[4] == 2  # курс — ЯВНЫЙ Excel перекрыл карточку
    assert record[10] == 1  # место — из Excel
    assert record[11] == student_id

    # Карточка-мастер не мутировалась: курс остался 1.
    master = storage.get_student_by_id(student_id)
    assert (master['sex'], master['institute'], master['group_name'], master['course']) == ('М', 'ИСЭиУ', 'G-101', '1')


def test_event_import_place_empty_awaits_result(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Место': ''}])
    token = import_session_token(response)
    # Bulk берёт строки, зафиксированные рендером предпросмотра.
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    # Пустое место → position = 0 — та же семантика «ждёт результата»,
    # что у ручного добавления участника события.
    assert records[0][10] == 0


def test_event_import_course_required_when_not_autofilled(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    # Карточка без курса: пусто и в файле, и в карточке → ошибка строки.
    storage.create_student('Безкурсов Студент Студентович', 'М', '', '', '')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Безкурсов Студент Студентович'}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'ошибочные: 1' in page.text
    assert 'Курс обязателен' in page.text


def test_event_import_custom_fields(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_custom_field('discipline', 'Дисциплина', 'text', False, True, True, True, 0)
    storage.create_custom_field('score', 'Очки', 'number', False, True, True, True, 1)
    storage.create_custom_field('req', 'Обязательное', 'text', True, True, True, True, 2)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')

    # Один Student, разные дисциплины — ДВЕ корректные записи; обязательное
    # custom-поле пусто и нечисловое числовое поле → ошибки валидации строк.
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Дисциплина': '100 м', 'Место': 1, 'Очки': 10},
            {'ФИО': 'Иванов Иван Иванович', 'Дисциплина': '200 м', 'Место': 3, 'Очки': 'шесть'},
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'ошибочные: 2' in page.text
    assert 'Поле &#34;Обязательное&#34; обязательно' in page.text
    # Числовое custom-поле — существующая семантика parse_custom_field_value
    # (сырой ValueError, как у импорта записей и ручного ввода).
    assert 'could not convert string to float' in page.text

    # Правка строк: заполняем обязательное поле — обе уходят в «готовы».
    for row_number in ('2', '3'):
        response = post_event_import_action(
            event_import_client,
            event_id,
            token,
            f'row/{row_number}/edit',
            {
                'full_name': 'Иванов Иван Иванович',
                'sex': 'М',
                'institute': 'ИСЭиУ',
                'group': 'G-101',
                'course': '1',
                'position': '1' if row_number == '2' else '3',
                'custom__discipline': '100 м' if row_number == '2' else '200 м',
                'custom__score': '10' if row_number == '2' else '6',
                'custom__req': 'да',
            },
        )
        assert response.status == 302
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 2' in page.text

    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 2.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 2
    assert all(record[11] == student_id for record in records)
    extra_values = [
        json.loads(row[0])
        for row in storage.connection.execute('SELECT extra_data FROM competitions ORDER BY id').fetchall()
    ]
    assert [values.get('discipline') for values in extra_values] == ['100 м', '200 м']
    assert [values.get('score') for values in extra_values] == ['10', '6']
    assert all(values.get('req') == 'да' for values in extra_values)


def test_event_import_duplicate_rows_warn_without_silent_drop(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},
            {'ФИО': 'иванов иван иванович', 'Курс': '1'},
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Полное равенство снимка (casefold) → обе строки требуют решения,
    # ничего не выбрасывается молча.
    assert 'требуют решения: 2' in page.text
    assert 'Одинаковые строки в файле: строки 2, 3' in page.text
    assert 'готовы к добавлению: 0' in page.text
    # Дубль не входит в bulk.
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 0.' in unquote_plus(response.headers['location'])
    assert event_records_snapshot(storage) == []


def test_event_import_existing_participation_is_informational(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    save_reconcile_record(storage, 'Иванов Иван Иванович', datetime(2026, 5, 10), position=5)
    # Существующая запись — участник того же события: P2 — связь по
    # calendar_event_id (пресет тоже совпадает, но состав участников
    # читается по ссылке).
    storage.connection.execute(
        "UPDATE competitions SET name = 'Забег 2026', date = '2026-05-10T00:00:00', "
        "date_to = '2026-05-11T00:00:00', "
        "sport = 'Бег', level = 'внутривузовские', student_ref_id = ?, calendar_event_id = ?",
        (student_id, event_id),
    )
    storage.connection.commit()
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Не блокировка: строка «готова», участие показано рядом.
    assert 'готовы к добавлению: 1' in page.text
    assert 'уже участвует: место 5' in page.text
    # Несколько участий одного Student — нормальный случай: bulk проходит.
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 1.' in unquote_plus(response.headers['location'])
    assert len(event_records_snapshot(storage)) == 2


def test_event_import_row_actions_select_add_skip(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')

    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'Найти студента' in page.text

    # Явный выбор карточки (поиск/POST select-student) → строка «готова»,
    # автозаполнение из неё.
    response = post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(student_id)}
    )
    assert 'Строка 2: выбрана карточка #%d — Иванов Иван Иванович.' % student_id in unquote_plus(
        response.headers['location']
    )
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'выбрана вручную' in page.text

    # Несуществующая/неактивная карточка — отказ без изменений сессии.
    response = post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': '99999'}
    )
    assert 'Студент не найден.' in unquote_plus(response.headers['location'])

    # Добавление одной строки: участие + student_ref_id + аудит.
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    assert 'Строка 2: участник добавлен.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] == student_id
    assert records[0][2] == 'ИСЭиУ'  # автозаполнение из выбранной карточки
    assert audit_details(storage, 'event_participant_imported') == [
        {'event_id': event_id, 'student_ref_id': student_id}
    ]
    page = get_event_preview(event_import_client, event_id, token)
    assert 'добавлен · карточка #%d' % student_id in re.sub(r'\s+', ' ', page.text)

    # Повторное действие над разобранной строкой — отказ.
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    assert 'Строка 2 уже разобрана.' in unquote_plus(response.headers['location'])
    assert len(event_records_snapshot(storage)) == 1


def test_event_import_row_skip_action(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1}])
    token = import_session_token(response)
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/skip')
    assert 'Строка 2 пропущена.' in unquote_plus(response.headers['location'])
    page = get_event_preview(event_import_client, event_id, token)
    assert 'пропущен' in page.text
    assert event_records_snapshot(storage) == []


def test_event_import_row_create_student_admin_only(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [{'ФИО': 'Ещё Один Студент', 'Пол': 'Ж', 'Институт': 'ИСИ', 'Группа': 'Т-100', 'Курс': 3, 'Место': 5}],
    )
    token = import_session_token(response)

    # Editor: создание карточек — admin-only (конвенция Phase 1/2.5).
    # Сессии предпросмотра привязаны к владельцу — у editor своя загрузка.
    editor_response = upload_event_xlsx(
        event_import_client,
        event_id,
        [{'ФИО': 'Ещё Один Студент', 'Пол': 'Ж', 'Институт': 'ИСИ', 'Группа': 'Т-100', 'Курс': 3, 'Место': 5}],
        role='editor',
    )
    assert editor_response.status == 302
    assert editor_response.headers['location'].startswith(f'/calendar/{event_id}/participants/import/preview/')
    editor_token = import_session_token(editor_response)
    response = post_event_import_action(
        event_import_client,
        event_id,
        editor_token,
        'row/2/create-student',
        {'full_name': 'Ещё Один Студент', 'sex': 'Ж', 'institute': 'ИСИ', 'group': 'Т-100', 'course': '3'},
        role='editor',
    )
    assert response.status == 403
    page = get_event_preview(event_import_client, event_id, editor_token, role='editor')
    assert page.status == 200
    # Строка валидна и без карточки (готова), но создание карточки — admin-only.
    assert 'готовы к добавлению: 1' in page.text
    assert 'Создать нового Student' not in page.text
    assert 'Создать карточку и добавить участника' not in page.text
    assert 'href="/admin/people/' not in page.text
    post_event_import_action(event_import_client, event_id, editor_token, 'discard', role='editor')

    # Admin: одна операция «карточка + связанный участник».
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/create-student',
        {'full_name': 'Ещё Один Студент', 'sex': 'Ж', 'institute': 'ИСИ', 'group': 'Т-100', 'course': '3'},
    )
    assert 'студент «Ещё Один Студент» создан, участник добавлен.' in unquote_plus(response.headers['location'])
    new_student = storage.connection.execute("SELECT id FROM students WHERE full_name = 'Ещё Один Студент'").fetchone()[
        'id'
    ]
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] == new_student
    assert records[0][10] == 5  # место — из Excel
    assert {'student_id': new_student, 'full_name': 'Ещё Один Студент', 'source': 'event-participant-import'} in (
        audit_details(storage, 'student_created')
    )


def test_event_import_bulk_commit_rechecks_candidates(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text

    # Между рендером и кликом появился полный тёзка — состав «готовых»
    # изменился: весь батч отменяется, ничего не вставляется.
    storage.create_student('Иванов Иван Иванович', 'Ж', 'ИМИ', 'СБ-202', '3')
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Список карточек изменился — обновите страницу и проверьте строки.' in unquote_plus(
        response.headers['location']
    )
    assert event_records_snapshot(storage) == []


def test_event_import_bulk_commit_rollback_on_failure(event_import_client: SanicTestClient, monkeypatch):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    storage.create_student('Петров Пётр Петрович', 'Ж', 'ИСИ', 'ПГС-101', '2')
    # Смешанный батч: со связью, без карточки (0 кандидатов) и строка с
    # ошибкой валидации (обязательное custom-поле пусто) — в батч НЕ входит.
    storage.create_custom_field('req', 'Обязательное', 'text', True, True, True, True, 0)
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1, 'Обязательное': 'есть'},
            {'ФИО': 'Петров Пётр Петрович', 'Курс': 2, 'Обязательное': 'есть'},
            {'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1, 'Обязательное': 'есть'},
            {'ФИО': 'Безполеный Гражданин', 'Курс': 1, 'Обязательное': ''},
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 3' in page.text
    assert 'ошибочные: 1' in page.text

    def failing_import(*args, **kwargs):
        raise sqlite3.OperationalError('simulated failure')

    monkeypatch.setattr(SQLiteAdapter, 'import_competitions', failing_import)
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert response.status == 500
    monkeypatch.undo()
    # Ни записи, ни справочники не изменились — откат всего батча.
    assert event_records_snapshot(storage) == []
    # Ошибочная строка осталась неразобранной (не потеряна молча).
    page = get_event_preview(event_import_client, event_id, token)
    assert 'ошибочные: 1' in page.text
    assert 'готовы к добавлению: 3' in page.text


def test_event_import_finish_and_discard(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},
            {'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1},
        ],
    )
    token = import_session_token(response)

    # Обе строки «готовы» (без карточки — тоже), но завершение требует
    # разобрать их явно: у Козлова нет кандидатов — строку пропускаем ДО
    # bulk, чтобы батч зафиксировал только строку Иванова.
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 2' in page.text
    response = post_event_import_action(event_import_client, event_id, token, 'finish')
    assert 'Завершение недоступно: не разобрано строк: 2.' in unquote_plus(response.headers['location'])
    response = post_event_import_action(event_import_client, event_id, token, 'row/3/skip')
    assert 'Строка 3 пропущена.' in unquote_plus(response.headers['location'])

    # Bulk берёт строки, зафиксированные рендером предпросмотра.
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    page = get_event_preview(event_import_client, event_id, token)
    assert 'Кнопка «Завершить импорт» станет активной' not in page.text
    assert 'disabled>Завершить импорт' not in page.text

    response = post_event_import_action(event_import_client, event_id, token, 'finish')
    assert response.status == 302
    assert response.headers['location'].startswith(f'/calendar/{event_id}')
    assert 'Импорт завершён: добавлено 1, пропущено 1.' in unquote_plus(response.headers['location'])
    completed = audit_details(storage, 'event_participants_import_completed')
    assert completed == [
        {
            'event_id': event_id,
            'event_name': 'Забег 2026',
            'counters': {'added': 1, 'skipped': 1, 'errors': 0, 'total': 2},
        }
    ]
    assert token not in event_import_sessions

    # Discard: сессия удаляется, добавленные участники остаются.
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    response = post_event_import_action(event_import_client, event_id, token, 'discard')
    assert 'Импорт отменён. Добавленные участники (1) сохранены.' in unquote_plus(response.headers['location'])
    assert token not in event_import_sessions
    assert len(event_records_snapshot(storage)) == 2


def test_event_import_relations_event_preset(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(
        storage,
        name='Кросс весны',
        date='2026-03-01T00:00:00',
        date_to=None,
        level='межвузовские',
        sport='Лёгкая атлетика',
    )
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Место': 2}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    record = records[0]
    # Пресет события скопирован в запись (однодневное событие: date_to NULL).
    assert record[5] == 'Лёгкая атлетика'
    assert record[6] == '2026-03-01T00:00:00'
    assert record[7] is None
    assert record[8] == 'межвузовские'
    assert record[9] == 'Кросс весны'
    # P2: bulk-импорт пишет явную ссылку на событие.
    assert (
        storage.connection.execute('SELECT calendar_event_id FROM competitions').fetchone()['calendar_event_id']
        == event_id
    )
    # Участник появился в списке события, счётчик обновился.
    participants = storage.list_calendar_event_participants(event_id)
    assert [participant['student_name'] for participant in participants] == ['Иванов Иван Иванович']


def test_event_import_post_routes_require_csrf(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    suffixes = (
        'bulk-commit',
        'finish',
        'discard',
        'row/2/skip',
        'row/2/add',
        'row/2/edit',
        'row/2/select-student',
        'row/2/clear-student',
        'row/2/create-student',
    )
    for suffix in suffixes:
        _, response = event_import_client.post(
            f'/calendar/{event_id}/participants/import/preview/{token}/{suffix}',
            headers=get_auth_headers(),
            data={},
            allow_redirects=False,
        )
        assert response.status == 403, suffix
    # И сама загрузка файла тоже.
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants/import', headers=get_auth_headers(), data={}, allow_redirects=False
    )
    assert response.status == 403
    assert event_records_snapshot(storage) == []


def test_event_import_unknown_token_redirects_to_upload(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = get_event_preview(event_import_client, event_id, 'no-such-token')
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло. Загрузите файл заново.' in unquote_plus(response.headers['location'])


def test_event_import_expired_session_swept(event_import_client: SanicTestClient, monkeypatch):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    # TTL истёк: сессия выметается при следующем обращении.
    event_import_sessions[token]['created_at'] -= EVENT_IMPORT_SESSION_TTL_SECONDS + 1
    response = get_event_preview(event_import_client, event_id, token)
    assert response.status == 302
    assert 'Время сессии предпросмотра истекло.' in unquote_plus(response.headers['location'])
    assert token not in event_import_sessions


# --- Ремонтный цикл 1 (QA FAIL): квота легаси, явный выбор тёзки,
# достижимые дубли, XSS превью. ---


def test_athletes_search_legacy_not_displaced_by_full_student_limit(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    # Полный лимит совпадающих карточек (8 из 10) + легаси-человек без карточки.
    for index in range(10):
        storage.create_student(f'Сидоров Сидор {index:02d}', 'М', '', '', '')
    save_reconcile_record(storage, 'Сидоркин Сидор Сидорович', datetime(2026, 1, 10))
    _, response = event_import_client.get('/api/athletes/search?q=Сид', headers=get_auth_headers())
    assert response.status == 200
    kinds = [item['kind'] for item in response.json]
    # Легаси НЕ вытесняется: у каждого вида своя квота (до 8), карточки —
    # раньше; иначе на главной (клиент фильтрует карточки) список пустел бы.
    assert kinds.count('student') == 8
    assert kinds.count('legacy') == 1
    assert kinds[0] == 'student'
    assert kinds[-1] == 'legacy'
    assert response.json[-1]['name'] == 'Сидоркин Сидор Сидорович'


def test_event_import_selected_namesake_becomes_ready_and_bulk(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИСИ', 'ПГС-201', '1')
    second = storage.create_student('Сидоров Сидор Сидорович', 'Ж', 'ИМИ', 'СБ-202', '2')

    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Сидоров Сидор Сидорович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Тёзки больше НЕ блокируют: строка «готова» БЕЗ авто-подбора (явного
    # решения нет) — участие уйдёт без связи, карточки показаны списком.
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert 'требуют решения: 0' in page.text
    assert 'Добавить 1 готовых' in page.text
    assert 'Найдено нескольких подходящих студентов' in collapsed
    assert 'найден автоматически' not in collapsed
    # Confirm-текст предупреждает о строках без карточки.
    assert 'Добавить 1 готовых участников (без карточки — 1)?' in collapsed

    # Явный выбор карточки → строка по-прежнему «готова», но уже со связью.
    response = post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(second)}
    )
    assert 'Строка 2: выбрана карточка #%d — Сидоров Сидор Сидорович.' % second in unquote_plus(
        response.headers['location']
    )
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Добавить 1 готовых' in page.text
    assert f'выбрана вручную · карточка #{second}' in collapsed
    assert '(без карточки —' not in collapsed

    # Bulk вставляет ровно готовую строку — с выбранной карточкой.
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 1.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] == second
    assert records[0][2] == 'ИМИ'  # автозаполнение из ВЫБРАННОЙ карточки тёзки


def test_event_import_bulk_counter_matches_inserted_rows(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИСИ', 'ПГС-201', '1')
    storage.create_student('Сидоров Сидор Сидорович', 'Ж', 'ИМИ', 'СБ-202', '2')

    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},  # единственный кандидат → готова, со связью
            {'ФИО': 'Сидоров Сидор Сидорович', 'Курс': 1},  # тёзки без выбора → готова, ref NULL
            {'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1},  # 0 кандидатов → готова, ref NULL
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Смешанный батч: карточка необязательна — все три строки «готовы»,
    # счётчик кнопки равен числу вставляемых строк.
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 3' in page.text
    assert 'требуют решения: 0' in page.text
    assert 'Добавить 3 готовых' in page.text
    assert 'Добавить 3 готовых участников (без карточки — 2)?' in collapsed

    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 3.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 3
    assert records[0][0] == 'Иванов Иван Иванович'
    assert records[0][11] == student_id
    # Тёзки без выбора и строка без кандидатов — без связи (сопоставить позже).
    assert records[1][0] == 'Сидоров Сидор Сидорович'
    assert records[1][11] is None
    assert records[2][0] == 'Козлов Козьма Козьмич'
    assert records[2][11] is None
    # Неразобранных строк не осталось: ничего не потеряно молча.
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 0' in page.text
    assert 'требуют решения: 0' in page.text


def test_event_import_duplicate_addable_after_sibling_resolved(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')

    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},
            {'ФИО': 'иванов иван иванович', 'Курс': '1'},
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Оба дубля не разобраны → требуют решения (создать обе нельзя вслепую).
    assert 'требуют решения: 2' in page.text
    assert 'Одинаковые строки в файле: строки 2, 3' in page.text

    # Одну пропустили → вторая разблокирована и добавляема (§17:
    # «оставить одну» достижима из UI).
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/skip')
    assert 'Строка 2 пропущена.' in unquote_plus(response.headers['location'])
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Добавить 1 готовых' in page.text
    # Предупреждение о дубле остаётся информацией (строки 2, 3).
    assert 'Одинаковые строки в файле: строки 2, 3' in page.text

    response = post_event_import_action(event_import_client, event_id, token, 'row/3/add')
    assert 'Строка 3: участник добавлен.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] == student_id

    # «Создать обе»: свежая сессия — первую добавляем, вторая после этого
    # разблокируется и тоже добавляется.
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {'ФИО': 'Иванов Иван Иванович', 'Курс': 1},
            {'ФИО': 'иванов иван иванович', 'Курс': '1'},
        ],
    )
    token = import_session_token(response)
    post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    response = post_event_import_action(event_import_client, event_id, token, 'row/3/add')
    assert 'Строка 3: участник добавлен.' in unquote_plus(response.headers['location'])
    assert len(event_records_snapshot(storage)) == 3


def test_event_import_preview_xss_rendered_as_text(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    payload = '<script>alert("xss")</script>'
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [{'ФИО': payload, 'Институт': '<img src=x onerror=alert(1)>', 'Курс': 1}],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert page.status == 200
    # Значения файла рендерятся как текст: сырой payload в разметке отсутствует.
    assert payload not in page.text
    assert '<img src=x onerror=' not in page.text
    assert '&lt;script&gt;alert(' in page.text
    assert '&lt;/script&gt;' in page.text
    assert '&lt;img src=x onerror=alert(1)&gt;' in page.text

    # Flash с ФИО из файла (создание карточки) отображается как текст.
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/create-student',
        {'full_name': payload, 'sex': '', 'institute': '', 'group': '', 'course': '1'},
    )
    assert response.status == 302
    _, flashed = event_import_client.get(response.headers['location'], headers=get_auth_headers())
    assert flashed.status == 200
    assert payload not in flashed.text
    assert '&lt;script&gt;alert(' in flashed.text


# --- Event Participant Import UX v2: курс опционален (field_settings),
# url-поля вне импорта участников, «Найти студента» вместо сырого № карточки. ---


def test_event_import_course_optional_accepts_name_only_row(event_import_client: SanicTestClient):
    """Курс опционален (настройка field_settings): строка только с ФИО
    валидна и без курса в карточке Student, bulk проходит. Дефолт
    (required) уже покрыт test_event_import_course_required_when_not_autofilled."""
    storage = app.ctx.storage
    storage.update_field_settings({'course': ('number', False)})
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович'}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Курс обязателен' not in page.text
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][4] == 0  # пустой опциональный курс → 0 (как Место)
    assert records[0][11] == student_id


def test_event_import_course_optional_autofills_from_student(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    storage.update_field_settings({'course': ('number', False)})
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '3')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович'}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    # Курс подтянут из карточки (Excel пуст) — автозаполнение не зависит от
    # обязательности поля.
    assert f'поля из карточки #{student_id}: Пол, Институт, Группа, Курс' in re.sub(r'\s+', ' ', page.text)
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert event_records_snapshot(storage)[0][4] == 3


def test_event_import_course_optional_excel_overrides_student(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    storage.update_field_settings({'course': ('number', False)})
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '3')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 5}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    # Явный Excel-курс не входит в автозаполненные из карточки поля.
    assert f'поля из карточки #{student_id}: Пол, Институт, Группа, Курс' not in re.sub(r'\s+', ' ', page.text)
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert records[0][4] == 5  # ЯВНЫЙ Excel-курс перекрыл карточку


# --- Url-поля вне импорта участников события (UX v2). ---


def test_event_import_excludes_url_custom_field(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    storage.create_custom_field('discipline', 'Дисциплина', 'text', False, True, True, True, 0)
    storage.create_custom_field('link', 'Ссылка на соревнование', 'url', False, True, True, True, 1)

    # Шаблон импорта участников: обычное custom-поле есть, url-поле — нет.
    _, response = event_import_client.get(
        f'/calendar/{event_id}/participants/import/template', headers=get_auth_headers()
    )
    assert response.status == 200
    header = get_xlsx_headers(response.body)
    assert 'Дисциплина' in header
    assert 'Ссылка на соревнование' not in header

    # Файл с колонкой url-поля: колонка становится «неизвестной» → штатное
    # предупреждение; строка валидна, значение не попадает ни в предпросмотр,
    # ни в extra_data участия.
    response = upload_event_xlsx(
        event_import_client,
        event_id,
        [
            {
                'ФИО': 'Иванов Иван Иванович',
                'Курс': 1,
                'Дисциплина': '100 м',
                'Ссылка на соревнование': 'https://sport.example.com/race',
            }
        ],
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Неизвестные колонки файла игнорируются: Ссылка на соревнование.' in page.text
    assert 'https://sport.example.com/race' not in page.text
    assert '<th>Ссылка на соревнование</th>' not in page.text
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    extra = json.loads(storage.connection.execute('SELECT extra_data FROM competitions').fetchone()[0])
    assert extra == {'discipline': '100 м'}  # ключа url-поля нет

    # Само url-поле остаётся в системе: список custom-полей админки.
    _, fields_page = event_import_client.get('/admin/fields', headers=get_auth_headers())
    assert fields_page.status == 200
    assert 'Ссылка на соревнование' in fields_page.text


def test_event_import_required_url_field_does_not_block_rows(event_import_client: SanicTestClient):
    """КРИТИЧНО: фильтр доходит до валидации — ОБЯЗАТЕЛЬНОЕ url-поле не
    блокирует строки импорта участников (в обычном реестре оно осталось бы
    обязательным)."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_custom_field('link', 'Ссылка на соревнование', 'url', True, True, True, True, 0)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Поле &#34;Ссылка на соревнование&#34; обязательно' not in page.text
    response = post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    assert 'Добавлено участников из Excel: 1.' in unquote_plus(response.headers['location'])
    extra = json.loads(storage.connection.execute('SELECT extra_data FROM competitions').fetchone()[0])
    assert extra == {}


def test_registry_import_and_export_keep_url_custom_field(client: SanicTestClient):
    """Url-поля исключены ТОЛЬКО из импорта участников события: обычный
    импорт записей валидирует и хранит их значения как раньше."""
    url_field = CustomField(
        field_id=1,
        key='link',
        label='Ссылка на соревнование',
        field_type='url',
        required=False,
        show_in_table=True,
        show_in_export=True,
        show_in_template=True,
    )
    app.ctx.storage.get_custom_fields.return_value = [url_field]
    app.ctx.storage.import_competitions.reset_mock()
    app.ctx.storage.get_competitions.return_value = []
    response = upload_xlsx(client, [make_import_row(**{'Ссылка на соревнование': 'https://sport.example.com/race'})])
    assert response.status == 200
    saved = app.ctx.storage.import_competitions.call_args[0][0][0]
    assert saved.extra_data == {'link': 'https://sport.example.com/race'}

    # Читатели extra_data используют .get: участия из импорта участников
    # события ключа url-поля не имеют — экспорт отдаёт пустое значение, а не
    # падает.
    competition = build_competition(make_import_row(), custom_fields=[])
    row = competition_to_export_row(competition, [url_field])
    assert row['Ссылка на соревнование'] == ''

    app.ctx.storage.get_custom_fields.return_value = []
    app.ctx.storage.get_competitions.return_value = []


# --- «Найти студента»: поиск карточки для строки предпросмотра (UX v2). ---


def get_event_search(
    client: SanicTestClient,
    event_id: int,
    token: str,
    row_number: int,
    query: str,
    *,
    role: str = 'admin',
):
    """GET предпросмотра с параметрами поиска find/q."""
    _, response = client.get(
        f'/calendar/{event_id}/participants/import/preview/{token}?{urlencode({"find": row_number, "q": query})}',
        headers=get_auth_headers(role=role),
        allow_redirects=False,
    )
    return response


def test_event_import_row_search_select_writes_student_ref(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    # Кнопка поиска есть, ручного ввода сырого номера карточки больше нет.
    assert 'Найти студента' in page.text
    assert '№ карточки' not in page.text
    assert 'Выбрать существующего' not in page.text
    assert 'type="number"' not in page.text

    # Подстрока по ФИО (casefold); collapse открыт для найденной строки.
    page = get_event_search(event_import_client, event_id, token, 2, 'иван')
    assert page.status == 200
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'По запросу «иван» найдено: 1.' in collapsed
    assert 'collapse show" id="import-row-2-search"' in collapsed
    assert f'<a href="/admin/people/{student_id}">Иванов Иван Иванович</a>' in collapsed
    assert f'#{student_id} · ИСЭиУ · G-101 · курс 1' in collapsed
    assert f'Выбрать карточку #{student_id}' in collapsed
    # Форма выбора результата — POST select-student с csrf (переиспользованный
    # макрос); бейдж «ФИО» у результатов поиска не выводится.
    assert '/row/2/select-student' in page.text
    assert '>ФИО</span>' not in collapsed

    # Регистр запроса не важен (casefold).
    page = get_event_search(event_import_client, event_id, token, 2, 'ИВАНОВ иВАН')
    assert 'По запросу «ИВАНОВ иВАН» найдено: 1.' in re.sub(r'\s+', ' ', page.text)

    # Выбор результата поиска: строка «готова», участие получает student_ref_id.
    response = post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(student_id)}
    )
    assert 'Строка 2: выбрана карточка #%d — Иванов Иван Иванович.' % student_id in unquote_plus(
        response.headers['location']
    )
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] == student_id


def test_event_import_row_search_by_alias_and_namesakes(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    alias_id = storage.create_student('Николаев Николай Николаевич', 'М', 'ИСИ', 'ПГС-201', '1')
    storage.add_student_alias(alias_id, 'Коля')
    first = storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИСИ', 'ПГС-201', '1')
    second = storage.create_student('Сидоров Сидор Сидорович', 'Ж', 'ИМИ', 'СБ-202', '2')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)

    # Подстрока псевдонима: бейдж «псевдоним» с исходным написанием в title.
    page = get_event_search(event_import_client, event_id, token, 2, 'коля')
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'По запросу «коля» найдено: 1.' in collapsed
    assert f'Выбрать карточку #{alias_id}' in collapsed
    assert 'title="Совпадение по псевдониму «Коля»">псевдоним</span>' in collapsed

    # Полные тёзки — два отдельных результата с разными карточками.
    page = get_event_search(event_import_client, event_id, token, 2, 'сидоров')
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'По запросу «сидоров» найдено: 2.' in collapsed
    assert f'#{first} · ИСИ · ПГС-201 · курс 1' in collapsed
    assert f'#{second} · ИМИ · СБ-202 · курс 2' in collapsed


def test_event_import_row_search_inactive_and_empty_result(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    storage.set_student_active(student_id, False)
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)
    # Неактивная карточка не находится — пустой результат.
    page = get_event_search(event_import_client, event_id, token, 2, 'иван')
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'По запросу «иван» ничего не найдено.' in collapsed
    assert 'По запросу «иван» найдено' not in collapsed
    assert f'Выбрать карточку #{student_id}' not in collapsed


def test_event_import_row_search_limit_and_truncation(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    for index in range(9):
        storage.create_student(f'Сидоров Сидор {index:02d}', 'М', 'ИСИ', 'ПГС-201', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)
    page = get_event_search(event_import_client, event_id, token, 2, 'сидоров')
    collapsed = re.sub(r'\s+', ' ', page.text)
    # Лимит 8 + признак усечения; 9-я карточка не показана.
    assert 'По запросу «сидоров» найдено: 8.' in collapsed
    assert 'Показаны первые 8 — уточните запрос.' in collapsed
    assert 'Сидоров Сидор 08' not in collapsed
    assert 'Сидоров Сидор 07' in collapsed


def test_event_import_row_search_ignores_invalid_params(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)

    # Пустой/пробельный q — не поиск: collapse доступен, блока результатов нет.
    for query in ('', '   '):
        page = get_event_search(event_import_client, event_id, token, 2, query)
        assert page.status == 200
        assert 'По запросу' not in page.text

    # Несуществующая строка — параметр find игнорируется целиком.
    page = get_event_search(event_import_client, event_id, token, 999, 'иван')
    assert page.status == 200
    assert 'По запросу' not in page.text
    assert 'collapse show' not in page.text
    # Не-число в find — то же молчаливое игнорирование (не 400).
    _, page = event_import_client.get(
        f'/calendar/{event_id}/participants/import/preview/{token}?find=abc&q=%D0%B8%D0%B2%D0%B0%D0%BD',
        headers=get_auth_headers(),
        allow_redirects=False,
    )
    assert page.status == 200
    assert 'По запросу' not in page.text

    # Разобранная (skipped) строка больше не ищется.
    post_event_import_action(event_import_client, event_id, token, 'row/2/skip')
    page = get_event_search(event_import_client, event_id, token, 2, 'иван')
    assert page.status == 200
    assert 'По запросу' not in page.text


def test_event_import_row_search_available_to_editor(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(
        event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}], role='editor'
    )
    token = import_session_token(response)
    page = get_event_search(event_import_client, event_id, token, 2, 'иван', role='editor')
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'По запросу «иван» найдено: 1.' in collapsed
    # Editor: ФИО результатом — текстом (без ссылки на карточку), выбор есть,
    # создание карточки по-прежнему только admin (подсказка больше не нужна —
    # строка и без карточки готова).
    assert f'<a href="/admin/people/{student_id}">' not in collapsed
    assert f'Выбрать карточку #{student_id}' in collapsed
    assert 'Создать нового Student' not in collapsed
    assert 'Создать карточку и добавить участника' not in collapsed
    assert 'Создать карточку студента может только администратор.' not in collapsed


def test_event_import_edit_full_name_resets_explicit_selection(event_import_client: SanicTestClient):
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '1')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Новый Человек Человечович', 'Курс': 2}])
    token = import_session_token(response)
    post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(student_id)}
    )
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text

    # Правка БЕЗ изменения ФИО (другой курс) — явный выбор сохраняется.
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/edit',
        {
            'full_name': 'Новый Человек Человечович',
            'sex': 'М',
            'institute': '',
            'group': '',
            'course': '3',
            'position': '',
        },
    )
    assert response.status == 302
    page = get_event_preview(event_import_client, event_id, token)
    assert 'готовы к добавлению: 1' in page.text
    assert event_import_sessions[token]['rows'][0]['student_id'] == student_id

    # Правка с изменением ФИО (регистр не в счёт — casefold) — не сбрасывает.
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/edit',
        {
            'full_name': 'Новый Человек Человечович'.upper(),
            'sex': 'М',
            'institute': '',
            'group': '',
            'course': '3',
            'position': '',
        },
    )
    assert response.status == 302
    assert event_import_sessions[token]['rows'][0]['student_id'] == student_id

    # Реальное изменение ФИО — сброшены ОБА решения (явный выбор и «без
    # карточки»): строка остаётся «готовой» (0 кандидатов → без связи),
    # устаревшая связь не переносится молча.
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/edit',
        {
            'full_name': 'Другой Человек Человечович',
            'sex': 'М',
            'institute': '',
            'group': '',
            'course': '3',
            'position': '',
        },
    )
    assert response.status == 302
    assert event_import_sessions[token]['rows'][0]['student_id'] is None
    assert event_import_sessions[token]['rows'][0]['student_forced_none'] is False
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'требуют решения: 0' in collapsed
    assert 'готовы к добавлению: 1' in collapsed
    assert 'выбрана вручную' not in collapsed
    assert 'После правки ФИО кандидаты пересчитаются' in collapsed

    # После сброса повторный поиск и выбор работают.
    page = get_event_search(event_import_client, event_id, token, 2, 'иван')
    assert 'По запросу «иван» найдено: 1.' in re.sub(r'\s+', ' ', page.text)
    post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(student_id)}
    )
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    assert 'Строка 2: участник добавлен.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert records[0][11] == student_id


# --- Импорт участников: карточка студента НЕ обязательна. Авто-подбор
# единственного точного кандидата, «Без карточки», сопоставление позже
# существующим механизмом, ручная форма по настройкам полей. ---


def test_event_import_auto_preselect_single_candidate(event_import_client: SanicTestClient):
    """Единственный точный кандидат подбирается автоматически: строка «готова»,
    карточка видна первой строкой ячейки («найден автоматически»), пустые поля
    файла заполняются из неё; явная смена карточки переключает источник."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    first = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '2')
    second = storage.create_student('Иванов Иван Иванович Другой', 'Ж', 'ИМИ', 'СБ-202', '3')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович'}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert f'найден автоматически · карточка #{first}' in collapsed
    assert f'<a href="/admin/people/{first}">Иванов Иван Иванович</a>' in collapsed
    assert f'#{first} · ИСЭиУ · G-101 · курс 2' in collapsed
    assert 'выбрана вручную' not in collapsed
    # Автозаполнение пустых полей файла из подобранной карточки.
    assert f'поля из карточки #{first}: Пол, Институт, Группа, Курс' in collapsed

    # Явная смена карточки: происхождение — «выбрана вручную», автозаполнение
    # пересчиталось на новую карточку.
    response = post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(second)}
    )
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert f'выбрана вручную · карточка #{second}' in collapsed
    assert f'поля из карточки #{second}: Пол, Институт, Группа, Курс' in collapsed
    assert f'поля из карточки #{first}' not in collapsed


def test_event_import_namesakes_ready_without_preselect(event_import_client: SanicTestClient):
    """Полные тёзки (≥2): строка «готова» БЕЗ авто-подбора — по умолчанию
    участие уходит без связи (ref NULL); явный выбор работает."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    first = storage.create_student('Сидоров Сидор Сидорович', 'М', 'ИСИ', 'ПГС-201', '1')
    storage.create_student('Сидоров Сидор Сидорович', 'Ж', 'ИМИ', 'СБ-202', '2')
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Сидоров Сидор Сидорович', 'Курс': 1}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert 'найден автоматически' not in collapsed
    assert page.text.count('Выбрать карточку #') == 2
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] is None  # без явного выбора — без связи

    # Явный выбор тёзки — участие со связью выбранной карточки (рендер
    # предпросмотра между выбором и bulk фиксирует состав «готовых»).
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Сидоров Сидор Сидорович', 'Курс': 1}])
    token = import_session_token(response)
    post_event_import_action(event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(first)})
    page = get_event_preview(event_import_client, event_id, token)
    assert f'выбрана вручную · карточка #{first}' in re.sub(r'\s+', ' ', page.text)
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 2
    assert records[1][11] == first
    assert records[1][2] == 'ИСИ'  # автозаполнение из выбранной карточки


def test_event_import_no_candidate_bulk_and_reconcile_later(event_import_client: SanicTestClient):
    """0 кандидатов: строка «готова», bulk добавляет ref NULL; запись видна в
    сопоставлении (/admin/people/reconcile) с кандидатом по ФИО и привязывается
    СУЩЕСТВУЮЩИМ механизмом — новой сущности не появилось."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(
        event_import_client, event_id, [{'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1, 'Место': 2}]
    )
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Совпадений с активными карточками нет — можно сопоставить позже.' in collapsed
    assert 'Добавить 1 готовых участников (без карточки — 1)?' in collapsed
    post_event_import_action(event_import_client, event_id, token, 'bulk-commit')
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] is None
    # Аудит допускает NULL-связь.
    assert audit_details(storage, 'event_participant_imported') == [{'event_id': event_id, 'student_ref_id': None}]

    # Карточку создали ПОСЛЕ импорта: запись без связи видна в сопоставлении
    # (list_unlinked_competitions), кандидат находится по ФИО.
    student_id = storage.create_student('Козлов Козьма Козьмич', 'М', 'ИСИ', 'ПГС-101', '1')
    record_id = storage.connection.execute(
        "SELECT id FROM competitions WHERE student_name = 'Козлов Козьма Козьмич'"
    ).fetchone()['id']
    unlinked = storage.list_unlinked_competitions(search='Козлов')
    assert [record['id'] for record in unlinked] == [record_id]
    assert [candidate['student_id'] for candidate in storage.find_student_candidates('Козлов Козьма Козьмич')] == [
        student_id
    ]

    # Привязка — существующий роут сопоставления (admin, CSRF).
    headers = get_auth_headers()
    _, response = event_import_client.post(
        '/admin/people/reconcile/link',
        headers=headers,
        data={
            **csrf_for(headers),
            'record_ids[]': str(record_id),
            'student_id': str(student_id),
            'back': '/admin/people/reconcile',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    linked = storage.connection.execute('SELECT student_ref_id FROM competitions WHERE id = ?', (record_id,)).fetchone()
    assert linked[0] == student_id


def test_event_import_clear_student_flow(event_import_client: SanicTestClient):
    """clear-student: сброс автоподобранной карточки — участие без связи;
    решение переживает GET, доступно модераторам и сбрасывается правкой ФИО."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    student_id = storage.create_student('Иванов Иван Иванович', 'М', 'ИСЭиУ', 'G-101', '2')
    # Курс задан файлом: без автозаполнения из карточки строка остаётся валидной.
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 2}])
    token = import_session_token(response)
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert f'найден автоматически · карточка #{student_id}' in collapsed
    assert 'Без карточки' in collapsed  # кнопка сброса у подобранной карточки

    response = post_event_import_action(event_import_client, event_id, token, 'row/2/clear-student')
    assert 'Строка 2: карточка сброшена — участие будет добавлено без связи.' in unquote_plus(
        response.headers['location']
    )
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    # Автоподбор и автозаполнение подавлены; строка остаётся «готовой».
    assert 'готовы к добавлению: 1' in page.text
    assert 'Без карточки — участие будет без связи, сопоставить можно позже.' in collapsed
    assert 'найден автоматически' not in collapsed
    assert f'поля из карточки #{student_id}' not in collapsed
    row = event_import_sessions[token]['rows'][0]
    assert (row['student_id'], row['student_forced_none']) == (None, True)

    # Решение переживает повторный GET предпросмотра (живёт в сессии).
    page = get_event_preview(event_import_client, event_id, token)
    assert 'Без карточки — участие будет без связи, сопоставить можно позже.' in re.sub(r'\s+', ' ', page.text)

    # Модератор (editor) тоже может сбросить карточку строки.
    editor_response = upload_event_xlsx(
        event_import_client, event_id, [{'ФИО': 'Иванов Иван Иванович', 'Курс': 2}], role='editor'
    )
    editor_token = import_session_token(editor_response)
    response = post_event_import_action(
        event_import_client, event_id, editor_token, 'row/2/clear-student', role='editor'
    )
    assert response.status == 302
    post_event_import_action(event_import_client, event_id, editor_token, 'discard', role='editor')

    # Viewer не может сбросить карточку (модераторская операция).
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants/import/preview/{token}/row/2/clear-student',
        headers=get_auth_headers(role='viewer'),
        data={},
        allow_redirects=False,
    )
    assert response.status == 403

    # Правка ФИО сбрасывает и решение «без карточки»: свежий подбор для
    # нового ФИО (кандидатов нет — строка готова без связи).
    response = post_event_import_action(
        event_import_client,
        event_id,
        token,
        'row/2/edit',
        {
            'full_name': 'Другой Человек Человечович',
            'sex': 'М',
            'institute': '',
            'group': '',
            'course': '3',
            'position': '',
        },
    )
    assert response.status == 302
    row = event_import_sessions[token]['rows'][0]
    assert (row['student_id'], row['student_forced_none']) == (None, False)
    page = get_event_preview(event_import_client, event_id, token)
    collapsed = re.sub(r'\s+', ' ', page.text)
    assert 'готовы к добавлению: 1' in page.text
    assert 'Без карточки — участие будет без связи' not in collapsed
    assert 'Совпадений с активными карточками нет — можно сопоставить позже.' in collapsed

    # Повторный выбор карточки снимает «без карточки» (select-student).
    post_event_import_action(
        event_import_client, event_id, token, 'row/2/select-student', {'student_id': str(student_id)}
    )
    row = event_import_sessions[token]['rows'][0]
    assert (row['student_id'], row['student_forced_none']) == (student_id, False)


def test_event_import_add_row_without_card_writes_null_ref(event_import_client: SanicTestClient):
    """Одиночное добавление «Добавить» строки без карточки (0 кандидатов):
    участие создаётся с ref NULL, аудит — с student_ref_id: null."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    response = upload_event_xlsx(event_import_client, event_id, [{'ФИО': 'Козлов Козьма Козьмич', 'Курс': 1}])
    token = import_session_token(response)
    response = post_event_import_action(event_import_client, event_id, token, 'row/2/add')
    assert 'Строка 2: участник добавлен.' in unquote_plus(response.headers['location'])
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][11] is None
    assert audit_details(storage, 'event_participant_imported') == [{'event_id': event_id, 'student_ref_id': None}]
    page = get_event_preview(event_import_client, event_id, token)
    assert 'добавлен</span>' in re.sub(r'\s+', ' ', page.text)


def manual_form_input(body: str, input_id: str) -> str:
    """Разметка поля ручной формы добавления участника по id."""
    match = re.search(rf'<(?:input|select)[^>]*id="{input_id}"[^>]*>', body)
    assert match is not None, input_id
    return match.group(0)


def test_calendar_event_manual_form_course_required_default(event_import_client: SanicTestClient):
    """Дефолт настроек: курс обязателен (HTML required + серверная проверка),
    место опционально с подсказкой «дописать позже», пол — select."""
    storage = app.ctx.storage
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers())
    assert response.status == 200
    body = response.text
    assert 'required' in manual_form_input(body, 'participant-fio')
    assert 'aria-label="ФИО"' in manual_form_input(body, 'participant-fio')
    course_input = manual_form_input(body, 'participant-course')
    assert 'required' in course_input
    assert 'type="number"' in course_input
    assert 'min="1"' in course_input
    position_input = manual_form_input(body, 'participant-position')
    assert 'required' not in position_input
    assert 'title="Место можно дописать позже"' in position_input
    # Пол — select с пустым «не указан» (резолвер script.js пишет value).
    sex_input = manual_form_input(body, 'participant-sex')
    assert sex_input.startswith('<select')
    assert 'required' not in sex_input
    assert '<option value="">не указан</option>' in body
    assert '<option value="М">М</option>' in body

    # Серверная проверка дефолта: без курса — отказ (как раньше).
    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Безкурсов Без Безович',
            'student_sex': '',
            'institute': '',
            'group': '',
            'course': '',
            'position': '',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Курс обязателен' in unquote_plus(response.headers['location'])
    assert event_records_snapshot(storage) == []


def test_calendar_event_manual_form_course_optional(event_import_client: SanicTestClient):
    """Курс опционален (field_settings): нет HTML required, добавление без
    курса проходит серверно (курс → 0, как место)."""
    storage = app.ctx.storage
    storage.update_field_settings({'course': ('number', False)})
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers())
    assert response.status == 200
    course_input = manual_form_input(response.text, 'participant-course')
    assert 'required' not in course_input
    assert 'type="number"' in course_input

    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Безкурсов Без Безович',
            'student_sex': '',
            'institute': '',
            'group': '',
            'course': '',
            'position': '',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][4] == 0  # пустой опциональный курс → 0
    assert records[0][11] is None  # без карточки — без связи


def test_calendar_event_manual_form_course_text_type(event_import_client: SanicTestClient):
    """Курс-текст (кейс «Выпускник»): type=text с подсказкой, обязательность
    по настройке; место-текст — без числовых атрибутов."""
    storage = app.ctx.storage
    storage.update_field_settings({'course': ('text', True), 'position': ('text', False)})
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers())
    assert response.status == 200
    course_input = manual_form_input(response.text, 'participant-course')
    assert 'type="text"' in course_input
    assert 'required' in course_input
    assert 'placeholder="например, 2"' in course_input
    assert 'min=' not in course_input
    position_input = manual_form_input(response.text, 'participant-position')
    assert 'type="text"' in position_input
    assert 'required' not in position_input
    assert 'title="Место можно дописать позже"' in position_input

    # Текстовый курс сохраняется как текст.
    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Выпускников Вып Выпович',
            'student_sex': '',
            'institute': '',
            'group': '',
            'course': 'Выпускник 2025/26',
            'position': '',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    records = event_records_snapshot(storage)
    assert len(records) == 1
    assert records[0][4] == 'Выпускник 2025/26'


def test_calendar_event_manual_form_position_required(event_import_client: SanicTestClient):
    """Место обязательно по настройке: HTML required, подсказки «дописать
    позже» нет; сервер отвергает пустое место."""
    storage = app.ctx.storage
    storage.update_field_settings({'position': ('number', True)})
    event_id = make_calendar_event(storage)
    _, response = event_import_client.get(f'/calendar/{event_id}', headers=get_auth_headers())
    position_input = manual_form_input(response.text, 'participant-position')
    assert 'required' in position_input
    assert 'Место можно дописать позже' not in position_input

    headers = get_auth_headers()
    _, response = event_import_client.post(
        f'/calendar/{event_id}/participants',
        headers=headers,
        data={
            **csrf_for(headers),
            'student_name': 'Безместов Без Безович',
            'student_sex': '',
            'institute': '',
            'group': '',
            'course': '1',
            'position': '',
        },
        allow_redirects=False,
    )
    assert response.status == 302
    assert 'Поле "Место" обязательно' in unquote_plus(response.headers['location'])
    assert event_records_snapshot(storage) == []


# --- Обратная совместимость: экспорты/очередь/сопоставление. ---


def test_exports_do_not_leak_student_ref_id(event_import_client: SanicTestClient):
    # Служебная колонка не попадает в Excel-выгрузки: формат выгрузок —
    # часть контракта импорта/экспорта (шаблон импорта записей её не знает).
    competition = build_competition(
        {
            'ФИО': 'Тестов Тест Тестович',
            'Пол': 'М',
            'Институт': 'ИСИ',
            'Группа': 'ПГС-101',
            'Вид спорта': 'Бег',
            'Дата': '15.03.2026',
            'Уровень соревнований': 'внутривузовские',
            'Название соревнований': 'Кубок',
            'Место': 1,
            'Курс': 2,
        },
        custom_fields=[],
    )
    competition.student_ref_id = 42
    # Wave 1 P1: служебные поля participation-identity тоже не отдаются.
    competition.discipline = 'Бег 100 м'
    competition.result = '11.2'
    competition.calendar_event_id = 7
    row = competition_to_export_row(competition, [])
    assert 'student_ref_id' not in row
    assert 'discipline' not in row
    assert 'result' not in row
    assert 'calendar_event_id' not in row
    assert set(row) == {
        'Код студента',
        'ФИО',
        'Пол',
        'Институт',
        'Группа',
        'Курс',
        'Вид спорта',
        'Дата',
        'Уровень соревнований',
        'Название соревнований',
        'Место',
    }


def test_import_queue_payload_round_trip_with_ref_field(client: SanicTestClient):
    # Конфликт импорта записей: model_dump → payload → model_validate.
    # Новое None-поле не ломает round-trip (существующие записи в очереди
    # без ключа тоже валидируются — поле имеет дефолт). Wave 1 P1: то же
    # для discipline/result/calendar_event_id.
    competition = build_competition(
        {
            'ФИО': 'Тестов Тест Тестович',
            'Пол': 'М',
            'Институт': 'ИСИ',
            'Группа': 'ПГС-101',
            'Вид спорта': 'Бег',
            'Дата': '15.03.2026',
            'Уровень соревнований': 'внутривузовские',
            'Название соревнований': 'Кубок',
            'Место': 1,
            'Курс': 2,
        },
        custom_fields=[],
    )
    payload = competition.model_dump(mode='json', by_alias=False)
    assert payload['student_ref_id'] is None
    assert payload['discipline'] is None
    assert payload['result'] is None
    assert payload['calendar_event_id'] is None
    validated = Competition.model_validate(payload)
    assert validated.student_ref_id is None
    assert validated.discipline is None
    assert validated.result is None
    assert validated.calendar_event_id is None
    legacy_payload = {
        key: value
        for key, value in payload.items()
        if key not in {'student_ref_id', 'discipline', 'result', 'calendar_event_id'}
    }
    legacy_validated = Competition.model_validate(legacy_payload)
    assert legacy_validated.student_ref_id is None
    assert legacy_validated.discipline is None
    assert legacy_validated.result is None
    assert legacy_validated.calendar_event_id is None
