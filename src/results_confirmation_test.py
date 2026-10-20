"""REP-05 synthetic migration, authorization, classification and Reports oracle."""
import sqlite3
from datetime import date
from datetime import datetime
from io import BytesIO
from urllib.parse import parse_qs
from urllib.parse import urlencode
from urllib.parse import urlsplit

import pandas as pd
import pytest
from openpyxl import load_workbook

from src import main_test as helpers
from src.conftest import make_report_record
from src.conftest import seed_fixture_levels
from src.main import app
from src.main_test import csrf_for
from src.main_test import get_auth_headers
from src.main_test import get_report_cells
from src.main_test import get_report_footer
from src.main_test import get_report_xlsx_rows
from src.records import build_competition
from src.routes.calendar import build_event_participant_competition
from src.routes.calendar import event_import_row_record
from src.routes.calendar import event_participant_preset
from src.routes.calendar import parse_event_participant_rows
from src.routes.reports import build_report_event_options
from src.settings import settings
from src.storage.sqlite import SQLiteAdapter


client = helpers.client
TODAY = date(2026, 10, 20)


def seed_results_history(storage):
    """No production-derived identities or records; also usable for browser QA."""
    seed_fixture_levels(storage, 'Тестовый уровень')
    ref = storage.create_student('Тестовый Атлет Альфа', 'М', 'Тестовый институт', 'СИН-24', '1')
    past = storage.create_calendar_event('Тестовый кубок', '2001-01-01', None, 'Тестовый уровень', 'Бег', [])
    pending = storage.create_calendar_event('Тестовый кубок', '2001-01-02', None, 'Тестовый уровень', 'Бег', [])
    future = storage.create_calendar_event('Тестовый план', '2099-01-01', None, 'Тестовый уровень', 'Бег', [])
    assert storage.set_calendar_event_results_confirmation(past, True, 1) is None
    # A legacy future event is confirmed by migration; its blank result is unknown.
    storage.connection.execute('UPDATE calendar_events SET results_confirmed = 1 WHERE id = ?', (future,))
    states = [(past, value) for value in (1, 2, 3, 4, 0, '', ' dNf ')]
    states += [(None, value) for value in (0, 'DNF', 1, 4)] + [(future, 0), (pending, 1)]
    for index, (event, position) in enumerate(states, 1):
        record = make_report_record(
            f'Тестовый Атлет {index}',
            'М',
            'Тестовый институт',
            'СИН-24',
            99,
            'Бег',
            datetime(2001, 1, 1),
            'Тестовый уровень',
            position,
            student_ref_id=ref if index <= 7 else None,
            extra={'legacy_result': 'DNF'},
        )
        record.calendar_event_id = event
        record.result = 'DNF'  # Informational only; never the place classifier.
        record.discipline = f'Тестовая дисциплина {index}'
        storage.save_competitions([record])
    rejected = record.model_copy(update={'position': 1, 'calendar_event_id': past})
    storage.save_competitions([rejected], review_status='pending')
    return past, pending, future


@pytest.fixture
def results_history(client, monkeypatch):
    monkeypatch.setattr(settings, 'auth_viewer_username', 'synthetic-viewer')
    monkeypatch.setattr(settings, 'auth_viewer_password', 'synthetic-viewer-password')
    monkeypatch.setattr('src.storage.events.today_in_timezone', lambda _: TODAY)
    monkeypatch.setattr('src.routes.calendar.today_in_timezone', lambda _: TODAY)
    storage = SQLiteAdapter(':memory:')
    storage.connection.create_function('app_today', 0, lambda: TODAY.isoformat())
    events = seed_results_history(storage)
    storage.get_user = helpers.fake_get_user
    previous = app.ctx.storage
    app.ctx.storage = storage
    try:
        yield client, storage, events
    finally:
        app.ctx.storage = previous
        storage.connection.close()


def ids(storage, **filters):
    count, rows, _ = storage.get_report_drilldown(metric='participations', **filters)
    assert count == len(rows)
    return {int(record.record_id) for record in rows}


def test_reports_classification_is_a_record_level_oracle(results_history):
    _, storage, (past, pending, _) = results_history
    assert ids(storage) == set(range(1, 13))
    assert ids(storage, position='<2') == {1, 10}
    assert ids(storage, position='<4') == {1, 2, 3, 10}
    assert ids(storage, position='>3') == {4, 5, 6, 11}
    assert ids(storage, position='0') == {7, 8, 9, 12}
    assert ids(storage, event_id=[str(pending)]) == set()
    summary = storage.get_report_summary()
    assert (summary.count_participation, summary.count_wins, summary.count_prizes) == (12, 2, 4)
    assert summary.count_events == 2 and summary.count_unlinked_participations == 4
    for key in ('student', 'group', 'institute', 'course', 'admission_year', 'sport', 'level', 'year'):
        rows = storage.get_filtered() if key == 'student' else storage.get_grouped_report(key)
        assert sum(row.count_participation for row in rows) == 12
        for row in rows:
            for metric, field in (
                ('participations', 'count_participation'),
                ('wins', 'count_wins'),
                ('prizes', 'count_prizes'),
            ):
                total, selected, _ = storage.get_report_drilldown(
                    metric=metric, slice_key=key, anchor_id=row.drilldown_record_id
                )
                assert total == len(selected) == getattr(row, field)
    # No history, zeros, free results, identity links or review flags rewritten.
    raw = storage.connection.execute('SELECT position, result, student_ref_id FROM competitions ORDER BY id').fetchall()
    assert [row['position'] for row in raw[:7]] == [1, 2, 3, 4, 0, '', ' dNf ']
    assert all(row['result'] == 'DNF' for row in raw)
    assert storage.get_calendar_event(past)['results_confirmed'] == 1


@pytest.mark.parametrize('position', ['', '<2', '<4', '>3', '0'])
@pytest.mark.parametrize(
    'slice_key', ['student', 'group', 'institute', 'course', 'admission_year', 'sport', 'level', 'year']
)
def test_html_xlsx_and_registry_parity(results_history, position, slice_key):
    browser, storage, _ = results_history
    query = urlencode({'period': 'all', 'position': position, 'group_by': slice_key})
    headers = get_auth_headers()
    _, screen = browser.get('/reports?' + query, headers=headers)
    _, export = browser.get('/export/report?' + query, headers=headers)
    _, registry = browser.get('/?report_metric=participations&' + query, headers=headers)
    assert screen.status == export.status == registry.status == 200
    column = {
        'student': 'ФИО',
        'group': 'Группа',
        'institute': 'Институт',
        'course': 'Курс',
        'admission_year': 'Год поступления',
        'sport': 'Вид спорта',
        'level': 'Уровень соревнований',
        'year': 'Год',
    }[slice_key]
    xlsx = get_report_xlsx_rows(export.content, column)
    width = 3 if slice_key == 'student' else 5
    assert [row[-width:] for row in get_report_cells(screen.text)] == [
        [str(value) for value in row[-width:]] for row in xlsx[1:-1]
    ]
    assert get_report_footer(screen.text)[-width:] == [str(value) for value in xlsx[-1][-width:]]
    assert len(ids(storage, position=position)) == storage.get_report_summary(position=position).count_participation
    assert 'data-column-key="position">0</td>' not in registry.text


def test_confirmation_admin_only_and_csrf(results_history):
    browser, storage, (_, pending, _) = results_history
    url = f'/calendar/{pending}/results-confirmation'
    for role in ('editor', 'viewer'):
        headers = get_auth_headers(role)
        _, response = browser.post(url, headers=headers, data={**csrf_for(headers), 'confirmed': '1'})
        assert response.status == 403
    athlete = {
        'cookie': f'{settings.auth_cookie_name}=' + helpers.create_auth_cookie_value(username='sportik', role='athlete')
    }
    _, response = browser.post(url, headers=athlete, data={**csrf_for(athlete), 'confirmed': '1'})
    assert response.status == 403
    headers = get_auth_headers()
    _, response = browser.post(url, headers=headers, data={'confirmed': '1'})
    assert response.status == 403
    assert not storage.get_calendar_event(pending)['results_confirmed']
    return_url = '/calendar?' + urlencode({'period': '2025-2026', 'sport': 'Бег'})
    _, response = browser.post(
        url,
        headers=headers,
        data={**csrf_for(headers), 'confirmed': '1', 'result_filter': 'without', 'return': return_url},
        allow_redirects=False,
    )
    assert response.status == 302
    navigation = parse_qs(urlsplit(response.headers['location']).query)
    assert navigation['result'] == ['without'] and navigation['return'] == [return_url]
    confirmed = storage.get_calendar_event(pending)
    assert confirmed['results_confirmed'] == 1 and confirmed['results_confirmed_by'] == 1
    assert datetime.fromisoformat(confirmed['results_confirmed_at']).utcoffset().total_seconds() == 0
    assert storage.get_report_summary().count_participation == 13
    assert storage.set_calendar_event_results_confirmation(pending, True, 999) is None
    assert storage.get_calendar_event(pending)['results_confirmed_at'] == confirmed['results_confirmed_at']
    assert storage.get_calendar_event(pending)['results_confirmed_by'] == 1
    # Existing editor update remains available and preserves event confirmation.
    storage.update_event_participation_result(13, pending, 'DNF', '12.5')
    storage.update_calendar_event(
        pending, 'Тестовый новый заголовок', '2001-01-02', None, 'Тестовый уровень', 'Бег', []
    )
    assert storage.get_calendar_event(pending)['results_confirmed_at'] == confirmed['results_confirmed_at']
    assert storage.get_report_summary(position='0', event_id=[str(pending)]).count_participation == 1
    _, response = browser.post(
        url, headers=headers, data={**csrf_for(headers), 'confirmed': '0'}, allow_redirects=False
    )
    assert response.status == 302
    event = storage.get_calendar_event(pending)
    assert (event['results_confirmed'], event['results_confirmed_at'], event['results_confirmed_by']) == (0, None, None)
    assert storage.get_report_summary().count_participation == 12


def test_confirmation_end_boundary_and_read_only_ui(results_history):
    browser, storage, (past, pending, future) = results_history
    today = storage.create_calendar_event('Тестовое сегодня', TODAY.isoformat(), None, 'Тестовый уровень', 'Бег', [])
    spanning = storage.create_calendar_event(
        'Тестовый период', '2001-01-01', TODAY.isoformat(), 'Тестовый уровень', 'Бег', []
    )
    for event_id in (today, spanning, future):
        assert storage.set_calendar_event_results_confirmation(event_id, True, 1) == 'not_finished'
    assert storage.set_calendar_event_results_confirmation(999999, True, 1) == 'not_found'
    for role in ('admin', 'editor', 'viewer'):
        _, response = browser.get(f'/calendar/{past}', headers=get_auth_headers(role))
        assert response.status == 200 and 'Результаты подтверждены' in response.text
        assert ('Снять подтверждение' in response.text) == (role == 'admin')
    _, response = browser.get(f'/calendar/{pending}', headers=get_auth_headers())
    assert 'Результаты ожидают подтверждения' in response.text and 'Подтвердить результаты' in response.text
    _, response = browser.get(f'/calendar/{today}', headers=get_auth_headers())
    assert 'Подтверждение доступно после окончания' in response.text
    assert 'value="1"' not in response.text.split('data-results-confirmation>', 1)[1].split('</div>', 1)[0]
    options = build_report_event_options(storage)
    namesakes = [option for option in options if option['label'] == 'Тестовый кубок']
    assert {option['value'] for option in namesakes} == {str(past), str(pending)}
    assert all('№' not in option['detail'] for option in options)


def test_legacy_migration_is_once_and_preserves_records(tmp_path):
    db = str(tmp_path / 'synthetic-legacy.sqlite3')
    storage = SQLiteAdapter(db)
    seed_results_history(storage)
    matching_name = 'Тестовый Legacy Кубок'
    storage.create_calendar_event(matching_name, '2001-03-01T00:00:00', None, '', 'Бег', [])
    legacy_record = make_report_record(
        'Тестовый Legacy Несвязанный', 'М', 'Тестовый институт', 'СИН-24', 99, 'Бег', datetime(2001, 3, 1), '', 1
    )
    legacy_record.name = matching_name
    storage.save_competitions([legacy_record])
    before = [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')]
    storage.connection.close()
    with sqlite3.connect(db) as legacy:
        for column in ('results_confirmed', 'results_confirmed_at', 'results_confirmed_by'):
            legacy.execute(f'ALTER TABLE calendar_events DROP COLUMN {column}')
    storage = SQLiteAdapter(db)
    assert all(event['results_confirmed'] == 1 for event in storage.list_calendar_events())
    assert storage.get_calendar_event(1)['results_confirmed_at'] is None
    assert storage.get_report_summary(name=legacy_record.student_name).count_participation == 1
    assert storage.get_report_summary(name=legacy_record.student_name).count_wins == 1
    assert [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')] == before
    assert storage.connection.execute('PRAGMA quick_check').fetchone()[0] == 'ok'
    assert storage.set_calendar_event_results_confirmation(1, False, 1) is None
    new = storage.create_calendar_event('Тестовое новое', '2001-02-01', None, '', '', [])
    storage.connection.close()
    storage = SQLiteAdapter(db)
    assert not storage.get_calendar_event(1)['results_confirmed']
    assert not storage.get_calendar_event(new)['results_confirmed']
    assert storage.get_calendar_event(2)['results_confirmed']
    assert storage.get_report_summary(name=legacy_record.student_name).count_unlinked_participations == 1
    assert storage.get_report_summary(name=legacy_record.student_name).count_wins == 1
    assert [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')] == before
    storage.connection.close()


def test_migration_failure_rolls_back_ddl_and_backfill():
    storage = SQLiteAdapter(':memory:')
    event = storage.create_calendar_event('Тестовая миграция', '2001-01-01', None, '', '', [])
    for column in ('results_confirmed', 'results_confirmed_at', 'results_confirmed_by'):
        storage.connection.execute(f'ALTER TABLE calendar_events DROP COLUMN {column}')
    storage.connection.commit()
    actual = storage.connection

    class FailSecondAlter:
        def execute(self, sql, *args):
            if 'ADD COLUMN results_confirmed_at' in sql:
                raise RuntimeError('Synthetic migration interruption')
            return actual.execute(sql, *args)

    storage.connection = FailSecondAlter()
    with pytest.raises(RuntimeError, match='interruption'):
        storage._migrate_event_results(set())
    storage.connection = actual
    columns = {row['name'] for row in actual.execute('PRAGMA table_info(calendar_events)')}
    assert not any(name.startswith('results_confirmed') for name in columns)
    storage._migrate_event_results(columns)
    assert storage.get_calendar_event(event)['results_confirmed']
    actual.close()


@pytest.mark.parametrize('value', ['DNF', 'dnf', ' dNf '])
def test_default_numeric_place_accepts_dnf_in_manual_and_excel(value):
    record = {
        'ФИО': 'Тестовый Атлет',
        'Пол': 'М',
        'Институт': 'Тестовый институт',
        'Группа': 'СИН-24',
        'Дата': '01.01.2001',
        'Вид спорта': 'Бег',
        'Уровень соревнований': '',
        'Название соревнований': 'Тестовый кубок',
        'Место': value,
    }
    competition = build_competition(record, custom_fields=[], manual_input=True)
    assert competition.position == 'DNF'
    output = BytesIO()
    pd.DataFrame([record]).to_excel(output, index=False)
    output.seek(0)
    loaded = pd.read_excel(output).fillna('').to_dict('records')[0]
    assert build_competition(loaded, custom_fields=[], manual_input=True).position == 'DNF'


def test_editor_dnf_entry_edit_and_import_paths(results_history):
    browser, storage, (past, _, _) = results_history
    headers = get_auth_headers('editor')
    form = dict(
        student_name='Тестовая Новая Запись',
        student_sex='Ж',
        institute='Тестовый институт',
        group='СИН-24',
        position=' dnf ',
        discipline='Тестовая новая дисциплина',
        result='12.5',
    )
    _, response = browser.post(
        f'/calendar/{past}/participants', headers=headers, data={**csrf_for(headers), **form}, allow_redirects=False
    )
    assert response.status == 302 and 'error=' not in response.headers['location']
    record = storage.connection.execute(
        'SELECT id, position FROM competitions WHERE student_name = ?', (form['student_name'],)
    ).fetchone()
    assert record['position'] == 'DNF'
    assert storage.get_calendar_event(past)['results_confirmed']
    form.update(position='2', date='01.01.2001', name='Ignored event name', sport='Бег', level='Тестовый уровень')
    _, response = browser.post(
        f'/competition/{record["id"]}', headers=headers, data={**csrf_for(headers), **form}, allow_redirects=False
    )
    assert response.status == 302
    assert storage.get_competition_by_id(record['id']).position == 2
    assert storage.get_calendar_event(past)['results_confirmed']
    form.update(student_name='Тестовый Реестр DNF', position='dnf')
    _, response = browser.post(
        '/competition', headers=headers, data={**csrf_for(headers), **form}, allow_redirects=False
    )
    assert response.status == 302
    assert (
        storage.connection.execute(
            'SELECT position FROM competitions WHERE student_name = ?', (form['student_name'],)
        ).fetchone()[0]
        == 'DNF'
    )
    dataframe = pd.DataFrame(
        [
            {
                'ФИО': 'Тестовый Импорт DNF',
                'Пол': 'М',
                'Институт': 'Тестовый институт',
                'Группа': 'СИН-24',
                'Место': ' dNf ',
                'Дисциплина': 'Тестовая импортная',
                'Результат': '42.5',
            }
        ]
    )
    output = BytesIO()
    dataframe.to_excel(output, index=False)
    output.seek(0)
    rows, unknown, _ = parse_event_participant_rows(pd.read_excel(output), [])
    assert not unknown and rows[0]['status'] == 'pending'
    imported, _, _ = event_import_row_record(
        storage, rows[0], event_participant_preset(storage.get_calendar_event(past)), None, []
    )
    competition, error = build_event_participant_competition(imported, [], storage, result=rows[0]['result'])
    assert error is None and competition.position == 'DNF' and competition.result == '42.5'
    competition.calendar_event_id = past
    storage.apply_event_participation_batch(past, [competition], owner_id=1)
    assert (
        storage.connection.execute(
            'SELECT position FROM competitions WHERE student_name = ?', ('Тестовый Импорт DNF',)
        ).fetchone()[0]
        == 'DNF'
    )


@pytest.mark.parametrize('custom_order', [('qa_count', 'qa_note'), ('qa_note', 'qa_count')])
def test_registry_xlsx_place_display_preserves_numeric_roundtrip(results_history, custom_order):
    browser, storage, _ = results_history
    before = [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')]
    definitions = {'qa_count': ('Тестовый счётчик', 'number'), 'qa_note': ('Тестовая заметка', 'text')}
    for index, key in enumerate(custom_order):
        label, kind = definitions[key]
        storage.create_custom_field(key, label, kind, False, True, True, True, index)
    # Custom numeric fields keep their existing exported value and format.
    storage.connection.execute('UPDATE competitions SET extra_data = ?', ('{"qa_count": "0", "qa_note": "Тест"}',))
    storage.connection.commit()
    before_export = [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')]
    _, response = browser.get('/export/index', headers=get_auth_headers())
    assert response.status == 200
    book = load_workbook(BytesIO(response.content), data_only=True)
    sheet = book.active
    headings = [cell.value for cell in sheet[1]]
    place = headings.index('Место') + 1
    other_zero = headings.index('Тестовый счётчик') + 1
    assert headings[-2:] == [definitions[key][0] for key in custom_order]
    place_cells = [row[place - 1] for row in list(sheet.rows)[1:]]
    zeros = [cell for cell in place_cells if cell.value == 0]
    assert zeros and all(cell.data_type == 'n' for cell in zeros)
    # In Excel's actual number-format contract the third section formats zero.
    assert all(cell.number_format.split(';')[2] == '"—"' for cell in zeros)
    assert any(cell.value == 'DNF' and cell.data_type == 's' for cell in place_cells)
    assert any(cell.value == 1 and cell.data_type == 'n' for cell in place_cells)
    assert all(
        sheet.cell(row, other_zero).value == '0' and sheet.cell(row, other_zero).number_format == 'General'
        for row in range(2, sheet.max_row + 1)
    )
    imported = pd.read_excel(BytesIO(response.content)).fillna('')
    imported_by_name = imported.set_index('ФИО')
    assert imported_by_name.loc['Тестовый Атлет 5', 'Место'] == 0
    assert imported_by_name.loc['Тестовый Атлет 9', 'Место'] == 'DNF'
    assert imported_by_name.loc['Тестовый Атлет 1', 'Место'] == 1
    reimported = build_competition(
        imported_by_name.loc['Тестовый Атлет 5'].to_dict() | {'ФИО': 'Тестовый Атлет 5'},
        custom_fields=[],
        manual_input=True,
    )
    assert reimported.position == 0
    reimported = build_competition(
        imported_by_name.loc['Тестовый Атлет 9'].to_dict() | {'ФИО': 'Тестовый Атлет 9'},
        custom_fields=[],
        manual_input=True,
    )
    assert reimported.position == 'DNF'
    assert [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')] == before_export
    assert [row[0] for row in before_export] == [row[0] for row in before]
    book.close()


@pytest.mark.parametrize('confirmed', [False, True])
@pytest.mark.parametrize('position', [1, 2, 3, 4, 0, 'DNF'])
@pytest.mark.parametrize('end', [None, '2001-01-03T00:00:00'])
def test_restarts_preserve_exact_matching_unlinked_history(tmp_path, confirmed, position, end):
    """A unique candidate is a suggestion, never startup authorization to link."""
    path = str(tmp_path / 'synthetic-unlinked-restarts.sqlite3')
    storage = SQLiteAdapter(path)
    seed_fixture_levels(storage, 'Тестовый уровень')
    event = storage.create_calendar_event(
        'Тестовый Совпадающий Кубок', '2001-01-01T00:00:00', end, 'Тестовый уровень', 'Бег', []
    )
    other = storage.create_calendar_event(
        'Тестовое Явно Выбранное Событие', '2001-02-01T00:00:00', None, 'Тестовый уровень', 'Бег', []
    )
    if confirmed:
        assert storage.set_calendar_event_results_confirmation(event, True, 1) is None
    record = make_report_record(
        'Тестовый Несвязанный Атлет',
        'М',
        'Тестовый Институт',
        'СИН-24',
        99,
        'Бег',
        datetime(2001, 1, 1),
        'Тестовый уровень',
        position,
        extra={'test_note': 'Тестовая история'},
    )
    record.name = 'Тестовый Совпадающий Кубок'
    record.date_to = datetime.fromisoformat(end) if end else None
    record.discipline, record.result = 'Тестовая дисциплина', 'Информационный результат'
    storage.save_competitions([record])
    # Already-linked rows must not be unlinked or retargeted, even when their
    # snapshot matches a different unique candidate and that event is unconfirmed.
    storage.save_competitions(
        [
            record.model_copy(update={'student_name': 'Тестовый Явно Связанный', 'calendar_event_id': event}),
            record.model_copy(update={'student_name': 'Тестовый Другая Связь', 'calendar_event_id': other}),
        ]
    )
    before_rows = [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')]
    before_events = [tuple(row) for row in storage.connection.execute('SELECT * FROM calendar_events ORDER BY id')]
    historical_name = 'Тестовый Несвязанный Атлет'
    baseline = {
        category: storage.get_report_summary(name=historical_name, position=category).model_dump()
        for category in ('', '<2', '<4', '>3', '0')
    }
    assert baseline['']['count_participation'] == 1
    assert baseline['']['count_wins'] == int(position == 1)
    assert baseline['']['count_prizes'] == int(isinstance(position, int) and 1 <= position <= 3)
    assert baseline['']['count_unlinked_participations'] == 1
    for _ in range(3):
        storage.connection.close()
        storage = SQLiteAdapter(path)
        assert [
            tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')
        ] == before_rows
        assert [
            tuple(row) for row in storage.connection.execute('SELECT * FROM calendar_events ORDER BY id')
        ] == before_events
        assert storage.get_competition_by_id(1).calendar_event_id is None
        assert storage.get_competition_by_id(2).calendar_event_id == event
        assert storage.get_competition_by_id(3).calendar_event_id == other
        assert storage.get_calendar_event(event)['results_confirmed'] == int(confirmed)
        for category, expected in baseline.items():
            assert storage.get_report_summary(name=historical_name, position=category).model_dump() == expected
        for metric, expected in (('participations', 1), ('wins', int(position == 1))):
            total, rows, _ = storage.get_report_drilldown(metric=metric, name=historical_name)
            assert total == len(rows) == expected
        preview = storage.calendar_link_backfill_preview()
        assert preview['counters']['matched'] == 1 and preview['counters']['already_linked'] == 2
        assert preview['matched'][0]['record_id'] == 1 and preview['matched'][0]['event_id'] == event
        assert storage.connection.execute('PRAGMA quick_check').fetchone()[0] == 'ok'
    storage.connection.close()


def test_manual_link_after_restart_is_explicit_and_respects_confirmation(tmp_path):
    path = str(tmp_path / 'synthetic-manual-link.sqlite3')
    storage = SQLiteAdapter(path)
    event = storage.create_calendar_event('Тестовый Ручной Кубок', '2001-01-01T00:00:00', None, '', 'Бег', [])
    record = make_report_record(
        'Тестовый Ручной Атлет', 'М', 'Тестовый Институт', 'СИН-24', 99, 'Бег', datetime(2001, 1, 1), '', 1
    )
    record.name = 'Тестовый Ручной Кубок'
    storage.save_competitions([record])
    storage.connection.close()
    storage = SQLiteAdapter(path)
    assert storage.get_competition_by_id(1).calendar_event_id is None
    assert (storage.get_report_summary().count_participation, storage.get_report_summary().count_wins) == (1, 1)
    counters, items = storage.apply_calendar_link_backfill()
    assert counters == {'matched': 1, 'linked': 1, 'skipped': 0}
    assert items == [{'record_id': 1, 'event_id': event, 'event_name': 'Тестовый Ручной Кубок'}]
    assert storage.get_competition_by_id(1).calendar_event_id == event
    assert storage.get_report_summary().count_participation == 0  # explicit link to unconfirmed event
    assert storage.set_calendar_event_results_confirmation(event, True, 1) is None
    assert (storage.get_report_summary().count_participation, storage.get_report_summary().count_wins) == (1, 1)
    rows_after_manual = [tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')]
    event_after_manual = storage.get_calendar_event(event)
    for _ in range(2):
        storage.connection.close()
        storage = SQLiteAdapter(path)
        assert [
            tuple(row) for row in storage.connection.execute('SELECT * FROM competitions ORDER BY id')
        ] == rows_after_manual
        assert storage.get_calendar_event(event) == event_after_manual
        assert (storage.get_report_summary().count_participation, storage.get_report_summary().count_wins) == (1, 1)
        assert storage.apply_calendar_link_backfill() == ({'matched': 0, 'linked': 0, 'skipped': 0}, [])
    storage.connection.close()
