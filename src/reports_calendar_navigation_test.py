"""RPT-10/CAL-02: fictional history, exact SQL membership and confidential GET navigation."""
import re
from html import unescape
from io import BytesIO
from urllib.parse import parse_qs
from urllib.parse import urlencode
from urllib.parse import urlsplit

import pytest
from openpyxl import load_workbook

from src.auth import reset_presence_tracking
from src.main_test import get_auth_headers
from src.report_context import context_args
from src.report_context import decode_report_context
from src.report_context import encode_report_context
from src.reports_metrics_test import client  # noqa: F401
from src.reports_metrics_test import metrics_history as history_fixture
from src.routes.reports import build_report_presentation
from src.routes.reports import build_report_workbook
from src.routes.reports import get_report_rows
from src.routes.reports import report_export_columns
from src.routes.reports import REPORT_SLICES


metrics_history = history_fixture


def hrefs(html):
    return [unescape(url) for url in re.findall(r'href="([^"]+)"', html)]


def calendar_link(html):
    return next(url for url in hrefs(html) if url.startswith('/calendar?report_context='))


def card_ids(html):
    return [int(value) for value in re.findall(r'class="calendar-comp-row__name" href="/calendar/(\d+)', html)]


def test_numbered_workbook_shares_all_slices_sorts_and_optional_columns(metrics_history):
    _, storage, _, _, _ = metrics_history
    summary = storage.get_report_summary()
    for slice_def in REPORT_SLICES:
        slice_key = slice_def['key']
        raw = get_report_rows(storage, slice_key, {})
        keys = ['count_participation', 'count_wins', 'count_prizes']
        keys += ['student_name', 'institute', 'group', 'course'] if slice_key == 'student' else ['slice_value']
        for key in ['', *keys]:
            for order in ['asc', 'desc']:
                args = context_args([('sort', key), ('order', order)])
                expected = [row for block in build_report_presentation(raw, slice_key, args) for row in block]
                for columns in (report_export_columns(slice_key), ['Участий']):
                    sheet = load_workbook(
                        build_report_workbook(
                            raw,
                            columns,
                            slice_key,
                            summary,
                            period_label='Всё время',
                            filter_chips=[],
                            slice_label=slice_def['label'],
                            sort_args=args,
                        )
                    ).active
                    header = next(cell.row for cell in sheet['A'] if cell.value == '№')
                    assert [sheet.cell(header, col).value for col in range(1, len(columns) + 2)] == ['№', *columns]
                    data = list(sheet.iter_rows(min_row=header + 1, max_row=header + len(expected)))
                    assert [row[0].value for row in data] == list(range(1, len(expected) + 1))
                    assert all(row[0].data_type == 'n' and isinstance(row[0].value, int) for row in data)
                    participation = columns.index('Участий') + 1
                    assert [row[participation].value for row in data] == [row.count_participation for row in expected]
                    assert sheet.cell(header + len(expected) + 1, 1).value is None
    empty = load_workbook(
        build_report_workbook(
            [], ['ФИО'], 'student', summary, period_label='Всё время', filter_chips=[], slice_label='Студент'
        )
    ).active
    header = next(cell.row for cell in empty['A'] if cell.value == '№')
    assert empty.cell(header + 1, 1).value is None and empty.cell(header + 1, 2).value == 'Итого'


def test_calendar_exact_snapshot_membership_private_links_and_no_writes(metrics_history):
    browser, storage, _, _, ids = metrics_history
    # Event metadata deliberately disagrees with the historical participation snapshot.
    storage.connection.execute(
        "UPDATE calendar_events SET date='2029-01-01T00:00:00',level='Иной уровень',sport='Иной спорт'"
    )
    storage.connection.commit()
    query = urlencode({'period': '2026-2027', 'sport': 'Бег', 'level': 'Тестовый уровень', 'name': 'Исторический'})
    _, report = browser.get('/reports?' + query, headers=get_auth_headers())
    assert report.status == 200
    link = calendar_link(report.text)
    assert 'Исторический' not in link and '%D0' not in link
    token = parse_qs(urlsplit(link).query)['report_context'][0]
    assert dict(decode_report_context(token))['name'] == 'Исторический'
    reset_presence_tracking()
    before = storage.connection.total_changes
    _, calendar = browser.get(link, headers=get_auth_headers())
    assert calendar.status == 200
    assert card_ids(calendar.text) == list(reversed(ids))
    assert 'В показателе: 2 · Показано: 2' in calendar.text
    assert storage.connection.total_changes == before
    links = hrefs(calendar.text)
    backlink = next(url for url in links if url.startswith('/reports?report_context='))
    for url in links:
        assert 'name=' not in url and 'custom__' not in url
    _, restored = browser.get(backlink, headers=get_auth_headers())
    assert restored.status == 200 and 'value="Исторический"' in restored.text
    _, again = browser.get(link, headers=get_auth_headers())
    assert card_ids(again.text) == card_ids(calendar.text)
    _, narrowed = browser.get(link + '&confirmed=no', headers=get_auth_headers())
    assert card_ids(narrowed.text) == [] and 'В показателе: 2 · Показано: 0' in narrowed.text
    event = next(url for url in links if re.match(r'/calendar/\d+\?', url))
    assert parse_qs(urlsplit(event).query)['return'] == [link]
    before = storage.connection.total_changes
    reset_presence_tracking()
    _, event_page = browser.get(event, headers=get_auth_headers())
    assert event_page.status == 200 and storage.connection.total_changes == before
    assert link in hrefs(event_page.text)


def test_source_chips_remove_one_repeated_value_without_losing_private_context(metrics_history):
    browser, _, _, _, _ = metrics_history
    entries = [
        ('period', 'all'),
        ('name', 'Исторический'),
        ('institute', 'Институт А'),
        ('institute', 'Институт Б'),
        ('sort', 'count_wins'),
        ('order', 'desc'),
    ]
    link = '/calendar?' + urlencode({'report_context': encode_report_context(entries)})
    _, page = browser.get(link, headers=get_auth_headers())
    assert page.status == 200
    removal = re.search(r'href="([^"]+)" aria-label="Убрать условие отчёта: Институт"', page.text).group(1)
    token = parse_qs(urlsplit(unescape(removal)).query)['report_context'][0]
    values = context_args(decode_report_context(token))
    assert values['institute'] == ['Институт Б']
    assert values['name'] == ['Исторический'] and values['sort'] == ['count_wins'] and values['order'] == ['desc']


def test_private_registry_metric_link_preserves_selection_and_safe_backlink(metrics_history):
    browser, _, _, _, _ = metrics_history
    _, report = browser.get(
        '/reports?' + urlencode({'period': 'all', 'name': 'Исторический'}), headers=get_auth_headers()
    )
    metric = next(
        url
        for url in hrefs(report.text)
        if url.startswith('/?report_context=') and 'report_metric=participations' in url
    )
    _, registry = browser.get(metric, headers=get_auth_headers())
    assert registry.status == 200 and 'name=' not in metric
    assert any(url.startswith('/reports?report_context=') for url in hrefs(registry.text))


@pytest.mark.parametrize(
    'suffix',
    [
        'report_context=',
        'report_context=broken',
        pytest.param('report_context=' + 'x' * 5601, id='oversized-context'),
        'report_context=broken&report_context=other',
        'status=unknown',
        'confirmed=unknown',
    ],
)
def test_invalid_calendar_context_never_broadens(metrics_history, suffix):
    browser, _, _, _, _ = metrics_history
    _, page = browser.get('/calendar?' + suffix, headers=get_auth_headers())
    assert page.status in (400, 413, 414, 431) and card_ids(page.text) == []


def test_stale_custom_context_and_long_report_are_visible_errors(metrics_history):
    browser, _, _, _, _ = metrics_history
    token = encode_report_context([('period', 'all'), ('custom__removed_field', 'fictional-private-value')])
    _, calendar = browser.get('/calendar?' + urlencode({'report_context': token}), headers=get_auth_headers())
    assert calendar.status == 400 and 'больше недоступно' in calendar.text
    _, report = browser.get('/reports?' + urlencode({'period': 'all', 'name': 'x' * 1025}), headers=get_auth_headers())
    assert report.status == 200 and 'слишком длинные условия' in report.text
    assert 'data-report-table' not in report.text


def test_context_is_authenticated_bounded_and_stateless():
    entries = [('period', 'all'), ('name', 'Полностью Вымышленный Студент'), ('custom__synthetic', 'fictional-private')]
    token = encode_report_context(entries)
    assert decode_report_context(token) == entries
    assert decode_report_context(token) == entries
    with pytest.raises(ValueError):
        decode_report_context(token[:30] + ('A' if token[30] != 'A' else 'B') + token[31:])
    for invalid in (
        [('unexpected', 'value')],
        [('name', 'x' * 1025)],
        [('name', 'a'), ('name', 'b')],
        [('sport', str(i)) for i in range(129)],
    ):
        with pytest.raises(ValueError):
            encode_report_context(invalid)


@pytest.mark.parametrize(
    'private_filters',
    [
        [('name', 'Исторический')],
        [('custom__synthetic_private', 'fictional-private-value')],
        [('name', 'Исторический'), ('custom__synthetic_private', 'fictional-private-value')],
    ],
)
def test_confidential_export_form_and_request_keep_context_sort_and_columns(metrics_history, private_filters):
    browser, storage, _, _, _ = metrics_history
    storage.create_custom_field(
        'synthetic_private', 'Тестовое персональное условие', 'text', False, True, True, True, 0
    )
    storage.connection.execute(
        "UPDATE competitions SET extra_data=json_set(extra_data,'$.synthetic_private','fictional-private-value')"
    )
    storage.connection.commit()
    entries = [('period', 'all'), ('sort', 'count_wins'), ('order', 'desc'), *private_filters]
    _, report = browser.get('/reports?' + urlencode(entries), headers=get_auth_headers())
    _, calendar = browser.get(calendar_link(report.text), headers=get_auth_headers())
    back = next(url for url in hrefs(calendar.text) if url.startswith('/reports?report_context='))
    _, restored = browser.get(back, headers=get_auth_headers())
    form = re.search(r'<form class="report-export-form".*?</form>', restored.text, re.S).group(0)
    hidden = [
        (unescape(key), unescape(value))
        for key, value in re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)">', form)
    ]
    assert len(hidden) == 1 and hidden[0][0] == 'report_context'
    assert context_args(decode_report_context(hidden[0][1])) == context_args(entries)
    export_url = '/export/report?' + urlencode([*hidden, ('columns', 'ФИО'), ('columns', 'Участий')])
    assert 'name=' not in export_url and 'custom__' not in export_url and 'fictional-private-value' not in export_url
    _, export = browser.get(export_url, headers=get_auth_headers())
    assert export.status == 200
    sheet = load_workbook(BytesIO(export.body)).active
    header = next(cell.row for cell in sheet['A'] if cell.value == '№')
    assert [sheet.cell(header, column).value for column in range(1, 4)] == ['№', 'ФИО', 'Участий']
    assert isinstance(sheet.cell(header + 1, 1).value, int)
