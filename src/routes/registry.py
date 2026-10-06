"""Маршруты реестра записей: главная, импорт/очередь конфликтов, шаблон,
выгрузка реестра, ручной ввод/правка/удаление записей, резолвер атлетов,
модерация и вложения. Перенесено из src/main.py без изменения поведения
(Architecture v1)."""
import asyncio
import logging
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable
from typing import Sequence
from urllib.parse import urlencode

import pandas as pd
from openpyxl.utils import get_column_letter
from sanic import redirect
from sanic import Request
from sanic import Sanic
from sanic import text
from sanic.response import json as json_response
from sanic.response import raw
from sanic_ext import render

from src.auth import athlete_identity_context
from src.auth import get_athlete_profile_defaults
from src.auth import get_current_user_id
from src.auth import log_audit_event
from src.auth import require_admin
from src.auth import require_moderator
from src.auth import require_writer
from src.auth import user_can_write
from src.auth import user_is_admin
from src.auth import user_is_athlete
from src.auth import user_is_moderator
from src.auth import user_owns_record
from src.auth import USER_ROLES
from src.education import derive_course
from src.files import ATTACHMENT_MAX_SIZE
from src.files import attachments_dir
from src.files import detect_attachment_type
from src.files import files_dir
from src.files import remove_attachment_record_dirs
from src.gto import build_gto_rows
from src.models.competition import Competition
from src.models.custom_field import CustomField
from src.records import BASE_FIELD_LABELS
from src.records import build_competition
from src.records import build_index_dataframe
from src.records import competition_duplicate_key
from src.records import competition_partial_key
from src.records import ensure_catalog_values
from src.records import event_participation_collision_custom_keys
from src.records import event_participation_effective_values
from src.records import event_participation_matched_duplicate_guard
from src.records import FIELD_TYPE_OPTIONS
from src.records import IMPORT_DISCIPLINE_COLUMN
from src.records import REQUIRED_IMPORT_COLUMNS
from src.records import select_export_custom_fields
from src.records import select_template_custom_fields
from src.settings import settings
from src.storage.sqlite import dedup_discipline
from src.storage.sqlite import SQLiteAdapter
from src.web import clean_str
from src.web import forbidden
from src.web import form_key_present
from src.web import format_date_range
from src.web import get_auth_user
from src.web import get_flash_args
from src.web import get_form_value
from src.web import get_param
from src.web import get_storage
from src.web import jinja_env
from src.web import unauthorized

logger = logging.getLogger(__name__)


def validate_import_columns(df: pd.DataFrame):
    missing_columns = [column for column in REQUIRED_IMPORT_COLUMNS if column not in df.columns]
    if missing_columns:
        raise ValueError(f'Missing required columns: {", ".join(missing_columns)}')


# Резолвер ФИО (/api/athletes/search): лимит вариантов на запрос ПО KIND —
# до 8 карточек Students и до 8 легаси-подсказок. Карточки идут раньше, но
# легаси-квота не вытесняется (см. search_athletes).
ATHLETE_SEARCH_LIMIT = 8


# Композиция главной по прототипу 02: серверные фильтры и пагинация реестра.
# per_page — из фиксированного набора; текущее поведение продукта показывало
# весь список сразу, поэтому дефолт — максимум набора.
INDEX_PER_PAGE_OPTIONS: Sequence[int] = (10, 25, 50, 100)


INDEX_DEFAULT_PER_PAGE = 100


INDEX_STATUS_FILTERS: Sequence[tuple[str, str]] = (
    ('approved', 'подтверждена'),
    ('pending', 'на проверке'),
    ('rejected', 'отклонена'),
)


INDEX_STATUS_KEYS = {key for key, _ in INDEX_STATUS_FILTERS}


def parse_index_multi_values(args: dict, key: str) -> list[str]:
    """Повторённые GET-параметры мульти-фильтра главной (institute/group/
    sport/level): точные значения справочников как есть, каждое strip,
    пустые отбрасываются, дубли схлопываются с сохранением порядка первого
    вхождения. Одиночное значение — список из одного (легаси-URL)."""
    values: list[str] = []
    seen: set[str] = set()
    for raw_value in args.get(key, []):
        value = str(raw_value).strip()
        if value and value not in seen:
            seen.add(value)
            values.append(value)
    return values


def parse_index_filters(args: dict) -> tuple[dict[str, object], str | None]:
    """Серверные фильтры главной (прототип 02) из GET-параметров.

    ФИО — подстрока (casefold на стороне storage); институт/группа/вид
    спорта/уровень — точные значения справочников, повторённые параметры =
    мульти-выбор (OR внутри фильтра, AND между фильтрами); даты —
    дд.мм.гггг (тот же формат, что в фильтрах отчёта). Возвращает сырые
    значения (для переотображения в форме) и текст ошибки при некорректной
    дате. Статус обрабатывает маршрут — он ролевой.
    """
    filters: dict[str, object] = {'name': (get_param(args, 'name') or '').strip()}
    for key in ('institute', 'group', 'sport', 'level'):
        filters[key] = parse_index_multi_values(args, key)
    for key in ('date_from', 'date_to'):
        raw_value = (get_param(args, key) or '').strip()
        if raw_value:
            try:
                datetime.strptime(raw_value, settings.date_format)
            except ValueError:
                return {}, f'Некорректная дата в фильтре ({key}): {raw_value}'
        filters[key] = raw_value
    filters['status'] = ''
    return filters, None


def parse_index_page(args: dict) -> int:
    """Номер страницы главной; пусто/нечисло — первая, меньше 1 — первая."""
    try:
        return max(1, int(get_param(args, 'page') or 1))
    except ValueError:
        return 1


def parse_index_per_page(args: dict) -> int:
    """Размер страницы главной; вне допустимого набора — дефолт."""
    raw_value = get_param(args, 'per_page') or ''
    try:
        value = int(raw_value)
    except ValueError:
        return INDEX_DEFAULT_PER_PAGE
    return value if value in INDEX_PER_PAGE_OPTIONS else INDEX_DEFAULT_PER_PAGE


def pager_items(current: int, total: int) -> list[dict]:
    """Номера страниц пагинации с разрывами: края (первая/последняя) и
    окно ±1 вокруг текущей, между непоследовательными номерами — «…»."""
    if total <= 0:
        return []
    wanted = {1, total}
    for offset in (-1, 0, 1):
        candidate = current + offset
        if 1 <= candidate <= total:
            wanted.add(candidate)
    items: list[dict] = []
    previous: int | None = None
    for number in sorted(wanted):
        if previous is not None and number - previous > 1:
            items.append({'gap': True})
        items.append({'page': number})
        previous = number
    return items


def build_index_query(filters: dict[str, object], per_page: int) -> str:
    """GET-параметры главной без page: для ссылок пагинации (page добавит
    шаблон/маршрут) и селекта «Показывать по» (он ставит per_page сам).
    Пустые значения и дефолтный per_page в URL не попадают — ссылки чище.
    doseq: мульти-фильтры — повторённые ключи (?institute=A&institute=B),
    без doseq пагинация и «Показывать по» теряли бы часть выбора."""
    params = {key: value for key, value in filters.items() if value}
    if per_page != INDEX_DEFAULT_PER_PAGE:
        params['per_page'] = str(per_page)
    return urlencode(params, doseq=True)


def index_multi_control(
    key: str,
    label: str,
    options: Sequence[str],
    selected: Sequence[str],
    institutes_by_group: dict[str, Sequence[str]] | None = None,
) -> dict:
    """Контекст одного мульти-фильтра главной (карточка фильтров).

    Опции рендера = справочник ∪ текущие выбранные: unknown-значение (его
    нет в справочнике — например, скрытое или удалённое) не ошибка и даёт
    0 записей, но обязано переотображаться выбранным. У опций групп —
    институты справочника (data-institutes для JS-сужения групп)."""
    known = set(options)
    values = [*options, *(value for value in selected if value not in known)]
    selected_set = set(selected)
    option_items = [
        {
            'value': value,
            'checked': value in selected_set,
            'unknown': value not in known,
            'institutes': list((institutes_by_group or {}).get(value, ())),
        }
        for value in values
    ]
    return {'key': key, 'label': label, 'options': option_items, 'selected': list(selected)}


def build_link_bindings(custom_fields: Sequence[CustomField]) -> dict[str, str]:
    """№24: карта «целевая колонка → ключ link-поля» для рендера таблицы.

    Целевая колонка показывает значение своей ячейки как гиперссылку на
    значение link-поля ЭТОЙ записи (если ссылка заполнена и это http/https).
    """
    bindings: dict[str, str] = {}
    for field in custom_fields:
        if field.field_type == 'url' and field.link_target and field.active:
            bindings[field.link_target] = field.key
    return bindings


def build_index_course_cells(competitions: Sequence[Competition]) -> dict[str, dict]:
    """Course/Education Phase A: ячейки колонки «Курс» реестра с производным
    курсом по снимку года поступления.

    У записи есть admission_year → показывается производный курс НА ДАТУ
    СОРЕВНОВАНИЯ (ok — число; future — «—» с подсказкой). Записи без снимка
    в словарь не попадают: шаблон показывает легаси competition.course как
    раньше (включая текстовые значения вроде «Выпускник 2025/26»).
    Расчёт живёт только в src.education.derive_course.
    """
    cells: dict[str, dict] = {}
    for competition in competitions:
        if competition.admission_year is None:
            continue
        derived = derive_course(competition.admission_year, competition.date.date())
        if derived['status'] == 'ok':
            cells[competition.record_id] = {'value': str(derived['course'])}
        elif derived['status'] == 'future':
            cells[competition.record_id] = {
                'value': '—',
                'title': 'Год поступления позже даты соревнования',
            }
    return cells


# Конфликт-режим импорта (№6/№8, docs/data-model-decisions.md «Конфликт-режим
# импорта: очередь на подтверждение»): разбора строка-кандидат показывается
# рядом с похожей существующей записью; решение — принять/пропустить.
QUEUE_FIELD_ORDER: Sequence[str] = (
    'student_name',
    'student_sex',
    'institute',
    'group',
    'course',
    'sport',
    'date',
    'date_to',
    'level',
    'name',
    'position',
    # Event Model, P5a: дисциплина кандидата видна при разборе конфликта
    # (разные дисциплины — не дубль) и участвует в diff аудита replace.
    'discipline',
)


# «discipline» НЕ добавляется в BASE_FIELD_LABELS: тот словарь уходит в
# настройки базовых полей (admin_fields_page), где дисциплины нет.
QUEUE_FIELD_LABELS: dict[str, str] = {**BASE_FIELD_LABELS, 'discipline': 'Дисциплина'}


def queue_candidate_view(payload: dict) -> list[tuple[str, str]]:
    """Поля кандидата из очереди как пары (название, значение) для страницы."""
    view = []
    for key in QUEUE_FIELD_ORDER:
        value = payload.get(key)
        if key in ('date', 'date_to') and value:
            try:
                value = format_date_range(
                    datetime.fromisoformat(payload['date']),
                    datetime.fromisoformat(payload['date_to']) if payload.get('date_to') else None,
                )
            except (TypeError, ValueError):
                pass
        view.append((QUEUE_FIELD_LABELS.get(key, key), '' if value is None else str(value)))
    return view


def queue_edit_values(payload: dict) -> dict:
    """Предзаполнение инлайн-формы ручного решения: значения кандидата,
    даты — в формате дд.мм.гггг (парсер ручного ввода понимает и диапазон
    «25-27.06.2026» в одном поле)."""
    date_display = ''
    if payload.get('date'):
        try:
            date_display = format_date_range(
                datetime.fromisoformat(payload['date']),
                datetime.fromisoformat(payload['date_to']) if payload.get('date_to') else None,
            )
        except (TypeError, ValueError):
            date_display = ''
    return {
        'student_name': '' if payload.get('student_name') is None else str(payload.get('student_name')),
        'student_sex': '' if payload.get('student_sex') is None else str(payload.get('student_sex')),
        'institute': '' if payload.get('institute') is None else str(payload.get('institute')),
        'group': '' if payload.get('group') is None else str(payload.get('group')),
        'course': '' if payload.get('course') is None else str(payload.get('course')),
        'sport': '' if payload.get('sport') is None else str(payload.get('sport')),
        'date': date_display,
        'level': '' if payload.get('level') is None else str(payload.get('level')),
        'name': '' if payload.get('name') is None else str(payload.get('name')),
        'position': '' if payload.get('position') is None else str(payload.get('position')),
        # Event Model, P5a: легаси-payload без ключа даёт пустую строку.
        'discipline': '' if payload.get('discipline') is None else str(payload.get('discipline')),
    }


def queue_field_changes(old: Competition, new: Competition) -> dict:
    """Изменения полей old→new для аудита замены существующей записи."""
    changes = {}
    for key in QUEUE_FIELD_ORDER:
        old_value = getattr(old, key, None)
        new_value = getattr(new, key, None)
        if key in ('date', 'date_to'):
            old_value = old_value.isoformat() if old_value else None
            new_value = new_value.isoformat() if new_value else None
        if old_value != new_value:
            changes[key] = {
                'old': '' if old_value is None else str(old_value),
                'new': '' if new_value is None else str(new_value),
            }
    return changes


def resolve_queue_entry(request: Request, entry_id: str):
    """Общая проверка действия над записью очереди: id и статус pending."""
    try:
        numeric_id = int(entry_id)
    except ValueError:
        return None, text(body='Invalid queue entry id', status=400)
    entry = get_storage(request.app).get_import_queue_entry(numeric_id)
    if entry is None:
        return None, text(body='Queue entry not found', status=404)
    if entry['status'] != 'pending':
        return None, text(body='Queue entry already resolved', status=409)
    return entry, None


def apply_template_date_format(buffer: BytesIO, columns: Sequence[str]) -> None:
    from openpyxl import load_workbook

    if 'Дата' not in columns:
        buffer.seek(0)
        return
    date_column_letter = get_column_letter(columns.index('Дата') + 1)
    buffer.seek(0)
    workbook = load_workbook(buffer)
    sheet = workbook.worksheets[0]
    for cell in sheet[date_column_letter]:
        cell.number_format = 'DD.MM.YYYY'
    sheet.column_dimensions[date_column_letter].width = 14
    buffer.seek(0)
    workbook.save(buffer)
    buffer.seek(0)


def import_record_value_is_empty(value) -> bool:
    """Пустая ячейка импорта: None/NaN или пробельная строка."""
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return str(value).strip() == ''


def autofill_record_from_known_athlete(
    record: dict,
    storage: SQLiteAdapter,
    cache: dict[str, dict],
) -> None:
    """№23доп (docs/feedback-live.md): автозаполнение пустых полей строки импорта.

    Для строки с пустыми Пол/Институт/Группа/Курс — данные известного атлета
    (профиль или последняя запись, точное ФИО). Заполненные поля строки НЕ
    перезаписываются: запись — исторический факт. Заполняется ДО
    build_competition и подсчёта дублей: дубль-ключ считается после
    автозаполнения, и пустой Курс проходит валидацию только после неё.
    """
    name = clean_str(record.get('ФИО', ''))
    if not name:
        return
    key = name.lower()
    if key not in cache:
        cache[key] = storage.find_athlete_fields(name)
    known = cache[key]
    if not known:
        return
    for column, result_key in (
        ('Пол', 'sex'),
        ('Институт', 'institute'),
        ('Группа', 'group'),
    ):
        if import_record_value_is_empty(record.get(column)) and known.get(result_key):
            record[column] = known[result_key]
    if import_record_value_is_empty(record.get('Курс')) and known.get('course'):
        record['Курс'] = known['course']


def build_import_competitions(df: pd.DataFrame, custom_fields, storage: SQLiteAdapter) -> list[Competition]:
    competitions = []
    known_athletes: dict[str, dict] = {}
    for index, row in df.iterrows():
        record = row.to_dict()
        try:
            autofill_record_from_known_athlete(record, storage, known_athletes)
            competitions.append(build_competition(record, custom_fields=custom_fields, storage=storage))
        except (TypeError, ValueError) as exc:
            raise ValueError(f'строка {index + 2}: {exc}') from exc
    return competitions


def split_import_competitions(
    competitions: Sequence[Competition],
    existing_competitions: Iterable[Competition],
) -> tuple[list[Competition], int, list[tuple[Competition, int | None]]]:
    """Разложить строки импорта на вставку / дубли / конфликты (№6/№8).

    Точный дубль (ФИО+дата+вид спорта+название) — пропуск как раньше.
    Похожая строка (совпадают ФИО+дата_начала, полный ключ другой) —
    НЕ вставляется и попадает в очередь: (конфликт, matched_record_id)
    похожей существующей записи. Конфликты между строками одного файла —
    тоже в очередь: первая строка вставляется, похожая на неё вторая —
    в очередь с matched_record_id=None (похожая запись появится после
    вставки первой).
    """
    existing_by_full_key = {competition_duplicate_key(comp): comp for comp in existing_competitions}
    existing_partial_keys = {competition_partial_key(comp) for comp in existing_competitions}
    new_competitions = []
    conflicts: list[tuple[Competition, int | None]] = []
    skipped_duplicates = 0
    seen_full_keys = set()
    inserted_partial_keys = set()
    for competition in competitions:
        full_key = competition_duplicate_key(competition)
        if full_key in existing_by_full_key or full_key in seen_full_keys:
            skipped_duplicates += 1
            continue
        partial_key = competition_partial_key(competition)
        if partial_key in existing_partial_keys or partial_key in inserted_partial_keys:
            matched = None
            if partial_key in existing_partial_keys:
                matched = next(
                    (
                        int(comp.record_id)
                        for comp in existing_competitions
                        if competition_partial_key(comp) == partial_key and comp.record_id
                    ),
                    None,
                )
            conflicts.append((competition, matched))
            continue
        seen_full_keys.add(full_key)
        inserted_partial_keys.add(partial_key)
        new_competitions.append(competition)
    return new_competitions, skipped_duplicates, conflicts


def build_import_summary(inserted: int, skipped_duplicates: int, conflicts: int, *, is_admin: bool) -> str:
    """Сводка импорта реестра (IA-редизайн 2026-09): admin видит адресата
    разбора, editor — что спорные строки переданы администратору (очередь
    ему недоступна)."""
    summary = f'Импортировано записей: {inserted}'
    if skipped_duplicates:
        summary += f'. Пропущено дублей: {skipped_duplicates}'
    if conflicts:
        if is_admin:
            summary += f'. На подтверждение: {conflicts}.'
        else:
            summary += f'. На подтверждение: {conflicts} — передано администратору.'
    return summary


def build_attachments_by_record(all_attachments: list[dict]) -> dict[int, list[dict]]:
    grouped: dict[int, list[dict]] = {}
    for attachment in all_attachments:
        grouped.setdefault(attachment['record_id'], []).append(attachment)
    return grouped


# --- P7: presence-based правка participation-полей записи (/competition/<id>) ---


def competition_update_discipline_collision_keys(custom_fields: Sequence[CustomField]) -> set[str]:
    """Ключи кастомов с label-коллизией «Дисциплина» (не «Результат»: тот
    живёт только в extra_data и восстанавливается общим циклом)."""
    return event_participation_collision_custom_keys(custom_fields) & {
        field.key for field in custom_fields if field.label.strip().casefold() == 'дисциплина'
    }


def competition_update_custom_values(
    request: Request,
    existing: Competition | None,
    custom_fields: Sequence[CustomField],
) -> tuple[dict[str, str], bool]:
    """label → значение кастома для правки записи + флаг «дисциплина пришла
    коллидирующим кастомом».

    P7 presence-based: custom__<key> отсутствует в форме — значение
    восстанавливается из existing.extra_data (лечит затирание кастомов
    правкой со страницы события, которая их не шлёт, и 400 «Поле
    обязательно» на required-кастоме). Коллизия label «Дисциплина»:
    непустой кастом из формы — это и есть значение дисциплины (двойная
    запись base+extra, P5a); ПУСТОЙ базовую колонку не затирает (легаси-
    строки, где база заполнена, а кастом пуст, правка главной не должна
    очищать), восстановление — отдельно, эффективным значением (см.
    competition_update_discipline_value).
    """
    collision_keys = competition_update_discipline_collision_keys(custom_fields)
    values: dict[str, str] = {}
    discipline_from_form_custom = False
    for field in custom_fields:
        form_key = f'custom__{field.key}'
        # Присутствие — включая пустое значение (очистка): form_key_present
        # читает сырое тело, request.form пустые значения теряет.
        field_in_form = form_key_present(request, form_key)
        if field_in_form and field.key in collision_keys:
            form_value = get_form_value(request, form_key)
            if clean_str(form_value):
                values[field.label] = form_value
                discipline_from_form_custom = True
            continue
        if field_in_form:
            values[field.label] = get_form_value(request, form_key)
        elif field.key not in collision_keys and existing is not None:
            values[field.label] = existing.extra_data.get(field.key, '')
    return values, discipline_from_form_custom


def competition_effective_discipline(competition_like: dict, custom_fields: Sequence[CustomField]) -> str:
    """Эффективная дисциплина записи/снимка (базовая колонка ?? legacy-фолбэк
    в extra_data коллидирующего кастома) — единая форма для существующей
    записи и пересобранной формы."""
    return event_participation_effective_values(competition_like, custom_fields)[0]


def competition_update_discipline_value(
    request: Request,
    existing: Competition | None,
    custom_fields: Sequence[CustomField],
    discipline_from_form_custom: bool,
) -> str:
    """Значение колонки «Дисциплина» для правки: 'discipline' из формы (пусто
    = NULL, осознанная очистка легальна — инпут правки со страницы события)
    > непустой коллидирующий кастом из формы > эффективное значение
    existing (путь O1 у главной тот же)."""
    if form_key_present(request, 'discipline'):
        return get_form_value(request, 'discipline')
    if existing is not None and not discipline_from_form_custom:
        return competition_effective_discipline(
            {'discipline': existing.discipline, 'result': None, 'extra_data': existing.extra_data},
            custom_fields,
        )
    return ''


def competition_update_result_value(request: Request, existing: Competition | None) -> str | None:
    """Результат после правки: 'result' из формы (пусто = NULL) или
    существующий, если форма поле не присылала (build_competition результат
    не читает — присвоение после build)."""
    if form_key_present(request, 'result'):
        return clean_str(get_form_value(request, 'result')) or None
    return existing.result if existing is not None else None


def competition_update_participation_guard_error(
    storage: SQLiteAdapter,
    existing: Competition,
    competition: Competition,
    custom_fields: Sequence[CustomField],
    record_id: int,
) -> str | None:
    """P7 matched-guard правки участия: смена дисциплины связанной с событием
    записи (эффективное значение с legacy-фолбэком изменилось) не должна
    нарушать инвариант «карточка + дисциплина уникальны в событии». Сама
    запись исключается (exclude_record_id): повтор собственной дисциплины —
    не конфликт. Путь только fetch (инлайн-правки главной/события): конфликт
    — 400 с текстом guard'а, клиент показывает его как есть."""
    new_discipline = competition_effective_discipline(
        {'discipline': competition.discipline, 'result': None, 'extra_data': competition.extra_data},
        custom_fields,
    )
    old_discipline = competition_effective_discipline(
        {'discipline': existing.discipline, 'result': None, 'extra_data': existing.extra_data},
        custom_fields,
    )
    if dedup_discipline(new_discipline) == dedup_discipline(old_discipline):
        return None
    event = storage.get_calendar_event(existing.calendar_event_id)
    if event is None:
        return None
    competition.student_ref_id = existing.student_ref_id
    return event_participation_matched_duplicate_guard(
        storage,
        event,
        competition,
        exclude_record_id=record_id,
    )


def competition_update_record(
    request: Request,
    existing: Competition | None,
    custom_fields: Sequence[CustomField],
) -> dict:
    """record для build_competition правки записи: базовые поля формы,
    event-owned restore для связанных записей (P2), кастомы и дисциплина
    presence-based (P7)."""
    record = {
        'ФИО': get_form_value(request, 'student_name'),
        'Пол': get_form_value(request, 'student_sex'),
        'Институт': get_form_value(request, 'institute'),
        'Группа': get_form_value(request, 'group'),
        'Вид спорта': get_form_value(request, 'sport'),
        'Дата': get_form_value(request, 'date'),
        'Уровень соревнований': get_form_value(request, 'level'),
        'Название соревнований': get_form_value(request, 'name'),
        'Место': get_form_value(request, 'position'),
        'Курс': get_form_value(request, 'course'),
    }
    custom_values, discipline_from_form_custom = competition_update_custom_values(request, existing, custom_fields)
    record.update(custom_values)
    if existing is not None and existing.calendar_event_id is not None:
        record['Название соревнований'] = existing.name
        record['Вид спорта'] = existing.sport
        record['Уровень соревнований'] = existing.level
        record['Дата'] = format_date_range(existing.date, existing.date_to)
    record['Дисциплина'] = competition_update_discipline_value(
        request, existing, custom_fields, discipline_from_form_custom
    )
    return record


def build_my_document_groups(documents: Sequence[dict]) -> list[dict]:
    """Документы атлета (личный кабинет, Event Documents), сгруппированные по
    событию с сохранением порядка storage (дата события DESC, затем
    sort_order/id документа): [{event_name, event_date_label, documents:
    [{id, calendar_event_id, title}]}]. Дата события — тем же форматером
    (format_date_range), что у календаря/страницы события."""
    groups: list[dict] = []
    group_by_event: dict[int, dict] = {}
    for document in documents:
        event_id = document['calendar_event_id']
        group = group_by_event.get(event_id)
        if group is None:
            date_from = datetime.fromisoformat(document['event_date'])
            date_to = datetime.fromisoformat(document['event_date_to']) if document.get('event_date_to') else None
            group = {
                'event_name': document.get('event_name') or '',
                'event_date_label': format_date_range(date_from, date_to),
                'documents': [],
            }
            group_by_event[event_id] = group
            groups.append(group)
        group['documents'].append(
            {
                'id': document['id'],
                'calendar_event_id': event_id,
                'title': document.get('title') or '',
            }
        )
    return groups


# C901 (осознанное подавление): mccabe суммирует сложность вложенных
# verbatim-хендлеров, перенесённых из main.py без изменений; разбиение
# register() — Architecture v2, не pre-merge gate.
def register(app: Sanic) -> None:  # noqa: C901
    @app.get('/')
    async def index(request: Request):
        storage = get_storage(request.app)
        custom_fields = storage.get_custom_fields()
        # P3 Runtime Identity: scope кабинета атлета — один helper (owner,
        # student_ref_id, легаси-хеши) + текущий режим dual/ref; список,
        # счётчик и has_unapproved согласованы — одни и те же параметры.
        identity = athlete_identity_context(request)
        identity_mode = storage.get_identity_mode() if identity is not None else 'dual'
        owner_filter = identity['user_id'] if identity is not None else None
        profile_hashes = identity['hashes'] if identity is not None else ()
        student_ref_filter = identity['student_ref_id'] if identity is not None else None
        # Event Documents (личный кабинет атлета): документы событий,
        # доступные карточке студента (all_participants — участие,
        # selected_students — маппинг). Не-атлеты и атлеты без стабильной
        # связи с карточкой — None: карточка «Мои документы» не рендерится,
        # лишних запросов нет.
        my_documents = None
        if student_ref_filter is not None:
            my_documents = build_my_document_groups(storage.list_documents_for_student(student_ref_filter))
        # ГТО (личный кабинет атлета): ступени и результаты карточки студента,
        # read-only, без форм. Только атлет со стабильной связью (student_ref_id)
        # и только при непустом списке — паттерн my_documents: пустой список
        # скрывает карточку, не-атлеты и атлеты без связи my_gto не получают.
        my_gto = None
        if student_ref_filter is not None:
            my_gto = build_gto_rows(storage.list_student_gto_records(student_ref_filter)) or None
        args = dict(request.args)

        # Серверная фильтрация реестра (прототип 02): GET-параметры рендерит,
        # фильтры живут в URL — ссылки шарятся. Некорректная дата — явная ошибка.
        index_filters, filter_error = parse_index_filters(args)
        if filter_error is not None:
            return text(body=filter_error, status=400)
        if user_is_moderator(request):
            raw_status = (get_param(args, 'status') or '').strip()
            if raw_status in INDEX_STATUS_KEYS:
                index_filters['status'] = raw_status

        per_page = parse_index_per_page(args)
        filter_kwargs = dict(
            owner_id=owner_filter,
            student_id_hashes=profile_hashes,
            student_ref_id=student_ref_filter,
            identity_mode=identity_mode,
            name=index_filters['name'],
            institute=index_filters['institute'],
            group=index_filters['group'],
            sport=index_filters['sport'],
            level=index_filters['level'],
            review_status=index_filters['status'],
            date_from=index_filters['date_from'],
            date_to=index_filters['date_to'],
        )
        total_count = storage.count_competitions_filtered(**filter_kwargs)
        page_count = max(1, -(-total_count // per_page))
        # Страница вне диапазона зажимается на последнюю (как в журнале аудита).
        page = min(parse_index_page(args), page_count)
        competitions = storage.get_competitions_page(
            limit=per_page,
            offset=(page - 1) * per_page,
            **filter_kwargs,
        )

        def page_url(page_number: int) -> str:
            query = build_index_query(index_filters, per_page)
            separator = '&' if query else ''
            return f'/?{query}{separator}page={page_number}'

        pager = []
        for item in pager_items(page, page_count):
            if item.get('gap'):
                pager.append({'gap': True})
            else:
                number = item['page']
                pager.append({'label': number, 'url': page_url(number), 'active': number == page})

        # Карточка фильтров (мульти-выбор): опции — справочники, группы —
        # плоский список + карта «группа → институты» для data-institutes
        # (JS сужает группы по выбранным институтам). Карточка вне
        # content-wrapper: карта в data-атрибуте таблицы ей недоступна,
        # дублируем на самой карточке.
        levels = storage.get_level_names()
        sport_options = storage.list_catalog('sport')
        institute_options = storage.list_catalog('institute')
        groups_by_institute = storage.get_group_options_by_institute()
        group_options = sorted({group for groups in groups_by_institute.values() for group in groups})
        institutes_by_group: dict[str, list[str]] = {}
        for institute, groups in groups_by_institute.items():
            for group_value in groups:
                institutes_by_group.setdefault(group_value, []).append(institute)

        return await render(
            template_name=jinja_env.get_template('index.html'),
            context={
                'request': request,
                'competitions': competitions,
                'custom_fields': custom_fields,
                # №24: карта «целевая колонка → ключ link-поля» для рендера.
                'link_bindings': build_link_bindings(custom_fields),
                # Course/Education Phase A: ячейки производного курса по
                # снимку года поступления (нет снимка — легаси course).
                'course_cells': build_index_course_cells(competitions),
                'admin_custom_fields': storage.get_custom_fields(include_inactive=True),
                'field_type_options': FIELD_TYPE_OPTIONS,
                'can_write': user_can_write(request),
                'can_import': user_is_moderator(request),
                'is_moderator': user_is_moderator(request),
                'is_athlete': user_is_athlete(request),
                'is_admin': user_is_admin(request),
                # Event Documents: карточка «Мои документы» (только атлет со
                # стабильной связью с карточкой студента и непустым списком).
                'my_documents': my_documents,
                # ГТО: read-only карточка кабинета (тот же паттерн).
                'my_gto': my_gto,
                'users': storage.list_users() if user_is_admin(request) else [],
                'user_roles': USER_ROLES,
                'levels': levels,
                'sport_options': sport_options,
                'institute_options': institute_options,
                'groups_by_institute': groups_by_institute,
                # Мульти-фильтры карточки: institute/group/sport/level.
                'index_multis': (
                    index_multi_control('institute', 'Институт', institute_options, index_filters['institute']),
                    index_multi_control('group', 'Группа', group_options, index_filters['group'], institutes_by_group),
                    index_multi_control('sport', 'Вид спорта', sport_options, index_filters['sport']),
                    index_multi_control('level', 'Уровень', levels, index_filters['level']),
                ),
                'index_groups_by_institute': groups_by_institute,
                # «Статус»-колонка видна, если у пользователя есть хоть одна
                # неподтверждённая запись — независимо от страницы и фильтра.
                'has_unapproved': storage.count_competitions_filtered(
                    owner_id=owner_filter,
                    student_id_hashes=profile_hashes,
                    student_ref_id=student_ref_filter,
                    identity_mode=identity_mode,
                    unapproved_only=True,
                )
                > 0,
                'attachments_by_record': build_attachments_by_record(storage.get_attachments()),
                'admin_levels': storage.list_levels() if user_is_admin(request) else [],
                'current_username': (get_auth_user(request) or {}).get('username'),
                'admin_message': get_param(args, 'admin_message'),
                'admin_error': get_param(args, 'admin_error'),
                # Композиция главной (прототип 02): фильтры, счётчик, пагинация.
                'index_filters': index_filters,
                'status_filter_options': INDEX_STATUS_FILTERS,
                'index_filter_query': urlencode(
                    {key: value for key, value in index_filters.items() if value}, doseq=True
                ),
                'index_is_filtered': any(index_filters.values()),
                'total_count': total_count,
                'shown_count': len(competitions),
                'page': page,
                'page_count': page_count,
                'per_page': per_page,
                'per_page_options': INDEX_PER_PAGE_OPTIONS,
                'pager_items': pager,
                'prev_url': page_url(page - 1) if page > 1 else None,
                'next_url': page_url(page + 1) if page < page_count else None,
            },
        )

    @app.get('/admin/import')
    async def admin_import_page(request: Request):
        auth_error = require_moderator(request)
        if auth_error is not None:
            return auth_error
        storage = get_storage(request.app)
        return await render(
            template_name=jinja_env.get_template('admin_import.html'),
            context={
                'request': request,
                'pending_queue_count': storage.count_import_queue('pending'),
                'is_admin': user_is_admin(request),
                **get_flash_args(request),
            },
        )

    @app.get('/admin/import-queue')
    async def admin_import_queue_page(request: Request):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error
        storage = get_storage(request.app)
        entries = []
        for entry in storage.list_import_queue('pending'):
            matched = None
            if entry['matched_record_id']:
                matched = storage.get_competition_by_id(int(entry['matched_record_id']))
            entries.append(
                {
                    **entry,
                    'candidate': queue_candidate_view(entry['payload']),
                    'edit_values': queue_edit_values(entry['payload']),
                    'matched': matched,
                }
            )
        return await render(
            template_name=jinja_env.get_template('admin_import_queue.html'),
            context={
                'request': request,
                'queue_entries': entries,
                **get_flash_args(request),
            },
        )

    @app.post('/admin/import-queue/<entry_id>/accept')
    async def accept_import_queue_entry(request: Request, entry_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error
        entry, error = resolve_queue_entry(request, entry_id)
        if error is not None:
            return error
        storage = get_storage(request.app)
        competition = Competition.model_validate(entry['payload'])

        # Строка вставляется только по явному решению админа; если пока она
        # лежала в очереди, такая запись появилась другим путём — не дублируем.
        duplicate_now = any(
            competition_duplicate_key(comp) == competition_duplicate_key(competition)
            for comp in storage.get_competitions()
        )
        if duplicate_now:
            storage.set_import_queue_status(entry['id'], 'skipped')
            log_audit_event(
                request,
                'import_conflict_resolved',
                {'decision': 'skipped', 'queue_id': entry['id'], 'reason': 'duplicate_already_exists'},
            )
            return redirect(to='/admin/import-queue?admin_message=Запись+уже+существует,+кандидат+пропущен')

        ensure_catalog_values(storage, [competition])
        storage.save_competitions([competition], review_status='approved', owner_id=entry['created_by'])
        storage.set_import_queue_status(entry['id'], 'accepted')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {
                'decision': 'accepted',
                'queue_id': entry['id'],
                'matched_record_id': entry['matched_record_id'],
            },
        )
        return redirect(to='/admin/import-queue?admin_message=Кандидат+принят+и+добавлен+в+реестр')

    @app.post('/admin/import-queue/<entry_id>/skip')
    async def skip_import_queue_entry(request: Request, entry_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error
        entry, error = resolve_queue_entry(request, entry_id)
        if error is not None:
            return error
        storage = get_storage(request.app)
        storage.set_import_queue_status(entry['id'], 'skipped')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {'decision': 'skipped', 'queue_id': entry['id'], 'matched_record_id': entry['matched_record_id']},
        )
        return redirect(to='/admin/import-queue?admin_message=Кандидат+пропущен')

    # Замена существующей записи данными кандидата (№25): историчность обеспечивает
    # аудит old→new по полям; record_id существующей сохраняется, её владелец и
    # review-статус не меняются (update_competition не трогает эти колонки).
    @app.post('/admin/import-queue/<entry_id>/replace')
    async def replace_import_queue_entry(request: Request, entry_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error
        entry, error = resolve_queue_entry(request, entry_id)
        if error is not None:
            return error
        if not entry['matched_record_id']:
            return text(body='Queue entry has no matched record', status=400)
        storage = get_storage(request.app)
        existing = storage.get_competition_by_id(int(entry['matched_record_id']))
        if existing is None:
            return text(body='Matched record not found', status=404)

        competition = Competition.model_validate(entry['payload'])
        ensure_catalog_values(storage, [competition])
        changes = queue_field_changes(existing, competition)
        storage.update_competition(int(entry['matched_record_id']), competition)
        storage.set_import_queue_status(entry['id'], 'replaced')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {
                'decision': 'replaced',
                'queue_id': entry['id'],
                'matched_record_id': entry['matched_record_id'],
                'changes': changes,
            },
        )
        return redirect(to='/admin/import-queue?admin_message=Существующая+запись+заменена+данными+кандидата')

    # Ручное решение (№25, №27в): админ правит поля кандидата инлайн, кроме ФИО
    # и даты — они задают конфликт и всегда берутся из payload очереди; после
    # сохранения кандидат вставляется как обычный accepted.
    @app.post('/admin/import-queue/<entry_id>/edit')
    async def edit_import_queue_entry(request: Request, entry_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error
        entry, error = resolve_queue_entry(request, entry_id)
        if error is not None:
            return error
        storage = get_storage(request.app)
        custom_fields = storage.get_custom_fields()
        # №27в: ФИО и дата задают конфликт с существующей записью — из формы не
        # берутся (подмена игнорируется), только из payload очереди. Остальные
        # поля приходят из формы. Дату форматируем в дд.мм.гггг (с диапазоном),
        # так её понимает ручной парсер build_competition.
        payload = entry['payload']
        record = {
            'ФИО': '' if payload.get('student_name') is None else str(payload['student_name']),
            'Пол': get_form_value(request, 'student_sex'),
            'Институт': get_form_value(request, 'institute'),
            'Группа': get_form_value(request, 'group'),
            'Вид спорта': get_form_value(request, 'sport'),
            'Дата': queue_edit_values(payload)['date'],
            'Уровень соревнований': get_form_value(request, 'level'),
            'Название соревнований': get_form_value(request, 'name'),
            'Место': get_form_value(request, 'position'),
            'Курс': get_form_value(request, 'course'),
            # Event Model, P5a: дисциплина правится вместе с остальными полями.
            'Дисциплина': get_form_value(request, 'discipline'),
        }
        try:
            competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
        except (TypeError, ValueError) as exc:
            return text(body=f'Invalid row data: {exc}', status=400)

        # Тот же анти-дубликат, что и у «Принять»: строка вставляется только
        # по явному решению админа.
        duplicate_now = any(
            competition_duplicate_key(comp) == competition_duplicate_key(competition)
            for comp in storage.get_competitions()
        )
        if duplicate_now:
            storage.set_import_queue_status(entry['id'], 'skipped')
            log_audit_event(
                request,
                'import_conflict_resolved',
                {'decision': 'skipped', 'queue_id': entry['id'], 'reason': 'duplicate_already_exists', 'edited': True},
            )
            return redirect(to='/admin/import-queue?admin_message=Запись+уже+существует,+кандидат+пропущен')

        ensure_catalog_values(storage, [competition])
        storage.save_competitions([competition], review_status='approved', owner_id=entry['created_by'])
        storage.set_import_queue_status(entry['id'], 'accepted')
        log_audit_event(
            request,
            'import_conflict_resolved',
            {
                'decision': 'accepted',
                'edited': True,
                'queue_id': entry['id'],
                'matched_record_id': entry['matched_record_id'],
            },
        )
        return redirect(to='/admin/import-queue?admin_message=Кандидат+принят+с+ручными+правками')

    @app.get('/template/empty.xlsx')
    async def export_empty_template(request: Request):
        auth_error = require_moderator(request)
        if auth_error is not None:
            return auth_error

        storage = get_storage(request.app)
        custom_fields = select_template_custom_fields(storage.get_custom_fields())
        # «Дисциплина» — опциональная фиксированная колонка шаблона (P5a):
        # колонки без неё продолжают импортироваться, «Результата» в шаблоне нет.
        columns = [*REQUIRED_IMPORT_COLUMNS, IMPORT_DISCIPLINE_COLUMN, *(field.label for field in custom_fields)]

        df = pd.DataFrame(columns=columns)
        buffer = BytesIO()
        df.to_excel(buffer, index=False)
        apply_template_date_format(buffer, columns)

        now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
        return raw(
            buffer.getvalue(),
            headers={
                'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                'content-disposition': f'attachment; filename="Шаблон_соревнования_{now_str}.xlsx"',
            },
        )

    @app.get('/export/index')
    async def export_index(request: Request):
        if user_is_athlete(request):
            return forbidden(request)
        storage = get_storage(request.app)
        competitions = storage.get_competitions()
        export_custom_fields = select_export_custom_fields(storage.get_custom_fields())
        df = await asyncio.to_thread(build_index_dataframe, competitions, export_custom_fields)

        now_str = datetime.utcnow().strftime('%d-%m-%Y_%H-%M-%S')
        filename = f'Отчет_{now_str}.xlsx'
        buffer = BytesIO()
        await asyncio.to_thread(df.to_excel, buffer, index=False)
        buffer.seek(0)

        return raw(
            buffer.getvalue(),
            headers={
                'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                'content-disposition': f'attachment; filename="{filename}"',
            },
        )

    @app.post('/')
    async def upload(request: Request):
        auth_error = require_moderator(request)
        if auth_error is not None:
            return auth_error

        upload_file = request.files.get('file')
        if not upload_file or not upload_file.body:
            return text(body='No file uploaded', status=400)

        storage = get_storage(request.app)
        custom_fields = storage.get_custom_fields()

        try:
            df = await asyncio.to_thread(pd.read_excel, io=upload_file.body)
            validate_import_columns(df)
        except ValueError as exc:
            return text(body=str(exc), status=400)

        try:
            competitions = await asyncio.to_thread(build_import_competitions, df, custom_fields, storage)
        except (TypeError, ValueError) as exc:
            return text(body=f'Invalid row data: {exc}', status=400)

        existing = await asyncio.to_thread(storage.get_competitions)
        new_competitions, skipped_duplicates, conflicts = split_import_competitions(competitions, existing)

        # Одна транзакция и один COMMIT на весь импорт, в отдельном потоке:
        # построчные коммиты справочников (WAL+fsync ~4 раза на строку) и запись
        # в event loop замораживали приложение на всё время импорта (QA №1).
        # Дубли отсеяны до транзакции; справочники storage пополняет значениями
        # из вставленных записей (решение 2026-09-13 по инциденту rename).
        await asyncio.to_thread(
            storage.import_competitions,
            new_competitions,
            owner_id=get_current_user_id(request),
        )

        # Конфликт-режим импорта (решение 2026-09-13, docs/data-model-decisions.md
        # «Конфликт-режим импорта: очередь на подтверждение»): похожие строки
        # не вставляются молча и не отклоняют импорт — ждут явного решения админа.
        created_by = get_current_user_id(request)
        for conflict, matched_record_id in conflicts:
            await asyncio.to_thread(
                storage.add_import_queue_entry,
                conflict.model_dump(mode='json', by_alias=False),
                matched_record_id=matched_record_id,
                created_by=created_by,
            )

        return text(
            body=build_import_summary(
                len(new_competitions),
                skipped_duplicates,
                len(conflicts),
                is_admin=user_is_admin(request),
            )
        )

    @app.get('/api/athletes/search')
    async def search_athletes(request: Request):
        # №23доп (docs/feedback-live.md): резолвер атлета. Подсказки по ФИО
        # раскрывают персональные данные других студентов — только admin/editor
        # (docs/data-model-decisions.md); атлету и наблюдателю — 403. У атлета
        # своя автоподстановка из собственного профиля (/api/profile).
        #
        # Источник — ОБЪЕДИНЕНИЕ: активные карточки студентов (таблица students,
        # находятся и без истории участий) + легаси-известные атлеты (профили и
        # последние записи). Каждый вариант помечен kind: 'student' | 'legacy';
        # при точном совпадении ФИО (strip + casefold) карточка Student приоритетна
        # — легаси-дубль не показывается, данные берутся из карточки. Разные
        # написания остаются обоими вариантами. Слияния/авто-связывания нет:
        # kind 'student' лишь позволяет вызывающей стороне записать student_ref_id
        # ПОСЛЕ явного выбора человеком.
        #
        # Лимит — ПО KIND (до ATHLETE_SEARCH_LIMIT карточек и до столько же
        # легаси-вариантов): карточки идут раньше, но НЕ вытесняют легаси-подсказки
        # при полном лимите совпадающих карточек — легаси-люди без карточек должны
        # оставаться досягаемыми и в резолвере главной (клиент фильтрует карточки).
        auth_error = require_moderator(request)
        if auth_error is not None:
            return auth_error
        query = str(request.args.get('q', '')).strip()
        if not query:
            return json_response([])

        storage = get_storage(request.app)

        def combined_search() -> list[dict]:
            students = storage.search_student_suggestions(query, limit=ATHLETE_SEARCH_LIMIT)
            legacy = storage.search_athletes(query, limit=ATHLETE_SEARCH_LIMIT)
            student_names = {(item['name'] or '').strip().casefold() for item in students}
            result = [
                {
                    'kind': 'student',
                    'student_id': item['student_id'],
                    'name': item['name'],
                    **{key: item[key] for key in ('sex', 'institute', 'group', 'course') if item.get(key)},
                }
                for item in students
            ]
            result.extend(
                {
                    'kind': 'legacy',
                    'student_id': None,
                    **{key: value for key, value in item.items() if key != 'name' and value},
                    'name': item['name'],
                }
                for item in legacy
                if (item['name'] or '').strip().casefold() not in student_names
            )
            return result

        athletes = await asyncio.to_thread(combined_search)
        return json_response(athletes)

    @app.post('/competition')
    async def add_competition(request: Request):
        auth_error = require_writer(request)
        if auth_error is not None:
            return auth_error

        storage = get_storage(request.app)
        custom_fields = storage.get_custom_fields()
        record = {
            'ФИО': get_form_value(request, 'student_name'),
            'Пол': get_form_value(request, 'student_sex'),
            'Институт': get_form_value(request, 'institute'),
            'Группа': get_form_value(request, 'group'),
            'Вид спорта': get_form_value(request, 'sport'),
            'Дата': get_form_value(request, 'date'),
            'Уровень соревнований': get_form_value(request, 'level'),
            'Название соревнований': get_form_value(request, 'name'),
            'Место': get_form_value(request, 'position'),
            'Курс': get_form_value(request, 'course'),
        }
        record.update({field.label: get_form_value(request, f'custom__{field.key}') for field in custom_fields})

        profile_defaults = get_athlete_profile_defaults(request, storage)
        for form_key, profile_key in (
            ('ФИО', 'student_name'),
            ('Пол', 'student_sex'),
            ('Институт', 'institute'),
            ('Группа', 'group'),
            ('Курс', 'course'),
        ):
            if not str(record.get(form_key, '')).strip() and profile_key in profile_defaults:
                record[form_key] = profile_defaults[profile_key]

        try:
            competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
        except (TypeError, ValueError) as exc:
            return text(body=f'Invalid row data: {exc}', status=400)

        ensure_catalog_values(storage, [competition])
        review_status = 'pending' if user_is_athlete(request) else 'approved'
        storage.save_competitions(
            [competition],
            review_status=review_status,
            owner_id=get_current_user_id(request),
        )
        return redirect(to='/')

    @app.post('/competition/<record_id>')
    async def update_competition(request: Request, record_id: str):
        auth_error = require_writer(request)
        if auth_error is not None:
            return auth_error

        try:
            numeric_id = int(record_id)
        except ValueError:
            return text(body='Invalid record id', status=400)

        storage = get_storage(request.app)
        existing_review = storage.get_competition_review(numeric_id)
        if existing_review is not None and not user_owns_record(request, existing_review):
            return forbidden(request)

        custom_fields = storage.get_custom_fields()
        # P2: снимок существующей записи — источник event-owned полей и
        # discipline/result (форма их не присылает). Записи со ссылкой на
        # событие календаря (calendar_event_id) этими полями не владеют:
        # название/вид спорта/уровень/дата берутся из записи-события (правка —
        # только со страницы события), подделка формы их изменить не может.
        # NULL-записи правятся прежним путём — все поля из формы.
        existing = storage.get_competition_by_id(numeric_id)
        record = competition_update_record(request, existing, custom_fields)

        try:
            competition = build_competition(record, custom_fields=custom_fields, manual_input=True, storage=storage)
        except (TypeError, ValueError) as exc:
            return text(body=f'Invalid row data: {exc}', status=400)

        # Результат — presence-based, как дисциплина (P1/P7): модельные None не
        # должны затирать сохранённые значения, пока форма полем не подтвердила
        # изменение/очистку.
        competition.result = competition_update_result_value(request, existing)

        if existing is not None and existing.calendar_event_id is not None:
            guard_error = competition_update_participation_guard_error(
                storage, existing, competition, custom_fields, numeric_id
            )
            if guard_error is not None:
                return text(body=guard_error, status=400)

        storage.update_competition(numeric_id, competition)
        # Роль решает статус после правки: модератор (admin/editor)
        # подтверждает, атлет отправляет на повторную проверку — независимо
        # от прежнего статуса записи.
        if user_is_moderator(request):
            storage.set_competition_review(numeric_id, 'approved')
            # Правка модератора = авто-подтверждение, фиксируем как record_approved.
            log_audit_event(request, 'record_approved', {'record_id': numeric_id, 'auto': True})
        else:
            storage.set_competition_review(numeric_id, 'pending')
        return redirect(to='/')

    @app.post('/competition/<record_id>/review/<decision>')
    async def review_competition(request: Request, record_id: str, decision: str):
        auth_error = require_moderator(request)
        if auth_error is not None:
            return auth_error

        if decision not in ('approve', 'reject'):
            return text(body='Unknown review decision', status=400)

        try:
            numeric_id = int(record_id)
        except ValueError:
            return text(body='Invalid record id', status=400)

        if get_storage(request.app).get_competition_review(numeric_id) is None:
            return text(body='Record not found', status=404)

        status = 'approved' if decision == 'approve' else 'rejected'
        comment = get_form_value(request, 'comment').strip() if decision == 'reject' else ''
        get_storage(request.app).set_competition_review(numeric_id, status, comment)
        log_audit_event(
            request,
            'record_approved' if decision == 'approve' else 'record_rejected',
            {'record_id': numeric_id, 'comment': comment},
        )
        return redirect(to='/')

    @app.post('/competition/<record_id>/delete')
    async def delete_competition(request: Request, record_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error

        try:
            numeric_id = int(record_id)
        except ValueError:
            return text(body='Invalid record id', status=400)

        storage = get_storage(request.app)
        # ФИО до удаления — в аудит удаления записи (после удаления взять неоткуда).
        student_name = storage.get_competition_student_name(numeric_id)
        # Удаление записи тянет за собой её вложения (M4): сначала транзакция БД
        # (строки вложений + запись), затем файлы — тот же порядок, что у
        # perform_wipe. Файлы удаляются best-effort: остаток каталога — только
        # предупреждение в лог, БД уже согласована.
        storage.delete_competition_with_attachments(numeric_id)
        remove_attachment_record_dirs([numeric_id])
        if (files_dir() / str(numeric_id)).exists():
            logger.warning('Attachment directory of record %s still exists after record deletion', numeric_id)
        log_audit_event(request, 'record_deleted', {'record_id': numeric_id, 'student_name': student_name})
        return redirect(to='/')

    @app.post('/competition/<record_id>/attachments')
    async def upload_attachment(request: Request, record_id: str):
        auth_error = require_writer(request)
        if auth_error is not None:
            return auth_error

        try:
            numeric_id = int(record_id)
        except ValueError:
            return text(body='Invalid record id', status=400)

        storage = get_storage(request.app)
        review = storage.get_competition_review(numeric_id)
        if review is None:
            return text(body='Record not found', status=404)
        if not user_owns_record(request, review):
            return forbidden(request)

        upload_file = request.files.get('file')
        if not upload_file or not upload_file.body:
            return text(body='No file uploaded', status=400)
        if len(upload_file.body) > ATTACHMENT_MAX_SIZE:
            return text(body='Файл больше 5 МБ', status=400)

        filename = (upload_file.name or 'attachment').rsplit('/', 1)[-1]
        content_type = detect_attachment_type(upload_file.body, filename)
        if content_type is None:
            return text(body='Допустимы только PDF, JPEG и PNG', status=400)

        extension = Path(filename).suffix.lower().lstrip('.') or 'bin'
        stored_name = f'{uuid.uuid4().hex}.{extension}'
        record_dir = attachments_dir() / str(numeric_id)
        record_dir.mkdir(parents=True, exist_ok=True)
        (record_dir / stored_name).write_bytes(upload_file.body)

        storage.create_attachment(
            record_id=numeric_id,
            filename=filename,
            stored_name=stored_name,
            content_type=content_type,
            size=len(upload_file.body),
            uploaded_by=get_current_user_id(request),
        )
        return text(body='Файл загружен')

    @app.get('/attachment/<attachment_id>')
    async def download_attachment(request: Request, attachment_id: str):
        if get_auth_user(request) is None:
            return unauthorized(request)

        try:
            numeric_attachment_id = int(attachment_id)
        except ValueError:
            return text(body='Invalid attachment id', status=400)

        storage = get_storage(request.app)
        attachment = storage.get_attachment(numeric_attachment_id)
        if attachment is None:
            return text(body='Not Found', status=404)
        review = storage.get_competition_review(attachment['record_id'])
        if review is not None and not user_owns_record(request, review):
            return forbidden(request)

        file_path = attachments_dir() / str(attachment['record_id']) / attachment['stored_name']
        if not file_path.is_file():
            return text(body='Not Found', status=404)

        return raw(
            await asyncio.to_thread(file_path.read_bytes),
            headers={
                'content-type': attachment['content_type'],
                'content-disposition': f'attachment; filename="{attachment["stored_name"]}"',
            },
        )

    @app.post('/attachment/<attachment_id>/delete')
    async def delete_attachment(request: Request, attachment_id: str):
        auth_error = require_admin(request)
        if auth_error is not None:
            return auth_error

        try:
            numeric_attachment_id = int(attachment_id)
        except ValueError:
            return text(body='Invalid attachment id', status=400)

        storage = get_storage(request.app)
        attachment = storage.get_attachment(numeric_attachment_id)
        if attachment is None:
            return text(body='Not Found', status=404)

        file_path = attachments_dir() / str(attachment['record_id']) / attachment['stored_name']
        file_path.unlink(missing_ok=True)
        storage.delete_attachment(numeric_attachment_id)
        return redirect(to='/')
